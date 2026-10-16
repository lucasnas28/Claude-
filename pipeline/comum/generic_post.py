"""Company-agnostic improvements to a model built on the standard mask.
usage: generic_post.py IN OUT [--realize]   (--realize turns 1Q26/2Q26 (AH/AI) into actuals)"""
import sys
import openpyxl
from openpyxl.utils import get_column_letter as CL
from openpyxl.formula.translate import Translator
from openpyxl.styles import PatternFill
from openpyxl.workbook.defined_name import DefinedName

IN, OUT = sys.argv[1], sys.argv[2]
REALIZE = '--realize' in sys.argv
wb = openpyxl.load_workbook(IN)
m, val = wb['Model'], wb['Valuation']
ANNUAL = set(range(8, 108, 5))
def tr(f, s, d): return Translator(f, origin=s).translate_formula(d)
log = []

# 1. realized quarters in AH/AI
if REALIZE:
    n = 0
    for dst, srcc in (('AH', 'AC'), ('AI', 'AD')):
        for rr in range(1, 359):
            if rr in (5, 6, 7, 8, 9, 10): continue
            s = m[f'{srcc}{rr}'].value
            if isinstance(s, str) and s.startswith('='):
                nf = tr(s, f'{srcc}{rr}', f'{dst}{rr}')
                if m[f'{dst}{rr}'].value != nf: m[f'{dst}{rr}'] = nf; n += 1
    m['AJ189'] = '=AI183+SUM(AJ187:AJ188)'
    for rr in (228, 229):
        for cc in ('AH', 'AI'):
            m[f'{cc}{rr}'] = None     # buyback / issuance projection is not used in realized quarters
    log.append(f'AH/AI (1Q26, 2Q26) em modo realizado: {n} celulas; AJ189 parte da divida LP real')

# 2. leverage on LTM EBITDA
def ltm(c):
    qs, k = [], c
    while len(qs) < 4 and k >= 4:
        if k not in ANNUAL: qs.append(CL(k))
        k -= 1
    return qs if len(qs) == 4 else None
for c in range(4, 108):
    if c in ANNUAL: continue
    L = CL(c); q = ltm(c)
    for rr, num in ((300, f'SUM({L}178,{L}179,{L}180,{L}183,{L}201,{L}204)'),
                    (302, f'(SUM({L}178,{L}179,{L}180,{L}183,{L}201,{L}204)-{L}110)'),
                    (303, f'(SUM({L}178,{L}179,{L}183,{L}204)-{L}110)')):
        if m.cell(rr, c).value is None: continue
        m.cell(rr, c).value = f'=IFERROR({num}/({"+".join(x + "35" for x in q)}),"")' if q else '=""'
m['C300'] = 'TOTAL DEBT / EBITDA LTM'; m['C302'] = 'NET DEBT / EBITDA LTM (com arrendamentos)'; m['C303'] = 'NET DEBT (EX IFRS 16) / EBITDA LTM'
log.append('alavancagem sobre EBITDA dos ultimos 12 meses (linhas 300, 302, 303)')

# 3. ROIC block: EBIT from row 39 and annual columns
for c in range(4, 108):
    if c in ANNUAL: continue
    L = CL(c)
    if m.cell(314, c).value == f'={L}96': m.cell(314, c).value = f'={L}39'
for c in sorted(ANNUAL):
    for rr in range(311, 322):
        if m.cell(rr, c).value is None and isinstance(m.cell(rr, c - 1).value, str):
            m.cell(rr, c).value = tr(m.cell(rr, c - 1).value, f'{CL(c - 1)}{rr}', f'{CL(c)}{rr}')
            m.cell(rr, c)._style = m.cell(rr, c - 1)._style
log.append('bloco ROIC: EBIT da linha 39 e colunas anuais')

# 4. Valuation: ROIC formulas typed as text, scenario selector, broken global names
for c in 'GHIJKL':
    if val[f'{c}50'].value == '#VALUE!': val[f'{c}50'] = tr('=R25/AVERAGE(Q46,R46)', 'R50', f'{c}50')
if 'scenario' not in val.defined_names:
    if val['E8'].value is None:
        val['E8'] = 2; val['E8'].fill = PatternFill('solid', fgColor='FFFFF2CC'); val['F8'] = '1 = Bull, 2 = Base, 3 = Bear'
    val.defined_names['scenario'] = DefinedName('scenario', attr_text='Valuation!$E$8', localSheetId=wb.sheetnames.index('Valuation'))
for nme in ('wacc', 'risk_free_rate', 'risk_premium', 'relevered_beta', 'totalcap'):
    if nme in wb.defined_names and wb.defined_names[nme].attr_text == '#REF!': del wb.defined_names[nme]
log.append('Valuation: ROIC G50:L50, seletor de cenario (E8), nomes globais quebrados removidos')

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('\n'.join(log))
