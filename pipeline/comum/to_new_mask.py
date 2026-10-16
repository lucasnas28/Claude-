"""Converte um modelo montado na máscara antiga (Modelo_2.0_proprio_Padrao_AJUSTADO-2028_corrigido) para a máscara
nova (Modelo_Base_credito_research_AJUSTADO_ratio): aba Control, Inputs com uma coluna de código e valores em mil,
Inputs (2) com flag de realizado e Model com trimestres históricos ligados direto na Inputs (2).

uso: to_new_mask.py MODELO_ANTIGO.xlsx (versão com fórmulas, ex. *_d.xlsx ou *_final.xlsx) MASCARA_NOVA.xlsx SAIDA.xlsx MOEDA ULTIMO_TRI
     MOEDA = R$ ou US$; ULTIMO_TRI = data do último trimestre realizado, ex. 2026-06-30

O que é feito:
- Inputs: cada linha de origem da Inputs antiga vira linha(s) da nova. Linha com um código e sinal +1 mantém o código;
  linha com sinal -1 ou vários códigos vira uma linha bruta sem código (fonte) + uma linha de lançamento por código
  (=sinal × linha bruta), rotulada. Linhas de memória, derivadas e checks vêm junto, com as fórmulas reescritas.
  Valores monetários passam de unidades para mil (a Inputs (2) da máscara divide por 1.000).
- Inputs (2): flag de realizado = 1 nos trimestres com dado; linhas novas para o código 44 (lucro reportado, base do
  check de DRE) e para os códigos 45+ (drivers), no mesmo padrão de fórmula da máscara.
- Model: abre o mesmo número de linhas do bloco de drivers (linha 18) e copia o bloco; linha 4 = lucro reportado.
- Control: companhia, ticker, moeda, unidade e último trimestre realizado.
O período 1T20-4T20 fica de fora: a grade da máscara nova começa no 1T21."""
import sys, os, re, copy, shutil, subprocess, tempfile, collections, datetime as dt
import openpyxl
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
from openpyxl.formula.translate import Translator

OLD, NEWM, OUT, CUR, LAST = sys.argv[1:6]
LAST = dt.date.fromisoformat(LAST)
HERE = os.path.dirname(os.path.abspath(__file__))
ow = openpyxl.load_workbook(OLD)
oi, o2, om = ow['Inputs'], ow['Inputs (2)'], ow['Model']
N = om.max_row - 358                                   # linhas do bloco de drivers (máscara antiga tem 358 linhas)
while N > 0 and all(om.cell(18 + N - 1, c).value is None for c in range(1, 40)): N -= 1

# períodos
PER_OLD = {}
for c in range(11, 80):
    h = oi.cell(4, c).value
    if isinstance(h, str) and re.fullmatch(r'[1-4]Q\d\d', h): PER_OLD[h] = c
tmp = tempfile.mkdtemp(); step = os.path.join(tmp, 's.xlsx')
shutil.copyfile(NEWM, step)
if N > 0:
    nxt = os.path.join(tmp, 's2.xlsx')
    subprocess.run([sys.executable, os.path.join(HERE, 'shift_rows.py'), step, nxt, 'Model', '18', str(N)], check=True, stdout=subprocess.DEVNULL)
    step = nxt
nw = openpyxl.load_workbook(step)
ni, n2, nm, ctl = nw['Inputs'], nw['Inputs (2)'], nw['Model'], nw['Control']
PER_NEW = {ni.cell(2, c).value: c for c in range(4, 34) if isinstance(ni.cell(2, c).value, str) and re.fullmatch(r'[1-4]Q\d\d', ni.cell(2, c).value)}
PERS = [p for p in PER_NEW if p in PER_OLD]
COLMAP = {L(PER_OLD[p]): L(PER_NEW[p]) for p in PERS}

# ------------------------------------------------------------------ Inputs: linhas antigas
R0 = 5
rows = []
for r in range(R0, oi.max_row + 1):
    desc = oi.cell(r, 8).value
    if desc is None and all(oi.cell(r, c).value is None for c in range(1, 7)):
        if any(oi.cell(r, PER_OLD[p]).value is not None for p in PER_OLD): pass
        else:
            rows.append(None); continue
    if isinstance(desc, str) and desc.startswith('CONVENCOES'): break
    pairs = [(oi.cell(r, 1 + 2 * i).value, oi.cell(r, 2 + 2 * i).value) for i in range(3)]
    pairs = [(c, s) for c, s in pairs if isinstance(c, int)]
    rows.append(dict(r=r, desc=desc, unid=oi.cell(r, 7).value, origem=oi.cell(r, 10).value, en=oi.cell(r, 9).value, pairs=pairs))
while rows and rows[-1] is None: rows.pop()

STD = {n2.cell(rr, 2).value: n2.cell(rr, 3).value for rr in range(4, 48) if isinstance(n2.cell(rr, 2).value, int)}
# limpa o modelo de linhas da máscara (4..44) e escreve a partir da 4
for rr in range(4, ni.max_row + 1):
    for c in range(1, 34): ni.cell(rr, c).value = None
plan = []          # (old_row or None, kind, code, sign)
for it in rows:
    if it is None: plan.append((None, 'blank', None, None)); continue
    ps = it['pairs']
    if len(ps) == 1 and ps[0][1] == 1: plan.append((it['r'], 'main', ps[0][0], 1))
    else:
        plan.append((it['r'], 'raw', None, None))
        for c, s in ps: plan.append((it['r'], 'post', c, s))
ROWOF = {}         # linha antiga -> linha nova (bruta/principal)
nr = 4
for o, kind, code, s in plan:
    if kind in ('main', 'raw'): ROWOF[o] = nr
    nr += 1
LASTNEW = nr - 1
SCAN = max(100, ((LASTNEW + 40) // 50 + 1) * 50)
FAT = oi['B1'].value if isinstance(oi['B1'].value, (int, float)) else 1e-6   # fator antigo para milhões
SCALE = {'moeda': FAT * 1e3, 'acoes': FAT * 1e3}  # -> mil (a Inputs (2) da máscara divide por 1.000)

def conv_formula(f):
    def sub(mm):
        cl, rw = mm.group(2), int(mm.group(4))
        ncl = COLMAP.get(cl, cl); nrw = ROWOF.get(rw, rw)
        return f'{mm.group(1)}{ncl}{mm.group(3)}{nrw}'
    f = re.sub(r'(\$?)([A-Z]{1,3})(\$?)(\d+)', sub, f)
    return f.replace(',-5)', ',-2)')
nr = 4
for o, kind, code, s in plan:
    if kind == 'blank': nr += 1; continue
    it = next(x for x in rows if x and x['r'] == o)
    sc = SCALE.get(it['unid'], 1e3)                   # não monetário: ×1000 para a Inputs (2) devolver o valor original
    src = f"{it['desc']}" if it['desc'] else ''
    if kind in ('main', 'raw'):
        ni.cell(nr, 1).value = (src + (f" · {it['origem']}" if it['origem'] else '') + (f" · unid. {it['unid']}" if it['unid'] else '')).strip(' ·')
        if kind == 'main': ni.cell(nr, 2).value = code; ni.cell(nr, 3).value = STD.get(code, src[:60])
        else: ni.cell(nr, 3).value = '[linha bruta, sem código; lançada nas linhas abaixo]'
        for p in PERS:
            v = oi.cell(o, PER_OLD[p]).value
            if isinstance(v, str) and v.startswith('='): ni.cell(nr, PER_NEW[p]).value = conv_formula(v)
            elif isinstance(v, (int, float)): ni.cell(nr, PER_NEW[p]).value = round(v * sc, 6)
    else:
        raw = ROWOF[o]
        ni.cell(nr, 1).value = f'(lançamento) código {code}, sinal {s:+d} sobre a linha bruta {raw}: {src[:90]}'
        ni.cell(nr, 2).value = code; ni.cell(nr, 3).value = STD.get(code, f'Driver / código {code}')
        for p in PERS:
            x = L(PER_NEW[p]); ni.cell(nr, PER_NEW[p]).value = f'=IF({x}{raw}="","",{s}*{x}{raw})' if s != 1 else f'=IF({x}{raw}="","",{x}{raw})'
    nr += 1
cur_lbl = 'US$' if CUR.upper().startswith('US') else 'R$'
ni['A1'] = f'Base histórica normalizada ({cur_lbl} mil) — convertida da máscara anterior; linhas de origem auditáveis'
ni['A3'] = 'Fonte: ver coluna A de cada linha'; ni['C3'] = f'Unidade: {cur_lbl} mil (monetário); demais unidades × 1.000'
ni.cell(LASTNEW + 2, 1).value = (f'Linhas de origem: 4 a {LASTNEW}. A Inputs (2) varre as linhas 4 a {SCAN}; as linhas {LASTNEW + 1} a {SCAN} estão livres. '
                                 'Linhas "(lançamento)" aplicam código e sinal sobre a linha bruta indicada (a máscara tem uma só coluna de código).')
for rr in range(4, LASTNEW + 1):
    for p in PERS: ni.cell(rr, PER_NEW[p]).number_format = '#,##0.0;(#,##0.0)'
ni.column_dimensions['A'].width = 90; ni.column_dimensions['C'].width = 45

# ------------------------------------------------------------------ Inputs (2)
for row in n2.iter_rows(min_row=4, max_row=n2.max_row):
    for c in row:
        if isinstance(c.value, str) and 'Inputs!' in c.value: c.value = c.value.replace('$100', f'${SCAN}')
n2['A2'] = f'{cur_lbl} milhões'
for p in PERS:
    y, q = 2000 + int(p[2:]), int(p[0])
    end = dt.date(y, 3 * q, [31, 30, 30, 31][q - 1])
    n2.cell(3, PER_NEW[p]).value = 1 if end <= LAST else 0
for c in range(4, 34):
    if n2.cell(2, c).value is not None and not re.fullmatch(r'[1-4]Q\d\d', str(n2.cell(2, c).value)):
        qs = [n2.cell(3, c - k).value for k in range(1, 5)]
        n2.cell(3, c).value = 1 if all(v == 1 for v in qs) else 0
TEMPL = 4
def code_row(rr, code, label):
    n2.cell(rr, 2).value, n2.cell(rr, 3).value = code, label
    for c in range(1, 34):
        a = n2.cell(TEMPL, c)
        if a.has_style: n2.cell(rr, c)._style = copy.copy(a._style)
        if c >= 4 and isinstance(a.value, str) and a.value.startswith('='):
            n2.cell(rr, c).value = Translator(a.value, origin=a.coordinate).translate_formula(n2.cell(rr, c).coordinate)
code_row(50, 44, 'Net income reported (lucro líquido reportado) — base do check de DRE [Model linha 4]')
# códigos 45+ nas mesmas linhas da Inputs (2) antiga (o bloco do Model referencia essas linhas)
drv = [(rr, o2.cell(rr, 2).value, o2.cell(rr, 3).value) for rr in range(51, o2.max_row + 1) if isinstance(o2.cell(rr, 2).value, int) and o2.cell(rr, 2).value >= 45]
for rr, code, lab in drv: code_row(rr, code, lab)
for rr in range(51, o2.max_row + 1):          # títulos e notas do bloco antigo
    if not isinstance(o2.cell(rr, 2).value, int) and o2.cell(rr, 3).value: n2.cell(rr, 3).value = o2.cell(rr, 3).value
# lucro reportado na linha 4 do Model (base do check de DRE)
for c in range(9, 36):
    lab = nm.cell(7, c).value
    if c - 5 < 4: continue
    if (c - 9) % 5 == 4: nm.cell(4, c).value = f'=IF(COUNT({L(c-4)}4:{L(c-1)}4)=4,SUM({L(c-4)}4:{L(c-1)}4),"")'
    else: nm.cell(4, c).value = f"=IF(ISNUMBER('Inputs (2)'!{L(c-5)}50),'Inputs (2)'!{L(c-5)}50,\"\")"
nm['C4'] = 'Net Income reportado (lucro líquido) — base do check de DRE'

# ------------------------------------------------------------------ bloco de drivers do Model (mesmas colunas e linhas)
for rr in range(18, 18 + N):
    nm.row_dimensions[rr].height = om.row_dimensions[rr].height
    for c in range(1, min(om.max_column, nm.max_column) + 1):
        a, b = om.cell(rr, c), nm.cell(rr, c)
        b.value = a.value
        if a.has_style: b._style = copy.copy(a._style)

# ------------------------------------------------------------------ correções da máscara (falhas herdadas da máscara original)
SH = lambda r: r + N if r >= 18 else r
FIX = []
for r in (183, 220):          # dívida LP e capital + APIC: 1T26/2T26 não estavam ligados à Inputs (2)
    src = nm.cell(SH(r), CI('AF'))
    for cl in ('AH', 'AI'):
        tgt = nm[f'{cl}{SH(r)}']
        tgt.value = Translator(src.value, origin=src.coordinate).translate_formula(tgt.coordinate)
    FIX.append(f'Model linha {SH(r)}: 1T26 e 2T26 ligados à Inputs (2), como os demais trimestres realizados')
n_un = 0
for c in range(1, nm.max_column + 1):     # (X259,X266) é união de células no Excel (#VALUE!): vira soma
    f = nm.cell(SH(196), c).value
    if isinstance(f, str) and re.search(r'\(([A-Z]{1,3}\d+),([A-Z]{1,3}\d+)\)', f):
        nm.cell(SH(196), c).value = re.sub(r'\(([A-Z]{1,3}\d+),([A-Z]{1,3}\d+)\)', r'(\1+\2)', f); n_un += 1
FIX.append(f'Model linha {SH(196)}: "(CFO,CFI)" era união de células (#VALUE!) — trocado por soma em {n_un} colunas')
for r in (228, 229, 230, 231):  # ações: dependem do preço (premissa vazia) -> IFERROR para não propagar #DIV/0!
    for c in range(1, nm.max_column + 1):
        f = nm.cell(SH(r), c).value
        if isinstance(f, str) and f.startswith('=') and 'IFERROR' not in f: nm.cell(SH(r), c).value = f'=IFERROR({f[1:]},"")'
FIX.append(f'Model linhas {SH(228)} a {SH(231)}: IFERROR (dependem do preço da ação, premissa vazia)')
# trimestres realizados 4T25, 1T26 e 2T26 (AF, AH, AI): a máscara deixou várias linhas sem a fórmula histórica (capex,
# imobilizado, intangível, DSO/DIO/DPO...) ou com a fórmula de projeção. Repete o padrão do 3T25 (coluna AE).
nfix = collections.Counter()
for r in range(8, nm.max_row + 1):
    if 18 <= r < 18 + N: continue
    for cl, scl in (('AF', 'AE'), ('AH', 'AE'), ('AI', 'AE'), ('AG', 'AB')):
        src = nm[f'{scl}{r}']
        if not (isinstance(src.value, str) and src.value.startswith('=')): continue
        tgt = nm[f'{cl}{r}']
        want = Translator(src.value, origin=src.coordinate).translate_formula(tgt.coordinate)
        if tgt.value != want:
            tgt.value = want; nfix[r] += 1
            if src.has_style and not tgt.has_style: tgt._style = copy.copy(src._style)
FIX.append(f'Model, 4T25/1T26/2T26 (e anual 2025 a partir de 2024): fórmula histórica do 3T25 aplicada em {sum(nfix.values())} células de {len(nfix)} linhas que estavam vazias ou com fórmula de projeção')
# linha 150 (intangível do roll-forward): no histórico a máscara aponta para o goodwill (145); a projeção usa o intangível (146)
for c in range(CI('I'), CI('AI') + 1):
    cl = openpyxl.utils.get_column_letter(c)
    nm[f'{cl}{SH(150)}'] = f'={cl}{SH(146)}'
FIX.append('Model linha 150 (intangível do roll-forward): histórico apontava para o goodwill (linha 145); passa a usar o intangível (146)')
# linha 189 (cronograma da dívida LP): a máscara tem 0 digitado nos trimestres 2 a 4 de cada ano projetado, o que zera a
# dívida LP projetada (linha 183 = linha 189). Realizados (1T26, 2T26) recebem o saldo do balanço; projetados, o roll-forward.
L_ = openpyxl.utils.get_column_letter
c = CI('AH')
while nm.cell(SH(183), c).value is not None:
    if c % 5 != 3:                                   # colunas anuais: M, R, W, AB, AG, AL, AQ...
        cl = L_(c); prev = L_(c - 1 if (c - 1) % 5 != 3 else c - 2)
        nm[f'{cl}{SH(189)}'] = (f'={cl}{SH(183)}' if c <= CI('AI')
                                else f'={prev}{SH(189)}+SUM({cl}{SH(187)}:{cl}{SH(188)})')
    c += 1
FIX.append(f'Model linha {SH(189)} (dívida LP no cronograma): zeros digitados nas projeções trocados pelo roll-forward; 1T26 e 2T26 = saldo do balanço')
# 1T21 e 2021 (colunas I e M): a grade começa no 1T21 e a coluna H (2020) está vazia, então tudo o que é variação de saldo
# (capex, intangível, capital de giro, DFC indireta) sairia igual ao saldo inteiro. Fica vazio quando não há o saldo anterior.
n_h = 0
for r in range(8, nm.max_row + 1):
    if 18 <= r < 18 + N: continue
    for cl in ('I', 'M'):
        f = nm[f'{cl}{r}'].value
        if not (isinstance(f, str) and f.startswith('=')) or 'IFERROR' in f: continue
        refs = sorted(set(re.findall(r"(?<![A-Z$!'])\$?H\$?(\d+)", f)))
        refs = [x for x in refs if int(x) >= SH(110)]
        if refs:
            nm[f'{cl}{r}'] = f'=IF(COUNT({",".join("H" + x for x in refs)})=0,"",{f[1:]})'; n_h += 1
FIX.append(f'Model 1T21/2021: {n_h} variações de saldo sem o saldo de 2020 ficam vazias (antes davam o saldo inteiro)')
nm['H5'] = dt.datetime(2020, 12, 31); nm['H5'].number_format = nm['I5'].number_format
FIX.append('Model H5: data de 31/12/2020 (base do DSO, DIO e DPO do 1T21)')
vs = nw['Valuation']
for row in vs.iter_rows():
    for x in row:
        if isinstance(x.value, str) and x.value.startswith('=') and 'IFERROR' not in x.value and re.search(r'\*\(1\+', x.value):
            x.value = f'=IFERROR({x.value[1:]},"")'
FIX.append('Valuation: IFERROR nas fórmulas de crescimento que dependem de premissa vazia')

# ------------------------------------------------------------------ Control e rótulos de moeda
name = oi['D1'].value or ''
mname = re.match(r'(.+?) \(([^,]+),', name)
ctl['B3'] = mname.group(1) if mname else name[:60]
ctl['B4'] = mname.group(2) if mname else ''
ctl['B5'] = 'USD' if cur_lbl == 'US$' else 'BRL'
ctl['B6'] = f'{cur_lbl} milhões'
ctl['B7'] = dt.datetime(LAST.year, LAST.month, LAST.day); ctl['B7'].number_format = 'dd/mm/yyyy'
ctl['D4'] = f'Last close ({cur_lbl}/share)'
ctl['B30'] = (name[:250] if name else 'ver Inputs')
ctl['B31'] = f'{cur_lbl} mil na Inputs → {cur_lbl} milhões via Inputs (2)'
# DSO, DIO e DPO projetados: sem custo ou receita no histórico a média dá "" e o saldo projetado vira #VALUE!; fica 0 dia
for r in (25, 26, 27):
    f = ctl[f'B{r}'].value
    if isinstance(f, str) and f.startswith('=IFERROR(') and f.endswith(',"")'):
        ctl[f'B{r}'] = f[:-3] + ',0)'
FIX.append('Control B25:B27 (DSO, DIO, DPO): sem histórico o padrão é 0 dia em vez de "" (evitava #VALUE! na projeção)')
nw.calculation.fullCalcOnLoad = True
nw.save(OUT)
print('correções:', *FIX, sep='\n  ')
print('ok', OUT, '| linhas Inputs', LASTNEW, '| varredura', SCAN, '| bloco drivers', N, '| códigos 45+', len(drv), '| períodos', PERS[0], '-', PERS[-1])
