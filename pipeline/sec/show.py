import openpyxl, collections, sys
wb=openpyxl.load_workbook(sys.argv[1],data_only=True); m=wb['Model']; i=wb['Inputs']
cols=list(range(4,36)); fm=lambda x: round(x) if isinstance(x,float) else x
rows=[int(x) for x in sys.argv[2].split(',')] if len(sys.argv)>2 else [1,2,3,4,15,20,24,35,36,40,89,97,101,102,103,104,163,214]
for r in rows: print(r,str(m.cell(r,3).value)[:22].ljust(22),[fm(m.cell(r,c).value) for c in cols])
for r in range(5,300):
    if i.cell(r,8).value and 'CHECK:' in str(i.cell(r,8).value): print('inp',r,i.cell(r,8).value[:40],set(i.cell(r,c).value for c in range(11,37)))
print('hist errors', collections.Counter((str(m.cell(r,3).value)[:25],m.cell(r,c).value) for r in range(1,400) for c in cols if isinstance(m.cell(r,c).value,str) and m.cell(r,c).value.startswith('#')))
