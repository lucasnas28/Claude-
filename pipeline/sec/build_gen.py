"""Construtor genérico SEC -> máscara padrão (Inputs / Inputs (2) / Model).
uso: build_gen.py CONFIG.py MASK OUT
CONFIG define: NAME, TICKER, CIK, DL (pasta), e opcionalmente listas de tags (REV, COGS, OPINC, PRETAX, TAX, NI, MI, INTEXP, DA, DIV), BSMAP extra, FACE_OVR."""
import sys, shutil, copy, json, runpy
sys.path.insert(0, __import__('os').path.dirname(__import__('os').path.abspath(__file__)))
from sx import Co, P, label
import openpyxl
from openpyxl.utils import get_column_letter as CL
from openpyxl.styles import Font

CFG = runpy.run_path(sys.argv[1]); MASK, OUT = sys.argv[2:4]
g = lambda k, d: CFG.get(k, d)
co = Co(CFG['DL'], CFG.get('MONTHMAP'))
REV = g('REV', ['Revenues', 'RevenueFromContractWithCustomerExcludingAssessedTax', 'SalesRevenueNet', 'RevenueFromContractWithCustomerIncludingAssessedTax'])
COGS = g('COGS', ['CostOfRevenue', 'CostOfGoodsAndServicesSold', 'CostOfGoodsSold', 'CostOfGoodsAndServiceExcludingDepreciationDepletionAndAmortization'])
OPINC = g('OPINC', ['OperatingIncomeLoss'])
PRETAX = g('PRETAX', ['IncomeLossFromContinuingOperationsBeforeIncomeTaxesExtraordinaryItemsNoncontrollingInterest', 'IncomeLossFromContinuingOperationsBeforeIncomeTaxesMinorityInterestAndIncomeLossFromEquityMethodInvestments',
                      'IncomeLossFromContinuingOperationsIncludingNoncontrollingInterestBeforeIncomeTaxesExtraordinaryItems', 'IncomeLossFromContinuingOperationsBeforeIncomeTaxesDomestic'])
TAX = g('TAX', ['IncomeTaxExpenseBenefit'])
NI = g('NI', ['NetIncomeLoss', 'NetIncomeLossAvailableToCommonStockholdersBasic', 'ProfitLoss'])
MI = g('MI', ['NetIncomeLossAttributableToNoncontrollingInterest'])
INTEXP = g('INTEXP', ['InterestExpense', 'InterestExpenseNonoperating', 'InterestExpenseDebt', 'InterestAndOtherFinancialCharges'])
DA = g('DA', ['DepreciationDepletionAndAmortization', 'DepreciationAmortizationAndImpairment', 'DepreciationAndAmortizationIncludingDiscontinuedOperations', 'DepreciationAmortizationAndAccretionNet', 'DepreciationAndAmortization', 'DepreciationAmortizationAndOther'])
DIV = g('DIV', ['PaymentsOfDividends', 'PaymentsOfDividendsCommonStock', 'PaymentsOfOrdinaryDividends'])

def qser(tags, orig=False):
    return {k: v for k, v in co.quarter(tags, orig=orig).items() if k in P}

shutil.copyfile(MASK, OUT)
wb = openpyxl.load_workbook(OUT)
inp, w2, m = wb['Inputs'], wb['Inputs (2)'], wb['Model']
PC = {p: 11 + i for i, p in enumerate(P)}
col = lambda p: CL(PC[p])
FIRSTC, LASTC = col(P[0]), col(P[-1])
HDR, R0 = 4, 5
bold = Font(bold=True)
FYM = {'4': 'dez', '1': 'mar', '2': 'jun', '3': 'set'}[co.fy_digit]
inp['A1'] = 'Fator p/ milhões:'; inp['B1'] = 0.000001
inp['D1'] = (f"{CFG['NAME']} ({CFG['TICKER']}, CIK {int(CFG['CIK'])}) - fonte: SEC EDGAR XBRL (instâncias dos 10-Q/10-K). Moeda: USD. Brutos em USD unidades (fator 0,000001 -> milhões). "
             f"Colunas = trimestres-CALENDÁRIO (cada fechamento vai para o fim de trimestre mais próximo); ano fiscal termina no trimestre-calendário de {FYM}. "
             "Trimestre fiscal 4 (e qualquer trimestre sem 3M publicado) = acumulado do ano menos o acumulado anterior, calculado na extração.")
for i, h in enumerate(['COD1', 'SINAL1', 'COD2', 'SINAL2', 'COD3', 'SINAL3', 'UNID', 'Conta (PT) / origem', 'Account (EN)', 'ORIGEM']):
    inp.cell(HDR, 1 + i, h).font = bold
for p in P:
    inp.cell(HDR, PC[p], p).font = bold
    inp.cell(2, PC[p], f'FQ{co.fq(p)}')
inp.cell(2, 10, 'Trimestre fiscal')

r = R0; pos = {}
def section(t):
    global r
    inp.cell(r, 8, t).font = bold; r += 1
def write(cods, desc, origem, vals=None, forms=None, key=None, unid='moeda', zero_fill=True, fmt='#,##0.0;(#,##0.0)'):
    global r
    for i, (c, s) in enumerate(cods): inp.cell(r, 1 + 2 * i, c); inp.cell(r, 2 + 2 * i, s)
    vals = {p: v for p, v in (vals or {}).items() if v is not None and p in P}
    if cods and zero_fill and not forms:
        vals.update({p: 0.0 for p in P if p not in vals})
    inp.cell(r, 7, unid); inp.cell(r, 8, desc); inp.cell(r, 10, origem)
    for p, x in vals.items(): inp.cell(r, PC[p], x)
    for p, f in (forms or {}).items(): inp.cell(r, PC[p], f)
    for p in P: inp.cell(r, PC[p]).number_format = fmt
    if key: pos[key] = r
    r += 1
def fx(expr):   # expr com {k} -> referência da linha k na mesma coluna
    return {p: '=' + expr.format(**{k: f'{col(p)}{v}' for k, v in pos.items()}) for p in P}

S = {k: qser(t) for k, t in dict(rev=REV, cogs=COGS, op=OPINC, pt=PRETAX, tax=TAX, ni=NI, mi=MI, ie=INTEXP).items()}
S['da'], S['div'] = qser(DA, orig=True), qser(DIV, orig=True)
if g('DA_FUNC', None):                            # regra específica da companhia (ver cfg)
    S['da'] = {k: v for k, v in CFG['DA_FUNC'](qser).items() if k in P}; DA = [CFG.get('DA_DESC', 'D&A (regra da companhia)')]
elif g('DA_SUM', None):                             # D&A = soma de linhas do fluxo de caixa (ex.: depreciação + amortização de intangíveis)
    parts = [qser(t, orig=True) for t in CFG['DA_SUM']]
    S['da'] = {p: sum(x[p] for x in parts) for p in P if all(p in x for x in parts)}
    DA = ['+'.join(t[0] for t in CFG['DA_SUM'])]   # fluxo de caixa: como divulgado originalmente (mesma cadeia de acumulados do ano)
if g('COGS_SUM', None):                          # custo = soma de linhas de custo da DRE (companhia não divulga o total)
    parts = [qser(t) for t in CFG['COGS_SUM']]
    S['cogs'] = {p: sum(x[p] for x in parts if p in x) for p in P if any(p in x for x in parts)}
    COGS = ['+'.join(t[0] for t in CFG['COGS_SUM'])]
has = lambda k: bool(S[k])
for mem_tag in g('COGS_DIMS', []):              # custo só aberto por produto/serviço (dimensional) em alguns períodos
    tag, mems = mem_tag
    parts = [co.dq(tag, [m]) for m in mems]
    for p in P:
        if p not in S['cogs'] and all(p in x for x in parts): S['cogs'][p] = sum(x[p] for x in parts)
OP_DESC = f'Lucro operacional [us-gaap:{OPINC[0]}]'
if len(S['op']) < len(S['rev']) // 2:      # companhia não divulga lucro operacional: receita - custos e despesas totais + juros (juros ficam no não operacional)
    ce = qser(['CostsAndExpenses', 'CostsAndExpensesExcludingInterest', 'BenefitsLossesAndExpenses', 'OperatingCostsAndExpenses'])
    if ce:
        S['op'] = {p: S['rev'][p] - ce[p] + S['ie'].get(p, 0) for p in ce if p in S['rev']}
        OP_DESC = 'Lucro operacional DERIVADO = receita - custos e despesas totais [us-gaap:CostsAndExpenses] + despesa de juros (a companhia não divulga lucro operacional)'
if g('OP_FUNC', None):                           # regra específica da companhia para o lucro operacional (ver cfg)
    S['op'] = {k: v for k, v in CFG['OP_FUNC'](qser).items() if k in P}; OP_DESC = CFG['OP_DESC']
section('DRE - trimestre isolado (USD). Brutos SEC XBRL; linhas "derivada" são fórmulas nesta aba')
write([(1, 1)], f'Receita [us-gaap:{"/".join(REV[:2])}]', 'fonte', S['rev'], key='rev')
write([(2, -1)], f'Custo da receita (divulgado positivo) [us-gaap:{"/".join(COGS[:2])}]', 'fonte', S['cogs'], key='cogs')
write([], OP_DESC, 'memoria', S['op'], key='op')
write([(3, -1)], 'Despesas operacionais = receita - custo - lucro operacional (P&D + SG&A + outras operacionais)', 'derivada', forms=fx('N({rev})-N({cogs})-N({op})'), key='opx')
write([], f'Lucro antes do IR [us-gaap:{PRETAX[0][:60]}]', 'memoria', S['pt'], key='pt')
write([(6, 1)], 'Resultado não operacional total = lucro antes do IR - lucro operacional', 'derivada', forms=fx('N({pt})-N({op})'), key='nonop')
if has('ie'):
    write([(5, -1), (6, 1)], f'Despesa de juros (divulgado positivo): COD1=5 lança; COD2=6 devolve ao não operacional, que já a contém [us-gaap:{"/".join(INTEXP[:2])}]', 'fonte', S['ie'], key='ie')
write([(7, -1)], f'Imposto de renda (divulgado positivo) [us-gaap:{TAX[0]}]', 'fonte', S['tax'], key='tax')
if has('mi'):
    write([(9, -1)], f'Lucro atribuído a não controladores (divulgado positivo) [us-gaap:{MI[0]}]', 'fonte', S['mi'], key='mi')
write([], f'Lucro líquido atribuído à companhia [us-gaap:{NI[0]}]', 'memoria', S['ni'], key='ni_raw')
mi_ref = 'N({mi})' if has('mi') else '0'
write([(8, 1)], 'Resultado após IR não classificado acima (equivalência patrimonial, operações descontinuadas): lucro - (LAIR - IR - minoritários); zero quando não há', 'derivada',
      forms=fx('N({ni_raw})-(N({pt})-N({tax})-' + mi_ref + ')'), key='oth')
write([(44, 1)], 'Lucro líquido (check IS, código 44)', 'derivada', forms=fx('N({ni_raw})'), key='ni')
section('DFC - trimestre isolado (desacumulado na extração)')
ADDB = g('DA_ADDBACK', 3)
write([(4, -1), (ADDB, 1)], f'D&A (fluxo de caixa). COD1=4 lança (positivo -> -1); COD2={ADDB} devolve ' + ('ao custo, onde a companhia apresenta a D&A' if ADDB == 2 else 'às despesas operacionais (a companhia não abre D&A entre custo e despesa)') + f' [us-gaap:{"/".join(DA[:2])}]', 'fonte', S['da'], key='da')
write([(10, -1)], f'Dividendos pagos [us-gaap:{DIV[0]}]', 'fonte', S['div'], key='div')

# ---------------------------------------------------------------- balanço de face
KNOWN = {
 'CashAndCashEquivalentsAtCarryingValue': 11, 'CashCashEquivalentsAndShortTermInvestments': 11, 'MarketableSecuritiesCurrent': 11, 'ShortTermInvestments': 11,
 'AvailableForSaleSecuritiesDebtSecuritiesCurrent': 11, 'DebtSecuritiesCurrent': 11, 'CashAndCashEquivalentsAtCarryingValueIncludingDiscontinuedOperations': 11,
 'CashCashEquivalentsRestrictedCashAndRestrictedCashEquivalents': 11, 'Cash': 11, 'AvailableForSaleSecuritiesCurrent': 11, 'HeldToMaturitySecuritiesCurrent': 11,
 'AccountsReceivableNetCurrent': 12, 'ReceivablesNetCurrent': 12, 'AccountsNotesAndLoansReceivableNetCurrent': 12, 'AccountsAndOtherReceivablesNetCurrent': 12,
 'InventoryNet': 13, 'InventoryFinishedGoodsNetOfReserves': 13, 'RetailRelatedInventoryMerchandise': 13, 'InventoryGross': 13,
 'PrepaidExpenseCurrent': 15, 'PrepaidExpenseAndOtherAssetsCurrent': 16, 'OtherAssetsCurrent': 16,
 'DeferredIncomeTaxAssetsNet': 14, 'DeferredTaxAssetsNetNoncurrent': 14, 'DeferredIncomeTaxesAndOtherAssetsNoncurrent': 22,
 'EquityMethodInvestments': 17, 'LongTermInvestments': 22,
 'PropertyPlantAndEquipmentNet': 18, 'PropertyPlantandEquipmentandFinanceLeaseRightofUseAssetafterAccumulatedDepreciationandAmortization': 18, 'PropertyPlantAndEquipmentAndFinanceLeaseRightOfUseAssetAfterAccumulatedDepreciationAndAmortization': 18,
 'OperatingLeaseRightOfUseAsset': 19, 'Goodwill': 20, 'IntangibleAssetsNetExcludingGoodwill': 21, 'FiniteLivedIntangibleAssetsNet': 21, 'IntangibleAssetsNetIncludingGoodwill': 21,
 'OtherAssetsNoncurrent': 22, 'OtherLiabilitiesCurrent': 26,
 'AccountsPayableCurrent': 23, 'AccountsPayableTradeCurrent': 23, 'AccountsPayableAndAccruedLiabilitiesCurrent': 23,
 'AccruedLiabilitiesCurrent': 24, 'EmployeeRelatedLiabilitiesCurrent': 24, 'AccruedCompensationCurrent': 24,
 'TaxesPayableCurrent': 25, 'AccruedIncomeTaxesCurrent': 25,
 'DebtCurrent': 27, 'LongTermDebtCurrent': 27, 'CommercialPaper': 27, 'ShortTermBorrowings': 27, 'LongTermDebtAndCapitalLeaseObligationsCurrent': 27, 'OtherShortTermBorrowings': 27,
 'OperatingLeaseLiabilityCurrent': 29,
 'LongTermDebtNoncurrent': 30, 'LongTermDebtAndCapitalLeaseObligations': 30, 'LongTermDebt': 30, 'SeniorLongTermNotes': 30,
 'OperatingLeaseLiabilityNoncurrent': 31, 'DeferredIncomeTaxLiabilitiesNet': 33, 'DeferredTaxLiabilitiesNoncurrent': 33,
 'OtherLiabilitiesNoncurrent': 34,
 'MinorityInterest': 35, 'RedeemableNoncontrollingInterestEquityCarryingAmount': 35,
 'CommonStockValue': 37, 'AdditionalPaidInCapital': 37, 'AdditionalPaidInCapitalCommonStock': 37, 'CommonStocksIncludingAdditionalPaidInCapital': 37, 'CommonStockSharesSubscribedButUnissuedSubscriptionsReceivable': 37,
 'TreasuryStockValue': 38, 'TreasuryStockCommonValue': 38, 'CommonStockSharesHeldInEmployeeTrust': 38,
 'AccumulatedOtherComprehensiveIncomeLossNetOfTax': 39, 'PreferredStockValueOutstanding': 40, 'RightToRecoverForCoveredLosses': 40,
 'RetainedEarningsAccumulatedDeficit': 41,
}
KNOWN.update(g('BSMAP', {}))
SIGN = {'RightToRecoverForCoveredLosses': -1, 'TreasuryStockValue': -1, 'TreasuryStockCommonValue': -1, 'CommonStockSharesHeldInEmployeeTrust': -1, 'TreasuryStockPreferredValue': -1}
SIGN_CI = {k.lower(): v for k, v in SIGN.items()}
SKIP = {'AssetsCurrent', 'Assets', 'LiabilitiesCurrent', 'Liabilities', 'LiabilitiesAndStockholdersEquity', 'StockholdersEquity', 'NoncurrentAssets',
        'StockholdersEquityIncludingPortionAttributableToNoncontrollingInterest', 'CommitmentsAndContingencies', 'LiabilitiesNoncurrent', 'AssetsNoncurrent', 'PreferredStockValue'}
SKIP -= set(g('UNSKIP', []))                        # cfg: linha de face que nesta companhia é conta própria (ex.: ações preferenciais conversíveis no PL)
SECTION_DEFAULT = {'ca': 16, 'nca': 22, 'cl': 26, 'ncl': 34, 'eq': 40}
FACE = {}
SUBTOT = set()
SECT = {}
DROPPED = []
for p in P:
    rows, a = co.face(p, g('FACE_OVR', {}).get(p))
    FACE[p] = (rows or [], a)
missing = [p for p in P if not FACE[p][0]]
tagcode, tagdesc, vals, totals = {}, {}, {}, {k: {} for k in ('Assets', 'Liabilities', 'StockholdersEquity', 'LiabilitiesAndStockholdersEquity', 'MinorityInterest_tot')}
order = []
for p in P:
    sec = 'ca'; seen_lc = False
    face_rows = FACE[p][0]
    sub = set()
    BOUND = {'Assets', 'AssetsCurrent', 'LiabilitiesCurrent', 'Liabilities', 'LiabilitiesAndStockholdersEquity'}
    for i, (n, v) in enumerate(face_rows):     # subtotal: igual à soma de 2-12 linhas imediatamente anteriores (sem cruzar totais de seção)
        if n in SKIP or abs(v) < 1: continue
        for j in range(2, 13):
            prev = [x for x in face_rows[max(0, i - j):i] if x[0] not in sub]
            if len(face_rows[max(0, i - j):i]) < j or any(x[0] in BOUND for x in prev): break
            if abs(sum(x[1] for x in prev) - v) < 0.5: sub.add(n); break
    if any(n.startswith('PropertyPlantAndEquipmentNet') or n.startswith('PropertyPlantAndEquipmentAndFinanceLease') for n, v in face_rows):
        # imobilizado aberto bruto na face (terrenos, prédios, equipamentos, obras) + líquido: fica só o líquido
        GROSS = ('Land', 'LandAndLandImprovements', 'BuildingsAndImprovementsGross', 'FixturesAndEquipmentGross', 'ConstructionInProgressGross', 'PropertyPlantAndEquipmentGross',
                 'MachineryAndEquipmentGross', 'LeaseholdImprovementsGross', 'PropertyPlantAndEquipmentOther')
        sub.update(n for n, v in face_rows if n.split('[')[0] in GROSS)
    if any(n == 'AccumulatedOtherComprehensiveIncomeLossNetOfTax' for n, v in face_rows):
        # componentes de OCI listados acima do total: fica só o total
        sub.update(n for n, v in face_rows if (n.startswith('AccumulatedOtherComprehensiveIncomeLoss') or n.startswith('Aoci')) and n != 'AccumulatedOtherComprehensiveIncomeLossNetOfTax')
    SUBTOT.update((p, n) for n in sub)
    # reconciliação por seção: se a soma da seção difere do subtotal divulgado por exatamente o valor de uma linha, essa linha não é de face
    secs, cur = {'ca': [], 'nca': [], 'cl': [], 'rest': []}, 'ca'
    tot = dict(face_rows)
    for n, v in face_rows:
        if n == 'AssetsCurrent': cur = 'nca'; continue
        if n == 'Assets': cur = 'cl'; continue
        if n == 'LiabilitiesCurrent': cur = 'rest'; continue
        if n in SKIP or n in sub or 'PerShare' in n or ('AccumulatedDepreciation' in n and 'after' not in n.lower()) or n.startswith('Allowance') or n.endswith('Shares'): continue
        kc = KNOWN.get(n, KNOWN.get(n.split('[')[0]))
        if cur == 'ca' and kc is not None and 17 <= kc <= 22: secs['nca'].append((n, v)); continue     # imobilizado antes do circulante (utilities)
        secs[cur].append((n, v))
    target = {'ca': tot.get('AssetsCurrent'), 'nca': (tot['Assets'] - tot['AssetsCurrent']) if 'Assets' in tot and 'AssetsCurrent' in tot else None,
              'cl': tot.get('LiabilitiesCurrent'),
              'rest': (tot['LiabilitiesAndStockholdersEquity'] - tot['LiabilitiesCurrent']) if 'LiabilitiesAndStockholdersEquity' in tot and 'LiabilitiesCurrent' in tot else None}
    if 'AssetsCurrent' not in tot and 'Assets' in tot: secs['ca'] += secs['nca']; secs['nca'] = []; target['ca'] = tot['Assets']; target['nca'] = None
    for k, items in secs.items():
        if target[k] is None: continue
        gap = sum(v for _, v in items) - target[k]
        if abs(gap) < 0.5: continue
        hit = [n for n, v in items if abs(v - gap) < 0.5 and KNOWN.get(n) is None]
        if hit: sub.add(hit[0]); DROPPED.append((p, hit[0], gap)); continue
        unk = [(n, v) for n, v in items if KNOWN.get(n) is None]
        pair = [(a_[0], b_[0]) for i_, a_ in enumerate(unk) for b_ in unk[i_ + 1:] if abs(a_[1] + b_[1] - gap) < 0.5]
        if pair: sub.update(pair[0]); DROPPED.append((p, '+'.join(pair[0]), gap))
    for n, v in face_rows:
        if n in sub: continue
        if n in totals: totals[n][p] = v
        if n == 'AssetsCurrent': sec = 'nca'; continue
        if n == 'Assets': sec = 'cl'; continue
        if n == 'LiabilitiesCurrent': sec = 'ncl'; continue
        if n in ('Liabilities',): sec = 'eq'; continue
        if n in SKIP or 'PerShare' in n or 'LiquidationPreference' in n or n.endswith('Shares') or ('AccumulatedDepreciation' in n and 'after' not in n.lower()) or n.startswith('Allowance') or 'AccumulatedAmortization' in n: continue
        if n in ('StockholdersEquity',): continue
        code = KNOWN.get(n, KNOWN.get(n.split('[')[0]))
        if code is None:                       # prefixos de PL (balanços de utility trazem o PL antes do passivo)
            for pre, cc in (('AccumulatedOtherComprehensiveIncome', 39), ('RetainedEarnings', 41), ('AdditionalPaidInCapital', 37), ('CommonStock', 37), ('TreasuryStock', 38),
                            ('MinorityInterest', 35), ('NoncontrollingInterest', 35), ('PreferredStock', 40)):
                if n.startswith(pre): code = cc; break
        if code is None: code = SECTION_DEFAULT[sec]
        # dívida LP marcada como LongTermDebt mas na seção circulante
        if code == 30 and sec in ('ca', 'cl') and n in ('LongTermDebt', 'LongTermDebtAndCapitalLeaseObligations'): code = 27
        if code in (11, 12, 13, 15) and sec == 'nca': code = 22
        key = (n, code)
        SECT[(p, n)] = sec
        if key not in tagcode: tagcode[key] = code; order.append(key)
        vals.setdefault(key, {})[p] = v * SIGN.get(n, SIGN_CI.get(n.lower().split('[')[0], 1))
# patrimônio sem "Liabilities" no documento: tudo após o passivo circulante/NC até o fim é tratado pela tabela KNOWN
section('BALANÇO DE FACE (USD, saldo no fim do trimestre; cada trimestre lido da instância do próprio 10-Q/10-K)')
akeys, lkeys = [], []
for key in sorted(order, key=lambda k: (tagcode[k], order.index(k))):
    n, code = key
    write([(code, 1)], f'{n}' + (' (divulgado positivo, lançado negativo)' if SIGN.get(n) else ''), 'fonte', vals[key], key='bs_' + n + '_' + str(code))
    (akeys if code <= 22 else lkeys).append('bs_' + n + '_' + str(code))
ABSENT = {8: 'Equivalência patrimonial', 9: 'Minoritários (DRE)', 11: 'Caixa', 12: 'Contas a receber', 13: 'Estoques', 14: 'IR diferido ativo', 15: 'Despesas antecipadas', 16: 'Outros ativos circulantes',
          17: 'Investimentos em coligadas', 18: 'Imobilizado', 19: 'Direito de uso', 20: 'Goodwill', 21: 'Intangíveis', 22: 'Outros ativos não circulantes', 23: 'Fornecedores',
          24: 'Despesas provisionadas', 25: 'Impostos a pagar', 26: 'Outros passivos circulantes', 27: 'Dívida CP', 28: 'Revolving', 29: 'Arrendamento CP', 30: 'Dívida LP',
          31: 'Arrendamento LP', 32: 'Dívida conversível', 33: 'IR diferido passivo', 34: 'Outros passivos não circulantes', 35: 'Participação de minoritários (balanço)',
          37: 'Capital + APIC', 38: 'Tesouraria', 39: 'Outros resultados abrangentes', 40: 'Outras contas do PL', 41: 'Lucros retidos'}
have = set()
for rr in range(R0, r):
    for cc in (1, 3, 5):
        if isinstance(inp.cell(rr, cc).value, int): have.add(inp.cell(rr, cc).value)
for code, nm in ABSENT.items():
    if code not in have:
        write([(code, 1)], f'[vazio] {nm}: não aparece como linha própria nas demonstrações (zero)', 'derivada', {}, key=f'zero{code}')
# arredondamento: algumas companhias publicam em milhões e a soma das linhas difere do total em poucos milhões; a diferença vira linha explícita (só se < 50 milhões)
def resid(keys, tot):
    out = {}
    for p in P:
        if p not in tot: continue
        sm = sum(vals[k2].get(p, 0) for k2 in [kk for kk in order if 'bs_' + kk[0] + '_' + str(kk[1]) in keys or ('sec_' + kk[0][6:] in keys and kk[0].startswith('__sec_'))])
        d_ = tot[p] - sm
        if 0.5 < abs(d_) < 5e7: out[p] = d_
    return out
SECRES = {}
for p in P:
    tot = dict(FACE[p][0]); rows_p = [(k, vals[k].get(p, 0)) for k in order]
    if 'AssetsCurrent' in tot:
        ca = sum(v for (n, c_), v in rows_p if c_ in (11, 12, 13, 15, 16) and SECT.get((p, n)) == 'ca')
        gap = tot['AssetsCurrent'] - ca
        if abs(gap) >= 0.5: SECRES.setdefault(('ca', 16), {})[p] = gap
    if 'LiabilitiesCurrent' in tot:
        cl = sum(v for (n, c_), v in rows_p if 23 <= c_ <= 29)
        gap = tot['LiabilitiesCurrent'] - cl
        if abs(gap) >= 0.5: SECRES.setdefault(('cl', 26), {})[p] = gap
    if 'AssetsCurrent' not in tot and 'Assets' in tot:       # balanço não classificado: resíduo do ativo total
        aa = sum(v for (n, c_), v in rows_p if c_ <= 22 and SECT.get((p, n)) in ('ca', 'nca'))
        gap = tot['Assets'] - aa
        if abs(gap) >= 0.5: SECRES.setdefault(('na', 22), {})[p] = gap
    if 'Liabilities' in tot and 'LiabilitiesAndStockholdersEquity' in tot:
        eq = sum(v for (n, c_), v in rows_p if 35 <= c_ <= 41)
        gap = tot['LiabilitiesAndStockholdersEquity'] - tot['Liabilities'] - eq
        if abs(gap) >= 0.5: SECRES.setdefault(('eq', 40), {})[p] = gap
        if 'LiabilitiesCurrent' in tot:
            ncl = sum(v for (n, c_), v in rows_p if 30 <= c_ <= 34)
            gap = tot['Liabilities'] - tot['LiabilitiesCurrent'] - ncl
            if abs(gap) >= 0.5: SECRES.setdefault(('ncl', 34), {})[p] = gap
for (sc, code), vv in SECRES.items():
    nm = {'ca': 'ativo circulante', 'cl': 'passivo circulante', 'eq': 'patrimônio líquido', 'ncl': 'passivo não circulante', 'na': 'ativo (balanço não classificado)'}[sc]
    write([(code, 1)], g('RESID_LABEL', {}).get(sc, f'Linhas de face do {nm} sem tag própria (subtotal divulgado - soma das linhas)'), 'derivada', vv, key=f'sec_{sc}')
    (akeys if code <= 22 else lkeys).append(f'sec_{sc}')
    for p, v in vv.items(): vals.setdefault(('__sec_' + sc, code), {})[p] = v
    order.append(('__sec_' + sc, code))
ra = resid(set(akeys), totals['Assets'])
rl = resid(set(lkeys), {p: v for p, v in totals['LiabilitiesAndStockholdersEquity'].items()})
if ra:
    write([(16, 1)], 'Diferença de arredondamento do ativo (total divulgado - soma das linhas; < US$ 50 mi)', 'derivada', ra, key='rnd_a'); akeys.append('rnd_a')
if rl:
    write([(26, 1)], 'Diferença de arredondamento do passivo + PL (total divulgado - soma das linhas; < US$ 50 mi)', 'derivada', rl, key='rnd_l'); lkeys.append('rnd_l')
section('BALANÇO - totais divulgados (memória)')
write([], 'Total do ativo [us-gaap:Assets]', 'memoria', totals['Assets'], key='ta')
lse = {p: totals['LiabilitiesAndStockholdersEquity'].get(p) for p in P}
write([], 'Total do passivo + PL [us-gaap:LiabilitiesAndStockholdersEquity]', 'memoria', lse, key='tle')
section('CHECKS (devem ser 0)')
def chk(desc, terms, total, key):
    global r
    inp.cell(r, 7, 'moeda'); inp.cell(r, 8, desc); inp.cell(r, 10, 'derivada')
    for p in P:
        e = '+'.join(f'N({col(p)}{pos[t]})' for t in terms)
        inp.cell(r, PC[p], f'=ROUND({e}-N({col(p)}{pos[total]}),-5)'); inp.cell(r, PC[p]).number_format = '#,##0;(#,##0)'
    pos[key] = r; r += 1
chk('CHECK: contas de ativo - ativo total (USD, arredondado a 100 mil)', akeys, 'ta', 'chk_a')
chk('CHECK: passivo + PL (contas) - passivo + PL total', lkeys, 'tle', 'chk_l')
inp.cell(r, 8, 'CHECK: receita - custo - despesas + não operacional - IR - minoritários + outros após IR - lucro (deve ser 0)'); inp.cell(r, 10, 'derivada')
for p in P:
    c = col(p)
    inp.cell(r, PC[p], f'=ROUND(N({c}{pos["rev"]})-N({c}{pos["cogs"]})-N({c}{pos["opx"]})+N({c}{pos["nonop"]})-N({c}{pos["tax"]})' +
             (f'-N({c}{pos["mi"]})' if 'mi' in pos else '') + f'+N({c}{pos["oth"]})-N({c}{pos["ni"]}),-5)')
pos['chk_is'] = r; r += 1
if missing: inp.cell(r, 8, f'[vazio] Sem 10-Q/10-K para: {", ".join(missing)}'); r += 1
LASTROW = r - 1
END = ((LASTROW + 60 + 9) // 10) * 10
conv = END + 3
inp.cell(conv, 8, 'CONVENCOES').font = bold
for i, t in enumerate([
    f'Linhas de origem: {R0} a {LASTROW}. As formulas da Inputs (2) varrem as linhas {R0} a {END}; as linhas {LASTROW + 1} a {END} estao livres para novas contas.',
    f'Periodos: colunas {FIRSTC} a {LASTC} = trimestres-calendário 1Q20 a 2Q26 (linha 2: trimestre fiscal correspondente).',
    'Brutos em USD unidades: fator da célula B1 = 0,000001 converte para USD milhões nas linhas UNID = moeda.',
    'ORIGEM: fonte = copiado da SEC; memoria = como divulgado, sem codigo; derivada = formula nesta aba ou desacumulado na extração.',
    'Balanço: contas de face sem mapeamento específico vão para "outros" do seu bloco (ativo circulante 16, não circulante 22, passivo circulante 26, não circulante 34, PL 40).']):
    inp.cell(conv + 1 + i, 8, t)
for c_, w in [(1, 6), (2, 7), (3, 6), (4, 7), (5, 6), (6, 7), (7, 8), (8, 95), (9, 20), (10, 10)]:
    inp.column_dimensions[CL(c_)].width = w
for p in P: inp.column_dimensions[col(p)].width = 13
inp.freeze_panes = inp.cell(R0, 11)

# ---------------------------------------------------------------- Inputs (2) e Model (mesmo cabeamento das outras companhias)
for i, h in enumerate(['1Q20', '2Q20', '3Q20', '4Q20', 2020]):
    c = 34 + i
    w2.cell(2, c, h); w2.cell(2, c)._style = copy.copy(w2.cell(2, 29 + i if i < 4 else 33)._style)
    w2.column_dimensions[CL(c)].width = w2.column_dimensions['AC'].width or 10
w2['B50'] = 44; w2['C50'] = 'Net income reported (lucro liquido) - memo p/ check IS [Model linha 4]'
for c in range(1, 39): w2.cell(50, c)._style = copy.copy(w2.cell(13, c)._style)
IN = 'Inputs!'
def rng(cl): return f'{IN}${cl}${R0}:${cl}${END}'
VAL = f'INDEX({IN}${FIRSTC}${R0}:${LASTC}${END},0,MATCH({{h}},{IN}${FIRSTC}${HDR}:${LASTC}${HDR},0))'
ESC = f'((({rng("G")}="moeda")+({rng("G")}="acoes"))*{IN}$B$1+(({rng("G")}<>"moeda")*({rng("G")}<>"acoes")))'
def f2(h, d_):
    v_ = VAL.format(h=h)
    cnt = '+'.join(f'SUMPRODUCT(({rng(c)}={d_})*ISNUMBER({v_}))' for c in 'ACE')
    parts = '+'.join(f'SUMPRODUCT(({rng(c)}={d_})*{rng(s)}*{ESC}*{v_})' for c, s in (('A', 'B'), ('C', 'D'), ('E', 'F')))
    return f'=IF(ISNA(MATCH({h},{IN}${FIRSTC}${HDR}:${LASTC}${HDR},0)),NA(),IF({cnt}=0,NA(),{parts}))'
rows2 = [rr for rr in range(4, 51) if isinstance(w2.cell(rr, 2).value, int)]
ANN2 = {8, 13, 18, 23, 28, 33, 38}
for rr in rows2:
    for c in [c for c in range(4, 38) if c not in ANN2]:
        assert w2.cell(rr, c).value is None, (rr, c)
        w2.cell(rr, c).value = f2(f'{CL(c)}$2', f'$B{rr}')
    w2.cell(rr, 38, f'=SUM(AH{rr}:AK{rr})' if rr <= 13 or rr == 50 else f'=AK{rr}')
    if rr == 50:
        for c, (a, b) in {8: ('D', 'G'), 13: ('I', 'L'), 18: ('N', 'Q'), 23: ('S', 'V'), 28: ('X', 'AA'), 33: ('AC', 'AF')}.items():
            w2.cell(rr, c, f'=SUM({a}{rr}:{b}{rr})')
    else:
        w2.cell(rr, 38)._style = copy.copy(w2.cell(rr, 33)._style)
        for i in range(4): w2.cell(rr, 34 + i)._style = copy.copy(w2.cell(rr, 29 + i)._style)
if w2['AG2'].value == 2025: w2['AG2'] = 2026
grey = Font(color='FF808080', bold=w2['AE2'].font.bold, name=w2['AE2'].font.name, size=w2['AE2'].font.sz)
for a in ('AE2', 'AF2', 'AG2'): w2[a].font = grey
w2['AN2'] = 'Sem dado na fonte: 3Q26, 4Q26 e 2026 (rotulo em cinza). 2020 fica nas colunas AH-AL.'
code2model = {1: 15, 2: 20, 3: 24, 4: 36, 5: 40, 6: 89, 7: 97, 8: 101, 9: 102, 10: 104,
              11: 110, 12: 111, 13: 115, 14: 119, 15: 120, 16: 121, 17: 127, 18: 132, 19: 142, 20: 145, 21: 146, 22: 155,
              23: 164, 24: 168, 25: 170, 26: 171, 27: 178, 28: 179, 29: 180, 30: 183, 31: 201, 32: 204, 33: 205, 34: 206,
              35: 214, 37: 220, 38: 233, 39: 234, 40: 235, 41: 236}
code2row2 = {w2.cell(rr, 2).value: rr for rr in rows2}
present = set()
for rr in range(R0, LASTROW + 1):
    for cc in (1, 3, 5):
        if isinstance(inp.cell(rr, cc).value, int): present.add(inp.cell(rr, cc).value)
ANN = set(range(8, 108, 5))
for code, mr in code2model.items():
    if code not in present or code not in code2row2: continue
    r2 = code2row2[code]
    for c in range(4, 33):
        if c in ANN: continue
        cell = m.cell(mr, c)
        if cell.value not in (None, 0):
            if not (isinstance(cell.value, str) and cell.value.startswith('=')): raise AssertionError((mr, c, cell.value))
        cell.value = f"=INDEX('Inputs (2)'!$D${r2}:$AL${r2},MATCH({CL(c)}$7,'Inputs (2)'!$D$2:$AL$2,0))"
for c in range(4, 33):
    if c in ANN: continue
    m.cell(4, c).value = f"=INDEX('Inputs (2)'!$D$50:$AL$50,MATCH({CL(c)}$7,'Inputs (2)'!$D$2:$AL$2,0))"
for c in range(4, 33):
    for rw in (41, 143, 181, 202):
        f = m.cell(rw, c).value
        if isinstance(f, str) and f.startswith('=') and 'IFERROR' not in f: m.cell(rw, c).value = f'=IFERROR({f[1:]},0)'
m['C4'] = 'Net Income reportado (lucro liquido) - base do check IS'
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
json.dump(dict(pos=pos, R0=R0, LASTROW=LASTROW, END=END), open(OUT + '.layout.json', 'w'))
unk = sorted(set(n for (n, c) in order if n not in KNOWN))
print('subtotais ignorados:', sorted(set(n for _, n in SUBTOT)), '| fora de face:', DROPPED)
print('Inputs', R0, LASTROW, 'END', END, '| codes', sorted(present), '| sem face:', missing, '| tags genéricas:', unk)
