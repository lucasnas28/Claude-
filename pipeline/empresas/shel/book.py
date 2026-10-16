"""Lê o Quarterly Databook da Shell (USD milhões, trimestre isolado)."""
import openpyxl
P = [f'{q}Q{y:02d}' for y in range(20, 27) for q in range(1, 5)][:26]
QLAB = {f'{q}Q{y:02d}': f'Q{q} 20{y:02d}' for y in range(20, 27) for q in range(1, 5)}


def hdr(ws, row):
    out = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row, c).value
        if isinstance(v, str) and v.startswith('Q') and ' ' in v:
            q, y = v.split()
            out[f'{q[1]}Q{y[2:]}'] = c
    return out


def line(ws, h, row):
    return {p: float(ws.cell(row, c).value) for p, c in h.items()
            if isinstance(ws.cell(row, c).value, (int, float))}


def find_row(ws, text, start=1):
    t = text.strip().lower()
    for r in range(start, ws.max_row + 1):
        b = str(ws.cell(r, 2).value or '').strip().lower()
        if b == t or b.rstrip(':') == t.rstrip(':'):
            return r
    raise KeyError(text)


def third_party(ws, h):
    """Primeira linha 'Third-party' imediatamente abaixo de 'Revenue:' (as reported)."""
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, 2).value or '').strip() == 'Revenue:':
            if str(ws.cell(r + 1, 2).value or '').strip() == 'Third-party':
                return line(ws, h, r + 1)
    raise KeyError('Third-party')


def load(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    is_ = wb['Statement of Income']; ih = hdr(is_, 7)
    bs = wb['Balance Sheet']; bh = hdr(bs, 6)
    cf = wb['Cash Flows']; ch = hdr(cf, 7)
    sh = wb['Shell']; shh = hdr(sh, 7)
    def IS(name): return line(is_, ih, find_row(is_, name))
    def BS(name): return line(bs, bh, find_row(bs, name))
    def CF(name): return line(cf, ch, find_row(cf, name))
    out = {
        'have': sorted(ih, key=lambda p: P.index(p) if p in P else 99),
        'rev': IS('Revenue'),
        'jv': IS('Share of profit/(loss) of joint ventures and associates'),
        'othinc': IS('Interest and other income/(expenses)'),
        'purch': IS('Purchases'),
        'prod': IS('Production and manufacturing expenses'),
        'sell': IS('Selling, distribution and administrative expenses'),
        'rd': IS('Research and development'),
        'expl': IS('Exploration'),
        'da': IS('Depreciation, depletion and amortisation'),
        'ie': IS('Interest expense'),
        'pt': IS('Income/(loss) before taxation'),
        'tax': IS('Taxation charge/ (credit)'),
        'ni_period': IS('Income/(loss) for the period'),
        'nci': IS('Income/(loss) attributable to non-controlling interest'),
        'ni': IS('Income/(loss) attributable to Shell plc shareholders'),
        'div': {p: abs(v) for p, v in CF('- Shell plc shareholders1').items()},
        'cfo': CF('Cash flow from operating activities'),
        'cfi': CF('Cash flow from investing activities'),
        'cff': CF('Cash flow from financing activities'),
        'capex': {p: abs(v) for p, v in CF('Capital expenditure').items()},
        'bs': {
            'gw': BS('Goodwill'), 'intan': BS('Other intangible assets'),
            'ppe': BS('Property, plant and equipment'),
            'jv': BS('Joint ventures and associates'),
            'sec': BS('Investments in securities'),
            'dta': BS('Deferred tax'),
            'ret_a': line(bs, bh, 16),
            'rec_nc': line(bs, bh, 17),
            'der_nca': line(bs, bh, 18),
            'inv': BS('Inventories'),
            'rec': line(bs, bh, 22),
            'der_ca': line(bs, bh, 23),
            'cash': BS('Cash and cash equivalents'),
            'hfs': line(bs, bh, 26),
            'ta': BS('Total assets'),
            'debt_nc': line(bs, bh, 33),
            'pay_nc': line(bs, bh, 34),
            'der_ncl': line(bs, bh, 35),
            'dtl': line(bs, bh, 36),
            'ret_l': line(bs, bh, 37),
            'decom_nc': line(bs, bh, 38),
            'debt_c': line(bs, bh, 41),
            'pay': line(bs, bh, 42),
            'der_cl': line(bs, bh, 43),
            'taxp': BS('Income taxes payable'),
            'decom_c': line(bs, bh, 45),
            'hfs_l': line(bs, bh, 47),
            'tl': BS('Total liabilities'),
            'eq': BS('Equity attributable to Shell plc shareholders'),
            'nci': BS('Non-controlling interest'),
            'te': BS('Total equity'),
            'tle': BS('Total liabilities and equity'),
        },
        'tp': {
            'ig': third_party(wb['IG '], hdr(wb['IG '], 7)),
            'up': third_party(wb['UP'], hdr(wb['UP'], 7)),
            'mk': third_party(wb['MK'], hdr(wb['MK'], 7)),
            'cp': third_party(wb['CP'], hdr(wb['CP'], 7)),
            'res': third_party(wb['RES'], hdr(wb['RES'], 7)),
        },
        'ae': {
            'ig': line(sh, shh, 66), 'up': line(sh, shh, 67), 'mk': line(sh, shh, 68),
            'cp': line(sh, shh, 72), 'res': line(sh, shh, 75), 'co': line(sh, shh, 76),
        },
    }
    # lucros operacionais derivados: receita - compras - opex caixa - exploração - D&A (JV e outras receitas ficam no não operacional)
    op = {}
    for p in out['rev']:
        op[p] = (out['rev'][p] - out['purch'].get(p, 0) - out['prod'].get(p, 0)
                 - out['sell'].get(p, 0) - out['rd'].get(p, 0) - out['expl'].get(p, 0)
                 - out['da'].get(p, 0))
    out['op'] = op
    return out
