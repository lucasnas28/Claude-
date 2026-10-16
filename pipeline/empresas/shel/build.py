"""Shell plc: máscara padrão a partir do Quarterly Databook do RI (USD milhões).
uso: build.py MASCARA SAIDA"""
import sys, os, shutil, json, copy
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from book import load, P
import openpyxl
from openpyxl.utils import get_column_letter as CL
from openpyxl.styles import Font

MASK, OUT = sys.argv[1], sys.argv[2]
CFG = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'cfg.py')
g = __import__('runpy').run_path(CFG)
D = load(g['BOOK'])
HAVE = set(D['have'])

shutil.copyfile(MASK, OUT)
wb = openpyxl.load_workbook(OUT)
inp, w2, m = wb['Inputs'], wb['Inputs (2)'], wb['Model']
PC = {p: 11 + i for i, p in enumerate(P)}
col = lambda p: CL(PC[p])
FIRSTC, LASTC = col(P[0]), col(P[-1])
HDR, R0 = 4, 5
bold = Font(bold=True)
inp['A1'] = 'Fator p/ milhões:'; inp['B1'] = 1
inp['D1'] = (f"{g['NAME']} ({g['TICKER']}, CIK {int(g['CIK'])}) - fonte: Quarterly Databook Q2 2026 do RI (shell.com/investors). "
             "Moeda: USD. Brutos já em USD milhões (fator = 1). Trimestre isolado (a soma dos quatro trimestres fecha o ano, arredondamento de US$ 1 mi). "
             "Série da fonte: 1T21 a 2T26. 1T20-4T20 ficam vazios: o databook vigente começa em Q1 2021. IFRS: DRE em etapa única, sem lucro operacional divulgado.")
for i, h in enumerate(['COD1', 'SINAL1', 'COD2', 'SINAL2', 'COD3', 'SINAL3', 'UNID', 'Conta (PT) / origem', 'Account (EN)', 'ORIGEM']):
    inp.cell(HDR, 1 + i, h).font = bold
for p in P:
    inp.cell(HDR, PC[p], p).font = bold
    inp.cell(2, PC[p], p)
inp.cell(2, 10, 'Trimestre')

r = R0; pos = {}
def section(t):
    global r
    inp.cell(r, 8, t).font = bold; r += 1
def write(cods, desc, origem, vals=None, forms=None, key=None, unid='moeda', fmt='#,##0.0;(#,##0.0)'):
    global r
    for i, (c, s) in enumerate(cods): inp.cell(r, 1 + 2 * i, c); inp.cell(r, 2 + 2 * i, s)
    vals = {p: v for p, v in (vals or {}).items() if v is not None and p in P and p in HAVE}
    inp.cell(r, 7, unid); inp.cell(r, 8, desc); inp.cell(r, 10, origem)
    for p, x in vals.items(): inp.cell(r, PC[p], x)
    forms = {p: f for p, f in (forms or {}).items() if p in HAVE}
    for p, f in forms.items(): inp.cell(r, PC[p], f)
    for p in P: inp.cell(r, PC[p]).number_format = fmt
    if key: pos[key] = r
    r += 1
def fx(expr):
    return {p: '=' + expr.format(**{k: f'{col(p)}{v}' for k, v in pos.items()}) for p in HAVE}

section('DRE - trimestre isolado (USD milhões). Fonte: aba Statement of Income do databook')
write([(1, 1)], 'Receita [Revenue] databook Statement of Income', 'fonte', D['rev'], key='rev')
write([(2, -1)], 'Compras / custo da receita (divulgado positivo) [Purchases]', 'fonte', D['purch'], key='cogs')
write([], 'Despesas de produção e fabricação [Production and manufacturing expenses] (memória; entram nas despesas operacionais)', 'memoria', D['prod'], key='prod')
write([], 'Despesas de vendas, distribuição e administrativas [Selling, distribution and administrative expenses]', 'memoria', D['sell'], key='sell')
write([], 'P&D [Research and development]', 'memoria', D['rd'], key='rd')
write([], 'Exploração [Exploration]', 'memoria', D['expl'], key='expl')
write([], 'Lucro operacional DERIVADO = receita − compras − produção − vendas/adm − P&D − exploração − D&A. A Shell não divulga EBIT; JV e outras receitas ficam no não operacional.', 'memoria', D['op'], key='op')
write([(3, -1)], 'Despesas operacionais = receita − custo − lucro operacional (produção, vendas/adm, P&D, exploração, D&A)', 'derivada', forms=fx('N({rev})-N({cogs})-N({op})'), key='opx')
write([], 'Lucro antes do IR [Income/(loss) before taxation]', 'memoria', D['pt'], key='pt')
write([(6, 1)], 'Resultado não operacional total = lucro antes do IR − lucro operacional (JV + outras receitas − juros)', 'derivada', forms=fx('N({pt})-N({op})'), key='nonop')
write([], 'Equivalência patrimonial [Share of profit/(loss) of joint ventures and associates] (memória; já está no código 6 via não operacional)', 'memoria', D['jv'], key='jv')
write([], 'Juros e outras receitas/(despesas) [Interest and other income/(expenses)] (memória; já está no código 6 via não operacional)', 'memoria', D['othinc'], key='othinc')
write([(5, -1), (6, 1)], 'Despesa de juros (divulgado positivo): COD1=5 lança; COD2=6 devolve ao não operacional, que já a contém [Interest expense]', 'fonte', D['ie'], key='ie')
write([(7, -1)], 'Imposto de renda (charge positivo / credit negativo) [Taxation charge/(credit)]', 'fonte', D['tax'], key='tax')
write([(9, -1)], 'Lucro atribuído a não controladores [Income/(loss) attributable to non-controlling interest]', 'fonte', D['nci'], key='mi')
write([], 'Lucro líquido atribuído à Shell plc [Income/(loss) attributable to Shell plc shareholders]', 'memoria', D['ni'], key='ni_raw')
write([(8, 1)], 'Resultado após IR não classificado: lucro − (LAIR − IR − minoritários); absorve arredondamento de US$ 1 mi', 'derivada',
      forms=fx('N({ni_raw})-(N({pt})-N({tax})-N({mi}))'), key='oth')
write([(44, 1)], 'Lucro líquido (check IS, código 44)', 'derivada', forms=fx('N({ni_raw})'), key='ni')

section('DFC - trimestre isolado (USD milhões). Dividendos: databook negativo, lançados em valor absoluto')
write([(4, -1), (3, 1)], 'D&A da DRE. COD1=4 lança; COD2=3 devolve às despesas operacionais (a D&A vem em linha própria, não no custo) [Depreciation, depletion and amortisation]', 'fonte', D['da'], key='da')
write([(10, -1)], 'Dividendos pagos aos acionistas da Shell plc (valor absoluto do fluxo) [Cash dividends paid to Shell plc shareholders]', 'fonte', D['div'], key='div')
write([], 'CFO [Cash flow from operating activities] memória', 'memoria', D['cfo'])
write([], 'CFI [Cash flow from investing activities] memória', 'memoria', D['cfi'])
write([], 'CFF [Cash flow from financing activities] memória', 'memoria', D['cff'])
write([], 'Capex (valor absoluto) [Capital expenditure] memória', 'memoria', D['capex'])

section('BALANÇO (USD milhões, saldo no fim do trimestre). IFRS: não circulante primeiro. Arrendamentos dentro de PPE e Debt.')
B = D['bs']
akeys, lkeys = [], []
def bs(code, key, desc):
    write([(code, 1)], desc, 'fonte', B[key], key='bs_' + key)
    (akeys if code <= 22 else lkeys).append('bs_' + key)
bs(20, 'gw', 'Goodwill')
bs(21, 'intan', 'Outros intangíveis [Other intangible assets]')
bs(18, 'ppe', 'Imobilizado (inclui direito de uso IFRS 16) [Property, plant and equipment]')
bs(17, 'jv', 'Investimentos em joint ventures e coligadas')
bs(22, 'sec', 'Investimentos em títulos [Investments in securities]')
bs(14, 'dta', 'IR diferido ativo [Deferred tax, não circulante]')
bs(22, 'ret_a', 'Ativo de benefícios de aposentadoria [Retirement benefits]')
bs(22, 'rec_nc', 'Contas a receber não circulantes [Trade and other receivables]')
bs(22, 'der_nca', 'Derivativos ativo não circulante')
bs(13, 'inv', 'Estoques [Inventories]')
bs(12, 'rec', 'Contas a receber circulantes [Trade and other receivables]')
bs(16, 'der_ca', 'Derivativos ativo circulante')
bs(11, 'cash', 'Caixa e equivalentes [Cash and cash equivalents]')
bs(16, 'hfs', 'Ativos classificados como mantidos para venda')
bs(30, 'debt_nc', 'Dívida não circulante (inclui arrendamentos) [Debt]')
bs(34, 'pay_nc', 'Fornecedores e outras contas não circulantes')
bs(34, 'der_ncl', 'Derivativos passivo não circulante')
bs(33, 'dtl', 'IR diferido passivo')
bs(34, 'ret_l', 'Passivo de benefícios de aposentadoria')
bs(34, 'decom_nc', 'Descomissionamento e outras provisões não circulantes')
bs(27, 'debt_c', 'Dívida circulante (inclui arrendamentos) [Debt]')
bs(23, 'pay', 'Fornecedores e outras contas circulantes')
bs(26, 'der_cl', 'Derivativos passivo circulante')
bs(25, 'taxp', 'Impostos a pagar [Income taxes payable]')
bs(26, 'decom_c', 'Descomissionamento e outras provisões circulantes')
bs(26, 'hfs_l', 'Passivos associados a ativos mantidos para venda')
bs(35, 'nci', 'Participação de não controladores (balanço)')
bs(41, 'eq', 'Patrimônio atribuível aos acionistas da Shell plc (capital, tesouraria e OCI não abertos na face condensada)')

# resíduo de arredondamento (< US$ 50 mi)
def resid(keys, tot, code, lado):
    out = {}
    for p in HAVE:
        if p not in tot: continue
        sm = sum(B[k].get(p, 0) for k in keys)
        dlt = tot[p] - sm
        if 0.4 < abs(dlt) < 50: out[p] = dlt
    if out:
        write([(code, 1)], f'Diferença de arredondamento do {lado} (total divulgado − soma das linhas; < US$ 50 mi)', 'derivada', out, key='rnd_' + lado[:3])
        (akeys if code <= 22 else lkeys).append('rnd_' + lado[:3])
resid(['gw','intan','ppe','jv','sec','dta','ret_a','rec_nc','der_nca','inv','rec','der_ca','cash','hfs'], B['ta'], 16, 'ativo')
resid(['debt_nc','pay_nc','der_ncl','dtl','ret_l','decom_nc','debt_c','pay','der_cl','taxp','decom_c','hfs_l','nci','eq'], B['tle'], 26, 'passivo + PL')

ABSENT = {15: 'Despesas antecipadas: não abertas na face condensada',
          19: 'Direito de uso: dentro do imobilizado (IFRS 16)', 24: 'Despesas provisionadas: dentro de fornecedores',
          28: 'Revolving: dentro da dívida', 29: 'Arrendamento CP: dentro da dívida circulante',
          31: 'Arrendamento LP: dentro da dívida não circulante', 32: 'Dívida conversível: não divulgada',
          37: 'Capital + APIC: não aberto (tudo no código 41)', 38: 'Tesouraria: não aberta (tudo no código 41)',
          39: 'Outros resultados abrangentes: não abertos (tudo no código 41)', 40: 'Outras contas do PL'}
have_codes = set()
for rr in range(R0, r):
    for cc in (1, 3, 5):
        if isinstance(inp.cell(rr, cc).value, int): have_codes.add(inp.cell(rr, cc).value)
# sem código numérico: linha vazia com código 15/19/24/... faria a Inputs (2) devolver #N/A (cnt=0) e quebraria Total Assets
for code, nm in ABSENT.items():
    if code not in have_codes:
        write([], f'[vazio] código {code}: {nm}', 'derivada', {})

section('BALANÇO - totais divulgados (memória)')
write([], 'Total do ativo [Total assets]', 'memoria', B['ta'], key='ta')
write([], 'Total do passivo + PL [Total liabilities and equity]', 'memoria', B['tle'], key='tle')
section('CHECKS (devem ser 0)')
def chk(desc, terms, total, key):
    global r
    inp.cell(r, 7, 'moeda'); inp.cell(r, 8, desc); inp.cell(r, 10, 'derivada')
    for p in HAVE:
        e = '+'.join(f'N({col(p)}{pos[t]})' for t in terms)
        inp.cell(r, PC[p], f'=ROUND({e}-N({col(p)}{pos[total]}),1)'); inp.cell(r, PC[p]).number_format = '#,##0.0;(#,##0.0)'
    pos[key] = r; r += 1
chk('CHECK: contas de ativo - ativo total (USD milhões)', akeys, 'ta', 'chk_a')
chk('CHECK: passivo + PL (contas) - passivo + PL total', lkeys, 'tle', 'chk_l')
inp.cell(r, 8, 'CHECK: receita - custo - despesas + não operacional - IR - minoritários + outros após IR - lucro (deve ser 0)'); inp.cell(r, 10, 'derivada')
for p in HAVE:
    c = col(p)
    inp.cell(r, PC[p], f'=ROUND(N({c}{pos["rev"]})-N({c}{pos["cogs"]})-N({c}{pos["opx"]})+N({c}{pos["nonop"]})-N({c}{pos["tax"]})-N({c}{pos["mi"]})+N({c}{pos["oth"]})-N({c}{pos["ni"]}),1)')
r += 1
inp.cell(r, 8, '[vazio] 1T20-4T20: o Quarterly Databook Q2 2026 começa em Q1 2021. Não interpolar.'); r += 1
inp.cell(r, 8, '[vazio] Direito de uso e arrendamentos: a face condensada IFRS junta no imobilizado e na dívida.'); r += 1
LASTROW = r - 1
END = ((LASTROW + 60 + 9) // 10) * 10
conv = END + 3
inp.cell(conv, 8, 'CONVENCOES').font = bold
for i, t in enumerate([
    f'Linhas de origem: {R0} a {LASTROW}. As formulas da Inputs (2) varrem as linhas {R0} a {END}; as linhas {LASTROW + 1} a {END} estao livres para novas contas.',
    f'Periodos: colunas {FIRSTC} a {LASTC} = 1Q20 a 2Q26. Primeiro trimestre com dado = 1Q21.',
    'Brutos em USD milhões: fator da célula B1 = 1. UNID = moeda.',
    'ORIGEM: fonte = databook; memoria = como divulgado, sem codigo; derivada = formula nesta aba.',
    'Balanço: PL dos controladores no código 41 (face condensada não abre capital/tesouraria/OCI).']):
    inp.cell(conv + 1 + i, 8, t)
for c_, w in [(1, 6), (2, 7), (3, 6), (4, 7), (5, 6), (6, 7), (7, 8), (8, 95), (9, 20), (10, 10)]:
    inp.column_dimensions[CL(c_)].width = w
for p in P: inp.column_dimensions[col(p)].width = 13
inp.freeze_panes = inp.cell(R0, 11)

# Inputs (2) e Model — mesmo cabeamento do build_gen, sem ligar 1Q20-4Q20 (sem dado)
for i, h in enumerate(['1Q20', '2Q20', '3Q20', '4Q20', 2020]):
    c = 34 + i
    w2.cell(2, c, h); w2.cell(2, c)._style = copy.copy(w2.cell(2, 29 + i if i < 4 else 33)._style)
    w2.column_dimensions[CL(c)].width = w2.column_dimensions['AC'].width or 10
w2['B50'] = 44; w2['C50'] = 'Net income reported (lucro liquido) - memo p/ check IS [Model linha 4]'
for c in range(1, 39): w2.cell(50, c)._style = copy.copy(w2.cell(13, c)._style)
IN = 'Inputs!'
def rng(cl): return f'{IN}${cl}${R0}:${cl}${END}'
VAL = f'INDEX({IN}${FIRSTC}${R0}:${LASTC}${END},0,MATCH({{h}},{IN}${FIRSTC}${HDR}:${LASTC}${HDR},0))'
ESC = f'((({rng("G")}="moeda")+({rng("G")}="acoes"))*{IN}$B$1+(({rng("G")}<>"moeda")*({rng("G")}<>"acoes")))'
def f2(h, d_):
    v_ = VAL.format(h=h)
    cnt = '+'.join(f'SUMPRODUCT(({rng(c)}={d_})*ISNUMBER({v_}))' for c in 'ACE')
    parts = '+'.join(f'SUMPRODUCT(({rng(c)}={d_})*{rng(s)}*{ESC}*{v_})' for c, s in (('A', 'B'), ('C', 'D'), ('E', 'F')))
    return f'=IF(ISNA(MATCH({h},{IN}${FIRSTC}${HDR}:${LASTC}${HDR},0)),NA(),IF({cnt}=0,NA(),{parts}))'
rows2 = [rr for rr in range(4, 51) if isinstance(w2.cell(rr, 2).value, int)]
ANN2 = {8, 13, 18, 23, 28, 33, 38}
for rr in rows2:
    for c in [c for c in range(4, 38) if c not in ANN2]:
        assert w2.cell(rr, c).value is None, (rr, c)
        w2.cell(rr, c).value = f2(f'{CL(c)}$2', f'$B{rr}')
    w2.cell(rr, 38, f'=SUM(AH{rr}:AK{rr})' if rr <= 13 or rr == 50 else f'=AK{rr}')
    if rr == 50:
        for c, (a, b) in {8: ('D', 'G'), 13: ('I', 'L'), 18: ('N', 'Q'), 23: ('S', 'V'), 28: ('X', 'AA'), 33: ('AC', 'AF')}.items():
            w2.cell(rr, c, f'=SUM({a}{rr}:{b}{rr})')
    else:
        w2.cell(rr, 38)._style = copy.copy(w2.cell(rr, 33)._style)
        for i in range(4): w2.cell(rr, 34 + i)._style = copy.copy(w2.cell(rr, 29 + i)._style)
if w2['AG2'].value == 2025: w2['AG2'] = 2026
grey = Font(color='FF808080', bold=w2['AE2'].font.bold, name=w2['AE2'].font.name, size=w2['AE2'].font.sz)
for a in ('AE2', 'AF2', 'AG2'): w2[a].font = grey
w2['AN2'] = 'Sem dado na fonte: 1Q20-4Q20 (databook começa em 1Q21); 3Q26, 4Q26 e 2026 (rotulo em cinza).'
code2model = {1: 15, 2: 20, 3: 24, 4: 36, 5: 40, 6: 89, 7: 97, 8: 101, 9: 102, 10: 104,
              11: 110, 12: 111, 13: 115, 14: 119, 15: 120, 16: 121, 17: 127, 18: 132, 19: 142, 20: 145, 21: 146, 22: 155,
              23: 164, 24: 168, 25: 170, 26: 171, 27: 178, 28: 179, 29: 180, 30: 183, 31: 201, 32: 204, 33: 205, 34: 206,
              35: 214, 37: 220, 38: 233, 39: 234, 40: 235, 41: 236}
code2row2 = {w2.cell(rr, 2).value: rr for rr in rows2}
present = set()
for rr in range(R0, LASTROW + 1):
    for cc in (1, 3, 5):
        if isinstance(inp.cell(rr, cc).value, int): present.add(inp.cell(rr, cc).value)
ANN = set(range(8, 108, 5))
SKIP2020 = {4, 5, 6, 7}   # D-G = 1Q20-4Q20, sem dado
for code, mr in code2model.items():
    if code not in present or code not in code2row2: continue
    r2 = code2row2[code]
    for c in range(4, 33):
        if c in ANN or c in SKIP2020: continue
        cell = m.cell(mr, c)
        if cell.value not in (None, 0):
            if not (isinstance(cell.value, str) and cell.value.startswith('=')): raise AssertionError((mr, c, cell.value))
        cell.value = f"=INDEX('Inputs (2)'!$D${r2}:$AL${r2},MATCH({CL(c)}$7,'Inputs (2)'!$D$2:$AL$2,0))"
for c in range(4, 33):
    if c in ANN or c in SKIP2020: continue
    m.cell(4, c).value = f"=INDEX('Inputs (2)'!$D$50:$AL$50,MATCH({CL(c)}$7,'Inputs (2)'!$D$2:$AL$2,0))"
for c in range(4, 33):
    for rw in (41, 143, 181, 202):
        f = m.cell(rw, c).value
        if isinstance(f, str) and f.startswith('=') and 'IFERROR' not in f: m.cell(rw, c).value = f'=IFERROR({f[1:]},0)'
m['C4'] = 'Net Income reportado (lucro liquido) - base do check IS'
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
json.dump(dict(pos=pos, R0=R0, LASTROW=LASTROW, END=END), open(OUT + '.layout.json', 'w'))
print('Inputs', R0, LASTROW, 'END', END, '| codes', sorted(present), '| periodos', sorted(HAVE, key=P.index))
