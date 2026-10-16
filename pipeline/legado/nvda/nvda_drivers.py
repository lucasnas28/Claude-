import sys, re, copy, json

import openpyxl
from openpyxl.utils import get_column_letter as CL
from openpyxl.formula.translate import Translator
from openpyxl.styles import Font

IN, OUT, LAYF, B0, NROWS = sys.argv[1], sys.argv[2], sys.argv[3], int(sys.argv[4]), int(sys.argv[5])
lay = json.load(open(LAYF)); pos, R0, OLD_END, LAST0 = lay['pos'], lay['R0'], lay['END'], lay['LASTROW']
wb = openpyxl.load_workbook(IN)
inp, w2, m = wb['Inputs'], wb['Inputs (2)'], wb['Model']
P = [f'{q}Q{y:02d}' for y in range(20, 27) for q in range(1, 5)][:26]
SQ = {p: f'{p[0]}T{p[2:]}' for p in P}
PC = {p: 11 + i for i, p in enumerate(P)}
bold = Font(bold=True)
ANN = set(range(8, 108, 5))
def tr(f, s, d): return Translator(f, origin=s).translate_formula(d)
sys.path.insert(0, __import__('os').environ.get('PIPE_DATA', 'dados')+'/nvda')
from nv import qser
def merge(*ds):
    out = {}
    for d_ in ds:
        for k, v in d_.items(): out.setdefault(k, v)
    return out
RV, OI = 'Revenues', 'OperatingIncomeLoss'
NOTE = 'trimestre-calendário; o 4T (nov-jan, FQ4) = ano fiscal do 10-K menos 9M do 10-Q de out'
ITEMS = [
 ('p_dc', f'Receita Data Center (USD; {NOTE})', 'moeda', qser(RV, 'nvda:DataCenterMember'), '#,##0', 'sum'),
 ('p_cmp', 'Data Center - Compute (USD; aberto de 1T23 a 4T25)', 'moeda', qser(RV, 'nvda:ComputeMember'), '#,##0', 'sum'),
 ('p_net', 'Data Center - Networking (USD; aberto de 1T23 a 4T25)', 'moeda', qser(RV, 'nvda:NetworkingMember'), '#,##0', 'sum'),
 ('p_hyp', 'Data Center - Hyperscale (USD; nova abertura do FY27, com 1T25 e 2T25 reapresentados)', 'moeda', qser(RV, 'nvda:HyperscaleMember'), '#,##0', 'sum'),
 ('p_aic', 'Data Center - AI Clouds, Industrial & Enterprise (USD; nova abertura do FY27)', 'moeda', qser(RV, 'nvda:AICloudsIndustrialEnterpriseMember'), '#,##0', 'sum'),
 ('p_gam', 'Receita Gaming (USD; até 4T25)', 'moeda', qser(RV, 'nvda:GamingMember'), '#,##0', 'sum'),
 ('p_pv', 'Receita Professional Visualization (USD; até 4T25)', 'moeda', qser(RV, 'nvda:ProfessionalVisualizationMember'), '#,##0', 'sum'),
 ('p_aut', 'Receita Automotive (USD; até 4T25)', 'moeda', qser(RV, 'nvda:AutomotiveMember'), '#,##0', 'sum'),
 ('p_oem', 'Receita OEM & Other (USD; até 4T25)', 'moeda', qser(RV, 'nvda:OEMAndOtherMember'), '#,##0', 'sum'),
 ('p_edg', 'Receita Edge Computing (USD; plataforma nova do FY27 que reúne o que não é Data Center)', 'moeda', qser(RV, 'nvda:EdgeComputingMember'), '#,##0', 'sum'),
 ('s_cn', 'Segmento Compute & Networking - receita (USD)', 'moeda', merge(qser(RV, 'nvda:ComputeAndNetworkingSegmentMember'), qser(RV, 'nvda:ComputeAndNetworkingMember')), '#,##0', 'sum'),
 ('s_gr', 'Segmento Graphics - receita (USD)', 'moeda', merge(qser(RV, 'nvda:GraphicsSegmentMember'), qser(RV, 'nvda:GraphicsMember')), '#,##0', 'sum'),
 ('o_cn', 'Segmento Compute & Networking - lucro operacional (USD)', 'moeda', merge(qser(OI, 'nvda:ComputeAndNetworkingSegmentMember'), qser(OI, 'nvda:ComputeAndNetworkingMember')), '#,##0', 'sum'),
 ('o_gr', 'Segmento Graphics - lucro operacional (USD)', 'moeda', merge(qser(OI, 'nvda:GraphicsSegmentMember'), qser(OI, 'nvda:GraphicsMember')), '#,##0', 'sum'),
 ('g_us', 'Receita por local de faturamento - Estados Unidos (USD; como divulgado originalmente: os 10-Qs do FY26 reapresentaram a abertura geográfica de 3T24-2T25 e o modelo mantém a base original para o 4T = ano - 9M ficar coerente)', 'moeda', qser(RV, 'country:US', orig=True), '#,##0', 'sum'),
 ('g_tw', 'Receita por local de faturamento - Taiwan (USD)', 'moeda', qser(RV, 'country:TW', orig=True), '#,##0', 'sum'),
 ('g_ch', 'Receita por local de faturamento - China, inclui Hong Kong (USD)', 'moeda', qser(RV, 'nvda:ChinaIncludingHongKongMember', orig=True), '#,##0', 'sum'),
]
# ---------------- Inputs memo rows
oldconv = OLD_END + 3
conv_texts = [inp.cell(rr, 8).value for rr in range(oldconv + 1, oldconv + 12) if inp.cell(rr, 8).value]
for rr in range(oldconv, oldconv + 12): inp.cell(rr, 8).value = None
r = LAST0 + 2
inp.cell(r, 8, 'DRIVERS DE RECEITA (codigos 45+): receita por plataforma de mercado, segmentos operacionais (receita e lucro operacional) e receita por local de faturamento - SEC XBRL (instâncias dos 10-Q/10-K; USD, fator 0,000001).').font = bold; r += 1
code = 45; row_of = {}
for key, label, unid, vals, fmt, ann in ITEMS:
    inp.cell(r, 1, code); inp.cell(r, 2, 1); inp.cell(r, 7, unid)
    inp.cell(r, 8, label); inp.cell(r, 10, 'fonte')
    for p in P:
        if vals.get(p) is not None: inp.cell(r, PC[p], float(vals[p]))
        inp.cell(r, PC[p]).number_format = fmt
    row_of[key] = (r, code); code += 1; r += 1
inp.cell(r, 8, 'Notas: valores dimensionais lidos das instâncias XBRL; o trimestre nov-jan (FQ4) é o ano fiscal menos 9M, calculado na extração. Em 2026 (FY27) a NVIDIA trocou a abertura por plataforma: Data Center passou a abrir Hyperscale e AI Clouds/Industrial/Enterprise, e Gaming, ProViz, Automotive e OEM viraram Edge Computing; o 3T25 e o 4T25 ainda não foram reapresentados no formato novo.'); r += 1
LAST = r - 1
END = ((LAST + 30 + 9) // 10) * 10
conv = END + 3
inp.cell(conv, 8, 'CONVENCOES').font = bold
conv_texts[0] = f'Linhas de origem: {R0} a {LAST}. As formulas da Inputs (2) varrem as linhas {R0} a {END}; as linhas {LAST + 1} a {END} estao livres para novas contas.'
conv_texts.append(f'Codigos 45 a {code - 1}: drivers de receita (memoria); alimentam o bloco logo abaixo da receita liquida no Model.')
for i, t in enumerate(conv_texts): inp.cell(conv + 1 + i, 8, t)
rx = re.compile(r'(Inputs!\$[A-Z]{1,2}\$\d+:\$[A-Z]{1,2}\$)' + str(OLD_END) + r'\b')
for row in w2.iter_rows():
    for c in row:
        if isinstance(c.value, str) and 'Inputs!' in c.value:
            c.value = rx.sub(lambda g: g.group(1) + str(END), c.value)

# ---------------- Inputs (2) rows
row2 = {}
w2['C52'] = 'Drivers de receita (memo) - codigos 45+ [Model, bloco abaixo da receita liquida]'; w2['C52'].font = bold
annual2 = {8: ('D', 'G'), 13: ('I', 'L'), 18: ('N', 'Q'), 23: ('S', 'V'), 28: ('X', 'AA'), 33: ('AC', 'AF'), 38: ('AH', 'AK')}
for i, (key, label, unid, vals, fmt, ann) in enumerate(ITEMS):
    rr = 53 + i; row2[key] = rr
    w2.cell(rr, 2, row_of[key][1]); w2.cell(rr, 3, label)
    for c in range(1, 39): w2.cell(rr, c)._style = copy.copy(w2.cell(4, c)._style)
    f2fmt = '#,##0.0' if unid == 'moeda' else fmt
    for c in range(4, 38):
        if c in annual2: continue
        w2.cell(rr, c, tr(w2.cell(4, c).value, f'{CL(c)}4', f'{CL(c)}{rr}'))
        w2.cell(rr, c).number_format = f2fmt
    for c, (a, b) in annual2.items():
        w2.cell(rr, c, f'=SUM({a}{rr}:{b}{rr})' if ann == 'sum' else (f'={b}{rr}' if ann == 'last' else None))
        w2.cell(rr, c).number_format = f2fmt

# ---------------- Model block
COLS = list(range(4, 36))
def link(key, c): return f"IFERROR(INDEX('Inputs (2)'!$D${row2[key]}:$AL${row2[key]},MATCH({CL(c)}$7,'Inputs (2)'!$D$2:$AL$2,0)),\"\")"
def label_cell(rr, text, b=False):
    c = m.cell(rr, 3, text); c._style = copy.copy(m['C15']._style)
    if b: c.font = Font(bold=True, name=m['C15'].font.name, size=m['C15'].font.sz)
def hdr_row(rr, text):
    m.cell(rr, 3, text)._style = copy.copy(m['C13']._style)
    for c in COLS: m.cell(rr, c, f'={CL(c)}7')._style = copy.copy(m.cell(13, c)._style)
def fill(rr, fn, fmt='#,##0', annual='sum'):
    for c in COLS:
        cell = m.cell(rr, c); cell._style = copy.copy(m.cell(15, c)._style)
        a, b = CL(max(c - 4, 1)), CL(max(c - 1, 1))
        if c in ANN and annual == 'sum': cell.value = f'=IF(COUNT({a}{rr}:{b}{rr})=4,SUM({a}{rr}:{b}{rr}),"")'
        elif c in ANN and annual == 'last': cell.value = f'=IF(COUNT({b}{rr})=0,"",{b}{rr})'
        elif c in ANN and annual == 'blank': cell.value = None
        else: cell.value = fn(c)
        cell.number_format = fmt
def prevq(c):
    p = c - 1
    if p in ANN: p -= 1
    return p if p >= 4 else None
def yoy(src):
    def f(c):
        if c - 5 < 4: return '=""'
        return f'=IFERROR({CL(c)}{src}/{CL(c - 5)}{src}-1,"")'
    return f
def ratio(a, b, k=1):
    return lambda c: f'=IFERROR({CL(c)}{a}/{CL(c)}{b}' + (f'*{k}' if k != 1 else '') + ',"")'
for rr in range(B0, B0 + NROWS):
    for c in range(1, 36): assert m.cell(rr, c).value is None, (rr, c)
R = {}
rr = B0
COGS = 20 + NROWS
def line(key, lab, fmt='#,##0', annual='sum', fn=None, b=False):
    global rr
    R[key] = rr; fill(rr, fn or (lambda c, k=key: '=' + link(k, c)), fmt, annual); label_cell(rr, lab, b); rr += 1
def sumf(*keys):
    return lambda c: '=IF(COUNT(' + ','.join(f'{CL(c)}{R[k]}' for k in keys) + ')=0,"",' + '+'.join(f'N({CL(c)}{R[k]})' for k in keys) + ')'
hdr_row(rr, 'Receita por plataforma de mercado - como a NVIDIA divulga (USD milhões)'); rr += 1
line('p_dc', 'Data Center', b=True)
line('p_cmp', '   Compute (até 4T25)'); line('p_net', '   Networking (até 4T25)')
line('p_hyp', '   Hyperscale (nova abertura)'); line('p_aic', '   AI Clouds, Industrial & Enterprise (nova abertura)')
line('p_gam', 'Gaming (até 4T25)'); line('p_pv', 'Professional Visualization (até 4T25)'); line('p_aut', 'Automotive (até 4T25)'); line('p_oem', 'OEM & Other (até 4T25)')
line('p_edg', 'Edge Computing (1T25-2T25 reapresentados; soma só a partir de 1T26)')
def totf(c):
    k = lambda x: f'{CL(c)}{R[x]}'
    return (f'=IF(COUNT({k("p_dc")})=0,"",N({k("p_dc")})+IF(COUNT({k("p_gam")},{k("p_pv")},{k("p_aut")},{k("p_oem")})>0,'
            f'N({k("p_gam")})+N({k("p_pv")})+N({k("p_aut")})+N({k("p_oem")}),N({k("p_edg")})))')
line('tot', 'Receita total (Data Center + Gaming/ProViz/Auto/OEM; Edge Computing onde a abertura antiga não existe)', annual='calc', fn=totf, b=True)
line('chk', 'Check vs Net Revenues (linha 15)', annual='calc', fn=lambda c: f'=IF(COUNT({CL(c)}{R["tot"]})=0,"",ROUND({CL(c)}15-{CL(c)}{R["tot"]},1))')
line('y_dc', '   % YoY Data Center', '0.0%', 'calc', yoy(R['p_dc']))
line('s_dc', '   % Data Center na receita', '0.0%', 'calc', ratio(R['p_dc'], R['tot']))
rr += 1
hdr_row(rr, 'Segmentos operacionais (USD milhões)'); rr += 1
line('s_cn', '   Receita Compute & Networking'); line('s_gr', '   Receita Graphics')
line('s_chk', 'Check segmentos vs Net Revenues (linha 15)', annual='calc', fn=lambda c: f'=IF(COUNT({CL(c)}{R["s_cn"]},{CL(c)}{R["s_gr"]})=0,"",ROUND({CL(c)}15-N({CL(c)}{R["s_cn"]})-N({CL(c)}{R["s_gr"]}),1))')
line('o_cn', '   Lucro operacional Compute & Networking'); line('o_gr', '   Lucro operacional Graphics')
line('m_cn', '   Margem operacional Compute & Networking', '0.0%', 'calc', ratio(R['o_cn'], R['s_cn']))
line('m_gr', '   Margem operacional Graphics', '0.0%', 'calc', ratio(R['o_gr'], R['s_gr']))
rr += 1
hdr_row(rr, 'Receita por local de faturamento do cliente (USD milhões)'); rr += 1
line('g_us', '   Estados Unidos'); line('g_tw', '   Taiwan'); line('g_ch', '   China (inclui Hong Kong)')
line('g_ot', '   Outros países (receita total - os três)', annual='calc', fn=lambda c: f'=IF(COUNT({CL(c)}{R["g_us"]})=0,"",{CL(c)}15-N({CL(c)}{R["g_us"]})-N({CL(c)}{R["g_tw"]})-N({CL(c)}{R["g_ch"]}))')
line('y_ch', '   % YoY China', '0.0%', 'calc', yoy(R['g_ch']))
line('p_ch', '   % China na receita', '0.0%', 'calc', ratio(R['g_ch'], R['tot']))
END_BLOCK = rr
assert END_BLOCK <= B0 + NROWS, END_BLOCK
# annual columns for 'calc' ratio rows that need annual formula: fill() already put fn for 'calc'? ensure annual uses the same formula
for c in COLS:
    if c in ANN:
        for k in ('l_pct', 'l_pm', 'l_yoy', 'v_chk', 'v_dr', 'v_pm', 'v_yoy', 'ref_res', 'ref_m', 'rec', 'mb', 'mb_gap', 'conv', 'r_yoy'):
            pass
for c in COLS:
    m.cell(17, c, f'=IF(COUNT({CL(c)}{R["chk"]})=0,"",{CL(c)}{R["chk"]})')._style = copy.copy(m.cell(15, c)._style)
m['C17'] = 'Check Receita (linha 15 - soma das plataformas)'
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('Inputs rows to', LAST, 'END', OLD_END, '->', END, '| Model block', B0, '->', END_BLOCK - 1)
json.dump(R, open(OUT + '.rows.json', 'w'))
