import sys, shutil, copy, json
sys.path.insert(0, __import__('os').environ.get('PIPE_DATA', 'dados')+'/nvda')
from nv import series, first, label, inst_series, face
from bsmap import BSM
import openpyxl
from openpyxl.utils import get_column_letter as CL
from openpyxl.styles import Font

MASK, OUT = sys.argv[1:3]
P = [f'{q}Q{y:02d}' for y in range(20, 27) for q in range(1, 5)][:26]
SQ = {p: f'FQ{p[0]} FY{(int(p[2:]) + 1) % 100:02d}' for p in P}
shutil.copyfile(MASK, OUT)
wb = openpyxl.load_workbook(OUT)
inp, w2, m = wb['Inputs'], wb['Inputs (2)'], wb['Model']
PC = {p: 11 + i for i, p in enumerate(P)}
col = lambda p: CL(PC[p])
FIRSTC, LASTC = col(P[0]), col(P[-1])
HDR, R0 = 4, 5
bold = Font(bold=True)
inp['A1'] = 'Fator p/ milhões:'; inp['B1'] = 0.000001
inp['D1'] = ('NVIDIA Corp. (NVDA, CIK 1045810) - fonte: SEC EDGAR XBRL (companyfacts API e instâncias dos 10-Q/10-K). Moeda: USD. Brutos em USD unidades (fator 0,000001 -> milhões). '
             'Colunas = trimestres-CALENDÁRIO; o ano fiscal NVIDIA termina no fim de janeiro (FQ1 = fev-abr = 1T calendário; FQ4 = nov-jan = 4T calendário). Trimestre nov-jan (FQ4) = ano fiscal (10-K) menos 9M (10-Q de out), por fórmula.')
for i, h in enumerate(['COD1', 'SINAL1', 'COD2', 'SINAL2', 'COD3', 'SINAL3', 'UNID', 'Conta (PT) / origem', 'Account (EN)', 'ORIGEM']):
    inp.cell(HDR, 1 + i, h).font = bold
for p in P: inp.cell(HDR, PC[p], p).font = bold; inp.cell(2, PC[p], SQ[p])
inp.cell(2, 10, 'Trimestre fiscal NVIDIA')

r = R0; pos = {}
def section(t):
    global r
    inp.cell(r, 8, t).font = bold; r += 1
def write(cods, desc, origem, vals=None, forms=None, key=None, unid='moeda', zero_fill=True, fmt='#,##0.0;(#,##0.0)', avail=None):
    global r
    for i, (c, s) in enumerate(cods): inp.cell(r, 1 + 2 * i, c); inp.cell(r, 2 + 2 * i, s)
    vals = {p: v for p, v in (vals or {}).items() if v is not None}
    if cods and zero_fill and not forms:
        miss = [p for p in P if p not in vals and (avail is None or p in avail)]
        if miss: vals.update({p: 0.0 for p in miss})
    inp.cell(r, 7, unid); inp.cell(r, 8, desc); inp.cell(r, 10, origem)
    for p, x in vals.items(): inp.cell(r, PC[p], x)
    for p, f in (forms or {}).items(): inp.cell(r, PC[p], f)
    for p in P: inp.cell(r, PC[p]).number_format = fmt
    if key: pos[key] = r
    r += 1

def V(dct): return {p: x[0] for p, x in dct.items() if p in P}
def ACC(dct): return {p: x[0] for p, x in dct.items()}
section('DRE - BRUTO COMO DIVULGADO (USD). 3M = trimestre do 10-Q (FQ1-FQ3 = 1T, 2T, 3T calendário); 12M = ano fiscal do 10-K (na coluna do 4T calendário); 9M = acumulado do 10-Q de out (coluna do 3T)')
ISL = [('rev', 'Receita (revenue)', ['Revenues'], [(1, 1)]),
       ('cogs', 'Custo da receita (divulgado positivo)', ['CostOfRevenue'], [(2, -1)]),
       ('opx', 'Despesas operacionais totais = P&D + SG&A (+ custo de rescisão da aquisição da Arm no FQ1 FY23) (divulgado positivo)', ['OperatingExpenses'], [(3, -1)]),
       ('nonop', 'Resultado não operacional total = juros recebidos - juros pagos + outras (inclui ganhos/perdas em participações)', ['NonoperatingIncomeExpense'], [(6, 1)]),
       ('iexp', 'Despesa de juros (divulgado positivo): COD1=5 lança a despesa; COD2=6 devolve ao não operacional, que já a contém', ['InterestExpense', 'InterestExpenseNonoperating'], [(5, -1), (6, 1)]),
       ('iinc', 'Receita de juros (memória; dentro do não operacional)', ['InvestmentIncomeInterest'], []),
       ('tax', 'Imposto de renda (divulgado positivo)', ['IncomeTaxExpenseBenefit'], [(7, -1)]),
       ('ni', 'Lucro líquido', ['NetIncomeLoss'], [(44, 1)])]
RAWR = {}
for k, lab_, tags, cods in ISL:
    s3 = {p: v for p, v in V(first(tags, '3M')).items() if p[0] != '4'}
    tg = '/'.join(tags)
    write([], f'{lab_} - 3M (10-Q) [us-gaap:{tg}]', 'memoria', s3, key=f'r3_{k}')
    write([], f'{lab_} - ano fiscal 12M (10-K; na coluna do 4T)', 'memoria', V(first(tags, '12M')), key=f'r12_{k}')
    write([], f'{lab_} - 9M acumulado (10-Q de out; na coluna do 3T)', 'memoria', V(first(tags, '9M')), key=f'r9_{k}')
section('DRE TRIMESTRALIZADA (derivada por fórmula): 4T calendário (nov-jan) = ano fiscal - 9M')
def qf(k):
    out = {}
    for p in P:
        c = col(p)
        if p[0] == '4':
            out[p] = f'=N({c}{pos["r12_" + k]})-N({CL(PC[p] - 1)}{pos["r9_" + k]})'
        else:
            out[p] = f'={c}{pos["r3_" + k]}'
    return out
for k, lab_, tags, cods in ISL:
    write(cods, f'{lab_} - trimestre', 'derivada', forms=qf(k), key=f'q_{k}')
section('DFC - ACUMULADO NO ANO FISCAL (USD): 3M (abr), 6M (jul), 9M (out), 12M (jan)')
def ytd(tags):
    out = {}
    for kind, qq in (('3M', '1'), ('6M', '2'), ('9M', '3'), ('12M', '4')):
        for p, v in ACC(first(tags, kind)).items():
            if p[0] == qq: out[p] = v
    return out
for k, lab_, tags in [('da', 'Depreciação e amortização', ['DepreciationDepletionAndAmortization', 'DepreciationAndAmortization']), ('div', 'Dividendos pagos', ['PaymentsOfDividends'])]:
    y = ytd(tags)
    write([], f'{lab_} - acumulado no ano fiscal [us-gaap:{"/".join(tags)}]', 'memoria', {p: v for p, v in y.items() if p in P}, key=f'y_{k}')
def cf_f(k):
    out = {}
    for p in P:
        c = col(p)
        if p[0] == '1': out[p] = f'=N({c}{pos["y_" + k]})'
        else: out[p] = f'=N({c}{pos["y_" + k]})-N({CL(PC[p] - 1)}{pos["y_" + k]})'
    return out
write([(4, -1), (3, 1)], 'D&A trimestral (desacumulada). COD1=4 lança (positivo -> -1); COD2=3 devolve às despesas operacionais (a NVIDIA não abre D&A entre custo e despesa)', 'derivada', forms=cf_f('da'), key='q_da')
write([(10, -1)], 'Dividendos pagos - trimestre (desacumulado)', 'derivada', forms=cf_f('div'), key='q_div')
section('BALANÇO (USD, saldo no fim do trimestre; cada trimestre lido da instância XBRL do próprio 10-Q/10-K; 4T25 da coluna comparativa do 10-Q de abr/26, que separa títulos de dívida e ações listadas)')
FACE = face()
for code, sg, tags, desc, key in BSM:
    vals = {}
    for p in P:
        d_ = dict(FACE[p][0])
        for t in tags:
            if t in d_: vals[p] = d_[t]; break
    write([(code, sg)], f'{desc} [us-gaap/nvda:{"/".join(tags)}]', 'fonte', vals, key=key)
section('BALANÇO - totais divulgados (memória)')
write([], 'Total do ativo [us-gaap:Assets]', 'memoria', {p: dict(FACE[p][0])['Assets'] for p in P}, key='ta')
write([], 'Total do passivo [us-gaap:Liabilities]', 'memoria', {p: dict(FACE[p][0])['Liabilities'] for p in P}, key='tl')
write([], 'Total do PL [us-gaap:StockholdersEquity]', 'memoria', {p: dict(FACE[p][0])['StockholdersEquity'] for p in P}, key='te')
section('CHECKS (devem ser 0)')
def chk(desc, terms, totals, key, sg=None):
    global r
    inp.cell(r, 7, 'moeda'); inp.cell(r, 8, desc); inp.cell(r, 10, 'derivada')
    for p in P:
        e = '+'.join(f'{"-" if sg and sg[i] < 0 else ""}N({col(p)}{t})' for i, t in enumerate(terms)).replace('+-', '-') + ''.join(f'-N({col(p)}{t})' for t in totals)
        inp.cell(r, PC[p], f'=ROUND({e},-5)'); inp.cell(r, PC[p]).number_format = '#,##0;(#,##0)'
    pos[key] = r; r += 1
akeys = ['cash', 'mkt', 'ar', 'inv', 'oca', 'eqs', 'ppe', 'rou', 'gw', 'intg', 'dta', 'nmk', 'onca']
lkeys = ['ap', 'acc', 'std', 'ltd', 'lease', 'oncl']
ekeys = ['cs', 'apic', 'tsy', 'aoci', 're']
SGN = {k: sg for c, sg, t, d_, k in BSM}
chk('CHECK: contas de ativo - ativo total (USD, arredondado a 100 mil)', [pos[k] for k in akeys], [pos['ta']], 'chk_a')
chk('CHECK: contas de passivo - passivo total', [pos[k] for k in lkeys], [pos['tl']], 'chk_l')
chk('CHECK: contas de PL - PL total', [pos[k] for k in ekeys], [pos['te']], 'chk_e', [SGN[k] for k in ekeys])
inp.cell(r, 8, 'CHECK: receita - custo - despesas + não operacional - IR - lucro (deve ser 0)'); inp.cell(r, 10, 'derivada')
for p in P:
    c = col(p); inp.cell(r, PC[p], f'=ROUND(N({c}{pos["q_rev"]})-N({c}{pos["q_cogs"]})-N({c}{pos["q_opx"]})+N({c}{pos["q_nonop"]})-N({c}{pos["q_tax"]})-N({c}{pos["q_ni"]}),-5)')
r += 1
for c_, t in [(8, 'Equity in affiliates: inexistente na DRE (ganhos/perdas em participações ficam em outras receitas, cod 6)'), (9, 'Minority interest: inexistente'),
              (15, 'Prepaid: junto com outros ativos circulantes (cod 16)'), (17, 'Equity in affiliates: participações ficam em ativos circulantes (ações listadas) e outros ativos NC (non-marketable)'),
              (25, 'IR a pagar CP: dentro de despesas provisionadas (cod 24)'), (26, 'Outros passivos circulantes: dentro de despesas provisionadas (cod 24)'), (29, 'Arrendamento CP: dentro de despesas provisionadas (cod 24)'),
              (33, 'IR diferido passivo: dentro de outros passivos não circulantes'), (38, 'Tesouraria: só até out/21; a NVIDIA cancelou as ações em tesouraria no FQ4 FY22 (a partir daí, zero)')]:
    inp.cell(r, 8, f'[vazio] cod {c_} {t}'); r += 1
LASTROW = r - 1
END = ((LASTROW + 60 + 9) // 10) * 10
conv = END + 3
inp.cell(conv, 8, 'CONVENCOES').font = bold
for i, t in enumerate([
    f'Linhas de origem: {R0} a {LASTROW}. As formulas da Inputs (2) varrem as linhas {R0} a {END}; as linhas {LASTROW + 1} a {END} estao livres para novas contas.',
    f'Periodos: colunas {FIRSTC} a {LASTC} = trimestres-calendário 1Q20 a 2Q26 (linha 2: trimestre fiscal NVIDIA correspondente).',
    'Brutos em USD unidades: fator da célula B1 = 0,000001 converte para USD milhões nas linhas UNID = moeda.',
    'ORIGEM: fonte = copiado da planilha; memoria = como divulgado, sem codigo; derivada = formula nesta aba.']):
    inp.cell(conv + 1 + i, 8, t)
for c_, w in [(1, 6), (2, 7), (3, 6), (4, 7), (5, 6), (6, 7), (7, 8), (8, 95), (9, 20), (10, 10)]:
    inp.column_dimensions[CL(c_)].width = w
for p in P: inp.column_dimensions[col(p)].width = 13
inp.freeze_panes = inp.cell(R0, 11)

# ---------------------------------------------------------------- Inputs (2) and Model links (same as the other companies)
for i, h in enumerate(['1Q20', '2Q20', '3Q20', '4Q20', 2020]):
    c = 34 + i
    w2.cell(2, c, h); w2.cell(2, c)._style = copy.copy(w2.cell(2, 29 + i if i < 4 else 33)._style)
    w2.column_dimensions[CL(c)].width = w2.column_dimensions['AC'].width or 10
w2['B50'] = 44; w2['C50'] = 'Net income reported (lucro liquido) - memo p/ check IS [Model linha 4]'
for c in range(1, 39): w2.cell(50, c)._style = copy.copy(w2.cell(13, c)._style)
IN = 'Inputs!'
def rng(cl): return f'{IN}${cl}${R0}:${cl}${END}'
VAL = f'INDEX({IN}${FIRSTC}${R0}:${LASTC}${END},0,MATCH({{h}},{IN}${FIRSTC}${HDR}:${LASTC}${HDR},0))'
ESC = f'((({rng("G")}="moeda")+({rng("G")}="acoes"))*{IN}$B$1+(({rng("G")}<>"moeda")*({rng("G")}<>"acoes")))'
def f2(h, d_):
    v_ = VAL.format(h=h)
    cnt = '+'.join(f'SUMPRODUCT(({rng(c)}={d_})*ISNUMBER({v_}))' for c in 'ACE')
    parts = '+'.join(f'SUMPRODUCT(({rng(c)}={d_})*{rng(s)}*{ESC}*{v_})' for c, s in (('A', 'B'), ('C', 'D'), ('E', 'F')))
    return f'=IF(ISNA(MATCH({h},{IN}${FIRSTC}${HDR}:${LASTC}${HDR},0)),NA(),IF({cnt}=0,NA(),{parts}))'
rows2 = [rr for rr in range(4, 51) if isinstance(w2.cell(rr, 2).value, int)]
ANN2 = {8, 13, 18, 23, 28, 33, 38}
for rr in rows2:
    for c in [c for c in range(4, 38) if c not in ANN2]:
        assert w2.cell(rr, c).value is None, (rr, c)
        w2.cell(rr, c).value = f2(f'{CL(c)}$2', f'$B{rr}')
    w2.cell(rr, 38, f'=SUM(AH{rr}:AK{rr})' if rr <= 13 or rr == 50 else f'=AK{rr}')
    if rr == 50:
        for c, (a, b) in {8: ('D', 'G'), 13: ('I', 'L'), 18: ('N', 'Q'), 23: ('S', 'V'), 28: ('X', 'AA'), 33: ('AC', 'AF')}.items():
            w2.cell(rr, c, f'=SUM({a}{rr}:{b}{rr})')
    else:
        w2.cell(rr, 38)._style = copy.copy(w2.cell(rr, 33)._style)
        for i in range(4): w2.cell(rr, 34 + i)._style = copy.copy(w2.cell(rr, 29 + i)._style)
if w2['AG2'].value == 2025: w2['AG2'] = 2026
grey = Font(color='FF808080', bold=w2['AE2'].font.bold, name=w2['AE2'].font.name, size=w2['AE2'].font.sz)
for a in ('AE2', 'AF2', 'AG2'): w2[a].font = grey
w2['AN2'] = 'Sem dado na fonte: 3Q26, 4Q26 e 2026 (rotulo em cinza). 2020 fica nas colunas AH-AL.'

code2model = {20: 145, 1: 15, 2: 20, 3: 24, 4: 36, 5: 40, 6: 89, 7: 97, 8: 101, 9: 102, 10: 104,
              11: 110, 12: 111, 13: 115, 14: 119, 15: 120, 16: 121, 17: 127, 18: 132, 19: 142, 21: 146, 22: 155,
              23: 164, 24: 168, 25: 170, 26: 171, 27: 178, 29: 180, 30: 183, 31: 201, 33: 205, 34: 206,
              35: 214, 37: 220, 38: 233, 39: 234, 40: 235, 41: 236}
code2row2 = {w2.cell(rr, 2).value: rr for rr in rows2}
present = set()
for rr in range(R0, LASTROW + 1):
    for cc in (1, 3, 5):
        if isinstance(inp.cell(rr, cc).value, int): present.add(inp.cell(rr, cc).value)
ANN = set(range(8, 108, 5))
for code, mr in code2model.items():
    if code not in present: continue
    r2 = code2row2[code]
    for c in range(4, 33):
        if c in ANN: continue
        cell = m.cell(mr, c)
        assert cell.value in (None, 0), (mr, c, cell.value)
        cell.value = f"=INDEX('Inputs (2)'!$D${r2}:$AL${r2},MATCH({CL(c)}$7,'Inputs (2)'!$D$2:$AL$2,0))"
for c in range(4, 33):
    if c in ANN: continue
    m.cell(4, c).value = f"=INDEX('Inputs (2)'!$D$50:$AL$50,MATCH({CL(c)}$7,'Inputs (2)'!$D$2:$AL$2,0))"
for c in range(4, 33):
  for rw in (41, 143, 181, 202):
    f = m.cell(rw, c).value
    if isinstance(f, str) and f.startswith('=') and 'IFERROR' not in f: m.cell(rw, c).value = f'=IFERROR({f[1:]},0)'
m['C4'] = 'Net Income reportado (lucro liquido) - base do check IS'
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
json.dump(dict(pos=pos, R0=R0, LASTROW=LASTROW, END=END), open(OUT + '.layout.json', 'w'))
print('Inputs', R0, LASTROW, 'END', END, '| codes present', sorted(present))
