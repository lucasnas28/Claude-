import sys, shutil, copy, re, datetime as dt, json
import openpyxl
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI
from openpyxl.formula.translate import Translator
from openpyxl.styles import Font

MASK, SRC, OUT = sys.argv[1:4]
src = openpyxl.load_workbook(SRC, read_only=True, data_only=True)

# ---------------------------------------------------------------- periods
P = [f'{q}Q{y:02d}' for y in range(20, 27) for q in range(1, 5)][:26]       # 1Q20..2Q26
def src_label(p): return f'{p[0]}T{p[2:]}'
FLEURY_P = P[:P.index('1Q23') + 1]           # contabil Fleury, pre-combination
COMB_P = P[P.index('2Q23'):]                 # contabil combinada (Pardini from May/23)

def sheet_table(name, hdr_row):
    rows = list(src[name].iter_rows(values_only=True))
    hdr = [str(x).strip() if x is not None else None for x in rows[hdr_row - 1]]
    return rows, hdr
TAB = {k: sheet_table(*v) for k, v in {
    'F_DRE': ('Fleury DRE', 8), 'F_BP': ('Fleury BP', 8), 'F_DFC': ('Fleury DFC', 8),
    'C_DRE': ('Combinada DRE', 5), 'C_BP': ('Combinada BP', 5), 'C_DFC': ('Combinada DFC', 5)}.items()}

def col_of(tab, p, kind):
    rows, hdr = TAB[tab]
    lab = src_label(p)
    cands = [lab]
    if tab in ('C_DRE', 'C_DFC') and lab == '2T23': cands = ['2T23 Contábil']
    if tab == 'F_BP' and p[0] == '4': cands = [lab, '20' + p[2:]]          # 4T empty -> 31/12 in annual column
    for c in cands:
        if c in hdr: return hdr.index(c), c
    raise KeyError((tab, p))

def value(tab, row, p):
    rows, hdr = TAB[tab]
    r = rows[row - 1]
    i, used = col_of(tab, p, None)
    v = r[i]
    if v is None and tab == 'F_BP' and p[0] == '4':
        i2 = hdr.index('20' + p[2:]); v = r[i2]; used = '20' + p[2:]
    return v, used

def label_of(tab, row): return str(TAB[tab][0][row - 1][0]).strip()

def series(frow, crow, kind):
    """kind: DRE/BP/DFC. Fleury sheet rows for 1Q20..1Q23, Combinada rows from 2Q23."""
    out = {}
    for p in P:
        if p in FLEURY_P:
            if frow is None: continue
            v, _ = value('F_' + kind, frow, p)
        else:
            if crow is None: continue
            v, _ = value('C_' + kind, crow, p)
        if isinstance(v, (int, float)): out[p] = float(v)
    return out

# (fleury_row, combinada_row, code, sign)
DRE = [
 (15, 13, 1, 1), (18, 16, 2, 1), (19, 17, 2, 1), (20, 18, 2, 1), (22, 20, 2, 1),
 (21, 19, 4, 1),
 (35, 33, 3, 1), (None, 36, 3, 1), (38, 37, 3, 1), (39, 38, 3, 1),
 (36, 34, 4, 1), (37, 35, 4, 1),
 (40, 39, 8, 1), (45, 44, 6, 1), (None, 46, 6, 1), (46, 45, 5, 1),
 (50, 50, 7, 1), (55, 55, 9, 1),
]
DRE_MEMO = [(10, 7, 'raw_rb'), (17, 15, 'raw_cogs'), (30, 28, 'raw_gp'), (33, 31, 'raw_opex'), (42, 41, 'raw_ebit'),
            (48, 48, 'raw_pt'), (56, 56, 'raw_ni'), (65, 68, 'raw_ebitda')]
DFC = [(63, 59, 10, 1)]
DFC_MEMO = [(13, 10, 'raw_da')]
BP = [
 (12, 9, 11), (13, 10, 11), (14, 11, 12), (15, 12, 13), (18, 16, 15), (19, 17, 14),
 (16, 13, 16), (None, 14, 16), (17, 15, 16), (20, 18, 16), (21, 19, 16), (22, 20, 16), (23, 21, 16),
 (28, 26, 22), (29, 27, 22), (30, 28, 22), (31, 29, 22), (None, 30, 22), (32, 31, 22), (33, 32, 14),
 (34, 33, 11), (35, 34, 22), (36, 35, 22), (37, 36, 22),
 (41, 40, 17), (42, 41, 18), (43, 42, 21), (44, 43, 19),
 (51, 50, 27), (None, 51, 27), (52, 52, 29), (53, 53, 26), (54, 54, 23), (55, 55, 24), (None, 56, 25),
 (56, 57, 25), (57, 58, 33), (58, 59, 25), (59, 60, 26), (60, 61, 26), (None, 62, 26), (61, 63, 26),
 (65, 67, 30), (None, 68, 30), (66, 69, 31), (67, 70, 34), (68, 71, 33), (None, 72, 34), (None, 73, 34),
 (69, 74, 34), (70, 75, 34), (71, 76, 34), (72, 77, 34), (73, 78, 34), (74, 79, 34), (75, 80, 34),
 (80, 85, 37), (81, 86, 37), (82, 87, 37), (83, 88, 40), (84, 89, 40), (85, 90, 40), (86, 91, 41),
 (87, 92, 40), (None, 93, 41), (None, 94, 39), (None, 95, 40), (88, 96, 38), (90, 98, 35),
]
BP_MEMO = [(24, 22, 'tca'), (47, 46, 'ta'), (62, 64, 'tcl'), (77, 82, 'tl'), (91, 99, 'te'), (93, 101, 'tle')]

def lab2(tab, fr, cr):
    a = f'Fleury {tab}: {label_of("F_" + tab, fr)} (l.{fr})' if fr else f'Fleury {tab}: [nao existe]'
    b = f'Combinada {tab}: {label_of("C_" + tab, cr)} (l.{cr})' if cr else ''
    return a + ' | ' + b

# ---------------------------------------------------------------- workbook
shutil.copyfile(MASK, OUT)
wb = openpyxl.load_workbook(OUT)
inp, w2, m, val = wb['Inputs'], wb['Inputs (2)'], wb['Model'], wb['Valuation']
PC = {p: 11 + i for i, p in enumerate(P)}
col = lambda p: CL(PC[p])
FIRSTC, LASTC = col(P[0]), col(P[-1])
HDR, R0 = 4, 5
bold = Font(bold=True)

inp['A1'] = 'Fator p/ milhões:'; inp['B1'] = 0.001
inp['D1'] = ('Fleury S.A. (FLRY3) - fonte: "Dados Historicos - 2T26.xlsm" (RI Fleury), consolidado contabil. '
             'Abas "Fleury" ate 1T23; abas "Combinada" (colunas contabeis) a partir do 2T23. Moeda: BRL. Unidade dos brutos: R$ mil. '
             'Fator converte p/ R$ milhoes.')
for i, h in enumerate(['COD1', 'SINAL1', 'COD2', 'SINAL2', 'COD3', 'SINAL3', 'UNID', 'Conta (PT) / origem', 'Account (EN)', 'ORIGEM']):
    inp.cell(HDR, 1 + i, h).font = bold
inp.cell(2, 10, 'Aba de origem'); inp.cell(3, 10, 'Coluna lida (BP)')
for p in P:
    inp.cell(HDR, PC[p], p).font = bold
    inp.cell(2, PC[p], 'Fleury' if p in FLEURY_P else 'Combinada')
    _, used = value('F_BP' if p in FLEURY_P else 'C_BP', 12 if p in FLEURY_P else 9, p)
    inp.cell(3, PC[p], used)

r = R0
pos = {}
def section(t):
    global r
    inp.cell(r, 8, t).font = bold; r += 1
def write(code, sign, desc, en, origem, vals, key=None, unid='moeda'):
    global r
    if code is not None:
        miss = [p for p in P if p not in vals]
        if miss:
            vals = dict(vals); vals.update({p: 0.0 for p in miss})
            desc += f' [0 em {miss[0]}..{miss[-1]} ({len(miss)} per.): linha inexistente/vazia na fonte; total divulgado confirma]'
    if code is not None: inp.cell(r, 1, code); inp.cell(r, 2, sign)
    inp.cell(r, 7, unid); inp.cell(r, 8, desc); inp.cell(r, 9, en); inp.cell(r, 10, origem)
    for p, x in vals.items(): inp.cell(r, PC[p], x)
    for p in P: inp.cell(r, PC[p]).number_format = '#,##0;(#,##0)'
    if key: pos[key] = r
    r += 1

section('DRE (R$ mil, trimestre isolado como divulgado; custos e despesas ja negativos na fonte)')
for fr, cr, code, s in DRE:
    write(code, s, lab2('DRE', fr, cr), '', 'fonte', series(fr, cr, 'DRE'), key=f'dre_{fr}_{cr}')
section('DRE - subtotais divulgados (memoria, sem codigo)')
for fr, cr, k in DRE_MEMO:
    write(None, None, lab2('DRE', fr, cr), '', 'memoria', series(fr, cr, 'DRE'), key=k)
section('DFC (R$ mil, trimestre isolado)')
for fr, cr, code, s in DFC:
    write(code, s, lab2('DFC', fr, cr), 'Dividends / interest on equity paid', 'fonte', series(fr, cr, 'DFC'), key='div')
for fr, cr, k in DFC_MEMO:
    write(None, None, lab2('DFC', fr, cr) + ' - memoria p/ conferir o cod 4', '', 'memoria', series(fr, cr, 'DFC'), key=k)
section('BALANCO (R$ mil, saldo de fim de trimestre; 4T21 e 4T22 da aba Fleury lidos da coluna anual)')
for fr, cr, code in BP:
    write(code, 1, lab2('BP', fr, cr), '', 'fonte', series(fr, cr, 'BP'), key=f'bp_{fr}_{cr}')
section('BALANCO - totais divulgados (memoria, sem codigo)')
for fr, cr, k in BP_MEMO:
    write(None, None, lab2('BP', fr, cr), '', 'memoria', series(fr, cr, 'BP'), key=k)
section('CHECKS (devem ser 0)')
def chk(desc, terms, total, key):
    global r
    inp.cell(r, 7, 'moeda'); inp.cell(r, 8, desc); inp.cell(r, 10, 'derivada')
    for p in P:
        inp.cell(r, PC[p], '=' + '+'.join(f'{col(p)}{t}' for t in terms) + f'-{col(p)}{total}')
        inp.cell(r, PC[p]).number_format = '#,##0;(#,##0)'
    pos[key] = r; r += 1
dre_rows = [pos[f'dre_{fr}_{cr}'] for fr, cr, *_ in DRE]
bp_rows = {c: [pos[f'bp_{fr}_{cr}'] for fr, cr, cc in BP if cc == c] for c in range(11, 42)}
chk('CHECK: soma das linhas codificadas 1-9 da DRE - lucro liquido do periodo', dre_rows, pos['raw_ni'], 'chk_ni')
chk('CHECK: D&A da DRE (cod 4) - D&A da DFC', [pos['dre_21_19'], pos['dre_36_34'], pos['dre_37_35']], pos['raw_da'], 'chk_da_tmp')
# D&A in DRE is negative, DFC positive -> sum + DFC must be 0; rewrite as a plus
rr = pos['chk_da_tmp']
inp.cell(rr, 8, 'CHECK: D&A da DRE (cod 4, negativa) + D&A da DFC (positiva)')
for p in P:
    inp.cell(rr, PC[p], f'={col(p)}{pos["dre_21_19"]}+{col(p)}{pos["dre_36_34"]}+{col(p)}{pos["dre_37_35"]}+{col(p)}{pos["raw_da"]}')
chk('CHECK: ativo (cod 11-22) - total do ativo', sum((bp_rows[c] for c in range(11, 23)), []), pos['ta'], 'chk_a')
chk('CHECK: passivo (cod 23-34) - total do passivo', sum((bp_rows[c] for c in range(23, 35)), []), pos['tl'], 'chk_l')
chk('CHECK: PL (cod 35-41) - total do PL', sum((bp_rows[c] for c in range(35, 42)), []), pos['te'], 'chk_e')
for c, t in [(20, 'Goodwill: o agio esta dentro do Intangivel na fonte (cod 21); nao aberto no balanco'),
             (28, 'Short term debt (revolving): sem linha propria; emprestimos e debentures CP estao no cod 27'),
             (32, 'Convertible debt: inexistente'), (36, 'Convertible Preferred Stock: inexistente')]:
    inp.cell(r, 8, f'[vazio] cod {c} {t}'); r += 1
# reported NI for the IS check (code 44, memo)
NI_ROW = pos['raw_ni']
inp.cell(NI_ROW, 1, 44); inp.cell(NI_ROW, 2, 1)
inp.cell(NI_ROW, 8, inp.cell(NI_ROW, 8).value + ' - COD 44 alimenta o check IS do Model (linha 4)')
LASTROW = r - 1
END = ((LASTROW + 30 + 9) // 10) * 10
conv = END + 3
inp.cell(conv, 8, 'CONVENCOES').font = bold
for i, t in enumerate([
    f'Linhas de origem: {R0} a {LASTROW}. As formulas da Inputs (2) varrem as linhas {R0} a {END}; as linhas {LASTROW + 1} a {END} estao livres para novas contas.',
    f'Periodos: colunas {FIRSTC} a {LASTC} (1Q20 a 2Q26 = 1T20 a 2T26). Linha 2 = aba de origem; linha 3 = coluna lida no balanco.',
    'Brutos em R$ mil, como na fonte. A Inputs (2) multiplica pelo fator da celula B1 (0,001) e entrega R$ milhoes.',
    'Fleury DRE/BP/DFC ate 1T23 (contabil pre-combinacao). Combinada a partir do 2T23, coluna "2T23 Contabil" na DRE/DFC. As colunas pro forma da Combinada (2022 a 1T23, "2T23" e "2023") nao foram usadas.',
    'ORIGEM: fonte = copiado da planilha de fundamentos; memoria = subtotal divulgado, sem codigo; derivada = check por formula.',
    'Uma linha da Inputs pode juntar a conta da aba Fleury e a da aba Combinada quando sao a mesma conta; a descricao mostra as duas linhas de origem.']):
    inp.cell(conv + 1 + i, 8, t)
for c, w in [(1, 6), (2, 7), (3, 6), (4, 7), (5, 6), (6, 7), (7, 8), (8, 95), (9, 30), (10, 10)]:
    inp.column_dimensions[CL(c)].width = w
for p in P: inp.column_dimensions[col(p)].width = 13
inp.freeze_panes = inp.cell(R0, 11)

# ---------------------------------------------------------------- Inputs (2)
ext = ['1Q20', '2Q20', '3Q20', '4Q20', 2020]
for i, h in enumerate(ext):
    c = 34 + i
    w2.cell(2, c, h)
    w2.cell(2, c)._style = copy.copy(w2.cell(2, 29 + i if i < 4 else 33)._style)
    w2.column_dimensions[CL(c)].width = w2.column_dimensions['AC'].width or 10
w2['B50'] = 44; w2['C50'] = 'Net income reported (lucro liquido do periodo) - memo p/ check IS [Model linha 4]'
for c in range(1, 39): w2.cell(50, c)._style = copy.copy(w2.cell(13, c)._style)
IN = 'Inputs!'
def rng(cl): return f'{IN}${cl}${R0}:${cl}${END}'
VAL = f'INDEX({IN}${FIRSTC}${R0}:${LASTC}${END},0,MATCH({{h}},{IN}${FIRSTC}${HDR}:${LASTC}${HDR},0))'
ESC = f'((({rng("G")}="moeda")+({rng("G")}="acoes"))*{IN}$B$1+(({rng("G")}<>"moeda")*({rng("G")}<>"acoes")))'
def f2(h, d):
    v_ = VAL.format(h=h)
    cnt = '+'.join(f'SUMPRODUCT(({rng(c)}={d})*ISNUMBER({v_}))' for c in 'ACE')
    parts = '+'.join(f'SUMPRODUCT(({rng(c)}={d})*{rng(s)}*{ESC}*{v_})' for c, s in (('A', 'B'), ('C', 'D'), ('E', 'F')))
    return f'=IF(ISNA(MATCH({h},{IN}${FIRSTC}${HDR}:${LASTC}${HDR},0)),NA(),IF({cnt}=0,NA(),{parts}))'
code_rows = [rr for rr in range(4, 51) if isinstance(w2.cell(rr, 2).value, int)]
annual_cols = {8, 13, 18, 23, 28, 33, 38}
qcols = [c for c in range(4, 38) if c not in annual_cols]
n2 = 0
for rr in code_rows:
    for c in qcols:
        assert w2.cell(rr, c).value is None, (rr, c)
        w2.cell(rr, c).value = f2(f'{CL(c)}$2', f'$B{rr}'); n2 += 1
    if rr <= 13 or rr == 50: w2.cell(rr, 38, f'=SUM(AH{rr}:AK{rr})')
    else: w2.cell(rr, 38, f'=AK{rr}')
    if rr == 50:
        for c, (a, b) in {8: ('D', 'G'), 13: ('I', 'L'), 18: ('N', 'Q'), 23: ('S', 'V'), 28: ('X', 'AA'), 33: ('AC', 'AF')}.items():
            w2.cell(rr, c, f'=SUM({a}{rr}:{b}{rr})')
    else:
        w2.cell(rr, 38)._style = copy.copy(w2.cell(rr, 33)._style)
        for i in range(4): w2.cell(rr, 34 + i)._style = copy.copy(w2.cell(rr, 29 + i)._style)
assert w2['AG2'].value == 2025
w2['AG2'] = 2026                                   # template typo
grey = Font(color='FF808080', bold=w2['AE2'].font.bold, name=w2['AE2'].font.name, size=w2['AE2'].font.sz)
for a in ('AE2', 'AF2', 'AG2'): w2[a].font = grey
w2['AN2'] = 'Sem dado na fonte: 3Q26, 4Q26 e 2026 (rotulo em cinza). 2020 fica nas colunas AH-AL.'

# ---------------------------------------------------------------- Model links
code2model = {1: 15, 2: 20, 3: 24, 4: 36, 5: 40, 6: 89, 7: 97, 8: 101, 9: 102, 10: 104,
              11: 110, 12: 111, 13: 115, 14: 119, 15: 120, 16: 121, 17: 127, 18: 132, 19: 142, 21: 146, 22: 155,
              23: 164, 24: 168, 25: 170, 26: 171, 27: 178, 29: 180, 30: 183, 31: 201, 33: 205, 34: 206,
              35: 214, 37: 220, 38: 233, 39: 234, 40: 235, 41: 236}
code2row2 = {w2.cell(rr, 2).value: rr for rr in code_rows}
mann = {8: ('D', 'G'), 13: ('I', 'L'), 18: ('N', 'Q'), 23: ('S', 'V'), 28: ('X', 'AA'), 33: ('AC', 'AF')}
nm = 0
pre_zero = []
for code, mr in code2model.items():
    r2 = code2row2[code]
    for c in range(4, 33):
        if c in mann: continue
        cell = m.cell(mr, c)
        if cell.value is not None:
            assert cell.value == 0, (mr, c, cell.value); pre_zero.append(cell.coordinate)
        cell.value = f"=INDEX('Inputs (2)'!$D${r2}:$AL${r2},MATCH({CL(c)}$7,'Inputs (2)'!$D$2:$AL$2,0))"; nm += 1

# ---------------------------------------------------------------- template fixes carried over from the NVIDIA work
def tr(f, s, d): return Translator(f, origin=s).translate_formula(d)
for colx, rws in {'I': [134], 'AF': [16, 21, 25, 37, 41, 98, 105, 112, 116, 122, 134, 143, 156, 165, 169, 172, 181, 202, 207]}.items():
    ref = 'N' if colx == 'I' else 'AE'
    for rr in rws:
        assert m[f'{colx}{rr}'].value is None
        m[f'{colx}{rr}'] = tr(m[f'{ref}{rr}'].value, f'{ref}{rr}', f'{colx}{rr}')
pat = re.compile(r"Model!(\$?)([A-Z]{1,3})(\$?)(\d+)")
for row in val.iter_rows(min_row=16, max_row=56, min_col=4, max_col=18):
    for c in row:
        if isinstance(c.value, str) and 'Model!' in c.value:
            c.value = pat.sub(lambda g: f"Model!{g.group(1)}{CL(CI(g.group(2)) + 10)}{g.group(3)}{g.group(4)}", c.value)
val['G24'] = val['G24'].value.replace('Model!AL99', 'Model!AL98')
for row in val.iter_rows(min_row=18, max_row=56, min_col=19, max_col=22):
    for c in row: c.value = None
for a in ('S16', 'T16', 'U16', 'V16'): val[a].font = Font(color='FF808080')
val['S15'] = 'Fora do horizonte do Model (projecao vai ate 2037)'; val['S15'].font = Font(color='FF808080')
val['D64'] = '=R50'; val['D67'] = '=D66*R55'; val['D68'] = '=D67*R54'; val['G57'] = '=SUM(G56:R56)'
val['D59'] = dt.datetime(2037, 12, 31)
val['Y18'] = '=((R18/L18)^(1/6))-1'; val['Y22'] = '=((R22/L22)^(1/6))-1'
val['X16'] = '=F16&"-"&L16'; val['Y16'] = '=L16&"-"&R16'
val['X49'] = '=F16'; val['Y49'] = '=L16'; val['Z49'] = '=R16'; val['Z50'] = '=R50'
# IS check against reported net income
m['C4'] = 'Net Income reportado (lucro liquido do periodo) - base do check IS'
m['C4']._style = copy.copy(m['C103']._style)
for c in range(4, 34):
    cell = m.cell(4, c)
    cell.value = (f'=SUM({mann[c][0]}4:{mann[c][1]}4)' if c in mann else
                  f"=INDEX('Inputs (2)'!$D$50:$AL$50,MATCH({CL(c)}$7,'Inputs (2)'!$D$2:$AL$2,0))")
    cell._style = copy.copy(m.cell(103, c)._style)
for c in range(4, m.max_column + 1):
    L = CL(c)
    if c >= 34 and m.cell(3, c).value is None: continue
    m.cell(3, c).value = f'=IF(COUNT({L}4)=0,"",ABS({L}103-{L}4))'
    if c < 34 and c not in mann: m.cell(3, c)._style = copy.copy(m['H3']._style)

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
json.dump(dict(pos=pos, R0=R0, LASTROW=LASTROW, END=END), open(OUT + '.layout.json', 'w'))
print('Inputs', R0, LASTROW, 'END', END, '| Inputs(2) formulas', n2, '| Model links', nm, '| pre-existing zeros overwritten', pre_zero)
