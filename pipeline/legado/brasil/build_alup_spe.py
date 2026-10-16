import sys, re, copy, json
sys.path.insert(0, __import__('os').environ.get('PIPE_DATA', 'dados')+'/work')
from alup_spe import wb as SRCWB, TRANS, GEN, net_rev, load
import openpyxl
from openpyxl.utils import get_column_letter as CL
from openpyxl.formula.translate import Translator
from openpyxl.styles import Font

IN, OUT, LAYF = sys.argv[1:4]
B0_ARG = int(sys.argv[4]) if len(sys.argv) > 4 else 361
lay = json.load(open(LAYF))
pos, R0, OLD_END, LAST0 = lay['pos'], lay['R0'], lay['END'], lay['LASTROW']
wb = openpyxl.load_workbook(IN)
inp, w2, m = wb['Inputs'], wb['Inputs (2)'], wb['Model']
P = [f'{q}Q{y:02d}' for y in range(20, 27) for q in range(1, 5)][:26]
PC = {p: 11 + i for i, p in enumerate(P)}
col = lambda p: CL(PC[p])
lab = lambda p: f'{p[0]}T{p[2:]}'
ANN = set(range(8, 108, 5))
bold = Font(bold=True)
def tr(f, s, d): return Translator(f, origin=s).translate_formula(d)

def series(sheet):
    out = {}
    for p in P:
        v = net_rev(sheet, lab(p))
        if v is not None: out[p] = v
    return out
NICE = {'Foz': 'Foz do Rio Claro', 'IJUÍ': 'Ijuí', 'Agua Limpa': 'Água Limpa', 'Energia dos Ventos até 4T14': 'Energia dos Ventos (até 4T14)'}
def title(sheet):
    return NICE.get(sheet, sheet)
def _title_unused(sheet):
    t = load(sheet)[3]
    mm = re.search(r'\(([^)]+)\)\s*$', t or '')
    return mm.group(1) if mm and mm.group(1) != sheet else sheet
def active(s): return any(abs(x) > 0 for x in s.values())

segT, segG = series('Consolidado - Transmissão'), series('Consolidado - Geração')
spe_T = [(s, series(s)) for s in TRANS]; spe_T = [(s, v) for s, v in spe_T if active(v)]
spe_G = [(s, series(s)) for s in GEN]; spe_G = [(s, v) for s, v in spe_G if active(v)]
assert 'TNE' in [s for s, _ in spe_T]

# ---------------------------------------------------------------- capture and clear the old conventions block first
oldconv = OLD_END + 3
conv_texts = [inp.cell(rr, 8).value for rr in range(oldconv + 1, oldconv + 12) if inp.cell(rr, 8).value]
for rr in range(oldconv, oldconv + 12): inp.cell(rr, 8).value = None
assert conv_texts and conv_texts[0].startswith('Linhas de origem'), conv_texts[:1]

# ---------------------------------------------------------------- Inputs rows (codes 45+ = memo)
r = LAST0 + 2
code = 45
items = []          # (key, code, label, inputs row)
def put(key, label_, vals, note=''):
    global r, code
    inp.cell(r, 1, code); inp.cell(r, 2, 1); inp.cell(r, 7, 'moeda'); inp.cell(r, 10, 'fonte')
    inp.cell(r, 8, label_ + note)
    for p, x in vals.items(): inp.cell(r, PC[p], x)
    for p in P: inp.cell(r, PC[p]).number_format = '#,##0;(#,##0)'
    items.append((key, code, label_, r)); code += 1; r += 1
inp.cell(r, 8, 'RECEITA LIQUIDA POR SEGMENTO E POR SPE (R$ mil; receita bruta + deducoes de cada aba, a 100%). Codigos 45+ = memoria, nao entram nos codigos 1-41').font = bold; r += 1
put('segT', 'Transmissao - consolidado do segmento | aba "Consolidado - Transmissao" (so a partir do 1T23)', segT)
put('segG', 'Geracao - consolidado do segmento | aba "Consolidado - Geracao" (so a partir do 1T23)', segG)
inp.cell(r, 8, 'Transmissao por SPE (aba de cada SPE)').font = bold; r += 1
for s, v in spe_T:
    put('T:' + s, f'Transmissao | {title(s)} | aba "{s}"', v, ' [controle compartilhado: entra no consolidado por equivalencia]' if s == 'TNE' else '')
inp.cell(r, 8, 'Geracao por usina (aba de cada empresa)').font = bold; r += 1
for s, v in spe_G:
    put('G:' + s, f'Geracao | {title(s)} | aba "{s}"', v)
LAST = r - 1
END = ((LAST + 30 + 9) // 10) * 10

# move conventions block (captured and cleared before writing new rows), widen Inputs (2) ranges
conv = END + 3
inp.cell(conv, 8, 'CONVENCOES').font = bold
conv_texts[0] = f'Linhas de origem: {R0} a {LAST}. As formulas da Inputs (2) varrem as linhas {R0} a {END}; as linhas {LAST + 1} a {END} estao livres para novas contas.'
conv_texts.append(f'Codigos 45 a {code - 1}: receita liquida por segmento e por SPE (memoria); alimentam o bloco no fim do Model.')
for i, t in enumerate(conv_texts): inp.cell(conv + 1 + i, 8, t)
rx = re.compile(r'(Inputs!\$[A-Z]{1,2}\$\d+:\$[A-Z]{1,2}\$)' + str(OLD_END) + r'\b')
for row in w2.iter_rows():
    for c in row:
        if isinstance(c.value, str) and 'Inputs!' in c.value:
            c.value = rx.sub(lambda g: g.group(1) + str(END), c.value)

# RAP table (cycle 2026/2027) stored below the conventions block, outside the formula range
rap_rows = [rw for rw in SRCWB['Anexo RAP '].iter_rows(values_only=True)]
RAP0 = conv + len(conv_texts) + 3
inp.cell(RAP0, 8, 'RAP CICLO 2026/2027 POR SPE (R$ milhoes; aba "Anexo RAP" da planilha de fundamentos) - fora do intervalo das formulas').font = bold
hdr = ['SPE', 'Indexador', 'RAP operacional', 'RAP implantacao', 'RAP total']
for i, h in enumerate(hdr): inp.cell(RAP0 + 1, 8 + i, h).font = bold
rap = []
for rw in rap_rows:
    v = [x for x in rw if x is not None]
    if len(v) >= 5 and isinstance(v[-1], (int, float)) and str(v[0]).strip() not in ('SPE',):
        name = str(v[0]).strip(); idx = str(v[1]).strip() if not isinstance(v[1], (int, float)) else ''
        nums = [float(x) if isinstance(x, (int, float)) else float(str(x).replace(',', '.')) if str(x).replace('.', '', 1).isdigit() else 0.0 for x in v[-3:]]
        rap.append((name, idx, *nums))
for i, (name, idx, a, b, c) in enumerate(rap):
    rr = RAP0 + 2 + i
    inp.cell(rr, 8, name); inp.cell(rr, 9, idx)
    for j, x in enumerate((a, b, c)):
        inp.cell(rr, 10 + j, x).number_format = '#,##0.0'
RAP_ROWS = (RAP0 + 2, RAP0 + 1 + len(rap))

# ---------------------------------------------------------------- Inputs (2) rows
row2 = {}
w2['C52'] = 'Receita liquida por segmento e por SPE (memo) - codigos 45+ [Model linhas 361+]'; w2['C52'].font = bold
annual2 = {8: ('D', 'G'), 13: ('I', 'L'), 18: ('N', 'Q'), 23: ('S', 'V'), 28: ('X', 'AA'), 33: ('AC', 'AF'), 38: ('AH', 'AK')}
for i, (key, cd, label_, _) in enumerate(items):
    rr = 53 + i; row2[key] = rr
    w2.cell(rr, 2, cd); w2.cell(rr, 3, label_.split(' | ')[0] + ' | ' + label_.split(' | ')[1] if ' | ' in label_ else label_)
    for c in range(1, 39): w2.cell(rr, c)._style = copy.copy(w2.cell(4, c)._style)
    for c in range(4, 38):
        if c in annual2: continue
        w2.cell(rr, c, tr(w2.cell(4, c).value, f'{CL(c)}4', f'{CL(c)}{rr}'))
    for c, (a, b) in annual2.items():
        w2.cell(rr, c, f'=SUM({a}{rr}:{b}{rr})')

# ---------------------------------------------------------------- Model block
B0 = B0_ARG
NROWS = 12 + len(spe_T) + 6 + len(spe_G) + 6 + len(rap) + 4
for rr in range(B0, B0 + NROWS):
    for c in range(1, 36): assert m.cell(rr, c).value is None, (rr, c)
COLS = list(range(4, 36))      # D..AI (through 2Q26)
def hdr_row(rr, text):
    m.cell(rr, 3, text)._style = copy.copy(m['C13']._style)
    for c in COLS: m.cell(rr, c, f'={CL(c)}7')._style = copy.copy(m.cell(13, c)._style)
def label_cell(rr, text, b=False):
    c = m.cell(rr, 3, text); c._style = copy.copy(m['C15']._style)
    if b: c.font = Font(bold=True, name=m['C15'].font.name, size=m['C15'].font.sz)
def link(key, c):
    return f"=IFERROR(INDEX('Inputs (2)'!$D${row2[key]}:$AL${row2[key]},MATCH({CL(c)}$7,'Inputs (2)'!$D$2:$AL$2,0)),\"\")"
def fill(rr, fn, fmt=None):
    for c in COLS:
        cell = m.cell(rr, c); cell._style = copy.copy(m.cell(15, c)._style)
        if c in ANN and fn is not None and fn.__name__ != 'mix':
            a, b = CL(c - 4), CL(c - 1); cell.value = f'=IF(COUNT({a}{rr}:{b}{rr})=4,SUM({a}{rr}:{b}{rr}),"")'
        else:
            cell.value = fn(c)
        if fmt: cell.number_format = fmt
rr = B0
hdr_row(rr, 'Receita liquida por segmento (R$ milhoes; segmentos so a partir do 1T23)'); rr += 2
SEG_T, SEG_G, SEG_H, SEG_TOT, SEG_CHK = rr, rr + 1, rr + 2, rr + 3, rr + 4
def lk_factory(key):
    def lk(c): return link(key, c)
    return lk
fill(SEG_T, lk_factory('segT')); label_cell(SEG_T, 'Transmissao')
fill(SEG_G, lk_factory('segG')); label_cell(SEG_G, 'Geracao')
def h(c): L = CL(c); return f'=IF(COUNT({L}{SEG_T},{L}{SEG_G})<2,"",{L}15-{L}{SEG_T}-{L}{SEG_G})'
fill(SEG_H, h); label_cell(SEG_H, 'Holding e eliminacoes entre segmentos (diferenca)')
def tot(c): L = CL(c); return f'=IF(COUNT({L}{SEG_T},{L}{SEG_G})<2,"",SUM({L}{SEG_T},{L}{SEG_G},{L}{SEG_H}))'
fill(SEG_TOT, tot); label_cell(SEG_TOT, 'Receita liquida consolidada (soma)', True)
def mix(c): L = CL(c); return f'=IFERROR({L}{SEG_T}/{L}15,"")'
fill(SEG_CHK, mix, '0.0%'); label_cell(SEG_CHK, 'Transmissao % da receita')
rr = SEG_CHK + 2

hdr_row(rr, 'Transmissao - receita liquida por SPE, a 100% (R$ milhoes)'); rr += 1
T0 = rr
for s, _ in spe_T:
    fill(rr, lk_factory('T:' + s)); label_cell(rr, '   ' + title(s) + (' (controle compartilhado)' if s == 'TNE' else '')); rr += 1
T1 = rr - 1
TSUM, TNE_R, TADJ, TCONS = rr, rr + 1, rr + 2, rr + 3
tne_row = T0 + [s for s, _ in spe_T].index('TNE')
def tsum(c): L = CL(c); return f'=IF(COUNT({L}{T0}:{L}{T1})=0,"",SUM({L}{T0}:{L}{T1}))'
fill(TSUM, tsum); label_cell(TSUM, 'Soma das SPEs (100%)', True)
def tne(c): L = CL(c); return f'=IF(COUNT({L}{tne_row})=0,"",-{L}{tne_row})'
fill(TNE_R, tne); label_cell(TNE_R, '(-) TNE (controle compartilhado, entra por equivalencia)')
def tadj(c): L = CL(c); return f'=IF(COUNT({L}{SEG_T})=0,"",{L}{SEG_T}-SUM({L}{TSUM},{L}{TNE_R}))'
fill(TADJ, tadj); label_cell(TADJ, 'Ajuste: eliminacoes e SPE sem dado na planilha (diferenca)')
def tcons(c): L = CL(c); return f'=IF(COUNT({L}{SEG_T})=0,"",SUM({L}{TSUM},{L}{TNE_R},{L}{TADJ}))'
fill(TCONS, tcons); label_cell(TCONS, 'Transmissao consolidado (fecha com o segmento)', True)
rr = TCONS + 2

hdr_row(rr, 'Geracao - receita liquida por usina, a 100% (R$ milhoes)'); rr += 1
G0 = rr
for s, _ in spe_G:
    fill(rr, lk_factory('G:' + s)); label_cell(rr, '   ' + title(s)); rr += 1
G1 = rr - 1
GSUM, GADJ, GCONS = rr, rr + 1, rr + 2
def gsum(c): L = CL(c); return f'=IF(COUNT({L}{G0}:{L}{G1})=0,"",SUM({L}{G0}:{L}{G1}))'
fill(GSUM, gsum); label_cell(GSUM, 'Soma das usinas (100%)', True)
def gadj(c): L = CL(c); return f'=IF(COUNT({L}{SEG_G})=0,"",{L}{SEG_G}-{L}{GSUM})'
fill(GADJ, gadj); label_cell(GADJ, 'Ajuste: comercializadora e eliminacoes de venda entre empresas (diferenca)')
def gcons(c): L = CL(c); return f'=IF(COUNT({L}{SEG_G})=0,"",{L}{GSUM}+{L}{GADJ})'
fill(GCONS, gcons); label_cell(GCONS, 'Geracao consolidado (fecha com o segmento)', True)
rr = GCONS + 2

m.cell(rr, 3, 'RAP ciclo 2026/2027 por SPE (R$ milhoes) - base para projetar a receita de transmissao')._style = copy.copy(m['C13']._style)
for i, h_ in enumerate(['Indexador', 'RAP operacional', 'RAP implantacao', 'RAP total']):
    m.cell(rr, 4 + i, h_)._style = copy.copy(m.cell(13, 4)._style)
rr += 1
for i in range(len(rap)):
    src_r = RAP_ROWS[0] + i
    label_cell(rr, f'=Inputs!H{src_r}')
    m.cell(rr, 4, f'=Inputs!I{src_r}')
    for j in range(3):
        cc = m.cell(rr, 5 + j, f'=Inputs!{CL(10 + j)}{src_r}'); cc._style = copy.copy(m.cell(15, 5)._style); cc.number_format = '#,##0.0'
    rr += 1

# row 17 check next to Net Revenues
for c in COLS:
    L = CL(c)
    m.cell(17, c, f'=IF(COUNT({L}{SEG_TOT})=0,"",{L}15-{L}{SEG_TOT})')._style = copy.copy(m.cell(15, c)._style)
m['C17'] = f'Check Receita (linha 15 - soma dos segmentos, linha {SEG_TOT})'
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('SPEs transmissao', len(spe_T), 'usinas', len(spe_G), 'RAP linhas', len(rap), '| Inputs ate', LAST, 'END', OLD_END, '->', END, '| Model', B0, '->', rr - 1)
print('rows', dict(SEG_T=SEG_T, TSUM=TSUM, TADJ=TADJ, GSUM=GSUM, GADJ=GADJ))
json.dump(dict(SEG_T=SEG_T, SEG_G=SEG_G, SEG_H=SEG_H, SEG_TOT=SEG_TOT, TSUM=TSUM, TNE_R=TNE_R, TADJ=TADJ, TCONS=TCONS, GSUM=GSUM, GADJ=GADJ, GCONS=GCONS, B0=B0, END=rr - 1), open(OUT + '.rows.json', 'w'))
