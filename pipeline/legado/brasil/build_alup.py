import sys, shutil, copy, json
import openpyxl
from openpyxl.utils import get_column_letter as CL
from openpyxl.styles import Font

MASK, SRC, OUT = sys.argv[1:4]
src = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
ROWS = list(src['Consolidado'].iter_rows(min_row=1, max_row=167, values_only=True))
HDR_SRC = [str(x).strip() if x is not None else None for x in ROWS[4]]

P = [f'{q}Q{y:02d}' for y in range(20, 27) for q in range(1, 5)][:26]       # 1Q20..2Q26
def lab(p): return f'{p[0]}T{p[2:]}'
def raw(n, p):
    v = ROWS[n - 1][HDR_SRC.index(lab(p))]
    return float(v) if isinstance(v, (int, float)) else None
# 3T23: the equity block of the source is shifted one row up (capital missing, each account one row above).
# Realigned below; capital = total equity - minority - other accounts (= 3.310.783, same as 2T23 and 4T23).
SHIFT_FIX = {('3Q23', 87): 86, ('3Q23', 88): 87, ('3Q23', 89): 88, ('3Q23', 90): 89, ('3Q23', 91): 90, ('3Q23', 92): 91}
def val(n, p):
    if p == '3Q23' and 66 <= n <= 83:          # non-current liabilities of 3T23 sit one row above in the source
        return raw(n - 1, p)
    if p == '3Q23' and n in range(86, 93):
        if n == 86:
            others = sum(raw(k, p) or 0 for k in range(86, 92))     # shifted values = accounts 87..92
            return raw(96, p) - raw(94, p) - others
        return raw(SHIFT_FIX[(p, n)], p) if (p, n) in SHIFT_FIX else None
    return raw(n, p)
def series(n):
    out = {}
    for p in P:
        v = val(n, p)
        if v is not None: out[p] = v
    return out
def label(n): return f'{str(ROWS[n - 1][1]).strip()} / {str(ROWS[n - 1][2]).strip()} (l.{n})'

# (source row, code)
DRE = [(r, 1) for r in range(104, 112)] + [(113, 1)] + \
      [(130, 2), (131, 2), (132, 2), (133, 4), (136, 3), (137, 4), (138, 3), (139, 8), (140, 3), (141, 3),
       (146, 5), (147, 6), (148, 6), (150, 6), (151, 6), (155, 7), (163, 9)]
DRE_MEMO = [(103, 'rb'), (129, 'cost'), (143, 'ebit_src'), (145, 'fe'), (149, 'fr'), (153, 'pt'), (161, 'ni_tot'), (165, 'ni')]
BP = [(10, 11), (11, 11), (12, 11), (13, 12), (14, 16), (15, 16), (16, 16), (17, 13), (18, 16), (19, 15), (20, 16), (21, 16),
      (22, 16), (23, 16), (24, 16),
      (27, 22), (28, 22), (29, 22), (30, 22), (31, 22), (32, 22), (33, 14), (34, 22), (35, 22), (36, 22), (37, 22), (38, 22),
      (39, 22), (40, 17), (41, 18), (42, 21),
      (48, 27), (49, 27), (50, 29), (51, 23), (52, 24), (53, 25), (54, 26), (55, 26), (56, 26), (57, 26), (58, 26), (59, 26),
      (60, 26), (61, 26), (62, 26), (63, 26),
      (66, 30), (67, 30), (68, 31), (69, 34), (70, 34), (71, 34), (72, 33), (73, 34), (74, 34), (75, 34), (76, 34), (77, 34),
      (78, 34), (79, 34), (80, 34), (81, 34), (82, 34), (83, 34),
      (86, 37), (87, 37), (88, 37), (89, 37), (90, 40), (91, 41), (92, 39), (94, 35)]
BP_MEMO = [(9, 'tca'), (26, 'tnca'), (44, 'ta'), (47, 'tcl'), (65, 'tncl'), (96, 'te'), (98, 'tle')]

shutil.copyfile(MASK, OUT)
wb = openpyxl.load_workbook(OUT)
inp, w2, m = wb['Inputs'], wb['Inputs (2)'], wb['Model']
PC = {p: 11 + i for i, p in enumerate(P)}
col = lambda p: CL(PC[p])
FIRSTC, LASTC = col(P[0]), col(P[-1])
HDR, R0 = 4, 5
bold = Font(bold=True)
inp['A1'] = 'Fator p/ milhões:'; inp['B1'] = 0.001
inp['D1'] = ('Alupar Investimento S.A. (ALUP11) - fonte: planilha de fundamentos do RI, aba "Consolidado" (IFRS, consolidado). '
             'Moeda: BRL. Unidade dos brutos: R$ mil. Fator converte p/ R$ milhoes.')
for i, h in enumerate(['COD1', 'SINAL1', 'COD2', 'SINAL2', 'COD3', 'SINAL3', 'UNID', 'Conta (PT) / origem', 'Account (EN)', 'ORIGEM']):
    inp.cell(HDR, 1 + i, h).font = bold
for p in P: inp.cell(HDR, PC[p], p).font = bold; inp.cell(2, PC[p], lab(p))
inp.cell(2, 10, 'Rotulo na fonte')

r = R0; pos = {}
def section(t):
    global r
    inp.cell(r, 8, t).font = bold; r += 1
def write(code, desc, origem, vals, key=None):
    global r
    if code is not None:
        inp.cell(r, 1, code); inp.cell(r, 2, 1)
        miss = [p for p in P if p not in vals]
        if miss:
            vals = dict(vals); vals.update({p: 0.0 for p in miss})
            desc += f' [0 em {len(miss)} per. ({miss[0]}..{miss[-1]}): celula vazia ou "-" na fonte; total divulgado confirma]'
    inp.cell(r, 7, 'moeda'); inp.cell(r, 8, desc); inp.cell(r, 10, origem)
    for p, x in vals.items(): inp.cell(r, PC[p], x)
    for p in P: inp.cell(r, PC[p]).number_format = '#,##0;(#,##0)'
    if key: pos[key] = r
    r += 1
section('DRE (R$ mil, trimestre isolado; aba Consolidado IFRS). Receita = componentes da receita bruta + deducoes (subtotal); custos e despesas negativos na fonte')
for n, code in DRE: write(code, label(n), 'fonte', series(n), key=f'd{n}')
section('DRE - subtotais divulgados (memoria)')
for n, k in DRE_MEMO: write(None, label(n), 'memoria', series(n), key=k)
section('BALANCO (R$ mil, saldo de fim de trimestre). 3T23: bloco do PL realinhado (ver nota na linha do capital social)')
for n, code in BP:
    d = label(n)
    if n == 86: d += ' [3T23 = PL total - minoritarios - demais contas, porque a fonte desloca o PL do 3T23 uma linha para cima]'
    elif n in range(87, 93) or 66 <= n <= 83: d += ' [3T23 lido da linha de cima na fonte (deslocamento de uma linha no 3T23)]'
    write(code, d, 'fonte', series(n), key=f'b{n}')
section('BALANCO - totais divulgados (memoria)')
for n, k in BP_MEMO: write(None, label(n), 'memoria', series(n), key=k)
section('CHECKS (devem ser 0)')
def chk(desc, terms, total, key, sign_total=-1):
    global r
    inp.cell(r, 7, 'moeda'); inp.cell(r, 8, desc); inp.cell(r, 10, 'derivada')
    for p in P:
        inp.cell(r, PC[p], '=' + '+'.join(f'{col(p)}{t}' for t in terms) + f'-{col(p)}{total}')
        inp.cell(r, PC[p]).number_format = '#,##0;(#,##0)'
    pos[key] = r; r += 1
code_rows = lambda lo, hi: [pos[f'b{n}'] for n, c in BP if lo <= c <= hi]
chk('CHECK: linhas codificadas 1-9 da DRE - lucro liquido Alupar (atribuivel)', [pos[f'd{n}'] for n, _ in DRE], pos['ni'], 'chk_ni')
chk('CHECK: ativo (cod 11-22) - ativo total', code_rows(11, 22), pos['ta'], 'chk_a')
chk('CHECK: passivo (cod 23-34) + PL (cod 35-41) - passivo total', code_rows(23, 41), pos['tle'], 'chk_le')
chk('CHECK: PL (cod 35-41) - PL + minoritarios', code_rows(35, 41), pos['te'], 'chk_e')
for c, t in [(10, 'Common Dividends: a aba Consolidado nao traz DFC; dividendos pagos nao preenchidos'),
             (20, 'Goodwill: nao aberto (dentro do intangivel, cod 21)'), (28, 'Short term debt (revolving): sem linha propria'),
             (32, 'Convertible debt: inexistente'), (36, 'Convertible Preferred Stock: inexistente')]:
    inp.cell(r, 8, f'[vazio] cod {c} {t}'); r += 1
inp.cell(r, 8, 'Nota: ativo contratual da concessao (CP e LP) mapeado em outros ativos (16/22); a variacao entra no fluxo de investimento do Model (linha 265). Para tratar como imobilizado, trocar o codigo para 18.'); r += 1
NI_ROW = pos['ni']; inp.cell(NI_ROW, 1, 44); inp.cell(NI_ROW, 2, 1)
inp.cell(NI_ROW, 8, inp.cell(NI_ROW, 8).value + ' - COD 44 alimenta o check IS do Model (linha 4)')
LASTROW = r - 1
END = ((LASTROW + 30 + 9) // 10) * 10
conv = END + 3
inp.cell(conv, 8, 'CONVENCOES').font = bold
for i, t in enumerate([
    f'Linhas de origem: {R0} a {LASTROW}. As formulas da Inputs (2) varrem as linhas {R0} a {END}; as linhas {LASTROW + 1} a {END} estao livres para novas contas.',
    f'Periodos: colunas {FIRSTC} a {LASTC} (1Q20 a 2Q26 = 1T20 a 2T26). A fonte tem historico desde 1T13; a grade do modelo comeca em 2020.',
    'Brutos em R$ mil, como na fonte. A Inputs (2) multiplica pelo fator da celula B1 (0,001) e entrega R$ milhoes.',
    'Resultado societario (IFRS). O release tambem mostra o regulatorio, que nao foi usado.',
    'ORIGEM: fonte = copiado da planilha de fundamentos; memoria = subtotal divulgado, sem codigo; derivada = check por formula.']):
    inp.cell(conv + 1 + i, 8, t)
for c, w in [(1, 6), (2, 7), (3, 6), (4, 7), (5, 6), (6, 7), (7, 8), (8, 95), (9, 30), (10, 10)]:
    inp.column_dimensions[CL(c)].width = w
for p in P: inp.column_dimensions[col(p)].width = 13
inp.freeze_panes = inp.cell(R0, 11)

# ---------------------------------------------------------------- Inputs (2)
for i, h in enumerate(['1Q20', '2Q20', '3Q20', '4Q20', 2020]):
    c = 34 + i
    w2.cell(2, c, h); w2.cell(2, c)._style = copy.copy(w2.cell(2, 29 + i if i < 4 else 33)._style)
    w2.column_dimensions[CL(c)].width = w2.column_dimensions['AC'].width or 10
w2['B50'] = 44; w2['C50'] = 'Net income reported (lucro liquido Alupar) - memo p/ check IS [Model linha 4]'
for c in range(1, 39): w2.cell(50, c)._style = copy.copy(w2.cell(13, c)._style)
IN = 'Inputs!'
def rng(cl): return f'{IN}${cl}${R0}:${cl}${END}'
VAL = f'INDEX({IN}${FIRSTC}${R0}:${LASTC}${END},0,MATCH({{h}},{IN}${FIRSTC}${HDR}:${LASTC}${HDR},0))'
ESC = f'((({rng("G")}="moeda")+({rng("G")}="acoes"))*{IN}$B$1+(({rng("G")}<>"moeda")*({rng("G")}<>"acoes")))'
def f2(h, d):
    v_ = VAL.format(h=h)
    cnt = '+'.join(f'SUMPRODUCT(({rng(c)}={d})*ISNUMBER({v_}))' for c in 'ACE')
    parts = '+'.join(f'SUMPRODUCT(({rng(c)}={d})*{rng(s)}*{ESC}*{v_})' for c, s in (('A', 'B'), ('C', 'D'), ('E', 'F')))
    return f'=IF(ISNA(MATCH({h},{IN}${FIRSTC}${HDR}:${LASTC}${HDR},0)),NA(),IF({cnt}=0,NA(),{parts}))'
rows2 = [rr for rr in range(4, 51) if isinstance(w2.cell(rr, 2).value, int)]
ANN2 = {8, 13, 18, 23, 28, 33, 38}
for rr in rows2:
    for c in [c for c in range(4, 38) if c not in ANN2]:
        assert w2.cell(rr, c).value is None, (rr, c)
        w2.cell(rr, c).value = f2(f'{CL(c)}$2', f'$B{rr}')
    w2.cell(rr, 38, f'=SUM(AH{rr}:AK{rr})' if rr <= 13 or rr == 50 else f'=AK{rr}')
    if rr == 50:
        for c, (a, b) in {8: ('D', 'G'), 13: ('I', 'L'), 18: ('N', 'Q'), 23: ('S', 'V'), 28: ('X', 'AA'), 33: ('AC', 'AF')}.items():
            w2.cell(rr, c, f'=SUM({a}{rr}:{b}{rr})')
    else:
        w2.cell(rr, 38)._style = copy.copy(w2.cell(rr, 33)._style)
        for i in range(4): w2.cell(rr, 34 + i)._style = copy.copy(w2.cell(rr, 29 + i)._style)
grey = Font(color='FF808080', bold=w2['AE2'].font.bold, name=w2['AE2'].font.name, size=w2['AE2'].font.sz)
for a in ('AE2', 'AF2', 'AG2'): w2[a].font = grey
w2['AN2'] = 'Sem dado na fonte: 3Q26, 4Q26 e 2026 (rotulo em cinza). 2020 fica nas colunas AH-AL.'

# ---------------------------------------------------------------- Model links
code2model = {1: 15, 2: 20, 3: 24, 4: 36, 5: 40, 6: 89, 7: 97, 8: 101, 9: 102,
              11: 110, 12: 111, 13: 115, 14: 119, 15: 120, 16: 121, 17: 127, 18: 132, 19: 142, 21: 146, 22: 155,
              23: 164, 24: 168, 25: 170, 26: 171, 27: 178, 29: 180, 30: 183, 31: 201, 33: 205, 34: 206,
              35: 214, 37: 220, 39: 234, 40: 235, 41: 236}
code2row2 = {w2.cell(rr, 2).value: rr for rr in rows2}
ANN = set(range(8, 108, 5))
present = {c for _, c in DRE + BP}
for code, mr in code2model.items():
    if code not in present: continue
    r2 = code2row2[code]
    for c in range(4, 33):
        if c in ANN: continue
        cell = m.cell(mr, c)
        assert cell.value in (None, 0), (mr, c, cell.value)
        cell.value = f"=INDEX('Inputs (2)'!$D${r2}:$AL${r2},MATCH({CL(c)}$7,'Inputs (2)'!$D$2:$AL$2,0))"
for c in range(4, 33):
    if c in ANN: continue
    m.cell(4, c).value = f"=INDEX('Inputs (2)'!$D$50:$AL$50,MATCH({CL(c)}$7,'Inputs (2)'!$D$2:$AL$2,0))"
m['C4'] = 'Net Income reportado (lucro liquido Alupar) - base do check IS'

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
json.dump(dict(pos=pos, R0=R0, LASTROW=LASTROW, END=END), open(OUT + '.layout.json', 'w'))
print('Inputs', R0, LASTROW, 'END', END)
