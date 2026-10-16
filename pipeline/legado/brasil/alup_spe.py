import openpyxl
F='/root/.claude/uploads/227f02e6-be63-5149-8c2b-85449cb2bc3d/afdd2cf9-5caa4edb-6f38-4961-b1bf-2db6017e61ec.xlsx'
wb=openpyxl.load_workbook(F,read_only=True,data_only=True)
names=wb.sheetnames
HOLD=names[names.index('Alupar Holding'):names.index('Alupar Colombia')+1]+['Reunidas - Holding']
TRANS=names[names.index('Transminas'):names.index('PALCA')+1]
GEN=[s for s in names[names.index('Foz'):names.index('Iracema')+1] if s!='Reunidas - Holding']
def load(s):
    rows=list(wb[s].iter_rows(min_row=1,max_row=200,values_only=True))
    import re
    hr=next((i for i,r in enumerate(rows) if sum(1 for x in r if x is not None and re.match(r'^\d[TQ]\d\d',str(x).strip()))>=2),None)
    if hr is None: return rows,[],{},'(sem cabecalho de periodos)'
    hdr=[re.sub(r'^(\d)Q(\d\d)$',r'\1T\2',str(x).strip()) if x else None for x in rows[hr]]
    labs={}
    for n,r in enumerate(rows):
        lab=r[1] if r[1] is not None else r[0]
        if lab is not None: labs.setdefault(str(lab).strip(),n)
    title=str(rows[0][1]).strip()
    return rows,hdr,labs,title
CACHE={}
def get(s,lab,p):
    if s not in CACHE: CACHE[s]=load(s)
    rows,hdr,labs,_=CACHE[s]
    if lab not in labs or p not in hdr: return None
    v=rows[labs[lab]][hdr.index(p)]
    return float(v) if isinstance(v,(int,float)) else None
def net_rev(s,p):
    a=get(s,'Receita operacional bruta',p); b=get(s,'Deduções da receita operacional',p)
    if a is None and b is None: return None
    return (a or 0)+(b or 0)
