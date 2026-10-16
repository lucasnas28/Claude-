import sys, re, copy, json
import openpyxl
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI
from openpyxl.formula.translate import Translator
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName

IN, SRC, OUT, LAYF = sys.argv[1:5]
lay = json.load(open(LAYF))
pos, R0, OLD_END, LAST0 = lay['pos'], lay['R0'], lay['END'], lay['LASTROW']
src = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
wb = openpyxl.load_workbook(IN)
inp, w2, m, val = wb['Inputs'], wb['Inputs (2)'], wb['Model'], wb['Valuation']
P = [f'{q}Q{y:02d}' for y in range(20, 27) for q in range(1, 5)][:26]
PC = {p: 11 + i for i, p in enumerate(P)}
col = lambda p: CL(PC[p])
FLEURY_P = P[:P.index('1Q23') + 1]
ANNUAL = set(range(8, 108, 5))
HISTQ = [c for c in range(4, 33) if c not in ANNUAL]
bold = Font(bold=True)
log = []
def tr(f, s, d): return Translator(f, origin=s).translate_formula(d)

# ---------------------------------------------------------------- source readers
def table(name):
    rows = list(src[name].iter_rows(values_only=True))
    for i, r in enumerate(rows):
        if any(str(x).strip() in ('1T23', '2T26') for x in r if x is not None):
            return rows, [str(x).strip() if x is not None else None for x in r]
T = {k: table(v) for k, v in {'F_OI': 'Fleury Outras Informações', 'C_OI': 'Combinada Outras Informações',
                               'F_DRE': 'Fleury DRE', 'C_DRE': 'Combinada DRE'}.items()}
def get(tab, row, p, contabil_2t23=True):
    rows, hdr = T[tab]
    lab = f'{p[0]}T{p[2:]}'
    if tab == 'C_DRE' and lab == '2T23': lab = '2T23 Contábil'
    if tab == 'C_OI' and lab == '2T23' and contabil_2t23: return None    # only pro forma available
    if lab not in hdr: return None
    v = rows[row - 1][hdr.index(lab)]
    return float(v) if isinstance(v, (int, float)) else None
def ser(frow, crow, ftab, ctab, contabil_2t23=True):
    out = {}
    for p in P:
        if p in FLEURY_P: v = get(ftab, frow, p) if frow else None
        else: v = get(ctab, crow, p, contabil_2t23) if crow else None
        if v is not None: out[p] = v
    return out

# ---------------------------------------------------------------- 1. acquisition payables become debt (company definition)
for key, code in (('bp_60_61', 27), ('bp_73_78', 30)):
    r = pos[key]
    assert inp.cell(r, 1).value in (26, 34)
    inp.cell(r, 1, code)
    inp.cell(r, 8, inp.cell(r, 8).value + f' -> cod {code} (divida, como no release: divida bruta inclui aquisicoes)')
    log.append(f'Inputs linha {r}: contas a pagar por aquisicoes -> cod {code}')

# ---------------------------------------------------------------- 2. new Inputs rows: revenue by business, deductions, volumes, shares
r = LAST0 + 2
def put(code, unid, desc, origem, vals, fmt='#,##0;(#,##0)', key=None):
    global r
    if code is not None: inp.cell(r, 1, code); inp.cell(r, 2, 1)
    inp.cell(r, 7, unid); inp.cell(r, 8, desc); inp.cell(r, 10, origem)
    for p, x in vals.items(): inp.cell(r, PC[p], x)
    for p in P: inp.cell(r, PC[p]).number_format = fmt
    if key: pos[key] = r
    r += 1
inp.cell(r, 8, 'RECEITA BRUTA POR NEGOCIO (R$ mil; aba "Outras Informacoes"). 2T23 vazio: a aba Combinada so traz o 2T23 pro forma. Codigos 45-55 = memoria, nao entram nos codigos 1-41').font = bold; r += 1
BU = [  # key, code, label, fleury OI row, combinada OI row
 ('bu_fl', 45, 'Marca Fleury', 13, 11), ('bu_sp', 46, 'Outras Marcas SP (release: Demais Marcas SP)', 14, 12),
 ('bu_rj', 47, 'Marcas Rio de Janeiro', 15, 25), ('bu_mg', 48, 'Marcas MG', 16, 37), ('bu_reg', 49, 'Marcas Regionais', 17, 50),
 ('bu_b2b', 50, 'B2B', 18, 91),
 ('bu_ne', 51, 'Novos Elos e Plataformas de Saude (Fleury: linha unica; Combinada: Novos Elos + Plataformas)', 19, 74),
 ('bu_pl', 51, 'Plataformas de saude (so na aba Combinada; soma no cod 51)', None, 90),
]
for key, code, lab, fr, cr in BU:
    put(code, 'moeda', f'{lab} | Fleury OI l.{fr} / Combinada OI l.{cr}', 'fonte', ser(fr, cr, 'F_OI', 'C_OI'), key=key)
put(None, 'moeda', 'Receita bruta total (Outras Informacoes) - memoria', 'memoria', ser(10, 8, 'F_OI', 'C_OI'), key='bu_rb')
inp.cell(r, 7, 'moeda'); inp.cell(r, 8, 'CHECK: soma dos negocios - receita bruta (deve ser 0)'); inp.cell(r, 10, 'derivada')
for p in P:
    inp.cell(r, PC[p], '=' + '+'.join(f'{col(p)}{pos[k]}' for k, *_ in BU) + f'-{col(p)}{pos["bu_rb"]}')
pos['chk_bu'] = r; r += 1
ded = [('ded_imp', 'Impostos', 12, 9), ('ded_glo', 'Glosas', 13, 10), ('ded_aba', 'Abatimentos (so na aba Combinada)', None, 11)]
for key, lab, fr, cr in ded:
    vals = ser(fr, cr, 'F_DRE', 'C_DRE')
    if fr is None: vals.update({p: 0.0 for p in FLEURY_P})
    put(52, 'moeda', f'Deducoes da receita - {lab} | Fleury DRE l.{fr} / Combinada DRE l.{cr}', 'fonte', vals, key=key)
r += 1
inp.cell(r, 8, 'VOLUME (mil exames) E ACOES. Volumes da aba Combinada incluem o lab-to-lab do Pardini (quebra de serie no 2T23)').font = bold; r += 1
put(53, 'un', 'Exames - Unidades de Atendimento (PSC), mil | Fleury OI l.27 / Combinada OI l.102', 'fonte', ser(27, 102, 'F_OI', 'C_OI'), key='ex_psc')
put(54, 'un', 'Exames - B2B, mil | Fleury OI l.28 / Combinada OI l.103', 'fonte', ser(28, 103, 'F_OI', 'C_OI'), key='ex_b2b')
put(None, 'un', 'Numero de acoes (unidades) | Fleury OI l.74 / Combinada OI l.293', 'memoria', ser(74, 293, 'F_OI', 'C_OI', False), '#,##0', key='sh_raw')
inp.cell(r, 7, 'un'); inp.cell(r, 8, 'Numero de acoes em milhoes (derivada) - cod 55 alimenta Model linha 230'); inp.cell(r, 10, 'derivada')
inp.cell(r, 1, 55); inp.cell(r, 2, 1)
for p in P:
    if inp.cell(pos['sh_raw'], PC[p]).value is not None:
        inp.cell(r, PC[p], f'={col(p)}{pos["sh_raw"]}/1000000')
    inp.cell(r, PC[p]).number_format = '#,##0.0'
pos['sh_mm'] = r; r += 1
put(56, 'un', 'Divida liquida / EBITDA LTM divulgado pela companhia (x) | Fleury OI l.72 / Combinada OI l.291', 'fonte', ser(72, 291, 'F_OI', 'C_OI'), '0.00', key='nd_co')
LAST = r - 1
END = ((LAST + 30 + 9) // 10) * 10

# move conventions block, widen Inputs (2) ranges
oldconv = OLD_END + 3
conv_texts = [inp.cell(rr, 8).value for rr in range(oldconv + 1, oldconv + 10) if inp.cell(rr, 8).value]
for rr in range(oldconv, oldconv + 10): inp.cell(rr, 8).value = None
conv = END + 3
inp.cell(conv, 8, 'CONVENCOES').font = bold
conv_texts[0] = f'Linhas de origem: {R0} a {LAST}. As formulas da Inputs (2) varrem as linhas {R0} a {END}; as linhas {LAST + 1} a {END} estao livres para novas contas.'
conv_texts.append('Codigos 45-56 (receita bruta por negocio, deducoes, volumes, acoes, alavancagem divulgada) sao memoria: alimentam o bloco no fim do Model; o cod 55 alimenta a linha 230.')
for i, t in enumerate(conv_texts): inp.cell(conv + 1 + i, 8, t)
rx = re.compile(r'(Inputs!\$[A-Z]{1,2}\$\d+:\$[A-Z]{1,2}\$)' + str(OLD_END) + r'\b')
for row in w2.iter_rows():
    for c in row:
        if isinstance(c.value, str) and 'Inputs!' in c.value:
            c.value = rx.sub(lambda g: g.group(1) + str(END), c.value)

# Inputs (2) rows for codes 45-56
w2['C52'] = 'Receita bruta por negocio, deducoes, volumes e acoes (memo) - codigos 45-56 [Model linhas 361+ e 230]'
w2['C52'].font = bold
row2 = {}
items = [(45, 'Marca Fleury'), (46, 'Outras Marcas SP'), (47, 'Marcas Rio de Janeiro'), (48, 'Marcas MG'), (49, 'Marcas Regionais'),
         (50, 'B2B'), (51, 'Novos Elos e Plataformas'), (52, 'Deducoes da receita'), (53, 'Exames PSC (mil)'), (54, 'Exames B2B (mil)'),
         (55, 'Acoes (milhoes)'), (56, 'Divida liquida/EBITDA LTM divulgado (x)')]
annual2 = {8: ('D', 'G'), 13: ('I', 'L'), 18: ('N', 'Q'), 23: ('S', 'V'), 28: ('X', 'AA'), 33: ('AC', 'AF'), 38: ('AH', 'AK')}
for i, (code, lab) in enumerate(items):
    rr = 53 + i; row2[code] = rr
    w2.cell(rr, 2, code); w2.cell(rr, 3, lab)
    for c in range(1, 39): w2.cell(rr, c)._style = copy.copy(w2.cell(4 if code <= 54 else 15, c)._style)
    for c in range(4, 38):
        if c in annual2: continue
        w2.cell(rr, c, tr(w2.cell(4, c).value, f'{CL(c)}4', f'{CL(c)}{rr}'))
    for c, (a, b) in annual2.items():
        w2.cell(rr, c, f'=SUM({a}{rr}:{b}{rr})' if code <= 54 else f'={b}{rr}')
    if code == 56:
        for c in range(4, 39): w2.cell(rr, c).number_format = '0.00'

# ---------------------------------------------------------------- 3. Model block: revenue by business, price x volume
B0 = 361
for rr in range(B0, B0 + 30):
    for c in range(1, 34): assert m.cell(rr, c).value is None, (rr, c)
m.cell(B0, 3, 'Receita bruta por negocio - como divulgado pela Fleury (R$ milhoes)')._style = copy.copy(m['C13']._style)
for c in range(4, 34): m.cell(B0, c, f'={CL(c)}7')._style = copy.copy(m.cell(13, c)._style)
def link(code, c): return f"=IFERROR(INDEX('Inputs (2)'!$D${row2[code]}:$AL${row2[code]},MATCH({CL(c)}$7,'Inputs (2)'!$D$2:$AL$2,0)),\"\")"
LINES = [
 (B0 + 2, 'sum', 'B2C - Unidades de Atendimento (PSC)', (B0 + 3, B0 + 7)),
 (B0 + 3, 45, '   Marca Fleury', None), (B0 + 4, 46, '   Outras Marcas SP', None), (B0 + 5, 47, '   Rio de Janeiro', None),
 (B0 + 6, 48, '   Minas Gerais', None), (B0 + 7, 49, '   Regionais', None),
 (B0 + 8, 50, 'B2B', None), (B0 + 9, 51, 'Novos Elos e Plataformas', None),
 (B0 + 10, 'gross', 'Receita bruta (soma dos negocios)', None),
 (B0 + 11, 52, '(-) Deducoes (impostos, glosas, abatimentos)', None),
 (B0 + 12, 'net', 'Receita liquida (receita bruta - deducoes)', None),
 (B0 + 13, 'chk', 'Check vs Net Revenues (linha 15)', None),
 (B0 + 15, 'hdr', 'Volume e receita media por exame (Combinada inclui lab-to-lab do Pardini: quebra no 2T23)', None),
 (B0 + 16, 53, 'Exames PSC (mil)', None),
 (B0 + 17, 'tk_psc', 'Receita bruta PSC por exame (R$)', None),
 (B0 + 18, 54, 'Exames B2B (mil)', None),
 (B0 + 19, 'tk_b2b', 'Receita bruta B2B por exame (R$)', None),
 (B0 + 21, 'hdr', 'Alavancagem', None),
 (B0 + 22, 56, 'Divida liquida / EBITDA LTM - divulgado pela companhia (x)', None),
 (B0 + 23, 'nd', 'Divida liquida / EBITDA LTM - modelo, linha 303 (x)', None),
]
GROSS, NET = B0 + 10, B0 + 12
for rr, kind, lab, extra in LINES:
    lc = m.cell(rr, 3, lab); lc._style = copy.copy(m['C15']._style)
    if kind in ('gross', 'net', 'sum', 'hdr'): lc.font = Font(bold=True, name=m['C15'].font.name, size=m['C15'].font.sz)
    if kind == 'hdr': continue
    for c in range(4, 34):
        L = CL(c); cell = m.cell(rr, c); cell._style = copy.copy(m.cell(15, c)._style)
        if isinstance(kind, int):
            if c in ANNUAL:
                a, b = CL(c - 4), CL(c - 1)
                cell.value = (f'=IF(COUNT({a}{rr}:{b}{rr})=4,SUM({a}{rr}:{b}{rr}),"")' if kind not in (56,) else f'=IF({b}{rr}="","",{b}{rr})')
            else:
                cell.value = link(kind, c)
            if kind == 56: cell.number_format = '0.00x'
        elif kind == 'sum':
            cell.value = f'=IF(COUNT({L}{extra[0]}:{L}{extra[1]})=0,"",SUM({L}{extra[0]}:{L}{extra[1]}))'
        elif kind == 'gross':
            cell.value = f'=IF(COUNT({L}{B0 + 2},{L}{B0 + 8},{L}{B0 + 9})=0,"",SUM({L}{B0 + 2},{L}{B0 + 8},{L}{B0 + 9}))'
        elif kind == 'net':
            cell.value = f'=IF(COUNT({L}{GROSS})=0,"",{L}{GROSS}+{L}{B0 + 11})'
        elif kind == 'chk':
            cell.value = f'=IF(COUNT({L}{NET})=0,"",{L}15-{L}{NET})'
        elif kind == 'tk_psc':
            cell.value = f'=IFERROR({L}{B0 + 2}/{L}{B0 + 16}*1000,"")'; cell.number_format = '#,##0.0'
        elif kind == 'tk_b2b':
            cell.value = f'=IFERROR({L}{B0 + 8}/{L}{B0 + 18}*1000,"")'; cell.number_format = '#,##0.0'
        elif kind == 'nd':
            cell.value = f'=IF(ISNUMBER({L}303),{L}303,"")'; cell.number_format = '0.00x'
for c in range(4, 34):
    assert m.cell(17, c).value is None
    m.cell(17, c, f'=IF(COUNT({CL(c)}{NET})=0,"",{CL(c)}15-{CL(c)}{NET})')._style = copy.copy(m.cell(15, c)._style)
m['C17'] = f'Check Receita (linha 15 - receita liquida pela soma dos negocios, linha {NET})'
# shares -> Model 230 (end-of-period shares)
for c in HISTQ:
    assert m.cell(230, c).value is None
    m.cell(230, c, f"=INDEX('Inputs (2)'!$D${row2[55]}:$AL${row2[55]},MATCH({CL(c)}$7,'Inputs (2)'!$D$2:$AL$2,0))")
m['AG230'] = '=AF230'; m['AG230'].fill = PatternFill(fill_type=None)
log.append('Model 230: acoes no fim do trimestre (historico) e AG230 = AF230')

# ---------------------------------------------------------------- 4. 1T26 and 2T26 become actuals in AH / AI
SKIP = {5, 6, 7, 8, 9, 10}
conv_rows = []
for dst, srcc in (('AH', 'AC'), ('AI', 'AD')):
    for rr in range(1, 359):
        if rr in SKIP: continue
        s = m[f'{srcc}{rr}'].value
        if isinstance(s, str) and s.startswith('='):
            nf = tr(s, f'{srcc}{rr}', f'{dst}{rr}')
            if m[f'{dst}{rr}'].value != nf:
                m[f'{dst}{rr}'] = nf; conv_rows.append(rr)
        elif s is None and rr in (230,):
            pass
m['AJ189'] = '=AI183+SUM(AJ187:AJ188)'
for rr in (228, 229):
    for cc in ('AH', 'AI'): m[f'{cc}{rr}'] = None     # share buyback/issuance projection not used: shares are actual
log.append(f'Model AH/AI (1T26, 2T26) em modo realizado: {len(conv_rows)} celulas; AJ189 parte da divida LP real do 2T26')

# ---------------------------------------------------------------- 5. leverage ratios on LTM EBITDA (company definition)
def ltm(c):
    qs = []
    k = c
    while len(qs) < 4 and k >= 4:
        if k not in ANNUAL: qs.append(CL(k))
        k -= 1
    return qs if len(qs) == 4 else None
for c in range(4, 108):
    if c in ANNUAL: continue
    L = CL(c); q = ltm(c)
    for rr, num in ((300, f'SUM({L}178,{L}179,{L}180,{L}183,{L}201,{L}204)'),
                    (302, f'(SUM({L}178,{L}179,{L}180,{L}183,{L}201,{L}204)-{L}110)'),
                    (303, f'(SUM({L}178,{L}179,{L}183,{L}204)-{L}110)')):
        if m.cell(rr, c).value is None: continue
        den = '+'.join(f'{x}35' for x in q) if q else None
        m.cell(rr, c).value = f'=IFERROR({num}/({den}),"")' if q else '=""'
m['C300'] = 'TOTAL DEBT / EBITDA LTM'; m['C302'] = 'NET DEBT / EBITDA LTM (com arrendamentos)'; m['C303'] = 'NET DEBT (EX IFRS 16) / EBITDA LTM'
log.append('Model 300/302/303: EBITDA dos ultimos 12 meses')

# ---------------------------------------------------------------- 6. ROIC block: EBIT row and annual columns
for c in range(4, 108):
    if c in ANNUAL: continue
    L = CL(c)
    if m.cell(314, c).value == f'={L}96': m.cell(314, c).value = f'={L}39'
for c in sorted(ANNUAL):
    if c > 107: continue
    for rr in range(311, 322):
        if m.cell(rr, c).value is None and isinstance(m.cell(rr, c - 1).value, str):
            m.cell(rr, c).value = tr(m.cell(rr, c - 1).value, f'{CL(c - 1)}{rr}', f'{CL(c)}{rr}')
            m.cell(rr, c)._style = copy.copy(m.cell(rr, c - 1)._style)
log.append('Model 311-321: EBIT da linha 39 (era lucro antes do IR) e colunas anuais preenchidas')

# ---------------------------------------------------------------- 7. Valuation fixes
for c in 'GHIJKL':
    assert val[f'{c}50'].value == '#VALUE!'
    val[f'{c}50'] = tr('=R25/AVERAGE(Q46,R46)', 'R50', f'{c}50')
assert val['E8'].value is None
val['E8'] = 2; val['E8'].fill = PatternFill('solid', fgColor='FFFFF2CC')
val['F8'] = '1 = Bull, 2 = Base, 3 = Bear'
val.defined_names['scenario'] = DefinedName('scenario', attr_text='Valuation!$E$8', localSheetId=3)
for n in ('wacc', 'risk_free_rate', 'risk_premium', 'relevered_beta', 'totalcap'):
    if n in wb.defined_names and wb.defined_names[n].attr_text == '#REF!': del wb.defined_names[n]
log.append('Valuation: G50:L50 com formula (eram texto "#VALUE!"), nome scenario criado (E8), 5 nomes globais quebrados removidos')

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
lay.update(pos=pos, END=END, LASTROW=LAST, row2=row2)
json.dump(lay, open(OUT + '.layout.json', 'w'))
print('\n'.join(log)); print('Inputs rows to', LAST, 'END', OLD_END, '->', END)
