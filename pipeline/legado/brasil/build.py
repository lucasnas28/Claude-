import sys, shutil, copy, datetime as dt
sys.path.insert(0, __import__('os').environ.get('PIPE_DATA', 'dados')+'/work')
from data import *
import openpyxl
from openpyxl.utils import get_column_letter as CL
from openpyxl.styles import Font, PatternFill, Alignment

SRC = sys.argv[1]
OUT = sys.argv[2]
shutil.copyfile(SRC, OUT)
wb = openpyxl.load_workbook(OUT)

# ---------------------------------------------------------------- data
P = PERIODS                      # 1Q20..2Q26 (calendar labels)
PC = {p: 11 + i for i, p in enumerate(P)}   # column K onward
LASTC = CL(PC[P[-1]])            # AJ
FIRSTC = CL(PC[P[0]])            # K

def q(*tags):  return ser(list(tags), 'q')
def ytd(*tags): return ser(list(tags), 'ytd')
def ins(*tags, only=None): return ser(list(tags), 'i', only)

# Non-marketable securities: custom tag nvda:NonMarketableSecurities is not in the
# companyfacts API; values typed from the balance sheet of each 10-Q (R4.htm).
NONMKT_CUSTOM = {'1Q26': (43364e6, '10-Q 2026-05-20 acc 0001045810-26-000052'),
                 '2Q26': (51157e6, '10-Q 2026-08-26 acc 0001045810-26-000075')}

rows = []   # each: dict(key, cods[(c,s)], unid, pt, en, origem, vals{p:num}|forms{p:str}, style)
def add(key=None, cods=(), unid='moeda', pt='', en='', origem='fonte', vals=None, forms=None, kind=None):
    rows.append(dict(key=key, cods=list(cods), unid=unid, pt=pt, en=en, origem=origem,
                     vals=vals or {}, forms=forms or {}, kind=kind))
def section(title): rows.append(dict(kind='section', pt=title))
def blank(): rows.append(dict(kind='blank'))

def v(sd): return {p: x for p, x in sd[0].items() if p in P}

# ---------- DRE raw
section('DRE - BRUTO COMO DIVULGADO (USD unidades). 1T-3T = trimestre (10-Q, 3M); coluna 4T = ANO FISCAL 12M (10-K)')
IS = [
 ('rev',  'Receita',                         'Revenue [us-gaap:Revenues]', q('Revenues')),
 ('cogs', 'Custo da receita',                'Cost of revenue [us-gaap:CostOfRevenue]', q('CostOfRevenue')),
 ('gp',   'Lucro bruto',                     'Gross profit [us-gaap:GrossProfit]', q('GrossProfit')),
 ('rd',   'P&D',                             'Research and development [us-gaap:ResearchAndDevelopmentExpense]', q('ResearchAndDevelopmentExpense')),
 ('sga',  'Despesas de vendas, gerais e adm.','Sales, general and administrative [us-gaap:SellingGeneralAndAdministrativeExpense]', q('SellingGeneralAndAdministrativeExpense')),
 ('opx',  'Total de despesas operacionais (inclui custo de rescisao Arm no 1T FY23)', 'Total operating expenses [us-gaap:OperatingExpenses]', q('OperatingExpenses')),
 ('oi',   'Lucro operacional',               'Operating income [us-gaap:OperatingIncomeLoss]', q('OperatingIncomeLoss')),
 ('ii',   'Receita de juros',                'Interest income [us-gaap:InvestmentIncomeInterest]', q('InvestmentIncomeInterest')),
 ('ie',   'Despesa de juros (divulgada positiva)', 'Interest expense [us-gaap:InterestExpense ate FY24; InterestExpenseNonoperating depois]', q('InterestExpense', 'InterestExpenseNonoperating')),
 ('non',  'Total outras receitas (despesas), liquido', 'Total other income (expense), net [us-gaap:NonoperatingIncomeExpense]', q('NonoperatingIncomeExpense')),
 ('pt',   'Lucro antes do IR',               'Income before income tax [us-gaap:IncomeLossFromContinuingOperationsBefore...]', q('IncomeLossFromContinuingOperationsBeforeIncomeTaxesExtraordinaryItemsNoncontrollingInterest')),
 ('tax',  'Imposto de renda (divulgado positivo)', 'Income tax expense [us-gaap:IncomeTaxExpenseBenefit]', q('IncomeTaxExpenseBenefit')),
 ('ni',   'Lucro liquido',                   'Net income [us-gaap:NetIncomeLoss]', q('NetIncomeLoss')),
]
for key, pt, en, sd in IS:
    add('raw_' + key, (), 'moeda', pt, en, 'memoria', v(sd))
blank()
section('DFC - BRUTO COMO DIVULGADO (USD unidades), ACUMULADO NO ANO FISCAL (3M, 6M, 9M, 12M)')
CF = [
 ('da',  'Depreciacao e amortizacao (acum.)', 'Depreciation and amortization, YTD [us-gaap:DepreciationAndAmortization ate 2T FY21; DepreciationDepletionAndAmortization depois]', ytd('DepreciationDepletionAndAmortization', 'DepreciationAndAmortization')),
 ('div', 'Dividendos pagos (acum.)',          'Dividends paid, YTD [us-gaap:PaymentsOfDividends]', ytd('PaymentsOfDividends')),
]
for key, pt, en, sd in CF:
    add('raw_' + key, (), 'moeda', pt, en, 'memoria', v(sd))
blank()

# ---------- derived quarterly flows
section('DRE / DFC TRIMESTRALIZADAS (derivadas por formula) - linhas que alimentam o modelo')
def is_forms(rawkey):
    return ('IS', rawkey)
def cf_forms(rawkey):
    return ('CF', rawkey)
add('q_rev',  [(1, 1)],  'moeda', 'Receita (trimestre)', 'Revenue, quarterly', 'derivada', forms=is_forms('raw_rev'))
add('q_cogs', [(2, -1)], 'moeda', 'Custo da receita (trimestre) - inclui a parcela de D&A alocada ao custo', 'Cost of revenue, quarterly', 'derivada', forms=is_forms('raw_cogs'))
add('q_opx',  [(3, -1)], 'moeda', 'Total de despesas operacionais (trimestre) - P&D + SG&A + outros', 'Total operating expenses, quarterly', 'derivada', forms=is_forms('raw_opx'))
add('q_da',   [(4, -1), (3, 1)], 'moeda', 'D&A (trimestre). COD1=4 lanca D&A; COD2=3 estorna o mesmo valor do SG&A and Other para deixa-lo caixa', 'D&A, quarterly (cod 4) and reclass out of opex (cod 3)', 'derivada', forms=cf_forms('raw_da'))
add('q_ie',   [(5, -1), (6, 1)], 'moeda', 'Despesa de juros (trimestre). COD1=5 lanca juros; COD2=6 devolve ao total de outras receitas, que ja o desconta', 'Interest expense, quarterly (cod 5) and add-back to other income total (cod 6)', 'derivada', forms=is_forms('raw_ie'))
add('q_non',  [(6, 1)],  'moeda', 'Total outras receitas (despesas), liquido (trimestre)', 'Total other income (expense), net, quarterly', 'derivada', forms=is_forms('raw_non'))
add('q_tax',  [(7, -1)], 'moeda', 'Imposto de renda (trimestre)', 'Income tax expense, quarterly', 'derivada', forms=is_forms('raw_tax'))
add('q_div',  [(10, -1)], 'moeda', 'Dividendos pagos (trimestre)', 'Dividends paid, quarterly', 'derivada', forms=cf_forms('raw_div'))
add('q_ni',   [], 'moeda', 'Lucro liquido (trimestre) - memoria p/ conferencia', 'Net income, quarterly (memo)', 'derivada', forms=is_forms('raw_ni'))
add('chk_ni', [], 'moeda', 'CHECK: receita - custo - despesas op. + outras receitas - IR - lucro divulgado (deve ser 0)', 'Check: NI rebuilt from codes 1-7 less reported NI', 'derivada', forms=('CHKNI',))
add(None, [], 'moeda', '[vazio] cod 8 Equity in Income of Affiliates: NVIDIA nao tem resultado de equivalencia patrimonial na DRE', '', 'memoria')
add(None, [], 'moeda', '[vazio] cod 9 Minority Interest: NVIDIA nao tem participacao de nao controladores', '', 'memoria')
blank()

# ---------- balance sheet
section('BALANCO PATRIMONIAL - ATIVO (USD unidades, saldo no fim do trimestre fiscal)')
def b(*tags, only=None): return v(ins(*tags, only=only))
P4Q25 = P[P.index('4Q25'):]
nonmkt = b('EquitySecuritiesFVNINoncurrent', only={'4Q24', '4Q25'})
for p, (val, _) in NONMKT_CUSTOM.items(): nonmkt[p] = val
add('cash', [(11, 1)], 'moeda', 'Caixa e equivalentes', 'Cash and cash equivalents [us-gaap:CashAndCashEquivalentsAtCarryingValue]', 'fonte', b('CashAndCashEquivalentsAtCarryingValue'))
add('mkt',  [(11, 1)], 'moeda', 'Titulos negociaveis de divida', 'Marketable (debt) securities [us-gaap:MarketableSecuritiesCurrent; DebtSecuritiesCurrent a partir do 4T FY26]', 'fonte', b('MarketableSecuritiesCurrent', 'DebtSecuritiesCurrent'))
add('mkteq',[(11, 1)], 'moeda', 'Titulos negociaveis de capital (linha propria no balanco so a partir do 4T FY26)', 'Marketable equity securities [us-gaap:EquitySecuritiesFvNi]', 'fonte', b('EquitySecuritiesFvNi', only=set(P4Q25)))
add('ar',   [(12, 1)], 'moeda', 'Contas a receber, liquido', 'Accounts receivable, net [us-gaap:AccountsReceivableNetCurrent]', 'fonte', b('AccountsReceivableNetCurrent'))
add('inv',  [(13, 1)], 'moeda', 'Estoques', 'Inventories [us-gaap:InventoryNet]', 'fonte', b('InventoryNet'))
add('prep', [(16, 1)], 'moeda', 'Despesas antecipadas e outros ativos circulantes (cod 16; cod 15 Prepaid nao e aberto separadamente)', 'Prepaid expenses and other current assets [us-gaap:PrepaidExpenseAndOtherAssetsCurrent]', 'fonte', b('PrepaidExpenseAndOtherAssetsCurrent'))
add('ppe',  [(18, 1)], 'moeda', 'Imobilizado, liquido', 'Property and equipment, net [us-gaap:PropertyPlantAndEquipmentNet]', 'fonte', b('PropertyPlantAndEquipmentNet'))
add('rou',  [(19, 1)], 'moeda', 'Ativos de arrendamento operacional', 'Operating lease assets [us-gaap:OperatingLeaseRightOfUseAsset]', 'fonte', b('OperatingLeaseRightOfUseAsset'))
add('gw',   [(20, 1)], 'moeda', 'Goodwill', 'Goodwill [us-gaap:Goodwill]', 'fonte', b('Goodwill'))
add('intg', [(21, 1)], 'moeda', 'Intangiveis, liquido', 'Intangible assets, net [us-gaap:IntangibleAssetsNetExcludingGoodwill]', 'fonte', b('IntangibleAssetsNetExcludingGoodwill'))
add('dta',  [(14, 1)], 'moeda', 'IR diferido ativo (nao circulante; vai no cod 14, unica linha de IR diferido ativo do modelo)', 'Deferred income tax assets [us-gaap:DeferredIncomeTaxAssetsNet]', 'fonte', b('DeferredIncomeTaxAssetsNet'))
add('nonmkt',[(22, 1)], 'moeda', 'Titulos nao negociaveis (linha propria so no 4T FY25, 4T FY26 e FY27; nos demais periodos estao dentro de Outros ativos)', 'Non-marketable securities [us-gaap:EquitySecuritiesFVNINoncurrent; nvda:NonMarketableSecurities em FY27, digitado do 10-Q R4]', 'fonte', nonmkt)
add('oth',  [(22, 1)], 'moeda', 'Outros ativos nao circulantes', 'Other assets [us-gaap:OtherAssetsNoncurrent; OtherAssets no 1T-2T FY21]', 'fonte', b('OtherAssetsNoncurrent', 'OtherAssets'))
add('tca',  [], 'moeda', 'Total do ativo circulante (divulgado)', 'Total current assets [us-gaap:AssetsCurrent]', 'memoria', b('AssetsCurrent'))
add('ta',   [], 'moeda', 'Total do ativo (divulgado)', 'Total assets [us-gaap:Assets]', 'memoria', b('Assets'))
add('chk_a',[], 'moeda', 'CHECK: soma das linhas do ativo menos total divulgado (deve ser 0)', 'Check: mapped assets less reported total', 'derivada', forms=('SUMCHK', ['cash','mkt','mkteq','ar','inv','prep','ppe','rou','gw','intg','dta','nonmkt','oth'], 'ta'))
add(None, [], 'moeda', '[vazio] cod 15 Prepaid expenses: NVIDIA divulga junto com outros ativos circulantes (cod 16)', '', 'memoria')
add(None, [], 'moeda', '[vazio] cod 17 Equity in affiliates: NVIDIA nao tem investimento por equivalencia patrimonial no balanco', '', 'memoria')
blank()

section('BALANCO PATRIMONIAL - PASSIVO E PL (USD unidades)')
add('ap',   [(23, 1)], 'moeda', 'Fornecedores', 'Accounts payable [us-gaap:AccountsPayableCurrent]', 'fonte', b('AccountsPayableCurrent'))
add('acc',  [(24, 1)], 'moeda', 'Provisoes e outros passivos circulantes (bruto; inclui IR a pagar e arrendamento CP, estornados abaixo)', 'Accrued and other current liabilities [us-gaap:AccruedLiabilitiesCurrent]', 'fonte', b('AccruedLiabilitiesCurrent'))
add('txp',  [(25, 1), (24, -1)], 'moeda', 'IR a pagar (nota de passivos circulantes). COD1=25 lanca; COD2=24 estorna de Accrued', 'Taxes payable [us-gaap:TaxesPayableCurrent] - reclass out of accrued', 'fonte', b('TaxesPayableCurrent'))
add('olc',  [(29, 1), (24, -1)], 'moeda', 'Arrendamento operacional CP (nota). COD1=29 lanca; COD2=24 estorna de Accrued', 'Operating lease liabilities, current [us-gaap:OperatingLeaseLiabilityCurrent] - reclass out of accrued', 'fonte', b('OperatingLeaseLiabilityCurrent'))
std = b('LongTermDebtCurrent', 'DebtCurrent')
for p in P: std.setdefault(p, 0)
add('std',  [(27, 1)], 'moeda', 'Divida de curto prazo (0 = linha nao apresentada no balanco do periodo; teste de subtotal confirma)', 'Short-term debt [us-gaap:LongTermDebtCurrent; DebtCurrent]', 'fonte', std)
add('ltd',  [(30, 1)], 'moeda', 'Divida de longo prazo', 'Long-term debt [us-gaap:LongTermDebtNoncurrent; LongTermDebt no 1T-2T FY21, sem parcela circulante]', 'fonte', b('LongTermDebtNoncurrent', 'LongTermDebt'))
add('oll',  [(31, 1)], 'moeda', 'Arrendamento operacional LP', 'Long-term operating lease liabilities [us-gaap:OperatingLeaseLiabilityNoncurrent]', 'fonte', b('OperatingLeaseLiabilityNoncurrent'))
add('olt',  [(34, 1)], 'moeda', 'Outros passivos de longo prazo (bruto; inclui IR diferido passivo, estornado abaixo)', 'Other long-term liabilities [us-gaap:OtherLiabilitiesNoncurrent]', 'fonte', b('OtherLiabilitiesNoncurrent'))
add('dtl',  [(33, 1), (34, -1)], 'moeda', 'IR diferido passivo (nota). COD1=33 lanca; COD2=34 estorna de Other LT liabilities', 'Deferred income tax liabilities [us-gaap:DeferredIncomeTaxLiabilitiesNet] - reclass out of other LT liabilities', 'fonte', b('DeferredIncomeTaxLiabilitiesNet'))
add('tl',   [], 'moeda', 'Total do passivo (divulgado)', 'Total liabilities [us-gaap:Liabilities]', 'memoria', b('Liabilities'))
add('chk_l',[], 'moeda', 'CHECK: soma das linhas do passivo menos total divulgado (deve ser 0)', 'Check: mapped liabilities less reported total', 'derivada', forms=('SUMCHK', ['ap','acc','std','ltd','oll','olt'], 'tl'))
add('cs',   [(37, 1)], 'moeda', 'Acoes ordinarias (valor nominal)', 'Common stock [us-gaap:CommonStockValue]', 'fonte', b('CommonStockValue'))
add('apic', [(37, 1)], 'moeda', 'Capital adicional integralizado', 'Additional paid-in capital [us-gaap:AdditionalPaidInCapital]', 'fonte', b('AdditionalPaidInCapital'))
ts = b('TreasuryStockValue')
for p in P: ts.setdefault(p, 0)
add('ts',   [(38, -1)], 'moeda', 'Acoes em tesouraria (divulgada positiva; 0 = acoes canceladas, linha nao apresentada a partir do 1T FY23)', 'Treasury stock [us-gaap:TreasuryStockValue]', 'fonte', ts)
add('aoci', [(39, 1)], 'moeda', 'Outros resultados abrangentes acumulados', 'Accumulated other comprehensive income (loss) [us-gaap:AccumulatedOtherComprehensiveIncomeLossNetOfTax]', 'fonte', b('AccumulatedOtherComprehensiveIncomeLossNetOfTax'))
add('re',   [(41, 1)], 'moeda', 'Lucros retidos', 'Retained earnings [us-gaap:RetainedEarningsAccumulatedDeficit]', 'fonte', b('RetainedEarningsAccumulatedDeficit'))
add('te',   [], 'moeda', 'Total do PL (divulgado)', 'Total shareholders equity [us-gaap:StockholdersEquity]', 'memoria', b('StockholdersEquity'))
add('chk_e',[], 'moeda', 'CHECK: soma das linhas do PL menos total divulgado (deve ser 0)', 'Check: mapped equity less reported total', 'derivada', forms=('EQCHK',))
add('tle',  [], 'moeda', 'Total do passivo + PL (divulgado)', 'Total liabilities and shareholders equity [us-gaap:LiabilitiesAndStockholdersEquity]', 'memoria', b('LiabilitiesAndStockholdersEquity'))
for c, t in [(26, 'Other current liabilities (non-debt): tudo que a NVIDIA divulga esta em Accrued and other (cod 24)'),
             (28, 'Short term debt (Revolving): NVIDIA nao tem linha de credito sacada'),
             (32, 'Convertible debt: sem divida conversivel no periodo'),
             (35, 'Minority interest: inexistente'), (36, 'Convertible Preferred Stock: preferenciais = 0 no balanco'),
             (40, 'Other equity account: inexistente')]:
    add(None, [], 'moeda', f'[vazio] cod {c} {t}', '', 'memoria')
blank()

section('ACOES (unidades de acoes como divulgado; ajuste por desdobramento 4:1 jul/21 e 10:1 jun/24)')
dil = q('WeightedAverageNumberOfDilutedSharesOutstanding'); bas = q('WeightedAverageNumberOfSharesOutstandingBasic'); eps = q('EarningsPerShareDiluted')
def factor(sd): return {p: split_factor(sd[1][p][1]) for p in P}
def drop4(dct): return {p: x for p, x in dct.items() if not p.startswith('4Q')}
add('raw_dil', [], 'acoes', 'Acoes diluidas, media ponderada, como divulgado (1T-3T: trimestre; 4T vazio, 10-K so traz o ano)', 'Diluted weighted average shares [us-gaap:WeightedAverageNumberOfDilutedSharesOutstanding]', 'memoria', drop4(v(dil)))
add('f_dil', [], 'fator', 'Fator de desdobramento aplicado (pela data de arquivamento do filing de onde veio o numero)', 'Split factor to current share basis', 'memoria', drop4(factor(dil)))
add('raw_bas', [], 'acoes', 'Acoes basicas, media ponderada, como divulgado (4T vazio)', 'Basic weighted average shares [us-gaap:WeightedAverageNumberOfSharesOutstandingBasic]', 'memoria', drop4(v(bas)))
add('f_bas', [], 'fator', 'Fator de desdobramento aplicado (basicas)', 'Split factor (basic)', 'memoria', drop4(factor(bas)))
add('adj_dil', [(42, 1)], 'acoes', 'Acoes diluidas, media ponderada, base atual (pos-desdobramentos)', 'Diluted weighted average shares, split-adjusted', 'derivada', forms=('MUL', 'raw_dil', 'f_dil'))
add('adj_bas', [(43, 1)], 'acoes', 'Acoes basicas, media ponderada, base atual (pos-desdobramentos)', 'Basic weighted average shares, split-adjusted', 'derivada', forms=('MUL', 'raw_bas', 'f_bas'))
fy_dil = {p: x for p, x in v(dil).items() if p.startswith('4Q')}
fy_f = {p: x for p, x in factor(dil).items() if p.startswith('4Q')}
add('fy_dil', [], 'acoes', 'Acoes diluidas, media do ANO FISCAL (10-K), como divulgado - coluna 4T', 'Diluted weighted average shares, fiscal year', 'memoria', fy_dil)
add('fy_f', [], 'fator', 'Fator de desdobramento (ano fiscal)', 'Split factor (FY)', 'memoria', fy_f)
add('raw_eps', [], 'por acao', 'LPA diluido, como divulgado (1T-3T trimestre; 4T = ano fiscal)', 'Diluted EPS [us-gaap:EarningsPerShareDiluted]', 'memoria', v(eps))
add('f_eps', [], 'fator', 'Fator de desdobramento (LPA)', 'Split factor (EPS)', 'memoria', factor(eps))
add('chk_eps', [], 'por acao', 'CHECK: lucro/acoes diluidas (base atual) menos LPA divulgado ajustado (1T-3T; arredondamento de 1 centavo e normal)', 'Check: NI / adj. diluted shares less adj. reported EPS', 'derivada', forms=('EPSCHK',))
add(None, [], 'acoes', '[vazio] End of period basic shares (Model linha 230): 10-Q nao traz saldo de acoes no fim do trimestre, so a capa (data posterior); nao preenchido', '', 'memoria')

# ---------------------------------------------------------------- write Inputs
ws = wb['Inputs']
HDR = 4; R0 = 5
ws['A1'] = 'Fator p/ milhões:'; ws['B1'] = 0.000001
ws['D1'] = 'NVIDIA Corp (CIK 1045810) - fonte: SEC EDGAR XBRL (companyfacts API, 10-Q/10-K). Moeda: USD. Unidade dos brutos: unidades (USD e acoes). Fator converte p/ milhoes.'
hdr = ['COD1', 'SINAL1', 'COD2', 'SINAL2', 'COD3', 'SINAL3', 'UNID', 'Conta (PT)', 'Account (EN)', 'ORIGEM']
for i, h in enumerate(hdr): ws.cell(HDR, 1 + i, h)
ws.cell(2, 10, 'Trimestre fiscal NVIDIA'); ws.cell(3, 10, 'Fim do trimestre fiscal')
for p in P:
    c = PC[p]
    ws.cell(HDR, c, p)
    ws.cell(2, c, fiscal(p))
    ws.cell(3, c, dt.date.fromisoformat(enddate(p))).number_format = 'dd/mm/yyyy'

r = R0
pos = {}
for row in rows:
    row['r'] = r
    if row.get('key'): pos[row['key']] = r
    r += 1
LASTROW = r - 1
END = ((LASTROW + 30 + 9) // 10) * 10     # range end of Inputs (2) formulas

def col(p): return CL(PC[p])
def quarter_group(p):
    yy = p[2:]; return [f'{k}Q{yy}' for k in range(1, 5) if f'{k}Q{yy}' in PC]

bold = Font(bold=True)
for row in rows:
    rr = row['r']
    if row.get('kind') == 'section':
        ws.cell(rr, 8, row['pt']).font = bold; continue
    if row.get('kind') == 'blank': continue
    for i, (c, s) in enumerate(row['cods']):
        ws.cell(rr, 1 + 2 * i, c); ws.cell(rr, 2 + 2 * i, s)
    ws.cell(rr, 7, row['unid']); ws.cell(rr, 8, row['pt']); ws.cell(rr, 9, row['en']); ws.cell(rr, 10, row['origem'])
    for p, x in row['vals'].items():
        ws.cell(rr, PC[p], x)
    f = row['forms']
    if not f: continue
    for p in P:
        cc = col(p); qn = int(p[0])
        if f[0] == 'IS':
            raw = pos[f[1]]
            if qn < 4: fx = f'={cc}{raw}'
            else:
                g = quarter_group(p); fx = f'={cc}{raw}-SUM({col(g[0])}{raw}:{col(g[2])}{raw})'
        elif f[0] == 'CF':
            raw = pos[f[1]]
            if qn == 1: fx = f'={cc}{raw}'
            else: fx = f'={cc}{raw}-{CL(PC[p]-1)}{raw}'
        elif f[0] == 'CHKNI':
            fx = f'={cc}{pos["q_rev"]}-{cc}{pos["q_cogs"]}-{cc}{pos["q_opx"]}+{cc}{pos["q_non"]}-{cc}{pos["q_tax"]}-{cc}{pos["q_ni"]}'
        elif f[0] == 'SUMCHK':
            fx = '=' + '+'.join(f'{cc}{pos[k]}' for k in f[1]) + f'-{cc}{pos[f[2]]}'
        elif f[0] == 'EQCHK':
            fx = f'={cc}{pos["cs"]}+{cc}{pos["apic"]}-{cc}{pos["ts"]}+{cc}{pos["aoci"]}+{cc}{pos["re"]}-{cc}{pos["te"]}'
        elif f[0] == 'MUL':
            if qn == 4: continue
            fx = f'={cc}{pos[f[1]]}*{cc}{pos[f[2]]}'
        elif f[0] == 'EPSCHK':
            if qn == 4: continue
            fx = f'=ROUND({cc}{pos["q_ni"]}/{cc}{pos["adj_dil"]}-{cc}{pos["raw_eps"]}/{cc}{pos["f_eps"]},4)'
        ws.cell(rr, PC[p], fx)

# number formats
for row in rows:
    if row.get('kind'): continue
    fmt = {'moeda': '#,##0;(#,##0)', 'acoes': '#,##0', 'fator': '0', 'por acao': '0.00;(0.00)'}[row['unid']]
    for p in P: ws.cell(row['r'], PC[p]).number_format = fmt
for cidx, w in [(1, 6), (2, 7), (3, 6), (4, 7), (5, 6), (6, 7), (7, 8), (8, 70), (9, 60), (10, 10)]:
    ws.column_dimensions[CL(cidx)].width = w
for p in P: ws.column_dimensions[col(p)].width = 15
for cidx in range(1, PC[P[-1]] + 1): ws.cell(HDR, cidx).font = bold
ws.freeze_panes = ws.cell(R0, 11)

# conventions block, generated from the same variables
conv = END + 3
ws.cell(conv, 8, 'CONVENCOES').font = bold
notes = [
 f'Linhas de origem: {R0} a {LASTROW}. As formulas da Inputs (2) varrem as linhas {R0} a {END}; as linhas {LASTROW + 1} a {END} estao livres para novas contas.',
 f'Periodos: colunas {FIRSTC} a {LASTC}, rotulo de trimestre-calendario. Cada trimestre fiscal NVIDIA vai para o trimestre-calendario em que majoritariamente cai (linha 2 mostra o trimestre fiscal; linha 3 a data de encerramento). Ex.: 1Q25 = 1T FY26 (fev-abr/2025).',
 'Brutos em USD unidades e acoes em unidades; a Inputs (2) multiplica pelo fator da celula B1 as linhas com UNID = moeda ou acoes. LPA e fatores passam sem fator.',
 'SINAL converte para a convencao do modelo (custos, despesas, impostos, D&A, juros e dividendos negativos). COD2/SINAL2 usam a mesma linha num segundo destino (reclassificacao/estorno).',
 'ORIGEM: fonte = copiado do XBRL; derivada = formula nesta aba (trimestralizacao, ajuste de desdobramento, checks); memoria = informativa, sem codigo.',
 'Valores do XBRL: para cada periodo usa-se o fato arquivado mais recentemente (ja reflete reapresentacoes e desdobramentos nas comparacoes posteriores).',
]
for i, t in enumerate(notes): ws.cell(conv + 1 + i, 8, t)

# ---------------------------------------------------------------- Inputs (2)
w2 = wb['Inputs (2)']
# extend header with 2020 after existing columns
ext = ['1Q20', '2Q20', '3Q20', '4Q20', 2020]
for i, h in enumerate(ext):
    c = 34 + i     # AH..AL
    w2.cell(2, c, h)
    w2.cell(2, c)._style = copy.copy(w2.cell(2, 29 + i if i < 4 else 33)._style)
    w2.column_dimensions[CL(c)].width = w2.column_dimensions['AC'].width or 10
# new share rows
w2['B48'] = 42; w2['C48'] = 'Diluted Shares outstanding (weighted avg, mm) [Model linha 222]'
w2['B49'] = 43; w2['C49'] = 'Weighted average basic shares outstanding (mm) [Model linha 231]'
for rr in (48, 49):
    for c in range(1, 39):
        src = w2.cell(47, c)
        w2.cell(rr, c)._style = copy.copy(src._style)
    for c in range(4, 39): w2.cell(rr, c).number_format = '#,##0'

IN = 'Inputs!'
def rng(cl): return f'{IN}${cl}${R0}:${cl}${END}'
VAL = f'INDEX({IN}${FIRSTC}${R0}:${LASTC}${END},0,MATCH({{h}},{IN}${FIRSTC}${HDR}:${LASTC}${HDR},0))'
ESC = f'((({rng("G")}="moeda")+({rng("G")}="acoes"))*{IN}$B$1+(({rng("G")}<>"moeda")*({rng("G")}<>"acoes")))'
def f2(h, d):
    val = VAL.format(h=h)
    cnt = '+'.join(f'SUMPRODUCT(({rng(c)}={d})*ISNUMBER({val}))' for c in 'ACE')
    parts = '+'.join(f'SUMPRODUCT(({rng(c)}={d})*{rng(s)}*{ESC}*{val})' for c, s in (('A', 'B'), ('C', 'D'), ('E', 'F')))
    return f'=IF(ISNA(MATCH({h},{IN}${FIRSTC}${HDR}:${LASTC}${HDR},0)),NA(),IF({cnt}=0,NA(),{parts}))'

code_rows = [r_ for r_ in range(4, 50) if isinstance(w2.cell(r_, 2).value, int)]
annual_cols = {8, 13, 18, 23, 28, 33, 38}
qcols = [c for c in range(4, 38) if c not in annual_cols]
n2 = 0
for r_ in code_rows:
    for c in qcols:
        cell = w2.cell(r_, c)
        assert cell.value is None, (r_, c, cell.value)
        cell.value = f2(f'{CL(c)}$2', f'$B{r_}'); n2 += 1
    # annual 2020 column AL, mirroring the existing annual logic
    if r_ <= 13: w2.cell(r_, 38, f'=SUM(AH{r_}:AK{r_})')
    elif r_ <= 47: w2.cell(r_, 38, f'=AK{r_}')
    if r_ <= 47:
        w2.cell(r_, 38)._style = copy.copy(w2.cell(r_, 33)._style)
        for i in range(4): w2.cell(r_, 34 + i)._style = copy.copy(w2.cell(r_, 29 + i)._style)
# grey label for grid periods without source data (3Q26, 4Q26, 2026)
grey = Font(color='FF808080', bold=w2['AE2'].font.bold, name=w2['AE2'].font.name, size=w2['AE2'].font.sz)
for a in ('AE2', 'AF2', 'AG2'): w2[a].font = grey
w2['AN2'] = 'Sem dado na fonte: 3Q26, 4Q26 e 2026 (rotulo em cinza). 2020 fica nas colunas AH-AL.'

# ---------------------------------------------------------------- Model links
m = wb['Model']
code2model = {1: 15, 2: 20, 3: 24, 4: 36, 5: 40, 6: 89, 7: 97, 10: 104,
              11: 110, 12: 111, 13: 115, 14: 119, 16: 121, 18: 132, 19: 142, 20: 145, 21: 146, 22: 155,
              23: 164, 24: 168, 25: 170, 27: 178, 29: 180, 30: 183, 31: 201, 33: 205, 34: 206,
              37: 220, 38: 233, 39: 234, 41: 236, 42: 222, 43: 231}
code2row2 = {w2.cell(r_, 2).value: r_ for r_ in code_rows}
mann = {8, 13, 18, 23, 28}   # H M R W AB
nm = 0
for code, mr in code2model.items():
    r2 = code2row2[code]
    for c in range(4, 33):
        if c in mann: continue
        if code in (42, 43) and m.cell(7, c).value and str(m.cell(7, c).value).startswith('="4Q'): continue
        cell = m.cell(mr, c)
        assert cell.value is None, (mr, c, cell.value)
        L = CL(c)
        cell.value = f"=INDEX('Inputs (2)'!$D${r2}:$AL${r2},MATCH({L}$7,'Inputs (2)'!$D$2:$AL$2,0))"
        nm += 1

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('Inputs rows', R0, LASTROW, 'END', END, 'Inputs(2) formulas', n2, 'Model links', nm)
import json
json.dump(dict(pos=pos, R0=R0, LASTROW=LASTROW, END=END, PC=PC), open(__import__('os').environ.get('PIPE_DATA', 'dados')+'/work/layout.json', 'w'))
