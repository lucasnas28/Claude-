"""Preenche a base IF v7 só com o IF.data (BCB): balanço, DRE (desacumulada), capital e carteira.
uso: build_ifd.py BASE OUT NOME TICKER PRU FIN SEGMENTO [FIN_PROXY_ATE]"""
import sys, shutil, datetime, re
sys.path.insert(0, __import__('os').environ.get('PIPE_DATA', 'dados')+'/bank')
from ifd import P, letter, lsum, name, new, rows, am
import openpyxl
from openpyxl.utils import get_column_letter as CL
BASE, OUT, NOME, TICK, PRU, FIN, SEG = sys.argv[1:8]
shutil.copyfile(BASE, OUT)
wb = openpyxl.load_workbook(OUT); inp = wb['Inputs']
COL = {p: 5 + i for i, p in enumerate(P)}
ROW = {inp.cell(r, 2).value: r for r in range(5, 149) if inp.cell(r, 2).value}
AV = [p for p in P if (letter(p, 4, PRU, 'z') if new(p) else letter(p, 4, PRU, 'j')) is not None and name(p, 2, PRU, r'^Ativo Total(?! Ajustado)') is not None]
written = 0
def put(code, vals, note):
    global written
    r = ROW[code]; inp.cell(r, 4).value = note
    for p, v in vals.items():
        if p in AV and v is not None: inp.cell(r, COL[p]).value = round(float(v), 6); written += 1
def empty(code, why): inp.cell(ROW[code], 4).value = f'[vazio] {why}'
SRC = f'IF.data/BCB, conglomerado prudencial {PRU}'
def L2(p, *lets): return lsum(p, 2, PRU, list(lets))
def N3(p, *pats):
    vs = [name(p, 3, PRU, pt) for pt in pats]
    return None if all(v is None for v in vs) else sum(v for v in vs if v is not None)
# ---------------- BALANÇO (rel. 2 ativo por letra - 2 layouts; rel. 3 passivo por nome)
A = {
 'disp': (lambda p: L2(p, 'a'), 'Disponibilidades (a)'),
 'aplic_mo': (lambda p: L2(p, 'b'), 'Aplicações interfinanceiras de liquidez TOTAL (b)'),
 'tvm_prop': (lambda p: L2(p, 'c'), 'TVM (c; até 4T24 TVM e derivativos juntos)'),
 'deriv_a': (lambda p: L2(p, 'd') if new(p) else None, 'Instrumentos derivativos (d; só 2025+)'),
 'cred_op': (lambda p: (L2(p, 'e1', 'e3', 'e4', 'f1', 'f3', 'h1') if new(p) else (L2(p, 'd1', 'e') or 0) - (L2(p, 'j') or 0)),
             'Operações de crédito brutas + arrendamento (até 4T24: d1 + e - j | 2025+: e1+e3+e4 + f1+f3 + h1 valores a receber de transações de pagamento pós-pago)'),
 'cred_tvm': (lambda p: L2(p, 'g1', 'g3', 'g4') if new(p) else None, 'Outras operações com características de concessão de crédito, valor bruto (g1+g3+g4; só 2025+)'),
 'cred_pdd': (lambda p: L2(p, 'e2', 'f2', 'g2', 'h2') if new(p) else L2(p, 'd2'), 'Provisão / perda esperada de crédito (até 4T24: d2 | 2025+: e2+f2+g2+h2)'),
 'out_a': (lambda p: L2(p, 'i') if new(p) else L2(p, 'f', 'g'), 'Outros créditos e outros ativos realizáveis (até 4T24: f + g | 2025+: i; inclui fiscais e relações interfinanceiras)'),
 'inv_col': (lambda p: L2(p, 'j') if new(p) else L2(p, 'h'), 'Ativo permanente TOTAL (investimentos + imobilizado + intangível: o IF.data não abre)'),
}
for code, (fn, desc) in A.items():
    put(code, {p: fn(p) for p in AV}, f'{SRC}, rel. 2 (ativo): {desc} · fonte · moeda (R$ convertido para R$ milhões)')
LIAB = {
 'dep_vista': ([r'^Depósitos à Vista'], 'Depósitos à vista'), 'dep_poup': ([r'^Depósitos de Poupança'], 'Depósitos de poupança'),
 'dep_inter': ([r'^Depósitos Interfinanceiros'], 'Depósitos interfinanceiros'), 'dep_prazo': ([r'^Depósitos a Prazo'], 'Depósitos a prazo'),
 'dep_out': ([r'^(Outros Depósitos|Depósitos Outros)', r'^Conta de Pagamento Pré-Paga'], 'Outros depósitos + conta de pagamento pré-paga'),
 'mo_prop': ([r'^Obrigações por Operações Compromissadas'], 'Obrigações por operações compromissadas TOTAL'),
 'div_emis': ([r'^Letras de Crédito Imobiliário', r'^Letras de Crédito do Agronegócio', r'^Letras Financeiras', r'^(Outros Recursos de Aceites|Demais instrumentos de dívida)'], 'LCI + LCA + LF + outros recursos de aceites/instrumentos de dívida'),
 'div_ext': ([r'^Obrigações por Títulos e Valores Mobiliários no Exterior'], 'Obrigações por TVM no exterior'),
 'div_sub': ([r'^Instrumentos de Dívida Elegíveis a Capital'], 'Instrumentos de dívida elegíveis a capital (só 2025+; antes dentro de outras obrigações)'),
 'emp_empr': ([r'^Obrigações por Empréstimos e Repasses'], 'Obrigações por empréstimos e repasses TOTAL'),
 'deriv_p': ([r'^Instrumentos Derivativos'], 'Instrumentos derivativos'),
 'out_p': ([r'^Outras Obrigações', r'^Resultados de Exercícios Futuros'], 'Outras obrigações (inclui fiscais, provisões e, até 4T24, dívida subordinada) + resultados de exercícios futuros'),
 'res_luc': ([r'^Patrimônio Líquido'], 'PATRIMÔNIO LÍQUIDO TOTAL do conglomerado (o IF.data não abre capital/reservas/minoritários): lançado aqui'),
}
for code, (pats, desc) in LIAB.items():
    put(code, {p: N3(p, *pats) for p in AV}, f'{SRC}, rel. 3 (passivo): {desc} · fonte · moeda')
for c_, why in [('aplic_di', 'dentro de aplic_mo'), ('aplic_rg', 'não se aplica'), ('aplic_pdd', 'não divulgada'), ('tvm_vinc', 'dentro de tvm_prop'), ('tvm_rg', 'não se aplica'),
                ('tvm_pdd', 'não divulgada'), ('rel_a', 'dentro de out_a'), ('fisc_a', 'dentro de out_a'), ('inv_out', 'dentro de inv_col'), ('imob', 'dentro de inv_col'), ('intang', 'dentro de inv_col'),
                ('mo_terc', 'dentro de mo_prop'), ('mo_livre', 'dentro de mo_prop'), ('div_coe', 'dentro de div_emis'), ('emp_rep', 'dentro de emp_empr'), ('rel_p', 'dentro de out_p'),
                ('prov_gar', 'dentro de out_p'), ('prov_tec', 'não se aplica'), ('prov_dem', 'dentro de out_p'), ('fisc_p', 'dentro de out_p'),
                ('cap_social', 'PL total em res_luc'), ('res_cap', 'idem'), ('ora', 'idem'), ('tesour', 'idem'), ('minor', 'IF.data não separa (PL total em res_luc)')]:
    empty(c_, why)
# ---------------- DRE (acumulada no semestre -> desacumulada por fórmula)
RAW = [
 ('rec_cred', ['a1', 'a2'], ['c', 'd', 'e'], 'Rendas de operações de crédito + arrendamento (até 4T24: a1+a2 | 2025+: c+d+e)'),
 ('rec_tvm', ['a3', 'a4', 'a5', 'b4'], ['b', 'i', 'j'], 'TVM + derivativos + câmbio/outros (até 4T24: a3+a4+a5+b4, TVM inclui aplicações interfinanceiras | 2025+: b+i+j)'),
 ('rec_aplic', ['a6'], ['a'], 'Aplicações (até 4T24: a6 compulsórias | 2025+: a aplicações interfinanceiras)'),
 ('desp_dep', ['b1'], ['g'], 'Despesas de captação (b1 | g)'),
 ('desp_div', [], ['h'], 'Instrumentos de dívida elegíveis a capital (só 2025+: h)'),
 ('desp_emp', ['b2', 'b3'], [], 'Empréstimos e repasses + arrendamento (até 4T24: b2+b3 | 2025+: dentro de captações)'),
 ('pdd_desp', ['b5'], ['f'], 'Provisão / perda esperada (b5 | f)'),
 ('serv', ['d1', 'd2'], ['l', 'm', 'n'], 'Serviços + tarifas (até 4T24: d1+d2 | 2025+: l+m+n, inclui transações de pagamento)'),
 ('pessoal', ['d3', 'i'], ['o', 'y'], 'Pessoal + participação nos lucros (d3+i | o+y); PLR também em pes_plr'),
 ('admin', ['d4'], ['p'], 'Despesas administrativas (d4 | p)'),
 ('desp_trib', ['d5'], ['r'], 'Despesas tributárias (d5 | r)'),
 ('res_part', ['d6'], ['s'], 'Resultado de participações (d6 | s)'),
 ('out_rec', ['d7'], ['t'], 'Outras receitas operacionais (d7 | t)'),
 ('out_desp', ['d8'], ['u', 'q'], 'Outras despesas operacionais (d8 | u + q perdas esperadas de outras operações)'),
 ('nao_op', ['f'], [], 'Resultado não operacional (até 4T24: f | 2025+: dentro de outras receitas/despesas)'),
 ('ir_corr', ['h'], ['x'], 'IR e CSLL TOTAL (h | x)'),
 ('serv_tarifas', ['d2'], ['m'], 'Tarifas bancárias (d2 | m)'),
 ('serv_out', ['d1'], ['l', 'n'], 'Outros serviços (d1 | l+n)'),
 ('pes_plr', ['i'], ['y'], 'Participação nos lucros (i | y)'),
]
DS = f'{SRC}, rel. 4 (DRE)'
r = 151; inp.cell(153, 3).value = None
inp.cell(r, 3).value = 'DRE IF.data ACUMULADA NO SEMESTRE (mar = 1T, jun = 1S, set = 3T, dez = 2S). Sem código; as linhas com código desacumulam por fórmula.'; r += 1
RAWROW = {}
for code, o, n, desc in RAW:
    inp.cell(r, 3).value = f'[acum. semestre] {desc}'; inp.cell(r, 4).value = f'{DS} · memoria · moeda'
    for p in AV:
        v = lsum(p, 4, PRU, n if new(p) else o) if (n if new(p) else o) else None
        if v is not None: inp.cell(r, COL[p]).value = round(v, 6); written += 1
    RAWROW[code] = r; r += 1
inp.cell(r, 3).value = '[acum. semestre] Lucro líquido (j | z) - conferência'; inp.cell(r, 4).value = f'{DS} · memoria · moeda'
for p in AV: inp.cell(r, COL[p]).value = round(letter(p, 4, PRU, 'z' if new(p) else 'j'), 6)
r += 1
inp.cell(r, 3).value = '[acum. semestre] Ativo total (rel. 2) - conferência do balanço'; inp.cell(r, 4).value = f'{SRC} · memoria · moeda'
for p in AV: inp.cell(r, COL[p]).value = round(name(p, 2, PRU, r'^Ativo Total(?! Ajustado)'), 6)
r += 1
for code, o, n, desc in RAW:
    rr = ROW[code]; inp.cell(rr, 4).value = f'{DS}: {desc} · derivada (desacumula a linha {RAWROW[code]}) · moeda'
    for p in AV:
        c, q = CL(COL[p]), int(p[0])
        prev = P[P.index(p) - 1] if P.index(p) > 0 else None
        if q in (1, 3) or prev not in AV: inp.cell(rr, COL[p]).value = f'={c}{RAWROW[code]}' if q in (1, 3) else None
        else: inp.cell(rr, COL[p]).value = f'={c}{RAWROW[code]}-{CL(COL[p] - 1)}{RAWROW[code]}'
for c_, why in [('rec_rg', 'não se aplica'), ('desp_prov_tec', 'não se aplica'), ('pdd_recup', 'IF.data não separa (líquida dentro da PDD)'), ('seguros', 'não se aplica'),
                ('prov_civ', 'dentro de outras despesas'), ('prov_trab', 'idem'), ('prov_fisc', 'idem'), ('ir_dif', 'IR total em ir_corr'), ('ll_minor', 'IF.data não separa')]:
    empty(c_, why)
# ---------------- CAPITAL e CARTEIRA
for code, pat, u in [('rwa', r'^Ativos Ponderados pelo Risco', 'moeda'), ('rwa_cred', r'^RWA para Risco de Crédito', 'moeda'), ('rwa_mkt', r'^RWA para Risco de Mercado', 'moeda'),
                     ('rwa_op', r'^RWA para Risco Operacional', 'moeda'), ('pr', r'^Patrimônio de Referência para Comparação', 'moeda'), ('nivel1', r'^Patrimônio de Referência Nível I', 'moeda'),
                     ('cet1', r'^Capital Principal para Comparação', 'moeda'), ('ib_rep', r'^Índice de Basileia', 'pct'), ('in1_rep', r'^Índice de Capital Nível I', 'pct'),
                     ('icp_rep', r'^Índice de Capital Principal', 'pct'), ('alav', r'^Razão de Alavancagem', 'pct')]:
    vals = {p: name(p, 5, PRU, pat) for p in AV}
    if u == 'pct': vals = {p: (v * 1e6 if v is not None else None) for p, v in vals.items()}
    put(code, vals, f'{SRC}, rel. 5 (capital): "{pat[1:]}" · fonte · {u}')
def cart(p, rel, pat, grupo=None):
    v = name(p, rel, PRU, pat, grupo, tipo=1) if new(p) else None
    if v is None and FIN != '-': v = name(p, rel, FIN, pat, grupo, tipo=2)
    return v
CC = f'IF.data/BCB carteira ativa (até 4T24 conglomerado financeiro {FIN}; 2025+ prudencial {PRU})'
pf = {p: cart(p, 11, r'^Total da Carteira de Pessoa F') for p in AV}; pj = {p: cart(p, 13, r'^Total da Carteira de Pessoa J') for p in AV}
put('cart_ampl', {p: (pf[p] or 0) + (pj[p] or 0) if (pf[p] is not None or pj[p] is not None) else None for p in AV}, f'{CC}: PF + PJ (carteira classificada) · fonte · moeda')
put('cart_pf', pf, f'{CC} rel.11 total PF · fonte · moeda'); put('cart_pj', pj, f'{CC} rel.13 total PJ · fonte · moeda')
prods = {}
for code, grp in [('cart_cartao', 'Cartão de Crédito'), ('cart_consig', 'Empréstimo com Consignação'), ('cart_veic', 'Veículos'), ('cart_imob', 'Habitação'), ('cart_rural', 'Rural')]:
    prods[code] = {p: cart(p, 11, r'^Total$', grp) for p in AV}; put(code, prods[code], f'{CC} rel.11 PF grupo "{grp}" · fonte · moeda')
put('cart_pf_out', {p: pf[p] - sum(prods[k][p] or 0 for k in prods) if pf[p] is not None else None for p in AV}, f'{CC}: PF total menos as modalidades acima · derivada · moeda')
mpe = {p: sum(cart(p, 14, f'^{k}$') or 0 for k in ('Micro', 'Pequena', 'Média')) if pj[p] is not None else None for p in AV}
gr = {p: cart(p, 14, r'^Grande$') for p in AV}
put('cart_mpe', mpe, f'{CC} rel.14 micro+pequena+média · fonte · moeda'); put('cart_grandes', gr, f'{CC} rel.14 grande · fonte · moeda')
put('cart_pj_out', {p: pj[p] - (mpe[p] or 0) - (gr[p] or 0) if pj[p] is not None else None for p in AV}, f'{CC}: PJ total menos porte · derivada · moeda')
put('prov_tot', {p: -(A['cred_pdd'][0](p) or 0) for p in AV}, f'{SRC}, rel. 2: provisão / perda esperada de crédito (saldo positivo) · derivada · moeda')
put('npl90', {p: name(p, 16, PRU, r'^Inadimplência$') for p in AV if new(p)}, f'{SRC}, rel. 16 "Inadimplência" (só desde 1T25) · fonte · moeda')
for c_, why in [('est1', 'IF.data não abre estágios'), ('est2', 'idem'), ('est3', 'idem'), ('prov_e1', 'idem'), ('prov_e2', 'idem'), ('prov_e3', 'idem'), ('npl15', 'não disponível no IF.data'),
                ('wo', 'idem'), ('reneg', 'idem'), ('mf_cli', 'idem'), ('mf_mkt', 'idem'), ('at_rent', 'idem'), ('pas_oner', 'idem'), ('lcr', 'idem'), ('nsfr', 'idem'),
                ('acoes', 'preencher (fora do IF.data)'), ('div_decl', 'idem'), ('jcp_decl', 'idem'), ('preco', 'idem'), ('funcion', 'idem'), ('agencias', 'idem'), ('clientes', 'idem')]:
    empty(c_, why)
for c_ in ['serv_cartoes', 'serv_adm_rec', 'serv_seg_corr', 'serv_merc_cap', 'pes_proventos', 'pes_benef', 'pes_out', 'adm_proc', 'adm_alug', 'adm_mkt', 'adm_terc', 'adm_deprec', 'adm_out']:
    empty(c_, 'o IF.data não abre')
MEMSTART = r
inp['E197'] = NOME; inp['E198'] = TICK; inp['E199'] = SEG
inp['E203'] = datetime.datetime(2026, 6, 30)
f0 = sys.argv[8] if len(sys.argv) > 8 else AV[0]; inp['E204'] = datetime.datetime(2000 + int(f0[2:]), int(f0[0]) * 3, {3: 31, 6: 30, 9: 30, 12: 31}[int(f0[0]) * 3])
inp['E203'].number_format = inp['E204'].number_format = 'dd/mm/yyyy'
inp['F1'] = (f'{NOME} ({TICK}). Fonte única: IF.data do Banco Central (API Olinda), conglomerado prudencial {PRU}, BR GAAP/COSIF: balanço, DRE (publicada acumulada no semestre e '
             f'desacumulada por fórmula nesta aba), capital regulatório e carteira. R$ milhões. Primeiro trimestre com dado: {f0}. Modo histórico: premissas e valuation vazios.')
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('written', written, 'periods', AV[0], AV[-1], len(AV), 'next free row', MEMSTART)
