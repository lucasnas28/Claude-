import sys, shutil, datetime, re
import openpyxl
BASE, OUT = sys.argv[1:3]
SH = openpyxl.load_workbook(__import__('os').environ.get('PIPE_DATA', 'dados')+'/abcb/dl/sh.xlsx', data_only=True)
P = [f'{q}Q{y:02d}' for y in range(20, 27) for q in range(1, 5)][:26]
MON = {1: 'Mar', 2: 'Jun', 3: 'Sep', 4: 'Dec'}
def keys(p): return {p, f'{p[0]}T{p[2:]}', f'{MON[int(p[0])]}/{p[2:]}'}
HDR = {}
def hcols(ws):
    if ws.title in HDR: return HDR[ws.title]
    best = {}
    for r in range(5, 12):
        d = {}
        for c in range(4, ws.max_column + 1):
            v = ws.cell(r, c).value
            if v is None: continue
            s = str(v).strip()
            for p in P:
                if s in keys(p): d[p] = c
        if len(d) > len(best): best = d
    HDR[ws.title] = best; return best
def ser(sheet, *rows, k=1.0):
    ws = SH[sheet]; hc = hcols(ws); out = {}
    for p, c in hc.items():
        vs = [ws.cell(r, c).value for r in rows]
        if all(not isinstance(v, (int, float)) for v in vs): continue
        out[p] = k * sum(v for v in vs if isinstance(v, (int, float)))
    return out
def splice(old, new): d = dict(old); d.update(new); return d

shutil.copyfile(BASE, OUT)
wb = openpyxl.load_workbook(OUT); inp = wb['Inputs']
COL = {p: 5 + i for i, p in enumerate(P)}
ROW = {inp.cell(r, 2).value: r for r in range(5, 149) if inp.cell(r, 2).value}
written = 0
def put(code, vals, note):
    global written
    r = ROW[code]; inp.cell(r, 4).value = note
    for p, v in vals.items():
        if p in COL and v is not None: inp.cell(r, COL[p]).value = round(float(v), 6); written += 1
def empty(code, why): inp.cell(ROW[code], 4).value = f'[vazio] {why}'
BO, BN = 'Balanço -Balance Sheet', 'Balanço - Balance Sheet'
DO, DN = 'DRE Contábil (Tri) - ISQua', 'DRE Contábil (Tri) - IS Quarter'
S = 'Séries Históricas 2T26 (RI ABC)'
def bal(code, old_rows, new_rows, desc):
    put(code, splice(ser(BO, *old_rows) if old_rows else {}, ser(BN, *new_rows) if new_rows else {}),
        f'{S}: até 4T24 aba "{BO}" l.{"+".join(map(str, old_rows)) or "-"} | 2025+ aba "{BN}" l.{"+".join(map(str, new_rows)) or "-"} · {desc} · fonte · moeda')
# ---------------- BALANÇO (contábil; até 4T24 Res. 2.682, desde 1T25 Res. 4.966)
bal('disp', [16], [16], 'Disponibilidades')
bal('aplic_mo', [17], [18], 'Aplicações interfinanceiras TOTAL (a fonte não separa mercado aberto x DI)')
bal('tvm_prop', [18], [19, 24, 27], 'TVM (até 4T24: TVM & derivativos juntos | 2025+: TVM custo amortizado + VJORA + VJR)')
bal('deriv_a', [], [28], 'Derivativos ativos (só 2025+; até 4T24 dentro de tvm_prop)')
bal('tvm_pdd', [], [25, 29], 'Perda esperada de TVM (VJORA e VJR; só 2025+)')
bal('rel_a', [19], [], 'Relações interfinanceiras (até 4T24; 2025+ dentro de outros ativos financeiros)')
bal('cred_op', [20], [20], 'Operações de crédito')
bal('cred_pdd', [], [22], 'Perda esperada de ativos ao custo amortizado (só 2025+; até 4T24 o balanço traz crédito líquido)')
bal('out_a', [21, 22], [21, 33, 34], 'Outros créditos/ativos (até 4T24: outros créditos + outros valores e bens, inclui fiscais | 2025+: outros ativos financeiros + não financeiros mantidos p/ venda + outros)')
bal('fisc_a', [], [31, 32], 'Ativos fiscais diferidos + correntes (só 2025+; antes dentro de out_a)')
bal('inv_out', [25], [], 'Investimentos')
bal('imob', [26], [35], 'Imobilizado de uso (2025+: imobilizado + intangível, linha única)')
bal('intang', [27, 28], [], 'Intangível + diferido (até 4T24; 2025+ dentro de imob)')
dep = splice(ser(BO, 36), ser(BN, 43)); FU = 'Captação - Funding'
vista, inter = ser(FU, 13), ser(FU, 18)
put('dep_vista', vista, f'{S} aba "{FU}" l.13 Depósitos à vista · fonte · moeda')
put('dep_inter', inter, f'{S} aba "{FU}" l.18 Depósitos interfinanceiros · fonte · moeda')
put('dep_prazo', {p: dep[p] - vista.get(p, 0) - inter.get(p, 0) for p in dep}, f'{S}: Depósitos do balanço (até 4T24 l.36 | 2025+ l.43) menos à vista e interfinanceiros da aba Captação · derivada · moeda')
bal('mo_prop', [37], [44], 'Captações no mercado aberto TOTAL')
bal('div_emis', [38], [45], 'Recursos de aceites e emissão de títulos (LF, LCA, LCI, COE; até 4T24 inclui LF subordinadas)')
bal('div_sub', [], [46], 'Dívidas subordinadas (só 2025+; até 4T24 dentro de aceites/outras obrigações)')
bal('emp_empr', [41], [47], 'Obrigações por empréstimos (exterior)')
bal('emp_rep', [42, 43], [48], 'Obrigações por repasses (país + exterior)')
bal('deriv_p', [], [50], 'Derivativos passivos (só 2025+; até 4T24 dentro de outras obrigações)')
bal('rel_p', [39, 40], [], 'Relações interfinanceiras + interdependências (até 4T24)')
bal('prov_gar', [], [52], 'Perda esperada de garantias e compromissos (só 2025+)')
bal('prov_dem', [], [53], 'Provisão para contingências (só 2025+)')
bal('fisc_p', [], [56, 57], 'Obrigações fiscais correntes + diferidas (só 2025+)')
bal('out_p', [44, 46], [54], 'Outras obrigações (até 4T24: inclui derivativos, fiscais, provisões e sub; + resultados de exercícios futuros | 2025+: diversos)')
bal('cap_social', [49], [59, 60], 'Capital social (+ aumento de capital em aprovação)')
bal('res_cap', [50], [61], 'Reservas de capital')
bal('res_luc', [51, 53], [62, 64], 'Reservas de lucros + lucros acumulados (até 4T24 a reserva vem somada às ações em tesouraria)')
bal('ora', [52], [63], 'Ajustes de avaliação patrimonial / ORA')
bal('tesour', [], [65], 'Ações em tesouraria (só 2025+; até 4T24 somada à reserva de lucros)')
bal('minor', [54], [66], 'Participação de não controladores')
for c_, why in [('aplic_di', 'dentro de aplic_mo'), ('aplic_rg', 'não se aplica'), ('aplic_pdd', 'não divulgada'), ('tvm_vinc', 'dentro de tvm_prop'), ('tvm_rg', 'não se aplica'),
                ('cred_tvm', 'títulos privados ficam em TVM no balanço (abertura na carteira expandida)'), ('inv_col', 'dentro de inv_out'), ('dep_poup', 'ABC não capta poupança'),
                ('dep_out', 'não divulgado'), ('mo_terc', 'dentro de mo_prop'), ('mo_livre', 'dentro de mo_prop'), ('div_ext', 'dentro de emp_empr/div_emis'), ('div_coe', 'dentro de div_emis'),
                ('prov_tec', 'não se aplica')]:
    empty(c_, why)
# ---------------- DRE (trimestral isolada)
def dre(code, old_rows, new_rows, desc):
    put(code, splice(ser(DO, *old_rows) if old_rows else {}, ser(DN, *new_rows) if new_rows else {}),
        f'{S}: até 4T24 aba "{DO}" l.{"+".join(map(str, old_rows)) or "-"} | 2025+ aba "{DN}" l.{"+".join(map(str, new_rows)) or "-"} · {desc} · fonte · moeda')
dre('rec_cred', [14], [15], 'Operações de crédito')
dre('rec_tvm', [15, 16, 17, 23], [16, 17, 18, 23], 'TVM + derivativos + câmbio + (até 4T24: venda/transferência de ativos | 2025+: variações cambiais líquidas)')
dre('rec_aplic', [18], [], 'Aplicações compulsórias (até 4T24)')
dre('desp_dep', [20], [20], 'Operações de captação no mercado')
dre('desp_emp', [21], [21], 'Empréstimos e repasses')
dre('pdd_desp', [22], [22], 'Provisão para créditos de liquidação duvidosa / perda esperada')
dre('serv', [26], [26], 'Receitas de prestação de serviços')
dre('pessoal', [27, 35], [27, 35], 'Despesas de pessoal + participações nos lucros (PLR incluída aqui para o lucro fechar; também em pes_plr)')
dre('admin', [28], [28], 'Outras despesas administrativas')
dre('desp_trib', [29], [29], 'Despesas tributárias')
dre('out_rec', [30], [30], 'Outras receitas/despesas operacionais (LÍQUIDO: a fonte não separa receitas x despesas)')
dre('nao_op', [32], [32], 'Resultado não operacional')
dre('ir_corr', [34], [34], 'IR e CSLL TOTAL (fonte não separa corrente x diferido)')
dre('ll_minor', [36], [36], 'Participação de acionistas não controladores')
dre('pes_plr', [35], [35], 'Participações nos lucros')
for c_, why in [('rec_rg', 'não se aplica'), ('desp_div', 'dentro de desp_dep'), ('desp_prov_tec', 'não se aplica'), ('pdd_recup', 'recuperações em linha de memória (aba Provisões); na DRE já estão líquidas'),
                ('seguros', 'não se aplica'), ('prov_civ', 'dentro de outras receitas/despesas'), ('prov_trab', 'idem'), ('prov_fisc', 'idem'), ('ir_dif', 'IR total em ir_corr'),
                ('out_desp', 'líquido em out_rec')]:
    empty(c_, why)
# ---------------- COMPLEMENTARES
CE_O, CE_N = 'Carteira Expandida - Exp Por', 'Carteira Expandida - Exp Port'
put('cart_ampl', splice(ser(CE_O, 29), ser(CE_N, 29)), f'{S}: Carteira de crédito expandida (empréstimos + títulos privados + garantias prestadas) · fonte · moeda')
put('cart_pj', splice(ser(CE_O, 29), ser(CE_N, 29)), f'{S}: carteira expandida (ABC só atende pessoa jurídica) · fonte · moeda')
put('cart_grandes', splice(ser(CE_O, 30), ser(CE_N, 30)), f'{S}: segmento Large Corporate / C&IB (faturamento > R$4 bi) · fonte · moeda')
put('cart_pj_out', splice(ser(CE_O, 31), ser(CE_N, 31)), f'{S}: segmento Corporate (faturamento R$500 mi a R$4 bi; até 2024 R$300 mi a R$4 bi) · fonte · moeda')
put('cart_mpe', splice(ser(CE_O, 32), ser(CE_N, 32)), f'{S}: segmento Middle (faturamento R$30 a R$500 mi) · fonte · moeda')
ST = 'Nível de Risco - Stages'
for code, rr in [('est1', 14), ('est2', 15), ('est3', 16), ('prov_e1', 30), ('prov_e2', 31), ('prov_e3', 32)]:
    put(code, ser(ST, rr), f'{S} aba "{ST}" l.{rr} (só desde 1T25, Res. 4.966) · fonte · moeda')
IO, IN_ = 'Indicadores - Port Indicators', 'Indicadores - Indicators'
put('prov_tot', splice(ser(IO, 20), ser(IN_, 20)), f'{S}: saldo de provisão TOTAL (2025+ inclui provisão prospectiva; desde 1T21) · fonte · moeda')
put('npl90', splice(ser(IO, 47), ser(IN_, 39)), f'{S}: NPL 90 dias - operações (parcelas vencidas e vincendas; desde 1T21) · fonte · moeda')
put('wo', ser(IN_, 87), f'{S} aba "{IN_}" l.87 Write-off (desde 1T24) · fonte · moeda')
put('reneg', ser(IN_, 92), f'{S} aba "{IN_}" l.92 Saldo de créditos renegociados (desde 1T24) · fonte · moeda')
CAP = 'Adequação de Capital - Capital'
for code, rr, u in [('pr', 11, 'moeda'), ('ib_rep', 15, 'pct'), ('in1_rep', 16, 'pct'), ('icp_rep', 17, 'pct'), ('rwa', 21, 'moeda'), ('rwa_cred', 22, 'moeda'), ('rwa_mkt', 23, 'moeda'), ('rwa_op', 24, 'moeda')]:
    put(code, ser(CAP, rr), f'{S} aba "{CAP}" l.{rr} · fonte · {u}')
SU = 'Sumário - Summary'
put('nivel1', ser(SU, 57), f'{S} aba "{SU}" l.57 Capital de Nível 1 (desde 1T21) · fonte · moeda')
put('cet1', ser(SU, 58), f'{S} aba "{SU}" l.58 Capital principal (desde 1T21) · fonte · moeda')
put('acoes', ser(SU, 22, k=1e-6), f'{S} aba "{SU}" l.22 Número de ações ex-tesouraria (convertido para milhões; desde 1T21) · fonte · un')
put('jcp_decl', ser(SU, 30), f'{S} aba "{SU}" l.30 JCP declarado (desde 1T21) · fonte · moeda')
put('preco', ser(SU, 25), f'{S} aba "{SU}" l.25 Cotação ABCB4 (R$; desde 1T21) · fonte · un')
put('funcion', ser(SU, 68), f'{S} aba "{SU}" l.68 Colaboradores (desde 1T21) · fonte · un')
put('clientes', ser(SU, 67, k=1e-3), f'{S} aba "{SU}" l.67 Clientes corporativos (convertido para mil; desde 1T21) · fonte · un')
NIM = 'NIM'
put('at_rent', ser(NIM, 16), f'{S} aba NIM l.16 Ativos remuneráveis MÉDIOS (a fonte publica a média, não o saldo de fim de período) · fonte · moeda')
for c_, why in [('cart_pf', 'ABC não atua com pessoa física'), ('cart_cartao', 'idem'), ('cart_consig', 'idem'), ('cart_veic', 'idem'), ('cart_imob', 'idem'), ('cart_rural', 'idem'),
                ('cart_pf_out', 'idem'), ('npl15', 'não divulgado'), ('mf_cli', 'a Cia divulga só a margem gerencial total (memória)'), ('mf_mkt', 'idem'), ('pas_oner', 'não divulgado'),
                ('alav', 'não está na planilha'), ('lcr', 'idem'), ('nsfr', 'idem'), ('div_decl', 'ABC remunera via JCP'), ('agencias', 'não divulgado'),
                ('serv_tarifas', 'a planilha não abre serviços'), ('serv_cartoes', 'idem'), ('serv_adm_rec', 'idem'), ('serv_seg_corr', 'idem'), ('serv_merc_cap', 'idem'), ('serv_out', 'idem'),
                ('pes_proventos', 'aba Eficiência traz visão gerencial (memória)'), ('pes_benef', 'idem'), ('pes_out', 'idem'), ('adm_proc', 'não aberto'), ('adm_alug', 'idem'),
                ('adm_mkt', 'idem'), ('adm_terc', 'idem'), ('adm_deprec', 'idem'), ('adm_out', 'idem')]:
    empty(c_, why)
# ---------------- MEMÓRIAS
r = 151; inp.cell(153, 3).value = None
inp.cell(r, 3).value = 'MEMÓRIAS DO RI (sem código; não entram no modelo)'; r += 1
MEM = [(DO, DN, 37, 37, 'Lucro líquido contábil (DRE contábil l.37) - conferência'), (SU, SU, 14, 14, 'Lucro líquido recorrente (Sumário l.14)'),
       (NIM, NIM, 13, 13, 'Margem financeira gerencial pré-provisão (aba NIM l.13)'), (NIM, NIM, 12, 12, 'NIM (% a.a.; aba NIM l.12)'),
       ('Provisões - Provisions (2.682)', 'Provisões - Provisions', 20, 15, 'Recuperações de crédito (abas Provisões)'),
       (CE_O, CE_N, 14, 14, 'Carteira expandida - empréstimos'), (CE_O, CE_N, 24, 19, 'Carteira expandida - títulos privados'), (CE_O, CE_N, 19, 24, 'Carteira expandida - garantias prestadas'),
       (FU, FU, 14, 14, 'Depósitos a prazo (aba Captação l.14)'), (FU, FU, 15, 15, 'LCA, LCI e COE (aba Captação l.15)'), (FU, FU, 16, 16, 'Letras financeiras (aba Captação l.16)'),
       (FU, FU, 17, 17, 'Letras financeiras subordinadas (aba Captação l.17)'), (FU, FU, 22, 22, 'LF subordinadas perpétuas (aba Captação l.22)'), (FU, FU, 23, 23, 'Captação com terceiros total (aba Captação l.23)'),
       ('Eficiência - Efficiency', 'Eficiência - Efficiency', 12, 12, 'Índice de eficiência (aba Eficiência l.12)')]
for so, sn, ro, rn, desc in MEM:
    vals = splice(ser(so, ro), ser(sn, rn)) if so != sn else ser(so, ro)
    inp.cell(r, 3).value = desc; inp.cell(r, 4).value = f'{S} · memoria'
    for p, v in vals.items():
        if p in COL: inp.cell(r, COL[p]).value = v
    r += 1
inp.cell(r, 3).value = f'Espaço livre até a linha 190 (linhas {r} a 190) para novas linhas de origem; qualquer linha com código existente é somada pela Inputs (2).'
# ---------------- PAINEL
inp['E197'] = 'Banco ABC Brasil S.A.'; inp['E198'] = 'ABCB4'
inp['E199'] = 'Banco de atacado (corporate) / BR GAAP (COSIF); Res. 2.682 até 4T24, Res. 4.966 desde 1T25'
inp['E203'] = datetime.datetime(2026, 6, 30); inp['E204'] = datetime.datetime(2020, 3, 31)
inp['E203'].number_format = inp['E204'].number_format = 'dd/mm/yyyy'
inp['F1'] = ('ABC Brasil (ABCB4). Fonte: Planilha de Séries Históricas 2T26 do RI (ri.abcbrasil.com.br > Central de Resultados): balanço e DRE contábeis (abas antigas até 4T24, '
             'Res. 2.682; abas novas desde 1T25, Res. 4.966), carteira expandida, estágios, indicadores, captação, capital, sumário. R$ milhões. Modo histórico: premissas e valuation vazios.')
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('values written', written, 'last memo row', r - 1, {s: sorted(HDR[s])[:1] + sorted(HDR[s])[-1:] for s in HDR})
