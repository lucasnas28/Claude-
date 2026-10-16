import sys, shutil, datetime
import openpyxl
from openpyxl.utils import get_column_letter as CL
BASE, OUT = sys.argv[1:3]
SH = openpyxl.load_workbook(__import__('os').environ.get('PIPE_DATA', 'dados')+'/brbi/dl/sh.xlsx', data_only=True)
P = [f'{q}Q{y:02d}' for y in range(20, 27) for q in range(1, 5)][:26]
END = {1: (3, 31), 2: (6, 30), 3: (9, 30), 4: (12, 31)}
def pkey(v):
    if isinstance(v, datetime.datetime): return f'{(v.month) // 3}Q{v.year % 100:02d}' if v.month in (3, 6, 9, 12) else None
    s = str(v).strip()
    if len(s) == 4 and s[1] == 'T' and s[0] in '1234': return f'{s[0]}Q{s[2:]}'
    if len(s) == 10 and s[2] == '/' and s[5] == '/':
        m, y = int(s[3:5]), int(s[6:])
        return f'{m // 3}Q{y % 100:02d}' if m in (3, 6, 9, 12) else None
    return None
def ser(sheet, hr, rows, k=1.0):
    ws = SH[sheet]; out = {}
    hc = {pkey(ws.cell(hr, c).value): c for c in range(2, ws.max_column + 1) if ws.cell(hr, c).value is not None and pkey(ws.cell(hr, c).value)}
    for p in P:
        c = hc.get(p)
        if c is None: continue
        vs = [ws.cell(r, c).value for r in rows]
        num = [v for v in vs if isinstance(v, (int, float))]
        if not num and all(v in (None, '-', 'N/A') for v in vs) and not any(v == '-' for v in vs): continue
        out[p] = k * sum(num)
    return out
shutil.copyfile(BASE, OUT)
wb = openpyxl.load_workbook(OUT); inp = wb['Inputs']
COL = {p: 5 + i for i, p in enumerate(P)}
ROW = {inp.cell(r, 2).value: r for r in range(5, 149) if inp.cell(r, 2).value}
written = 0
def put(code, vals, note):
    global written
    r = ROW[code]; inp.cell(r, 4).value = note
    for p, v in vals.items():
        if p in COL and v is not None: inp.cell(r, COL[p]).value = round(float(v), 6); written += 1
def empty(code, why): inp.cell(ROW[code], 4).value = f'[vazio] {why}'
S = 'Séries Históricas 2T26 (RI BR Partners)'
BP, DR = 'BP (Trimestral)', 'DRE Contábil (Trimestral)'
def b(code, rows, desc): put(code, ser(BP, 9, rows, 1e-3), f'{S} aba "{BP}" l.{"+".join(map(str, rows))} · {desc} · fonte · moeda (R$ mil convertido para R$ milhões)')
def d(code, rows, desc): put(code, {p: v for p, v in ser(DR, 10, rows, 1e-3).items() if P.index(p) >= 3}, f'{S} aba "{DR}" l.{"+".join(map(str, rows))} · {desc} · fonte · moeda (R$ mil convertido para R$ milhões)')
# BALANÇO consolidado (formato IFRS da holding; desde dez/20)
b('disp', [11], 'Caixa e equivalentes de caixa')
b('tvm_prop', [12, 16], 'Ativos financeiros a VJR (títulos públicos, privados, fundos) + a VJORA')
b('deriv_a', [19], 'Instrumentos financeiros derivativos (ativo)')
b('cred_op', [21], 'Operações de crédito (custo amortizado)')
b('out_a', [22, 23], 'Outros ativos financeiros ao custo amortizado + outros ativos')
b('fisc_a', [24], 'Ativo fiscal diferido'); b('inv_col', [25], 'Investimentos em controladas')
b('imob', [26], 'Imobilizado'); b('intang', [27], 'Intangíveis')
b('mo_prop', [34], 'Recursos de operações compromissadas')
b('dep_prazo', [35], 'Recursos de clientes (depósitos; a fonte não abre à vista/a prazo)')
b('div_emis', [36], 'Recursos de emissão de títulos')
b('deriv_p', [32, 38], 'Passivos financeiros a VJR + instrumentos financeiros derivativos')
b('out_p', [37, 39, 40], 'Outros passivos financeiros + fornecedores + valores a pagar')
fp = ser(BP, 9, [41, 42, 43], 1e-3); fp['2Q21'] = fp['2Q21'] + 36.152
put('fisc_p', fp, f'{S} aba "{BP}" l.41+42+43 · Impostos a recolher + passivo fiscal corrente + diferido. 2T21: a planilha deixou o diferido (l.43) em branco; usado o ITR/CVM (conta 2.02.03, R$ 36.152 mil) · fonte · moeda (R$ mil convertido para R$ milhões)')
b('cap_social', [48], 'Capital social'); b('res_cap', [49], 'Reserva de capital'); b('res_luc', [50, 52], 'Reserva de lucros + lucros acumulados'); b('ora', [51], 'Outros resultados abrangentes')
for c_, why in [('aplic_mo', 'não aberto (dentro de caixa/TVM)'), ('aplic_di', 'idem'), ('aplic_rg', 'não se aplica'), ('aplic_pdd', 'não divulgada'), ('tvm_vinc', 'dentro de tvm_prop'),
                ('tvm_rg', 'não se aplica'), ('tvm_pdd', 'não divulgada'), ('cred_tvm', 'títulos privados em tvm_prop (carteira nas complementares)'), ('cred_pdd', 'não aberta no balanço'),
                ('rel_a', 'não se aplica'), ('inv_out', 'dentro de inv_col'), ('dep_vista', 'dentro de dep_prazo'), ('dep_poup', 'não capta'), ('dep_inter', 'dentro de dep_prazo'), ('dep_out', 'idem'),
                ('mo_terc', 'dentro de mo_prop'), ('mo_livre', 'idem'), ('div_ext', 'dentro de div_emis'), ('div_coe', 'idem'), ('div_sub', 'dentro de div_emis'), ('emp_empr', 'não divulgado'),
                ('emp_rep', 'idem'), ('rel_p', 'não se aplica'), ('prov_gar', 'dentro de out_p'), ('prov_tec', 'não se aplica'), ('prov_dem', 'dentro de out_p'),
                ('tesour', 'não aberta'), ('minor', 'não há')]:
    empty(c_, why)
# DRE contábil (trimestral; formato IFRS)
d('rec_tvm', [12], 'Receitas de juros e ganhos em instrumentos financeiros TOTAL (formato IFRS: não separa crédito x TVM x derivativos)')
d('desp_dep', [13], 'Despesas de juros e perdas em instrumentos financeiros TOTAL (captação + perdas com instrumentos)')
d('serv', [17, 18, 19], 'Receitas de prestação de serviços (líquidas de serviços técnicos especializados + outras receitas, como no total da Cia)')
d('pessoal', [25], 'Despesas de pessoal'); d('admin', [26], 'Despesas administrativas'); d('desp_trib', [27], 'Despesas tributárias')
d('pdd_desp', [28], 'Perda por redução ao valor recuperável'); d('out_rec', [29], 'Outras receitas/despesas (LÍQUIDO)'); d('res_part', [34], 'Resultado de equivalência patrimonial')
d('ir_corr', [37], 'Tributos sobre lucros TOTAL')
for c_, why in [('rec_cred', 'dentro de rec_tvm (formato IFRS)'), ('rec_rg', 'não se aplica'), ('rec_aplic', 'dentro de rec_tvm'), ('desp_div', 'dentro de desp_dep'), ('desp_emp', 'idem'),
                ('desp_prov_tec', 'não se aplica'), ('pdd_recup', 'não divulgada'), ('seguros', 'não se aplica'), ('prov_civ', 'dentro de outras'), ('prov_trab', 'idem'), ('prov_fisc', 'idem'),
                ('out_desp', 'líquido em out_rec'), ('nao_op', 'não há'), ('ir_dif', 'IR total em ir_corr'), ('ll_minor', 'não há')]:
    empty(c_, why)
# COMPLEMENTARES (R$ milhões)
CA = 'Adequação de Capital'
for code, rr, u in [('pr', 11, 'moeda'), ('nivel1', 12, 'moeda'), ('cet1', 13, 'moeda'), ('rwa', 16, 'moeda'), ('ib_rep', 17, 'pct'), ('icp_rep', 18, 'pct')]:
    put(code, ser(CA, 10, [rr]), f'{S} aba "{CA}" l.{rr} · fonte · {u}')
inp.cell(ROW['in1_rep'], 4).value = f'{S}: Nível I / RWA (aba "{CA}" l.12 / l.16) · derivada · pct'
for p in ser(CA, 10, [12]):
    c = CL(COL[p]); inp.cell(ROW['in1_rep'], COL[p]).value = f'=IFERROR({c}{ROW["nivel1"]}/{c}{ROW["rwa"]},"")'
put('cart_ampl', ser('Carteira Expandida', 10, [11]), f'{S} aba "Carteira Expandida" l.11: carteira de títulos privados (crédito via títulos; é a carteira de crédito da Cia) · fonte · moeda')
put('cart_pj', ser('Carteira Expandida', 10, [11]), f'{S}: carteira de títulos privados (100% PJ) · fonte · moeda')
r = 151; inp.cell(153, 3).value = None
inp.cell(r, 3).value = 'MEMÓRIAS DO RI (sem código)'; r += 1
GER = 'DRE Gerencial (Trimestral)'
MEM = [(DR, 10, 38, 1e-3, 'Lucro líquido contábil (DRE contábil l.38; R$ mi) - conferência; 1T20-3T20 só aqui (balanço começa em dez/20)'), (BP, 9, 29, 1e-3, 'Total do ativo divulgado (BP l.29; R$ mi)'),
       (BP, 9, 56, 1e-3, 'Total do passivo e PL divulgado (BP l.56; R$ mi) - a fonte tem diferença em alguns trimestres'),
       (GER, 11, 14, 1, 'Receita gerencial - Investment Banking + CM (R$ mi)'), (GER, 11, 15, 1, 'Receita gerencial - Treasury Sales & Structuring (R$ mi)'),
       (GER, 11, 16, 1, 'Receita gerencial - Gestão de patrimônio (R$ mi)'), (GER, 11, 17, 1, 'Receita gerencial - Remuneração do capital (R$ mi)'),
       ('Captação', 10, 11, 1, 'Volume de captação (R$ mi)'), ('Eficiência & Remuneração', 11, 12, 1, 'Índice de eficiência'), ('Eficiência & Remuneração', 11, 13, 1, 'Índice de remuneração')]
for sh, hr, rr, k, desc in MEM:
    inp.cell(r, 3).value = desc; inp.cell(r, 4).value = f'{S} · memoria'
    for p, v in ser(sh, hr, [rr], k).items():
        if p in COL: inp.cell(r, COL[p]).value = v
    r += 1
inp.cell(r, 3).value = f'Espaço livre até a linha 190 (linhas {r} a 190) para novas linhas de origem.'
for c_, why in [('cart_pf', 'não atua com PF'), ('cart_cartao', 'idem'), ('cart_consig', 'idem'), ('cart_veic', 'idem'), ('cart_imob', 'idem'), ('cart_rural', 'idem'), ('cart_pf_out', 'idem'),
                ('cart_mpe', 'não aberto'), ('cart_grandes', 'não aberto'), ('cart_pj_out', 'não aberto'), ('est1', 'não divulgado'), ('est2', 'idem'), ('est3', 'idem'), ('prov_e1', 'idem'),
                ('prov_e2', 'idem'), ('prov_e3', 'idem'), ('prov_tot', 'idem'), ('npl90', 'idem'), ('npl15', 'idem'), ('wo', 'idem'), ('reneg', 'idem'), ('mf_cli', 'idem'), ('mf_mkt', 'idem'),
                ('at_rent', 'idem'), ('pas_oner', 'idem'), ('rwa_cred', 'só RWA total na planilha'), ('rwa_mkt', 'idem'), ('rwa_op', 'idem'), ('alav', 'não está na planilha'), ('lcr', 'idem'),
                ('nsfr', 'idem'), ('acoes', 'não está na planilha (preencher)'), ('div_decl', 'só anual na planilha (aba Distribuição de Dividendos)'), ('jcp_decl', 'idem'), ('preco', 'preencher'),
                ('funcion', 'não está na planilha'), ('agencias', 'não se aplica'), ('clientes', 'não está na planilha')]:
    empty(c_, why)
for c_ in ['serv_tarifas', 'serv_cartoes', 'serv_adm_rec', 'serv_seg_corr', 'serv_merc_cap', 'serv_out', 'pes_proventos', 'pes_benef', 'pes_plr', 'pes_out', 'adm_proc', 'adm_alug', 'adm_mkt', 'adm_terc', 'adm_deprec', 'adm_out']:
    empty(c_, 'a planilha não abre (receita gerencial por área nas memórias)')
inp['E197'] = 'BR Partners (BR Advisory Partners Participações S.A.)'; inp['E198'] = 'BRBI11'
inp['E199'] = 'Banco de investimento / demonstrações consolidadas da holding em formato IFRS'
inp['E203'] = datetime.datetime(2026, 6, 30); inp['E204'] = datetime.datetime(2020, 12, 31)
inp['E203'].number_format = inp['E204'].number_format = 'dd/mm/yyyy'
inp['F1'] = ('BR Partners (BRBI11). Fonte: Séries Históricas 2T26 do RI (ri.brpartners.com.br): balanço consolidado (desde dez/20), DRE contábil (desde 1T20), capital, carteira de títulos privados, '
             'receita gerencial por área. Balanço e DRE em R$ mil na fonte, convertidos para R$ milhões. Primeiro trimestre realizado: 4T20 (primeiro balanço disponível).')
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('written', written, 'memo end', r)
