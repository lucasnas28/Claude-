import sys, shutil, datetime
sys.path.insert(0, __import__('os').environ.get('PIPE_DATA', 'dados')+'/bank')
from ifd import P, letter, lsum, name, new, rows, am
import openpyxl
from openpyxl.utils import get_column_letter as CL

BASE, OUT = sys.argv[1:3]
RI = openpyxl.load_workbook(__import__('os').environ.get('PIPE_DATA', 'dados')+'/bpac/dl/pf.xlsx', data_only=True)
BS, HL, IS, BA = RI['InputSite_BalanceSheet'], RI['InputSite_Highlights'], RI['InputSite_IncomeStatement'], RI['InputSite_BusinessAreas ']
HC = {str(BS.cell(3, c).value).strip(): c for c in range(3, BS.max_column + 1) if BS.cell(3, c).value is not None}
def ri(ws, rr, p):
    v = ws.cell(rr, HC[f'{p[0]}Q 20{p[2:]}']).value
    return float(v) if isinstance(v, (int, float)) else None
PRU, FIN = 'C0080336', 'C0049944'

shutil.copyfile(BASE, OUT)
wb = openpyxl.load_workbook(OUT)
inp = wb['Inputs']
COL = {p: 5 + i for i, p in enumerate(P)}
assert [inp.cell(2, COL[p]).value for p in ('1Q20', '2Q26')] == ['1T20', '2T26']
ROW = {inp.cell(r, 2).value: r for r in range(5, 149) if inp.cell(r, 2).value}
written = 0
def put(code, vals, note):
    global written
    r = ROW[code]
    inp.cell(r, 4).value = note
    for p, v in vals.items():
        if v is not None:
            inp.cell(r, COL[p]).value = round(float(v), 6); written += 1
def empty(code, why):
    inp.cell(ROW[code], 4).value = f'[vazio] {why}'

# ---------------- BALANÇO: planilha do RI (consolidado, como no release; R$ milhoes)
SRC = 'Planilha Séries Históricas 2T26 (RI), aba InputSite_BalanceSheet'
def bs(*rr): return {p: sum(ri(BS, x, p) or 0 for x in rr) for p in P}
put('disp', bs(8), f'{SRC} l.8 · fonte · moeda')
put('aplic_mo', bs(9), f'{SRC} l.9 Aplicações interfinanceiras de liquidez (TOTAL: a fonte não separa mercado aberto x depósitos interfinanceiros) · fonte · moeda')
put('tvm_prop', bs(10), f'{SRC} l.10 TVM e instrumentos financeiros derivativos (TOTAL: a fonte não separa carteira própria, vinculados nem derivativos) · fonte · moeda')
put('rel_a', bs(11), f'{SRC} l.11 Relações interfinanceiras · fonte · moeda')
put('cred_op', bs(12), f'{SRC} l.12 Operações de crédito (saldo do balanço consolidado; provisão não aberta no balanço da planilha) · fonte · moeda')
put('out_a', bs(13, 14), f'{SRC} l.13 Outros créditos + l.14 Outros valores e bens (inclui ativos fiscais, não abertos) · fonte · moeda')
put('inv_col', bs(15), f'{SRC} l.15 Ativo permanente TOTAL (investimentos + imobilizado + intangível: a fonte não abre) · fonte · moeda')
put('dep_prazo', bs(19), f'{SRC} l.19 Depósitos TOTAL (a planilha não abre à vista/interfinanceiros/a prazo; abertura gerencial do funding nas memórias) · fonte · moeda')
put('mo_prop', bs(20), f'{SRC} l.20 Captações no mercado aberto TOTAL · fonte · moeda')
put('div_emis', bs(21), f'{SRC} l.21 Recursos de aceites e emissão de títulos (LF, LCA, LCI, exterior, COE: não abertos) · fonte · moeda')
put('div_sub', bs(25), f'{SRC} l.25 Dívida subordinada · fonte · moeda')
put('emp_empr', bs(23), f'{SRC} l.23 Obrigações por empréstimos e repasses TOTAL · fonte · moeda')
put('deriv_p', bs(24), f'{SRC} l.24 Instrumentos financeiros derivativos · fonte · moeda')
put('rel_p', bs(22), f'{SRC} l.22 Relações interfinanceiras · fonte · moeda')
put('out_p', bs(26, 27), f'{SRC} l.26 Outras obrigações + l.27 Resultados de exercícios futuros (inclui provisões e obrigações fiscais, não abertas) · fonte · moeda')
put('res_luc', bs(28), f'{SRC} l.28 Patrimônio líquido dos controladores TOTAL (a planilha não abre capital/reservas/ORA/tesouraria): lançado aqui · fonte · moeda')
put('minor', bs(29), f'{SRC} l.29 Participação de não controladores · fonte · moeda')
for c_, why in [('aplic_di', 'dentro de aplic_mo (fonte só traz o total)'), ('aplic_rg', 'banco sem seguradora no balanço consolidado'), ('aplic_pdd', 'não divulgada'),
                ('tvm_vinc', 'dentro de tvm_prop (fonte só traz o total)'), ('tvm_rg', 'não se aplica'), ('tvm_pdd', 'não divulgada'), ('deriv_a', 'dentro de tvm_prop (fonte junta TVM e derivativos)'),
                ('cred_tvm', 'não aberto no balanço da planilha'), ('cred_pdd', 'balanço da planilha não abre a provisão; saldo do conglomerado prudencial em prov_tot'),
                ('fisc_a', 'dentro de out_a'), ('inv_out', 'dentro de inv_col'), ('imob', 'dentro de inv_col (permanente total)'), ('intang', 'dentro de inv_col (permanente total)'),
                ('dep_vista', 'dentro de dep_prazo (depósitos total)'), ('dep_poup', 'BTG não capta poupança'), ('dep_inter', 'dentro de dep_prazo'), ('dep_out', 'dentro de dep_prazo'),
                ('mo_terc', 'dentro de mo_prop'), ('mo_livre', 'dentro de mo_prop'), ('div_ext', 'dentro de div_emis'), ('div_coe', 'dentro de div_emis'), ('emp_rep', 'dentro de emp_empr'),
                ('prov_gar', 'dentro de out_p'), ('prov_tec', 'não se aplica'), ('prov_dem', 'dentro de out_p'), ('fisc_p', 'dentro de out_p'),
                ('cap_social', 'PL total em res_luc (fonte não abre)'), ('res_cap', 'idem'), ('ora', 'idem'), ('tesour', 'idem')]:
    empty(c_, why)

# ---------------- DRE: IF.data (BCB), conglomerado prudencial BTG Pactual (C0080336). Acumulado no SEMESTRE -> desacumulado aqui
IFD = 'IF.data/BCB, conglomerado prudencial C0080336, rel. 4 (DRE)'
def acc(old, new_):
    out = {}
    for p in P:
        lets = new_ if new(p) else old
        out[p] = lsum(p, 4, PRU, lets) if lets else None
        if out[p] is not None: out[p] = out[p]
    return out
RAW = [  # code, old letters (até 4T24), new letters (2025+), descricao
 ('rec_cred', ['a1', 'a2'], ['c', 'd', 'e'], 'Rendas de operações de crédito + arrendamento (até 4T24: a1+a2 | 2025+: c+d+e, inclui outras operações com característica de crédito)'),
 ('rec_tvm', ['a3', 'a4', 'a5', 'b4'], ['b', 'i', 'j'], 'Rendas de TVM + derivativos + câmbio (até 4T24: a3+a4+a5+b4, TVM inclui aplicações interfinanceiras | 2025+: b+i+j)'),
 ('rec_aplic', ['a6'], ['a'], 'Aplicações (até 4T24: a6 rendas de aplicações compulsórias; as interfinanceiras estão em TVM | 2025+: a rendas de aplicações interfinanceiras)'),
 ('desp_dep', ['b1'], ['g'], 'Despesas de captação (até 4T24: b1 | 2025+: g; inclui depósitos, compromissadas e emissões, não abertos)'),
 ('desp_div', [], ['h'], 'Despesas de instrumentos de dívida elegíveis a capital (só 2025+: h)'),
 ('desp_emp', ['b2', 'b3'], [], 'Despesas de empréstimos e repasses + arrendamento (até 4T24: b2+b3 | 2025+: dentro de captações)'),
 ('pdd_desp', ['b5'], ['f'], 'Provisão/perda esperada de crédito (até 4T24: b5 | 2025+: f)'),
 ('serv', ['d1', 'd2'], ['l', 'm', 'n'], 'Rendas de prestação de serviços + tarifas (até 4T24: d1+d2 | 2025+: l+m+n, inclui transações de pagamento)'),
 ('pessoal', ['d3', 'i'], ['o', 'y'], 'Despesas de pessoal + participação nos lucros (até 4T24: d3+i | 2025+: o+y). PLR incluída aqui para o lucro fechar; também em pes_plr'),
 ('admin', ['d4'], ['p'], 'Despesas administrativas (até 4T24: d4 | 2025+: p)'),
 ('desp_trib', ['d5'], ['r'], 'Despesas tributárias (d5 | r)'),
 ('res_part', ['d6'], ['s'], 'Resultado de participações (d6 | s)'),
 ('out_rec', ['d7'], ['t'], 'Outras receitas operacionais (d7 | t)'),
 ('out_desp', ['d8'], ['u', 'q'], 'Outras despesas operacionais (d8 | 2025+: u + q perdas esperadas de outras operações)'),
 ('nao_op', ['f'], [], 'Resultado não operacional (até 4T24: f | 2025+: dentro de outras receitas/despesas)'),
 ('ir_corr', ['h'], ['x'], 'IR e CSLL TOTAL (h | x; IF.data não separa corrente x diferido)'),
 ('serv_tarifas', ['d2'], ['m'], 'Rendas de tarifas bancárias (d2 | m)'),
 ('serv_out', ['d1'], ['l', 'n'], 'Outras rendas de serviços (d1 | l+n)'),
 ('pes_plr', ['i'], ['y'], 'Participação nos lucros (i | y)'),
]
r = 151
inp.cell(153, 3).value = None
inp.cell(r, 3).value = 'DRE IF.data ACUMULADA NO SEMESTRE (como publicada pelo BCB: mar = 1T, jun = 1S, set = 3T, dez = 2S). Sem código: só memória. As linhas com código acima desacumulam por fórmula.'
r += 1
RAWROW = {}
for code, o, n, desc in RAW:
    vals = acc(o, n)
    inp.cell(r, 3).value = f'[acum. semestre] {desc}'
    inp.cell(r, 4).value = f'{IFD} · memoria · moeda (R$ convertido para R$ milhões)'
    for p, v in vals.items():
        if v is not None: inp.cell(r, COL[p]).value = round(v, 6); written += 1
    RAWROW[code] = r; r += 1
ni = {p: (letter(p, 4, PRU, 'z') if new(p) else letter(p, 4, PRU, 'j')) for p in P}
inp.cell(r, 3).value = '[acum. semestre] Lucro líquido (j | z) - memória para conferência'
inp.cell(r, 4).value = f'{IFD} · memoria · moeda'
for p, v in ni.items(): inp.cell(r, COL[p]).value = round(v, 6)
NIROW = r; r += 1
for code, o, n, desc in RAW:
    rr = ROW[code]
    inp.cell(rr, 4).value = f'{IFD}: {desc} · derivada (desacumula a linha {RAWROW[code]}) · moeda'
    for p in P:
        c, q = CL(COL[p]), int(p[0])
        inp.cell(rr, COL[p]).value = f'={c}{RAWROW[code]}' if q in (1, 3) else f'={c}{RAWROW[code]}-{CL(COL[p] - 1)}{RAWROW[code]}'
for c_, why in [('rec_rg', 'não se aplica (sem seguradora)'), ('desp_prov_tec', 'não se aplica'), ('pdd_recup', 'IF.data não separa recuperações (líquidas dentro da PDD)'),
                ('seguros', 'não se aplica'), ('prov_civ', 'dentro de outras despesas operacionais'), ('prov_trab', 'idem'), ('prov_fisc', 'idem'),
                ('ir_dif', 'IR total em ir_corr (IF.data não separa)'), ('ll_minor', 'lucro do IF.data já é o atribuível (= lucro contábil do release)')]:
    empty(c_, why)

# ---------------- DADOS COMPLEMENTARES
C5 = 'IF.data/BCB rel. 5 (capital), conglomerado prudencial'
for code, pat, unit in [('rwa', r'^Ativos Ponderados pelo Risco', 'moeda'), ('rwa_cred', r'^RWA para Risco de Crédito', 'moeda'), ('rwa_mkt', r'^RWA para Risco de Mercado', 'moeda'),
                        ('rwa_op', r'^RWA para Risco Operacional', 'moeda'), ('pr', r'^Patrimônio de Referência para Comparação', 'moeda'), ('nivel1', r'^Patrimônio de Referência Nível I', 'moeda'),
                        ('cet1', r'^Capital Principal para Comparação', 'moeda'), ('ib_rep', r'^Índice de Basileia', 'pct'), ('in1_rep', r'^Índice de Capital Nível I', 'pct'),
                        ('icp_rep', r'^Índice de Capital Principal', 'pct'), ('alav', r'^Razão de Alavancagem', 'pct')]:
    vals = {p: name(p, 5, PRU, pat) for p in P}
    if unit == 'pct': vals = {p: (v * 1e6 if v is not None else None) for p, v in vals.items()}
    put(code, vals, f'{C5}: "{pat[1:]}" (localizado pelo nome; as letras mudam entre layouts) · fonte · {unit}')
def cart(p, rel, pat, grupo=None):
    t, c = (1, PRU) if new(p) else (2, FIN)
    return name(p, rel, c, pat, grupo, tipo=t)
pf = {p: cart(p, 11, r'^Total da Carteira de Pessoa F') for p in P}
pj = {p: cart(p, 13, r'^Total da Carteira de Pessoa J') for p in P}
CC = 'IF.data/BCB carteira ativa (rel. 11/13/14; até 4T24 conglomerado financeiro C0049944, 2025+ prudencial)'
put('cart_ampl', {p: (pf[p] or 0) + (pj[p] or 0) if pf[p] is not None else None for p in P}, f'{CC}: PF + PJ (carteira classificada; a carteira gerencial ampliada do RI está nas memórias) · fonte · moeda')
put('cart_pf', pf, f'{CC} rel.11 total PF · fonte · moeda'); put('cart_pj', pj, f'{CC} rel.13 total PJ (inclui exterior) · fonte · moeda')
prods = {}
for code, grp in [('cart_cartao', 'Cartão de Crédito'), ('cart_consig', 'Empréstimo com Consignação'), ('cart_veic', 'Veículos'), ('cart_imob', 'Habitação'), ('cart_rural', 'Rural')]:
    prods[code] = {p: cart(p, 11, r'^Total$', grp) for p in P}
    put(code, prods[code], f'{CC} rel.11 PF grupo "{grp}" · fonte · moeda')
put('cart_pf_out', {p: pf[p] - sum(prods[k][p] or 0 for k in prods) if pf[p] is not None else None for p in P}, f'{CC}: PF total menos as modalidades acima (sem consignação, outros, exterior) · derivada · moeda')
mpe = {p: sum(cart(p, 14, f'^{k}$') or 0 for k in ('Micro', 'Pequena', 'Média')) if pj[p] is not None else None for p in P}
gr = {p: cart(p, 14, r'^Grande$') for p in P}
put('cart_mpe', mpe, f'{CC} rel.14 porte micro+pequena+média · fonte · moeda'); put('cart_grandes', gr, f'{CC} rel.14 porte grande · fonte · moeda')
put('cart_pj_out', {p: pj[p] - (mpe[p] or 0) - (gr[p] or 0) if pj[p] is not None else None for p in P}, f'{CC}: PJ total menos porte (indisponível + exterior + não individualizado) · derivada · moeda')
def prov(p):
    if new(p): return -sum(letter(p, 2, PRU, l) or 0 for l in ('e2', 'f2', 'g2', 'h2'))
    return -(letter(p, 2, PRU, 'd2') or 0) - (letter(p, 2, PRU, 'e4') or 0)
put('prov_tot', {p: prov(p) for p in P}, 'IF.data/BCB rel. 2, conglomerado prudencial: provisão/perda esperada de crédito (até 4T24 d2+e4 | 2025+ e2+f2+g2+h2), saldo positivo · fonte · moeda')
put('npl90', {p: name(p, 16, PRU, r'^Inadimplência$') for p in P if new(p)}, 'IF.data/BCB rel. 16 "Inadimplência" (só desde 1T25, Res. 4.966) · fonte · moeda')
HLS = 'Planilha Séries Históricas 2T26 (RI), aba InputSite_Highlights'
put('acoes', {p: ri(HL, 27, p) for p in P}, f'{HLS} l.27 BPAC11 total de UNITS ex-tesouraria (milhões; 1 unit = 1 ON + 2 PN) · fonte · un')
put('div_decl', {p: ri(HL, 20, p) for p in P}, f'{HLS} l.20 Dividendos · fonte · moeda')
put('jcp_decl', {p: ri(HL, 21, p) for p in P}, f'{HLS} l.21 Juros sobre capital próprio · fonte · moeda')
put('funcion', {p: ri(HL, 41, p) for p in P}, f'{HLS} l.41 Funcionários · fonte · un')
for c_, why in [('est1', 'IF.data não abre estágios; BTG não divulga na planilha'), ('est2', 'idem'), ('est3', 'idem'), ('prov_e1', 'idem'), ('prov_e2', 'idem'), ('prov_e3', 'idem'),
                ('npl15', 'não divulgado'), ('wo', 'não divulgado'), ('reneg', 'não divulgado'), ('mf_cli', 'BTG não divulga margem com clientes x mercado (receita por área nas memórias)'), ('mf_mkt', 'idem'),
                ('at_rent', 'não divulgado'), ('pas_oner', 'não divulgado'), ('lcr', 'não está na planilha'), ('nsfr', 'idem'), ('preco', 'não está na fonte (preencher)'),
                ('agencias', 'não se aplica'), ('clientes', 'não divulgado na planilha'), ('serv_cartoes', 'IF.data só separa tarifas x demais serviços'), ('serv_adm_rec', 'idem'),
                ('serv_seg_corr', 'idem'), ('serv_merc_cap', 'idem'), ('pes_proventos', 'IF.data não abre'), ('pes_benef', 'idem'), ('pes_out', 'idem'),
                ('adm_proc', 'IF.data não abre'), ('adm_alug', 'idem'), ('adm_mkt', 'idem'), ('adm_terc', 'idem'), ('adm_deprec', 'idem'), ('adm_out', 'idem')]:
    empty(c_, why)

# ---------------- MEMÓRIAS DO RI (sem código)
inp.cell(r, 3).value = 'MEMÓRIAS DO RI (sem código; não entram no modelo): visão gerencial do BTG e perímetro consolidado'; r += 1
MEM = [(IS, 21, 'Lucro líquido contábil (RI, Resumo contábil l.21) - confere com o IF.data'), (HL, 17, 'Lucro líquido ajustado (RI Highlights l.17)'),
       (HL, 51, 'Carteira de crédito gerencial ampliada (RI Highlights l.51)'), (HL, 39, 'AuM e AuA (R$ bilhões; RI Highlights l.39)'), (HL, 40, 'WuM (R$ bilhões; RI Highlights l.40)')]
MEM += [(IS, rr, f'Receita gerencial - {str(IS.cell(rr, 1).value).strip()} (RI Income Statement l.{rr})') for rr in range(24, 34)]
for ws, rr, desc in MEM:
    inp.cell(r, 3).value = desc; inp.cell(r, 4).value = f'Planilha Séries Históricas 2T26 (RI) · memoria · moeda'
    for p in P:
        v = ri(ws, rr, p)
        if v is not None: inp.cell(r, COL[p]).value = v
    r += 1
assert r <= 191, r
inp.cell(r, 3).value = f'Espaço livre até a linha 190 (linhas {r} a 190) para novas linhas de origem; qualquer linha com código existente é somada pela Inputs (2).' if r <= 190 else None

# ---------------- PAINEL
inp['E197'] = 'Banco BTG Pactual S.A.'; inp['E198'] = 'BPAC11'
inp['E199'] = 'Banco de investimento / BR GAAP (COSIF). DRE: conglomerado prudencial (IF.data); balanço: consolidado (RI)'
inp['E203'] = datetime.datetime(2026, 6, 30); inp['E204'] = datetime.datetime(2020, 3, 31)
inp['E203'].number_format = inp['E204'].number_format = 'dd/mm/yyyy'
inp['F1'] = ('BTG Pactual (BPAC11). Fontes: (1) Planilha Séries Históricas 2T26 do RI (ri.btgpactual.com > Apresentações e Planilhas): balanço consolidado, units, proventos, funcionários e memórias; '
             '(2) IF.data do Banco Central (API Olinda), conglomerado prudencial C0080336: DRE detalhada (lucro = lucro contábil do release em todos os trimestres), capital regulatório e carteira. '
             'Tudo em R$ milhões (IF.data convertido de R$). Modo histórico: premissas e valuation vazios.')
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('values written', written, 'last memo row', r - 1)
