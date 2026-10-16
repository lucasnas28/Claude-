import sys, copy
sys.path.insert(0, __import__('os').environ.get('PIPE_DATA', 'dados')+'/aure')
from cv import V, q, bs, dfc_q, PP, DA_PAT, DA_PRE, DIV_PAT, DIV_PRE
import openpyxl
from openpyxl.utils import get_column_letter as CL
from openpyxl.formula.translate import Translator
from openpyxl.styles import Font

IN, OUT = sys.argv[1:3]
wb = openpyxl.load_workbook(IN)
inp, w2, m, va = wb['Inputs'], wb['Inputs (2)'], wb['Model'], wb['Valuation']
ov = openpyxl.load_workbook(IN, data_only=True)['Inputs']
bold = Font(bold=True)

# ---------------------------------------------------------------- períodos da Inputs (linha 2, D..AG)
PC = {str(inp.cell(2, c).value): c for c in range(4, 34)}
QC = {p: PC[p] for p in PP}
ANN = {y: PC[str(2000 + y)] for y in range(21, 27)}

# pro forma antigo (R$ mil, RI) guardado como memória antes de apagar
pf_rev = {p: ov.cell(4, QC[p]).value for p in PP if isinstance(ov.cell(4, QC[p]).value, (int, float)) and p < '1Q25' and p[2:] in ('23', '24')}
pf_rev = {p: ov.cell(4, QC[p]).value for p in PP if p[2:] in ('23', '24') and isinstance(ov.cell(4, QC[p]).value, (int, float))}
style_txt = copy.copy(inp['A50']._style); style_num = copy.copy(inp['N50']._style)

for r in range(4, 101):
    for c in range(1, 34):
        inp.cell(r, c).value = None

r = 4
def head(t):
    global r
    inp.cell(r, 1, t).font = bold; r += 1
def row(desc, code, conta, vals, kind='flow', note=None):
    """vals: {periodo: valor em R$ mil}. kind flow -> anual = soma; stock -> anual = último trimestre."""
    global r
    inp.cell(r, 1, desc)._style = style_txt
    inp.cell(r, 2, code); inp.cell(r, 3, conta)
    for p in PP:
        v = vals.get(p)
        cell = inp.cell(r, QC[p])
        cell._style = style_num
        if v is not None: cell.value = round(float(v), 3)
    for y, c in ANN.items():
        qs = [QC[f'{k}Q{y:02d}'] for k in range(1, 5) if f'{k}Q{y:02d}' in QC]
        if len(qs) < 4: continue
        inp.cell(r, c)._style = style_num
        inp.cell(r, c).value = f'=SUM({CL(qs[0])}{r}:{CL(qs[-1])}{r})' if kind == 'flow' else f'={CL(qs[-1])}{r}'
    r += 1

def dre(code_cvm, neg=False):
    out = {}
    for p in PP:
        v = q('DRE', code_cvm, p)
        if v is not None: out[p] = -v if neg else v
    return out
CONTA = {1: 'Net Revenues', 2: 'COGS (cash)', 3: 'SG&A and Other (cash)', 4: 'D&A', 5: 'Interest Expense', 6: 'Nonoperating Income / (Loss)',
         7: 'Taxes', 8: 'Equity in Income of Affiliates, after tax (enter as +)', 9: 'Minority Interest Expense, after tax (enter as - )', 10: 'Common Dividends',
         11: 'Cash and equivalents (inc. investment securities)', 12: 'Accounts receivable', 13: 'Inventories', 14: 'Deferred income taxes', 15: 'Prepaid expenses',
         16: 'Other current assets', 17: 'Equity in affiliates', 18: 'PP&E', 19: 'Right-of-use-assets', 20: 'Goodwill', 21: 'Intangibles', 22: 'Other non-current assets',
         23: 'Accounts payable', 24: 'Accrued expenses', 25: 'Taxes payable', 26: 'Other current liabilities (non-debt)', 27: 'Current portion of long-term debt',
         28: 'Short term debt (Revolving credit facility)', 29: 'Short term operating leases', 30: 'Long term debt', 31: 'Long term operating leases', 32: 'Convertible debt',
         33: 'Deferred income taxes', 34: 'Other non-current liabilities', 35: 'Minority interest', 36: 'Convertible Preferred Stock', 37: 'Common Stock and APIC',
         38: 'Treasury stock (contra account)', 39: 'Comprehensive (accumulated) loss', 40: 'Other equity account', 41: 'Retained earnings', 44: 'Net income reported (controladores) - memo p/ check'}

# ---------------------------------------------------------------- D&A: DVA até 2T22 (DFC pelo método direto não abre D&A), DFC depois
DA = {}
for p in PP:
    if p in ('1Q21', '2Q21', '3Q21', '4Q21', '1Q22', '2Q22'):
        v = q('DVA', '7.04.01', p); DA[p] = -v if v is not None else None
    else:
        DA[p] = dfc_q(p, DA_PAT, DA_PRE)
DIV = {p: dfc_q(p, DIV_PAT, DIV_PRE) for p in PP}
def ctrl(p, kind):
    d = V.get(('DRE', p, kind), {})
    a, t = d.get('3.11.01', (0, 0))[1], d.get('3.11', (0, None))[1]
    mi = d.get('3.11.02', (0, 0))[1]
    return t if (a == 0 and mi == 0) else a            # 1T24-3T24: ITR sem a abertura (sem minoritários no período)
NI = {}
for p in PP:
    if p[0] != '4' and ('DRE', p, '3M') in V: NI[p] = ctrl(p, '3M')
    else:
        prev = f'{int(p[0]) - 1}Q{p[2:]}'
        NI[p] = ctrl(p, 'YTD') - ctrl(prev, 'YTD' if ('DRE', prev, 'YTD') in V else '3M')
head('DRE - CVM ITR/DFP consolidado (Auren Energia S.A., CVM 026620), R$ mil, trimestre isolado (4T = ano - 9M)')
row('CVM DRE 3.01 Receita de venda de bens e/ou serviços', 1, CONTA[1], dre('3.01'))
row('CVM DRE 3.02 Custo dos bens e serviços vendidos (inclui D&A)', 2, CONTA[2], dre('3.02'))
row('(+) estorno da D&A embutida no custo, relançada no código 4 (a CVM não separa D&A entre custo e despesa; toda a D&A sai do custo)', 2, CONTA[2], {p: v for p, v in DA.items() if v is not None})
row('CVM DRE 3.04 Despesas/receitas operacionais (G&A + outras receitas e despesas)', 3, CONTA[3], dre('3.04'))
row('(-) estorno da equivalência patrimonial (3.04.06), lançada no código 8', 3, CONTA[3], dre('3.04.06', neg=True))
row('D&A - DVA 7.04.01 até 2T22; DFC 6.01.01 (depreciação e amortização + amortização de mais-valia/ágio) a partir de 3T22', 4, CONTA[4], {p: -v for p, v in DA.items() if v is not None})
row('CVM DRE 3.06.02 Despesas financeiras', 5, CONTA[5], dre('3.06.02'))
row('CVM DRE 3.06.01 Receitas financeiras', 6, CONTA[6], dre('3.06.01'))
row('CVM DRE 3.08 IR e CS sobre o lucro', 7, CONTA[7], dre('3.08'))
row('CVM DRE 3.04.06 Resultado de equivalência patrimonial', 8, CONTA[8], dre('3.04.06'))
row('CVM DRE 3.11.02 Lucro atribuído a não controladores (sinal invertido)', 9, CONTA[9], dre('3.11.02', neg=True))
row('CVM DFC 6.03 Pagamento de dividendos e JCP (caixa)', 10, CONTA[10], {p: v for p, v in DIV.items() if v is not None})
row('CVM DRE 3.11.01 Lucro atribuído aos controladores (1T24-3T24: 3.11, pois a ITR não separou e não havia minoritários)', 44, CONTA[44], NI)
head('DRE - memória (sem código)')
row('CVM DRE 3.03 Resultado bruto', None, None, dre('3.03'))
row('CVM DRE 3.05 Resultado antes do resultado financeiro (EBIT)', None, None, dre('3.05'))
row('CVM DRE 3.07 Resultado antes dos tributos', None, None, dre('3.07'))
row('CVM DRE 3.11 Lucro consolidado do período', None, None, dre('3.11'))
row('Receita líquida PRO FORMA (Auren + AES Brasil) - Planilha Interativa da RI; base antiga do modelo, só comparação', None, None, pf_rev)

head('BALANÇO - CVM ITR/DFP consolidado, R$ mil, saldo no fim do trimestre')
BSMAP = [('BPA', '1.01.01', 11), ('BPA', '1.01.02', 11), ('BPA', '1.01.03', 12), ('BPA', '1.01.04', 13), ('BPA', '1.01.05', 16), ('BPA', '1.01.06', 16),
         ('BPA', '1.01.07', 15), ('BPA', '1.01.08', 16), ('BPA', '1.02.01', 22), ('BPA', '1.02.02', 17), ('BPA', '1.02.03', 18), ('BPA', '1.02.04', 21),
         ('BPP', '2.01.01', 24), ('BPP', '2.01.02', 23), ('BPP', '2.01.03', 25), ('BPP', '2.01.04', 27), ('BPP', '2.01.05', 26), ('BPP', '2.01.06', 26), ('BPP', '2.01.07', 26),
         ('BPP', '2.02.01', 30), ('BPP', '2.02.02', 34), ('BPP', '2.02.03', 33), ('BPP', '2.02.04', 34), ('BPP', '2.02.05', 34), ('BPP', '2.02.06', 34),
         ('BPP', '2.03.01', 37), ('BPP', '2.03.02', 37), ('BPP', '2.03.03', 41), ('BPP', '2.03.04', 41), ('BPP', '2.03.05', 41), ('BPP', '2.03.06', 39),
         ('BPP', '2.03.07', 39), ('BPP', '2.03.08', 39), ('BPP', '2.03.09', 35)]
ESTORNO = {'1.02.01': ('1.02.01.07', 14, 'IR e CS diferidos'), '1.02.03': ('1.02.03.02', 19, 'direito de uso'), '2.01.04': ('2.01.04.03', 29, 'arrendamento CP'),
           '2.02.01': ('2.02.01.03', 31, 'arrendamento LP'), '2.03.02': ('2.03.02.05', 38, 'ações em tesouraria')}
def name(s, c):
    ds = [V[k][c][0] for k in V if k[0] == s and c in V[k]]
    return ds[-1] if ds else c
for s, c, code in BSMAP:
    vals = {p: bs(s, c, p) for p in PP}
    if all(abs(v) < 0.5 for v in vals.values()): continue
    row(f'CVM {s} {c} {name(s, c)}', code, CONTA[code], vals, 'stock')
    if c in ESTORNO:
        sub, code2, lab = ESTORNO[c]
        sv = {p: bs(s, sub, p) for p in PP}
        if all(abs(v) < 0.5 for v in sv.values()): continue
        row(f'(-) estorno {lab} ({sub}), lançado no código {code2}', code, CONTA[code], {p: -v for p, v in sv.items()}, 'stock')
        row(f'CVM {s} {sub} {name(s, sub)}', code2, CONTA[code2], sv, 'stock')
head('BALANÇO - totais divulgados (memória)')
row('CVM BPA 1 Ativo total', None, None, {p: V[('BPA', p, 'I')]['1'][1] for p in PP}, 'stock')
row('CVM BPP 2.01 + 2.02 Passivo circulante + não circulante (reconstruído pelas contas)', None, None,
    {p: sum(bs('BPP', c, p) for _, c, k in BSMAP if c[:4] in ('2.01', '2.02')) for p in PP}, 'stock')
row('CVM BPP 2.03 Patrimônio líquido consolidado', None, None, {p: V[('BPP', p, 'I')]['2.03'][1] for p in PP}, 'stock')
LAST = r - 1
r += 1
for t in ['Notas:',
          '1) Fonte única: ITR/DFP consolidados da CVM (dados abertos), R$ mil; a Inputs (2) divide por 1.000 (R$ milhões). Trimestres de fluxo: 3M quando a ITR traz, senão acumulado do ano menos o trimestre anterior.',
          '2) 2023-2024 são o REPORTADO (sem AES Brasil até out/24). O pro forma da RI (Auren + AES) ficou só como memória acima e na linha 41 do Model (base do bloco de receita).',
          '3) 2T22: a ITR veio sem as contas-mãe 2.01.05, 2.01.06 e 2.02.04 e com 2.02.02 incompleta; os saldos foram reconstruídos pela soma das contas filhas (o total bate com o passivo total divulgado).',
          '4) Caixa (código 11) = caixa e equivalentes + aplicações financeiras. Fundo de liquidez/conta reserva de longo prazo fica em outros ativos (restrito).',
          '5) Fluxo de caixa de 1T21 sem saldo inicial (dez/20 fora da grade): as variações do 1T21 no Model não são interpretáveis.']:
    inp.cell(r, 1, t); r += 1
assert r <= 100, r
inp['A1'] = 'Base histórica (R$ mil) - AUREN ENERGIA S.A. (AURE3) - CVM ITR/DFP consolidado'
inp['A3'] = 'Fonte: CVM dados abertos (ITR/DFP consolidado, CVM 026620) - R$ mil, trimestre isolado'
inp['C3'] = 'Unidade original: R$ mil | a Inputs (2) divide por 1.000 -> R$ milhões'

# ---------------------------------------------------------------- Inputs (2): flags e linha do lucro reportado
for c in range(4, 34):
    h = str(w2.cell(2, c).value)
    if h in PP: w2.cell(3, c).value = 1
    elif h in ('2021', '2022', '2023', '2024', '2025'): w2.cell(3, c).value = 1
    else: w2.cell(3, c).value = 0
w2['B48'] = 44; w2['C48'] = 'Net income reported (controladores) - memo p/ check IS [Model linha 4]'
for c in range(1, 34):
    w2.cell(48, c)._style = copy.copy(w2.cell(13, c)._style)
    src = w2.cell(13, c).value
    if isinstance(src, str) and src.startswith('='):
        w2.cell(48, c).value = Translator(src, origin=f'{CL(c)}13').translate_formula(f'{CL(c)}48')

# ---------------------------------------------------------------- Model
QCOL = [c for c in range(9, 36) if c not in (13, 18, 23, 28, 33)]      # I..AI trimestres realizados
ACOL = [13, 18, 23, 28, 33]                                            # M R W AB AG (anuais 2021-2025)
REAL = sorted(QCOL + ACOL)
PROJ = list(range(36, 108))
def tr(f, src, dst): return Translator(f, origin=src).translate_formula(dst)
# 1) linha 4 = lucro reportado; linha 3 = check IS em todas as colunas realizadas; projeção sem check (não há reportado)
for c in REAL:
    m.cell(4, c).value = f"='Inputs (2)'!{CL(c - 5)}48"
    m.cell(3, c).value = f'=IF(COUNT({CL(c)}4,{CL(c)}136)<2,"",ABS({CL(c)}136-{CL(c)}4))'
    m.cell(4, c).number_format = m.cell(15, c).number_format
for c in PROJ:
    if m.cell(3, c).value is not None: m.cell(3, c).value = None
m['C4'] = 'Net Income reportado (controladores, CVM) - base do check IS'
# 2) fórmulas que faltavam nas colunas 4T25, 1T26, 2T26 (e similares): copiadas da coluna 1T25 (AC)
FILL = [16, 54, 58, 70, 74, 131, 145, 149, 162, 198]
for rr in FILL:
    base = m[f'AC{rr}'].value
    for c in QCOL:
        if c < 19: continue
        if m.cell(rr, c).value is None:
            m.cell(rr, c).value = tr(base, f'AC{rr}', f'{CL(c)}{rr}')
            m.cell(rr, c)._style = copy.copy(m[f'AC{rr}']._style)
# 249: %YoY da despesa com minoritários apontava para a linha da equivalência (134)
for c in QCOL + ACOL:
    p = c - 5 if c - 5 >= 9 else None
    m.cell(249, c).value = f'=IF(N({CL(p)}135)=0,"",{CL(c)}135/{CL(p)}135-1)' if p else None
# 3) imobilizado e intangível: colunas realizadas = saldo do balanço; capex e compra de intangível = variação + D&A
for c in REAL:
    L, P_ = CL(c), CL(c - 1)
    m.cell(170, c).value = f'={L}165'
    m.cell(183, c).value = f'={L}179'
    m.cell(182, c).value = 0
    if c in ACOL:
        a, b = CL(c - 4), CL(c - 1)
        m.cell(167, c).value = f'=SUM({a}167:{b}167)'; m.cell(181, c).value = f'=SUM({a}181:{b}181)'
        m.cell(297, c).value = f'=SUM({a}297:{b}297)'
    else:
        m.cell(167, c).value = f'={L}165-{P_}165-{L}69'          # D&A total (inclui amortização de intangível/mais-valia)
        m.cell(181, c).value = f'={L}179-{P_}179'
        m.cell(297, c).value = f'=-{L}181'
    for rr in (167, 181, 182, 183, 297): m.cell(rr, c)._style = copy.copy(m[f'AC{rr if rr != 297 else 294}']._style)
for c in PROJ:
    if m.cell(181, c).value is None and m.cell(170, c).value is not None:
        m.cell(181, c).value = 0; m.cell(297, c).value = f'=-{CL(c)}181'
# 4) dívida total (linha 91) para o custo da dívida da Control e para os juros projetados
for c in REAL + PROJ:
    if m.cell(15, c).value is None and c in PROJ: continue
    L = CL(c)
    m.cell(87, c).value = f'={L}211'; m.cell(88, c).value = f'={L}212'; m.cell(89, c).value = f'={L}216'; m.cell(90, c).value = f'={L}237'
    m.cell(91, c).value = f'=SUM({L}87:{L}90)'
    for rr in range(87, 92): m.cell(rr, c).number_format = '#,##0;(#,##0)'
# 5) equivalência e minoritários projetados: média dos últimos 4 trimestres realizados (antes: zero)
LAST4 = 'AE{0},AF{0},AH{0},AI{0}'
for rr in (134, 135):
    for c in PROJ:
        if m.cell(rr, c).value != '=0': continue
        if c % 5 == 3:                                   # coluna anual
            m.cell(rr, c).value = f'=SUM({CL(c - 4)}{rr}:{CL(c - 1)}{rr})'
        elif c == 36:
            m.cell(rr, c).value = f'=AVERAGE({LAST4.format(rr)})'
        else:
            prev = c - 1 if (c - 1) % 5 != 3 else c - 2
            m.cell(rr, c).value = f'={CL(prev)}{rr}'
# 6) bloco de receita: até 4T24 a base do build é pro forma; linha 41 guarda o pro forma e o check 17 só vale de 1T25 em diante
for p, v in pf_rev.items():
    c = QC[p] + 5
    m.cell(41, c).value = round(v / 1000, 3)
for c in (23, 28):
    m.cell(41, c).value = f'=SUM({CL(c - 4)}41:{CL(c - 1)}41)'
for c in range(9, 29):
    m.cell(17, c).value = None
m['C41'] = 'Receita líquida da base do build (pro forma Auren + CESP/VGE em 2021-1T22 e Auren + AES em 2023-4T24, releases/RI; reportada a partir de 1T25)'
m['C17'] = 'Check Receita (só de 1T25 em diante; até 4T24 o build é pro forma e a receita reportada é a CVM)'
# 6b) crescimento de 2025 contra 2024 pro forma (com AES), base comparável; é o que alimenta o driver de receita da Control
for c in range(29, 34):
    m.cell(16, c).value = f'=IFERROR({CL(c)}15/{CL(c - 5)}41-1,"")'
m['C16'] = '% YoY (2025: contra 2024 pro forma com AES, linha 41)'
# 7) linhas de driver ocultas que alimentam a Control: mostrar
for rr in (54, 58, 70, 74, 131, 145, 149, 198):
    m.row_dimensions[rr].hidden = False
# ---------------------------------------------------------------- Valuation: alíquota limitada a 0-100% (linha 348 do Model) para o NOPLAT
for c in range(4, 12):
    f = va.cell(24, c).value
    if isinstance(f, str) and 'Model!' in f and '131' in f:
        va.cell(24, c).value = f.replace('131', '348')
va['C24'] = 'Tax rate (ETR limitada a 0%-100%)'
va['D65'] = '=IF(D63="","",K25*(1+D63))'
va['D88'] = '=IF(OR(D86="",D87=""),"",D86*(1-D87))'
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('Inputs rows 4 ->', LAST)
