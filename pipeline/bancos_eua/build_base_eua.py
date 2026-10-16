"""Gera a base de bancos dos EUA (MODELO_IF_BANCOS_EUA) a partir da base brasileira MODELO_IF_BANCOS_BASE_v7.

Mantém a mecânica da v7 (Inputs -> Inputs (2) -> Model -> Valuation, checks, plug de funding, roll-forward do PL,
glossário e os quatro métodos de valuation) e troca o conteúdo:
  - plano de contas Cosif -> US GAAP (rótulos e códigos)
  - linhas novas: ações preferenciais, dividendo preferencial, lucro ao acionista comum, recompras,
    PL ordinário, MSR, TCE, segmentos de negócio, CET1, SLR, Tier 1 leverage, LCR, ROE ordinário, ROTCE e TBVPS
  - qualidade de crédito: CECL (nonaccrual, 90+ ainda acumulando, 30-89, net charge-offs) no lugar dos estágios da 4.966
  - macro: Fed funds, SOFR, Treasury 10 anos, CPI e PIB real (FRED) no lugar de Selic, CDI, IPCA e TLP
  - valuation: retorno total ao acionista (dividendos + recompras) no DDM, PL ordinário no lucro residual, P/TBV

uso: build_base_eua.py BASE_v7.xlsx SAIDA.xlsx PASTA_FRED
"""
import sys, os, csv, shutil, subprocess, tempfile, datetime as dt, statistics, copy
import openpyxl
from openpyxl.utils import get_column_letter as L

BASE, OUT, FRED = sys.argv[1:4]
HERE = os.path.dirname(os.path.abspath(__file__))
SHIFT = os.path.join(HERE, '..', 'comum', 'shift_rows.py')
tmp = tempfile.mkdtemp()
STEP = os.path.join(tmp, 'step0.xlsx')
shutil.copyfile(BASE, STEP)

# ------------------------------------------------------------------ inserções (coordenadas da v7; executadas de baixo para cima)
INS = {
    'Inputs (2)': [(146, 4), (107, 2), (65, 1)],
    'Inputs': [(302, 1), (299, 1), (248, 1)],
    'Model': [(202, 1), (194, 1), (163, 2), (160, 6), (142, 12), (137, 5), (128, 2), (85, 1), (77, 1)],
}
def fin(sheet, r):            # linha final de uma linha original
    return r + sum(n for at, n in INS[sheet] if r >= at)
def blk(sheet, at):           # primeira linha final do bloco inserido em 'at'
    return at + sum(n for a, n in INS[sheet] if a < at)

# ------------------------------------------------------------------ 1) rótulos e códigos (coordenadas originais)
wb = openpyxl.load_workbook(STEP)
inp, w2, m, val, glo = wb['Inputs'], wb['Inputs (2)'], wb['Model'], wb['Valuation'], wb['Glossário']

# código antigo -> (código novo, rótulo PT, rótulo EN)
CODES = {
 'disp': ('cash', 'Caixa e depósitos não remunerados em bancos', 'Cash and due from banks'),
 'aplic_mo': ('resale', 'Fed funds vendidos e compromissadas reversas', 'Fed funds sold and securities purchased under resale agreements'),
 'aplic_di': ('dep_banks', 'Depósitos remunerados em bancos (inclui Federal Reserve)', 'Deposits with banks'),
 'aplic_rg': ('sec_borrowed', 'Títulos tomados em empréstimo (securities borrowed)', 'Securities borrowed'),
 'aplic_pdd': ('liq_acl', '(Provisão para perdas — aplicações de liquidez)', '(Allowance — liquidity placements)'),
 'tvm_prop': ('trading', 'Ativos de trading (inclui derivativos a receber)', 'Trading assets (incl. derivative receivables)'),
 'tvm_vinc': ('afs', 'Títulos disponíveis para venda (AFS, valor justo)', 'Available-for-sale securities, at fair value'),
 'tvm_rg': ('htm', 'Títulos mantidos até o vencimento (HTM, custo amortizado)', 'Held-to-maturity securities'),
 'tvm_pdd': ('sec_acl', '(Provisão para perdas — títulos, quando à parte)', '(Allowance — securities)'),
 'deriv_a': ('deriv_a', 'Derivativos a receber (quando fora de trading)', 'Derivative receivables, if separate'),
 'cred_op': ('loans', 'Empréstimos, líquidos de receitas diferidas', 'Loans, net of unearned income'),
 'cred_tvm': ('loans_fv', 'Empréstimos a valor justo / mantidos para venda (quando à parte)', 'Loans at fair value / held for sale, if separate'),
 'cred_pdd': ('acl', '(Provisão para perdas de crédito — ACL, CECL)', '(Allowance for loan losses)'),
 'rel_a': ('receivables', 'Juros e contas a receber', 'Accrued interest and accounts receivable'),
 'fisc_a': ('dta', 'Impostos diferidos ativos (quando à parte)', 'Deferred tax assets, if separate'),
 'out_a': ('other_a', 'Outros ativos', 'Other assets'),
 'inv_col': ('eq_inv', 'Investimentos por equivalência patrimonial (quando à parte)', 'Equity-method investments, if separate'),
 'inv_out': ('inv_out', 'Outros investimentos (quando à parte)', 'Other investments, if separate'),
 'imob': ('premises', 'Imóveis, equipamentos e direito de uso', 'Premises and equipment'),
 'intang': ('gw_intang', 'Goodwill, MSRs e outros intangíveis', 'Goodwill, MSRs and other intangibles'),
 'dep_vista': ('dep_nib', 'Depósitos sem juros (noninterest-bearing)', 'Noninterest-bearing deposits'),
 'dep_poup': ('dep_ib', 'Depósitos com juros (interest-bearing)', 'Interest-bearing deposits'),
 'dep_inter': ('dep_x1', 'Depósitos — abertura adicional (quando a fonte separa)', 'Deposits — additional split'),
 'dep_prazo': ('dep_time', 'Depósitos a prazo (time deposits, quando à parte)', 'Time deposits, if separate'),
 'dep_out': ('dep_other', 'Outros depósitos (quando à parte)', 'Other deposits, if separate'),
 'mo_prop': ('repo', 'Fed funds comprados e compromissadas (repos)', 'Fed funds purchased and securities sold under repurchase agreements'),
 'mo_terc': ('st_borr', 'Captações de curto prazo (short-term borrowings)', 'Short-term borrowings'),
 'mo_livre': ('sec_loaned', 'Títulos emprestados (securities loaned, quando à parte)', 'Securities loaned, if separate'),
 'div_emis': ('ltd', 'Dívida de longo prazo (inclui subordinada)', 'Long-term debt'),
 'div_ext': ('vie', 'Participações emitidas por VIEs consolidadas', 'Beneficial interests issued by consolidated VIEs'),
 'div_coe': ('ltd_x', 'Outras emissões (quando à parte)', 'Other issuances, if separate'),
 'div_sub': ('sub_debt', 'Dívida subordinada (quando à parte)', 'Subordinated debt, if separate'),
 'emp_empr': ('fhlb', 'Adiantamentos do FHLB e outros empréstimos (quando à parte)', 'FHLB advances and other borrowings'),
 'emp_rep': ('borr_x', 'Outros empréstimos tomados (quando à parte)', 'Other borrowings, if separate'),
 'deriv_p': ('trading_l', 'Passivos de trading (inclui derivativos a pagar)', 'Trading liabilities (incl. derivative payables)'),
 'rel_p': ('brokers_l', 'Obrigações com corretoras e clientes (quando à parte)', 'Payables to brokers and customers, if separate'),
 'prov_gar': ('lrc_liab', 'Provisão para compromissos de crédito (quando à parte)', 'Allowance for lending-related commitments, if separate'),
 'prov_tec': ('ins_res', 'Reservas de seguros (quando aplicável)', 'Insurance reserves, if applicable'),
 'prov_dem': ('prov_other', 'Demais provisões (quando à parte)', 'Other provisions, if separate'),
 'fisc_p': ('dtl', 'Impostos a pagar e diferidos (quando à parte)', 'Taxes payable and deferred, if separate'),
 'out_p': ('other_l', 'Contas a pagar e outros passivos', 'Accounts payable and other liabilities'),
 'cap_social': ('common_stk', 'Ações ordinárias (valor nominal)', 'Common stock'),
 'res_cap': ('apic', 'Capital adicional integralizado (APIC) e outros', 'Additional paid-in capital and other'),
 'res_luc': ('ret_earn', 'Lucros retidos', 'Retained earnings'),
 'ora': ('aoci', 'Outros resultados abrangentes acumulados (AOCI)', 'Accumulated other comprehensive income'),
 'tesour': ('treasury', '(Ações em tesouraria, a custo)', '(Treasury stock, at cost)'),
 'minor': ('nci', 'Participação de não controladores', 'Noncontrolling interests'),
 'rec_cred': ('ii_loans', 'Juros de empréstimos', 'Interest income — loans'),
 'rec_tvm': ('ii_sec', 'Juros de títulos e ativos de trading', 'Interest income — securities and trading assets'),
 'rec_rg': ('ii_repo', 'Juros de compromissadas reversas e securities borrowed', 'Interest income — resale agreements and securities borrowed'),
 'rec_aplic': ('ii_other', 'Juros de depósitos em bancos e outros', 'Interest income — deposits with banks and other'),
 'desp_dep': ('ie_dep', '(Juros de depósitos)', '(Interest expense — deposits)'),
 'desp_div': ('ie_ltd', '(Juros de dívida de longo prazo e VIEs)', '(Interest expense — long-term debt and VIEs)'),
 'desp_emp': ('ie_st', '(Juros de compromissadas, curto prazo, trading e outros)', '(Interest expense — repos, short-term, trading and other)'),
 'desp_prov_tec': ('ie_other', '(Outras despesas de juros, quando à parte)', '(Other interest expense, if separate)'),
 'pdd_desp': ('provision', '(Provisão para perdas de crédito)', '(Provision for credit losses)'),
 'pdd_recup': ('recoveries', 'Recuperações (quando à parte; nos EUA a provisão já é líquida)', 'Recoveries, if separate'),
 'serv': ('fees', 'Receitas de tarifas e comissões', 'Fees and commissions'),
 'seguros': ('principal', 'Resultado de trading (principal transactions)', 'Principal transactions'),
 'pessoal': ('comp', '(Remuneração e benefícios)', '(Compensation expense)'),
 'admin': ('opex', '(Ocupação, tecnologia, serviços profissionais e marketing)', '(Occupancy, technology, professional services and marketing)'),
 'prov_civ': ('legal', '(Despesas jurídicas, quando à parte)', '(Legal expense, if separate)'),
 'prov_trab': ('prov_x1', '(Outras provisões, quando à parte)', '(Other provisions, if separate)'),
 'prov_fisc': ('prov_x2', '(Outras provisões, quando à parte)', '(Other provisions, if separate)'),
 'desp_trib': ('fdic', '(FDIC e contribuições regulatórias, quando à parte)', '(FDIC and regulatory assessments, if separate)'),
 'res_part': ('eq_income', 'Resultado de equivalência patrimonial (quando à parte)', 'Equity-method income, if separate'),
 'out_rec': ('other_inc', 'Outras receitas não-juros (inclui ganhos com títulos de investimento)', 'Other noninterest revenue (incl. investment securities gains)'),
 'out_desp': ('other_exp', '(Outras despesas não-juros)', '(Other noninterest expense)'),
 'nao_op': ('nonop', 'Resultado não operacional (raro nos EUA)', 'Non-operating result'),
 'ir_corr': ('tax', '(Imposto de renda)', '(Income tax expense)'),
 'ir_dif': ('tax_def', 'Imposto diferido (quando à parte)', 'Deferred tax, if separate'),
 'll_minor': ('nci_ni', '(Lucro atribuído a não controladores)', '(Net income to noncontrolling interests)'),
 'cart_ampl': ('loans_tot', 'Empréstimos totais', 'Total loans'),
 'cart_pf': ('loans_cons', 'Consumo ex-cartão (consumer ex-card)', 'Consumer, excluding credit card'),
 'cart_pj': ('loans_wh', 'Atacado (wholesale / commercial)', 'Wholesale / commercial'),
 'est1': ('loans_card', 'Cartão de crédito', 'Credit card'),
 'est2': ('nonaccrual', 'Empréstimos em nonaccrual', 'Nonaccrual loans'),
 'est3': ('npa', 'Ativos não performados (nonperforming assets)', 'Nonperforming assets'),
 'prov_e1': ('acl_cons', 'ACL — consumo ex-cartão', 'ACL — consumer ex-card'),
 'prov_e2': ('acl_lrc', 'ACL de compromissos de crédito (fora do balanço)', 'ACL — lending-related commitments'),
 'prov_e3': ('acl_secs', 'ACL de títulos (AFS e HTM)', 'ACL — investment securities'),
 'prov_tot': ('acl_tot', 'ACL total de empréstimos (saldo positivo)', 'Total allowance for loan losses'),
 'npl90': ('dpd90', 'Vencidos há 90+ dias ainda acumulando juros', '90+ days past due and still accruing'),
 'npl15': ('dpd30', 'Vencidos de 30 a 89 dias', '30-89 days past due'),
 'wo': ('nco', 'Baixas líquidas de recuperações (net charge-offs)', 'Net charge-offs'),
 'reneg': ('mods', 'Empréstimos modificados por dificuldade financeira', 'Loan modifications to borrowers in financial difficulty'),
 'mf_cli': ('nii_exmkts', 'NII ex-Markets (reportado)', 'NII excluding Markets (reported)'),
 'mf_mkt': ('nii_mkts', 'NII de Markets (reportado)', 'Markets NII (reported)'),
 'at_rent': ('iea', 'Ativos rentáveis (saldo fim de período)', 'Interest-earning assets, eop'),
 'pas_oner': ('ibl', 'Passivos onerosos (saldo fim de período)', 'Interest-bearing liabilities, eop'),
 'rwa': ('rwa', 'RWA — abordagem padronizada (Standardized)', 'RWA — Standardized'),
 'rwa_cred': ('rwa_cred', 'RWA de crédito', 'Credit RWA'),
 'rwa_mkt': ('rwa_mkt', 'RWA de mercado', 'Market RWA'),
 'rwa_op': ('rwa_op', 'RWA operacional (só na abordagem avançada)', 'Operational RWA (Advanced only)'),
 'pr': ('total_cap', 'Capital total (Standardized)', 'Total capital'),
 'nivel1': ('tier1', 'Capital Nível 1 (Tier 1)', 'Tier 1 capital'),
 'cet1': ('cet1', 'Capital principal (CET1)', 'Common equity Tier 1 capital'),
 'ib_rep': ('tcr_rep', 'Índice de capital total reportado — Standardized (%)', 'Total capital ratio — Standardized'),
 'in1_rep': ('t1r_rep', 'Índice Tier 1 reportado — Standardized (%)', 'Tier 1 ratio — Standardized'),
 'icp_rep': ('cet1r_rep', 'Índice CET1 reportado — Standardized (%)', 'CET1 ratio — Standardized'),
 'alav': ('t1lev', 'Tier 1 leverage ratio (%)', 'Tier 1 leverage ratio'),
 'lcr': ('lcr', 'LCR — liquidity coverage ratio (%)', 'Liquidity coverage ratio'),
 'nsfr': ('nsfr', 'NSFR — net stable funding ratio (%)', 'Net stable funding ratio'),
 'acoes': ('shares', 'Ações ordinárias em circulação no fim do período (milhões)', 'Common shares outstanding, eop (mn)'),
 'div_decl': ('div_common', 'Dividendos ordinários declarados no período', 'Common dividends declared'),
 'jcp_decl': ('buyback', 'Recompras de ações ordinárias no período (caixa)', 'Common stock repurchased (cash)'),
 'preco': ('price', 'Preço da ação no fechamento', 'Share price, eop'),
 'funcion': ('headcount', 'Número de funcionários', 'Headcount'),
 'agencias': ('branches', 'Número de agências', 'Branches'),
 'clientes': ('clients', 'Clientes (mil)', 'Clients (thousands)'),
 'cart_cartao': ('seg1_rev', 'Segmento 1 — receita líquida (ex.: Consumer & Community Banking)', 'Segment 1 — net revenue'),
 'cart_consig': ('seg2_rev', 'Segmento 2 — receita líquida (ex.: Commercial & Investment Bank)', 'Segment 2 — net revenue'),
 'cart_veic': ('seg3_rev', 'Segmento 3 — receita líquida (ex.: Asset & Wealth Management)', 'Segment 3 — net revenue'),
 'cart_imob': ('seg4_rev', 'Segmento 4 — receita líquida (ex.: Corporate)', 'Segment 4 — net revenue'),
 'cart_rural': ('seg_recon', 'Ajustes de reconciliação dos segmentos (FTE e outros)', 'Segment reconciling items'),
 'cart_pf_out': ('seg1_ni', 'Segmento 1 — lucro líquido', 'Segment 1 — net income'),
 'cart_mpe': ('seg2_ni', 'Segmento 2 — lucro líquido', 'Segment 2 — net income'),
 'cart_grandes': ('seg3_ni', 'Segmento 3 — lucro líquido', 'Segment 3 — net income'),
 'cart_pj_out': ('seg4_ni', 'Segmento 4 — lucro líquido', 'Segment 4 — net income'),
 'serv_tarifas': ('fee_ib', 'Investment banking fees', 'Investment banking fees'),
 'serv_cartoes': ('fee_lend_dep', 'Tarifas de empréstimos e depósitos', 'Lending- and deposit-related fees'),
 'serv_adm_rec': ('fee_am', 'Gestão de ativos (asset management fees)', 'Asset management fees'),
 'serv_seg_corr': ('fee_comm', 'Comissões e outras tarifas', 'Commissions and other fees'),
 'serv_merc_cap': ('fee_card', 'Cartão (card income, líquido de rewards)', 'Card income'),
 'serv_out': ('fee_mort', 'Hipotecas (mortgage fees) e demais tarifas', 'Mortgage fees and other'),
 'pes_proventos': ('comp_sal', '(Salários)', '(Salaries)'),
 'pes_benef': ('comp_ben', '(Benefícios)', '(Benefits)'),
 'pes_plr': ('comp_inc', '(Remuneração variável)', '(Incentive compensation)'),
 'pes_out': ('comp_oth', '(Demais despesas de pessoal)', '(Other compensation)'),
 'adm_proc': ('tech', '(Tecnologia, comunicações e equipamentos)', '(Technology, communications and equipment)'),
 'adm_alug': ('occupancy', '(Ocupação)', '(Occupancy)'),
 'adm_mkt': ('marketing', '(Marketing)', '(Marketing)'),
 'adm_terc': ('prof_serv', '(Serviços profissionais e terceirizados)', '(Professional and outside services)'),
 'adm_deprec': ('adm_da', '(Depreciação e amortização, quando à parte)', '(Depreciation and amortization, if separate)'),
 'adm_out': ('adm_oth', '(Demais despesas administrativas)', '(Other administrative expense)'),
}
NEWCODES = [  # linhas novas de origem na Inputs (área livre); (código, PT, unidade)
 ('pref', 'Ações preferenciais (valor contábil)', 'moeda'),
 ('div_pref', '(Dividendos preferenciais e alocação a títulos participantes)', 'moeda'),
 ('slr', 'Supplementary leverage ratio — SLR (%)', 'pct'),
 ('rwa_adv', 'RWA — abordagem avançada (Advanced)', 'moeda'),
 ('cet1r_adv', 'Índice CET1 reportado — Advanced (%)', 'pct'),
 ('msr', 'MSRs — direitos de serviço de hipotecas (dentro de goodwill e intangíveis)', 'moeda'),
]
for r in range(5, 150):
    c = inp.cell(r, 2).value
    if c in CODES:
        nc, pt, en = CODES[c]
        inp.cell(r, 2).value = nc; inp.cell(r, 3).value = pt
        unit = 'pct' if '(%)' in pt else ('un' if c in ('funcion', 'agencias', 'clientes') else ('acoes' if c == 'acoes' else ('por acao' if c == 'preco' else 'moeda')))
        inp.cell(r, 4).value = f'preencher: fonte e linha · fonte · {unit}'
for r in range(5, 196):
    c = w2.cell(r, 2).value
    if c in CODES:
        nc, pt, en = CODES[c]
        w2.cell(r, 2).value, w2.cell(r, 3).value, w2.cell(r, 4).value = nc, pt, en
# títulos e subtotais (Inputs e Inputs (2))
SUBT2 = {6: ('Aplicações de liquidez', 'Liquidity placements'), 11: ('Títulos e ativos de trading', 'Securities and trading assets'),
         17: ('Empréstimos', 'Loans'), 24: ('Investimentos', 'Investments'), 38: ('Compromissadas e captações de curto prazo', 'Repos and short-term funding'),
         42: ('Dívida de longo prazo e VIEs', 'Long-term debt and VIEs'), 47: ('Outros empréstimos tomados', 'Other borrowings'),
         59: ('Patrimônio líquido da companhia (ordinário + preferencial)', 'Total stockholders’ equity'),
         70: ('Receita de juros', 'Interest income'), 75: ('(Despesa de juros)', '(Interest expense)'),
         80: ('Receita líquida de juros (NII)', 'Net interest income'), 81: ('Provisão para perdas de crédito', 'Provision for credit losses'),
         84: ('NII após provisão', 'NII after provision'), 85: ('Receitas e despesas não-juros', 'Noninterest revenue and expense'),
         90: ('(Provisões jurídicas e outras, quando à parte)', '(Legal and other provisions)'), 98: ('Resultado operacional', 'Operating result'),
         100: ('Lucro antes do imposto de renda', 'Income before income taxes'), 101: ('Imposto de renda', 'Income tax'),
         104: ('LUCRO LÍQUIDO', 'NET INCOME'), 106: ('LUCRO LÍQUIDO ATRIBUÍVEL À COMPANHIA', 'NET INCOME ATTRIBUTABLE TO THE COMPANY'),
         110: ('Empréstimos e qualidade de crédito (CECL)', 'Loans and credit quality (CECL)'), 126: ('Margem financeira e base de ativos', 'NII and earning assets'),
         132: ('Capital regulatório e liquidez (Basileia III dos EUA)', 'Regulatory capital and liquidity (US Basel III)'),
         147: ('Ações, distribuições ao acionista e dados operacionais', 'Shares, capital return and operating data'),
         158: ('Segmentos de negócio (quando a companhia abre)', 'Business segments'), 169: ('Receitas de tarifas por origem', 'Fee revenue by source'),
         177: ('Despesas de pessoal por natureza', 'Compensation by nature'), 183: ('Despesas não-juros por natureza', 'Noninterest expense by nature'),
         194: ('Empréstimos totais – soma dos portfólios', 'Total loans less sum of portfolios'), 195: ('ACL do balanço – ACL total informada', 'Balance-sheet ACL less reported ACL')}
for r, (pt, en) in SUBT2.items():
    w2.cell(r, 3).value, w2.cell(r, 4).value = pt, en
for col in range(5, 60):
    x = L(col)
    w2.cell(194, col).value = f'=IFERROR(IF(N({x}112+{x}113+{x}114)=0,"",ROUND({x}111-({x}112+{x}113+{x}114),0)),"")'
    w2.cell(195, col).value = f'=IFERROR(IF(OR(N({x}120)=0,N({x}20)=0),"",ROUND(ABS({x}20)-ABS({x}120),0)),"")'
SECT_IN = {5: 'ATIVO', 27: 'PASSIVO', 50: 'PATRIMÔNIO LÍQUIDO', 58: 'DEMONSTRAÇÃO DO RESULTADO (US GAAP)', 85: 'DADOS COMPLEMENTARES'}
for r, t in SECT_IN.items(): inp.cell(r, 3).value = t if inp.cell(r, 3).value else inp.cell(r, 3).value

# Model: rótulos (linhas originais)
MLAB = {
 2: ('MODELO OPERACIONAL — BALANÇO, RESULTADO E INDICADORES (BANCO DOS EUA, US GAAP)', None),
 18: ('Caixa e depósitos não remunerados em bancos', 'Cash and due from banks'), 19: ('Aplicações de liquidez', 'Liquidity placements'),
 20: ('Fed funds vendidos e compromissadas reversas', 'Fed funds sold and resale agreements'), 21: ('Depósitos remunerados em bancos', 'Deposits with banks'),
 22: ('Títulos tomados em empréstimo (securities borrowed)', 'Securities borrowed'), 23: ('(Provisão para perdas — aplicações)', '(Allowance)'),
 24: ('Títulos e ativos de trading', 'Securities and trading assets'), 25: ('Ativos de trading', 'Trading assets'), 26: ('Títulos AFS (valor justo)', 'AFS securities'),
 27: ('Títulos HTM (custo amortizado)', 'HTM securities'), 28: ('(Provisão para perdas — títulos)', '(Allowance — securities)'),
 29: ('Derivativos a receber (quando fora de trading)', 'Derivative receivables'), 30: ('Empréstimos, líquidos de ACL', 'Loans, net of allowance'),
 31: ('Empréstimos', 'Loans'), 32: ('Empréstimos a valor justo / mantidos para venda', 'Loans at fair value / held for sale'),
 33: ('(Provisão para perdas de crédito — ACL)', '(Allowance for loan losses)'), 34: ('Juros e contas a receber', 'Accrued interest and accounts receivable'),
 35: ('Impostos diferidos ativos', 'Deferred tax assets'), 36: ('Outros ativos', 'Other assets'), 38: ('Investimentos por equivalência', 'Equity-method investments'),
 40: ('Imóveis e equipamentos', 'Premises and equipment'), 41: ('Goodwill, MSRs e outros intangíveis', 'Goodwill, MSRs and other intangibles'),
 45: ('Depósitos sem juros', 'Noninterest-bearing deposits'), 46: ('Depósitos com juros', 'Interest-bearing deposits'),
 47: ('Depósitos — abertura adicional', 'Deposits — additional split'), 48: ('Depósitos a prazo', 'Time deposits'), 49: ('Outros depósitos', 'Other deposits'),
 50: ('Compromissadas e captações de curto prazo', 'Repos and short-term funding'),
 51: ('Fed funds comprados e compromissadas (repos)  ← PLUG DE FUNDING', 'Repos (funding plug)'), 52: ('Captações de curto prazo', 'Short-term borrowings'),
 53: ('Títulos emprestados (securities loaned)', 'Securities loaned'), 54: ('Dívida de longo prazo e VIEs', 'Long-term debt and VIEs'),
 55: ('Dívida de longo prazo', 'Long-term debt'), 56: ('Participações emitidas por VIEs consolidadas', 'Beneficial interests issued by consolidated VIEs'),
 57: ('Outras emissões', 'Other issuances'), 58: ('Dívida subordinada (quando à parte)', 'Subordinated debt'), 59: ('Outros empréstimos tomados', 'Other borrowings'),
 60: ('Adiantamentos do FHLB e outros', 'FHLB advances and other'), 61: ('Outros empréstimos', 'Other borrowings'),
 62: ('Passivos de trading (inclui derivativos a pagar)', 'Trading liabilities'), 63: ('Obrigações com corretoras e clientes', 'Payables to brokers and customers'),
 64: ('Provisão para compromissos de crédito', 'Allowance for lending-related commitments'), 65: ('Reservas de seguros', 'Insurance reserves'),
 66: ('Demais provisões', 'Other provisions'), 67: ('Impostos a pagar e diferidos', 'Taxes payable and deferred'),
 68: ('Contas a pagar e outros passivos', 'Accounts payable and other liabilities'),
 71: ('Patrimônio líquido da companhia (ordinário + preferencial)', 'Total stockholders’ equity'), 72: ('Ações ordinárias (valor nominal)', 'Common stock'),
 73: ('Capital adicional integralizado (APIC) e outros', 'Additional paid-in capital'), 74: ('Lucros retidos', 'Retained earnings'),
 75: ('Outros resultados abrangentes acumulados (AOCI)', 'AOCI'), 76: ('(Ações em tesouraria, a custo)', '(Treasury stock)'),
 84: ('(–) Dividendos declarados (ordinários e preferenciais)', '(–) Dividends declared (common and preferred)'),
 86: ('(+/–) Variação de capital (ações, APIC e preferenciais)', '(+/–) Capital changes'),
 91: ('Receita de juros', 'Interest income'), 92: ('Juros de empréstimos', 'Loans'), 93: ('Juros de títulos e ativos de trading', 'Securities and trading assets'),
 94: ('Juros de compromissadas reversas e securities borrowed', 'Resale agreements and securities borrowed'), 95: ('Juros de depósitos em bancos e outros', 'Deposits with banks and other'),
 96: ('(Despesa de juros)', '(Interest expense)'), 97: ('(Depósitos)', '(Deposits)'), 98: ('(Dívida de longo prazo e VIEs)', '(Long-term debt and VIEs)'),
 99: ('(Compromissadas, curto prazo, trading e outros)', '(Repos, short-term, trading and other)'), 100: ('(Outras despesas de juros)', '(Other interest expense)'),
 101: ('Receita líquida de juros (NII)', 'Net interest income'), 102: ('Provisão para perdas de crédito', 'Provision for credit losses'),
 103: ('(Provisão para perdas de crédito)', '(Provision for credit losses)'), 104: ('Recuperações (quando à parte)', 'Recoveries'),
 105: ('NII após provisão', 'NII after provision'), 106: ('Receitas e despesas não-juros', 'Noninterest revenue and expense'),
 107: ('Receitas de tarifas e comissões', 'Fees and commissions'), 108: ('Resultado de trading (principal transactions)', 'Principal transactions'),
 109: ('(Remuneração e benefícios)', '(Compensation)'), 110: ('(Ocupação, tecnologia, serviços profissionais e marketing)', '(Occupancy, technology, professional services, marketing)'),
 111: ('(Provisões jurídicas e outras, quando à parte)', '(Legal and other provisions)'), 112: ('(Despesas jurídicas)', '(Legal expense)'),
 113: ('(Outras provisões)', '(Other provisions)'), 114: ('(Outras provisões)', '(Other provisions)'),
 115: ('(FDIC e contribuições regulatórias, quando à parte)', '(FDIC and regulatory assessments)'), 116: ('Resultado de equivalência patrimonial', 'Equity-method income'),
 117: ('Outras receitas não-juros', 'Other noninterest revenue'), 118: ('(Outras despesas não-juros)', '(Other noninterest expense)'),
 119: ('Resultado operacional', 'Operating result'), 120: ('Resultado não operacional', 'Non-operating result'),
 121: ('Lucro antes do imposto de renda', 'Income before income taxes'), 122: ('Imposto de renda', 'Income tax expense'),
 123: ('(Imposto de renda)', '(Income tax)'), 124: ('Imposto diferido (quando à parte)', 'Deferred tax'),
 125: ('LUCRO LÍQUIDO', 'NET INCOME'), 126: ('(Lucro atribuído a não controladores)', '(Noncontrolling interests)'),
 127: ('LUCRO LÍQUIDO ATRIBUÍVEL À COMPANHIA', 'NET INCOME ATTRIBUTABLE TO THE COMPANY'),
 136: ('PL da companhia — saldo médio', 'Total equity, average'),
 137: ('Receita líquida de juros (NII)', 'Net interest income'),
 138: ('NII ex-Markets (reportado; senão juros de empréstimos – juros de depósitos)', 'NII ex-Markets'),
 139: ('NII de Markets e tesouraria (por diferença)', 'Markets and treasury NII'),
 140: ('Receita líquida total (total net revenue)', 'Total net revenue'), 141: ('Despesas não-juros (noninterest expense)', 'Noninterest expense'),
 144: ('Empréstimos totais', 'Total loans'), 145: ('ACL de empréstimos', 'Allowance for loan losses'),
 146: ('Empréstimos em nonaccrual', 'Nonaccrual loans'), 147: ('Ativos não performados (NPA)', 'Nonperforming assets'),
 148: ('Vencidos há 90+ dias ainda acumulando juros', '90+ days past due, still accruing'), 149: ('Vencidos de 30 a 89 dias', '30-89 days past due'),
 150: ('Baixas líquidas (net charge-offs)', 'Net charge-offs'), 151: ('ACL de compromissos de crédito (fora do balanço)', 'ACL — lending-related commitments'),
 152: ('ACL de títulos (AFS e HTM)', 'ACL — securities'),
 155: ('RWA — abordagem padronizada (Standardized)', 'RWA — Standardized'), 156: ('Capital total', 'Total capital'),
 157: ('Capital Nível 1 (Tier 1)', 'Tier 1 capital'), 158: ('Índice de capital total', 'Total capital ratio'), 159: ('Índice Tier 1', 'Tier 1 ratio'),
 162: ('ROE — sobre o PL total da companhia (inclui preferenciais)', 'ROE on total equity'),
 165: ('NIM ex-Markets — sobre empréstimos médios (% a.a.)', 'NII ex-Markets over average loans'),
 166: ('NIM de Markets e tesouraria (% a.a.)', 'Markets and treasury NIM'),
 170: ('Índice de eficiência (overhead ratio)', 'Overhead ratio'), 171: ('Remuneração / receita líquida', 'Compensation over revenue'),
 172: ('Outras despesas administrativas / receita líquida', 'Admin over revenue'), 173: ('Cobertura da remuneração pelas tarifas', 'Fee coverage of compensation'),
 175: ('Receita líquida total — % YoY', 'Total net revenue growth YoY'), 176: ('NII — % YoY', 'NII growth YoY'),
 177: ('Empréstimos líquidos — % QoQ', 'Loan growth QoQ'), 178: ('Empréstimos líquidos — % YoY', 'Loan growth YoY'),
 181: ('Vencidos 90+ ainda acumulando (% dos empréstimos)', '90+ DPD still accruing'), 182: ('Vencidos de 30 a 89 dias (% dos empréstimos)', '30-89 DPD'),
 183: ('Nonaccrual (% dos empréstimos)', 'Nonaccrual ratio'), 184: ('NPA (% dos empréstimos)', 'NPA ratio'),
 185: ('Cobertura do nonaccrual (ACL / nonaccrual)', 'ACL over nonaccrual'),
 186: ('ACL total, incluindo compromissos (% dos empréstimos)', 'Total ACL incl. commitments over loans'),
 187: ('Cobertura de NPA (ACL / NPA)', 'ACL over NPA'), 188: ('ACL / empréstimos (%)', 'ACL over loans'),
 189: ('Custo do crédito — provisão / empréstimos médios (% a.a.)', 'Provision over average loans'),
 190: ('Taxa de net charge-offs (% dos empréstimos médios, a.a.)', 'Net charge-off rate'),
 192: ('Índice de capital total', 'Total capital ratio'), 193: ('Índice Tier 1', 'Tier 1 ratio'),
 199: ('Ações ordinárias em circulação (milhões)', 'Common shares outstanding (mn)'),
 203: ('Payout total — dividendos ordinários + recompras / lucro ao acionista comum', 'Total payout (dividends + buybacks)'),
}
for r, (pt, en) in MLAB.items():
    m.cell(r, 3).value = pt
    if en: m.cell(r, 4).value = en
m.cell(200, 3).value = '="Lucro por ação ordinária — EPS ("&Inputs!$E$200&")"'
m.cell(201, 3).value = '="Valor patrimonial por ação ordinária — BVPS ("&Inputs!$E$200&")"'
m.cell(202, 3).value = '="Dividendo ordinário por ação — DPS ("&Inputs!$E$200&")"'
for col in range(5, 60):     # cobertura: nonaccrual, ACL total incl. compromissos, NPA
    x = L(col)
    m.cell(183, col).value = f'=IF(N({x}146)=0,"",IFERROR({x}146/{x}144,0))'
    m.cell(184, col).value = f'=IF(N({x}147)=0,"",IFERROR({x}147/{x}144,0))'
    m.cell(185, col).value = f'=IF(N({x}146)=0,"",IFERROR({x}145/{x}146,0))'
    m.cell(186, col).value = f'=IFERROR(({x}145+{x}151)/{x}144,0)'
    m.cell(187, col).value = f'=IF(N({x}147)=0,"",IFERROR({x}145/{x}147,0))'
    m.cell(181, col).value = f'=IF(N({x}148)=0,"",IFERROR({x}148/{x}144,0))'
    m.cell(182, col).value = f'=IF(N({x}149)=0,"",IFERROR({x}149/{x}144,0))'

# Inputs: painel (rótulos)
PLAB = {
 197: 'Instituição', 199: 'Tipo de banco / norma contábil', 207: 'Ações ordinárias emitidas (milhões)', 208: 'Ações preferenciais (não entram na contagem ordinária)',
 209: '(-) Ações em tesouraria (milhões, negativo)', 210: 'Ações ordinárias em circulação (milhões)',
 222: 'Rendimento médio dos empréstimos (% a.a.)', 225: 'Recuperações (% da provisão; nos EUA a provisão é líquida)',
 236: 'Resultado de trading — crescimento % YoY', 239: 'Provisões jurídicas e outras (% da receita líquida)',
 240: 'FDIC e contribuições regulatórias (% da receita líquida)', 244: 'Alíquota efetiva de imposto de renda (%)',
 245: 'Participação de não controladores (% do lucro líquido)', 247: 'Payout de dividendos ordinários (% do lucro ao acionista comum)',
 249: 'Índice de capital total reportado (%)', 250: 'Índice Tier 1 reportado (%)',
 254: 'Vencidos 90+ ainda acumulando (% dos empréstimos)', 255: 'Vencidos de 30 a 89 dias (% dos empréstimos)',
 256: 'Cobertura dos vencidos 90+ pela ACL (%)', 257: 'ACL / empréstimos (%)',
 286: 'Resultado de trading — crescimento % YoY', 289: 'Provisões jurídicas e outras (% da receita líquida)',
 290: 'FDIC e contribuições regulatórias (% da receita líquida)', 294: 'Alíquota efetiva de imposto de renda (%)',
 298: 'Payout de dividendos ordinários (% do lucro ao acionista comum)', 300: 'Índice de capital total alvo (%)', 301: 'Tier 1 alvo (%)',
 306: 'Taxa livre de risco — Treasury 10 anos (% a.a.)', 307: 'Prêmio de risco de mercado — EUA (%)',
 309: 'Prêmio de tamanho / liquidez (%) — normalmente zero em large caps dos EUA',
}
for r, t in PLAB.items(): inp.cell(r, 3).value = t
inp['E200'] = 'US$'; inp['E201'] = 'milhões'
inp['F1'] = ('MODELO-BASE VAZIO — BANCOS DOS EUA (US GAAP). Preencha os valores nas linhas abaixo e as datas de primeiro e último trimestre no painel. '
             'Fonte preferencial: 10-Q/10-K (SEC XBRL) e earnings supplement da companhia. Macro histórico do FRED (St. Louis Fed).')

# macro (FRED) — histórico 1T20..último trimestre completo; colunas projetadas ficam vazias
def series(sid):
    out = []
    for row in csv.DictReader(open(os.path.join(FRED, sid + '.csv'))):
        v = row.get(sid)
        if v not in (None, '', '.'): out.append((dt.date.fromisoformat(row['observation_date']), float(v)))
    return out
SER = {s: series(s) for s in ('DFEDTARU', 'DFF', 'CPIAUCSL', 'GDPC1', 'DGS10', 'SOFR')}
def qbounds(y, q):
    a = dt.date(y, 3 * q - 2, 1); b = dt.date(y + (q == 4), (3 * q) % 12 + 1, 1) - dt.timedelta(days=1)
    return a, b
def in_q(s, y, q): a, b = qbounds(y, q); return [v for d, v in s if a <= d <= b]
def last_m(s, y, mth): return next((v for d, v in s if d.year == y and d.month == mth), None)
def macro(y, q):
    out = {}
    a, b = qbounds(y, q)
    if b > max(d for d, _ in SER['DFF']): return out
    ff = in_q(SER['DFEDTARU'], y, q); out[260] = ff[-1] / 100
    eff = in_q(SER['DFF'], y, q); out[261] = statistics.mean(eff) / 100
    out[262] = (1 + out[261]) ** 0.25 - 1
    c1, c0, c12 = last_m(SER['CPIAUCSL'], y, 3 * q), last_m(SER['CPIAUCSL'], y - (q == 1), (3 * q - 3) or 12), last_m(SER['CPIAUCSL'], y - 1, 3 * q)
    if c1 and c0: out[263] = c1 / c0 - 1
    if c1 and c12: out[264] = c1 / c12 - 1
    g1 = next((v for d, v in SER['GDPC1'] if d == a), None); g0 = next((v for d, v in SER['GDPC1'] if d == dt.date(y - 1, a.month, 1)), None)
    if g1 and g0: out[265] = g1 / g0 - 1
    t10 = in_q(SER['DGS10'], y, q); out[266] = t10[-1] / 100
    so = in_q(SER['SOFR'], y, q)
    if so: out[267] = statistics.mean(so) / 100
    return out
MACROLAB = {260: ('Fed funds — teto da meta no fim do período (% a.a.) [FRED DFEDTARU]', 'Fed funds target (upper), eop'),
            261: ('Fed funds efetivo — média do período (% a.a.) [FRED DFF]', 'Effective fed funds, period average'),
            262: ('Fed funds efetivo acumulado no trimestre (%) [derivado de DFF]', 'Effective fed funds, quarterly'),
            263: ('CPI no trimestre (%) [FRED CPIAUCSL]', 'CPI, quarterly'), 264: ('CPI 12 meses (%) [FRED CPIAUCSL]', 'CPI, 12m'),
            265: ('PIB real (% YoY) [FRED GDPC1]', 'Real GDP growth, YoY'), 266: ('Treasury 10 anos — fim do período (% a.a.) [FRED DGS10]', 'UST 10y, eop'),
            267: ('SOFR — média do período (% a.a.) [FRED SOFR]', 'SOFR, period average')}
for r, (pt, en) in MACROLAB.items():
    inp.cell(r, 3).value, inp.cell(r, 4).value = pt, en
    for col in range(5, 60):
        if (col - 5) % 5 != 4: inp.cell(r, col).value = None
for col in range(5, 60):
    if (col - 5) % 5 == 4: continue
    y = 2020 + (col - 5) // 5; q = (col - 5) % 5 + 1
    for r, v in macro(y, q).items(): inp.cell(r, col).value = round(v, 6)
inp.cell(259, 3).value = 'CENÁRIO MACROECONÔMICO — EUA (histórico do FRED; colunas projetadas são premissa do analista)'

# Valuation: rótulos
VLAB = {3: 'Ke por CAPM (Treasury 10 anos + beta × ERP dos EUA). Excess Return e DDM (dividendos + recompras) como métodos primários; P/BV justificado e múltiplos como checagem. Lucro e PL referem-se ao acionista ordinário.',
        16: 'Taxa livre de risco — Treasury 10 anos', 19: 'Prêmio de tamanho / liquidez (normalmente zero)', 29: 'Lucro líquido ao acionista comum',
        30: 'PL ordinário (final, ex-preferenciais)', 31: 'PL ordinário inicial', 32: 'ROE ordinário do exercício',
        33: 'Dividendos ordinários + recompras (retorno total ao acionista)', 34: 'Payout total (dividendos + recompras)',
        35: 'Ações ordinárias em circulação (milhões)', 44: 'Lucro líquido ao acionista comum', 45: '(–) Encargo de capital: Ke × PL ordinário inicial',
        52: 'PL ordinário na data-base', 58: 'Fluxo de retorno ao acionista (dividendos + recompras)', 59: 'Dividendos ordinários + recompras',
        60: 'Valor presente do retorno ao acionista', 62: 'Retorno ao acionista do último ano projetado', 63: 'Valor terminal por Gordon (payout total sustentável)'}
for r, t in VLAB.items(): val.cell(r, 3).value = t
val['D33'] = 'Dividends + buybacks'; val['D59'] = 'Dividends + buybacks'
for i, t in enumerate(['BAC', 'C', 'WFC', 'GS', 'MS', 'PNC', 'USB']): val.cell(82 + i, 3).value = t
val['F81'] = 'P/TBV (x)'; val['G81'] = 'ROTCE (%)'
val['C92'] = 'P/E alvo (x)'; val['C93'] = 'P/TBV alvo (x) — aplicado ao TBVPS'

wb.save(STEP)

# ------------------------------------------------------------------ 2) inserções de linhas
cur = STEP
for sheet, lst in INS.items():
    for at, n in lst:
        nxt = os.path.join(tmp, f'ins_{len(os.listdir(tmp))}.xlsx')
        subprocess.run([sys.executable, SHIFT, cur, nxt, sheet, str(at), str(n)], check=True, stdout=subprocess.DEVNULL)
        cur = nxt

# ------------------------------------------------------------------ 3) linhas novas e fórmulas (coordenadas finais)
wb = openpyxl.load_workbook(cur)
inp, w2, m, val, glo = wb['Inputs'], wb['Inputs (2)'], wb['Model'], wb['Valuation'], wb['Glossário']
FM = lambda r: fin('Model', r); F2 = lambda r: fin('Inputs (2)', r); FI = lambda r: fin('Inputs', r)
QC = [c for c in range(5, 60) if (c - 5) % 5 != 4]
FYC = [c for c in range(5, 60) if (c - 5) % 5 == 4]
def prevq(c):
    p = c - 1
    if p >= 5 and (p - 5) % 5 == 4: p -= 1
    return p if p >= 5 else None
def styl(ws, src, dst, c1=1, c2=61):
    for c in range(c1, c2 + 1):
        a, b = ws.cell(src, c), ws.cell(dst, c)
        if a.has_style: b._style = copy.copy(a._style)

# Inputs (2): linhas novas de código
R2 = {'pref': blk('Inputs (2)', 65), 'div_pref': blk('Inputs (2)', 107), 'nicommon': blk('Inputs (2)', 107) + 1,
      'slr': blk('Inputs (2)', 146), 'rwa_adv': blk('Inputs (2)', 146) + 1, 'cet1r_adv': blk('Inputs (2)', 146) + 2, 'msr': blk('Inputs (2)', 146) + 3}
from openpyxl.formula.translate import Translator
def copy_code_row(src, dst, code, pt, en):
    styl(w2, src, dst)
    w2.cell(dst, 2).value, w2.cell(dst, 3).value, w2.cell(dst, 4).value = code, pt, en
    for c in range(5, 60):
        f = w2.cell(src, c).value
        if isinstance(f, str) and f.startswith('='):
            w2.cell(dst, c).value = Translator(f, origin=w2.cell(src, c).coordinate).translate_formula(w2.cell(dst, c).coordinate)
copy_code_row(F2(64), R2['pref'], 'pref', 'Ações preferenciais', 'Preferred stock')
copy_code_row(F2(105), R2['div_pref'], 'div_pref', '(Dividendos preferenciais e alocação a títulos participantes)', '(Preferred dividends and participating securities)')
styl(w2, F2(106), R2['nicommon'])
w2.cell(R2['nicommon'], 3).value, w2.cell(R2['nicommon'], 4).value = 'LUCRO LÍQUIDO AO ACIONISTA COMUM', 'NET INCOME APPLICABLE TO COMMON'
for c in range(5, 60):
    x = L(c); w2.cell(R2['nicommon'], c).value = f'=IFERROR({x}{F2(106)}+N({x}{R2["div_pref"]}),"")' if (c - 5) % 5 != 4 else f'=SUM({L(c-4)}{R2["nicommon"]}:{L(c-1)}{R2["nicommon"]})'
copy_code_row(F2(144), R2['slr'], 'slr', 'Supplementary leverage ratio — SLR (%)', 'Supplementary leverage ratio')
copy_code_row(F2(133), R2['rwa_adv'], 'rwa_adv', 'RWA — abordagem avançada (Advanced)', 'RWA — Advanced')
copy_code_row(F2(144), R2['cet1r_adv'], 'cet1r_adv', 'Índice CET1 reportado — Advanced (%)', 'CET1 ratio — Advanced')
copy_code_row(F2(133), R2['msr'], 'msr', 'MSRs (dentro de goodwill e intangíveis)', 'Mortgage servicing rights')
for c in range(5, 60):        # PL da companhia inclui as preferenciais
    x = L(c)
    if (c - 5) % 5 != 4:
        w2.cell(F2(59), c).value = f'=IFERROR({x}{F2(60)}+{x}{F2(61)}+{x}{F2(62)}+{x}{F2(63)}+{x}{F2(64)}+N({x}{R2["pref"]}),"")'

# Inputs: códigos novos na área livre (linhas 151+)
r0 = 151
inp.cell(r0, 3).value = 'CONTAS ESPECÍFICAS DE BANCOS DOS EUA'
for i, (code, pt, unit) in enumerate(NEWCODES, 1):
    styl(inp, 6, r0 + i)
    inp.cell(r0 + i, 2).value, inp.cell(r0 + i, 3).value, inp.cell(r0 + i, 4).value = code, pt, f'preencher: fonte e linha · fonte · {unit}'
inp.cell(153, 3).value = None
free0 = r0 + len(NEWCODES) + 2
inp.cell(free0, 3).value = f'Espaço livre até a linha 190 (linhas {free0 + 1} a 190) para novas linhas de origem. Qualquer linha com um código existente é somada automaticamente pela Inputs (2).'

# Inputs: painel — payout ordinário e recompras (implícitos e premissas), CET1 alvo
RI_PAY, RI_BUY = FI(247), blk('Inputs', 248)
RD_PAY, RD_BUY = FI(298), blk('Inputs', 299)
RD_T1, RD_CET1 = FI(301), blk('Inputs', 302)
def implied(row, num_row):
    for c in range(5, 60):
        x = L(c)
        if inp.cell(RI_PAY, c).value is None and row != RI_PAY: continue
        if (inp.cell(RI_PAY, c).value or '').startswith('='):
            inp.cell(row, c).value = (f"=IFERROR(IF(OR(N('Inputs (2)'!{x}{R2['nicommon']})<=0,N('Inputs (2)'!{x}{num_row})=0),\"\","
                                      f"'Inputs (2)'!{x}{num_row}/'Inputs (2)'!{x}{R2['nicommon']}),\"\")")
styl(inp, RI_PAY, RI_BUY)
inp.cell(RI_BUY, 3).value, inp.cell(RI_BUY, 4).value = 'Recompras de ações (% do lucro ao acionista comum)', 'Buybacks over net income to common'
implied(RI_PAY, F2(149)); implied(RI_BUY, F2(150))
for src, dst, pt, en in [(RD_PAY, RD_BUY, 'Recompras de ações (% do lucro ao acionista comum)', 'Buybacks'), (RD_T1, RD_CET1, 'CET1 alvo (%)', 'Target CET1')]:
    styl(inp, src, dst); inp.cell(dst, 3).value, inp.cell(dst, 4).value = pt, en
    for c in FYC:
        f = inp.cell(src, c).value
        if isinstance(f, str): inp.cell(dst, c).value = Translator(f, origin=inp.cell(src, c).coordinate).translate_formula(inp.cell(dst, c).coordinate)

# Model: linhas novas
MR = {'pref': blk('Model', 77), 'buy': blk('Model', 85), 'prefdiv': blk('Model', 128), 'nicom': blk('Model', 128) + 1}
b = blk('Model', 137); MR.update(ceq=b, ceq_avg=b + 1, msr=b + 2, tce=b + 3, tce_avg=b + 4)
b = blk('Model', 142); MR.update(seg_hdr=b + 1, s1=b + 2, s2=b + 3, s3=b + 4, s4=b + 5, srec=b + 6, schk=b + 7, n1=b + 8, n2=b + 9, n3=b + 10, n4=b + 11)
b = blk('Model', 160); MR.update(cet1=b, cet1r=b + 1, cet1r_adv=b + 2, slr=b + 3, t1lev=b + 4, lcr=b + 5)
b = blk('Model', 163); MR.update(roe_c=b, rotce=b + 1)
MR['cet1r_ratio'] = blk('Model', 194); MR['tbvps'] = blk('Model', 202)

def put(row, lab_pt, lab_en, qf, fyf, style_from):
    styl(m, style_from, row)
    m.cell(row, 3).value, m.cell(row, 4).value = lab_pt, lab_en
    for c in range(5, 60):
        x = L(c); p = prevq(c)
        m.cell(row, c).value = qf(x, L(p) if p else None, c) if (c - 5) % 5 != 4 else fyf(x, L(c - 4), L(c - 1), c)
last = lambda r: (lambda x, a, z, c: f'={z}{r}')
ssum = lambda r: (lambda x, a, z, c: f'=SUM({a}{r}:{z}{r})')
same = lambda f: (lambda x, a, z, c: f(x, None, c))
def hist_or_prev(i2row, row, first=0):
    return lambda x, p, c: f"=IF({x}$14=1,'Inputs (2)'!{x}{i2row},{p}{row})" if p else f"=IF({x}$14=1,'Inputs (2)'!{x}{i2row},{first})"

put(MR['pref'], 'Ações preferenciais', 'Preferred stock', hist_or_prev(R2['pref'], MR['pref']), last(MR['pref']), FM(76))
put(MR['buy'], '(–) Recompras de ações ordinárias', '(–) Common stock repurchases',
    lambda x, p, c: f"=IF({x}$14=1,-N('Inputs (2)'!{x}{F2(150)}),-MAX(0,{x}{MR['nicom']})*Inputs!{x}${RD_BUY})", ssum(MR['buy']), FM(84))
put(MR['prefdiv'], '(Dividendos preferenciais e alocação a títulos participantes)', '(Preferred dividends)',
    lambda x, p, c: f"=IF({x}$14=1,N('Inputs (2)'!{x}{R2['div_pref']}),{p}{MR['prefdiv']})" if p else f"=IF({x}$14=1,N('Inputs (2)'!{x}{R2['div_pref']}),0)", ssum(MR['prefdiv']), FM(126))
put(MR['nicom'], 'LUCRO LÍQUIDO AO ACIONISTA COMUM', 'NET INCOME APPLICABLE TO COMMON', lambda x, p, c: f'={x}{FM(127)}+{x}{MR["prefdiv"]}',
    lambda x, a, z, c: f'={x}{FM(127)}+{x}{MR["prefdiv"]}', FM(127))
avg = lambda r: (lambda x, p, c: f'=AVERAGE({p}{r},{x}{r})' if p else f'={x}{r}')
avgfy = lambda r: (lambda x, a, z, c: f'=AVERAGE({a}{r}:{z}{r})')
put(MR['ceq'], 'PL ordinário (ex-preferenciais) — saldo final', 'Common equity, eop', lambda x, p, c: f'={x}{FM(71)}-{x}{MR["pref"]}', last(MR['ceq']), FM(136))
put(MR['ceq_avg'], 'PL ordinário — saldo médio', 'Common equity, average', avg(MR['ceq']), avgfy(MR['ceq_avg']), FM(136))
put(MR['msr'], 'MSRs (dentro de goodwill e intangíveis; não deduzidos do TCE)', 'MSRs', hist_or_prev(R2['msr'], MR['msr']), last(MR['msr']), FM(136))
put(MR['tce'], 'PL tangível ordinário (TCE) — saldo final', 'Tangible common equity, eop', lambda x, p, c: f'={x}{MR["ceq"]}-({x}{FM(41)}-{x}{MR["msr"]})', last(MR['tce']), FM(136))
put(MR['tce_avg'], 'TCE — saldo médio', 'Tangible common equity, average', avg(MR['tce']), avgfy(MR['tce_avg']), FM(136))

# segmentos
styl(m, FM(129), MR['seg_hdr']); m.cell(MR['seg_hdr'], 3).value = '="SEGMENTOS DE NEGÓCIO — histórico reportado ("&Inputs!$E$202&"; projeção fica a critério do analista)"'
m.cell(MR['seg_hdr'], 4).value = 'BUSINESS SEGMENTS'
segq = lambda i2: (lambda x, p, c: f"=IF({x}$14=1,IF(ISNUMBER('Inputs (2)'!{x}{i2}),'Inputs (2)'!{x}{i2},\"\"),\"\")")
segfy = lambda r: (lambda x, a, z, c: f'=IF(COUNT({a}{r}:{z}{r})=4,SUM({a}{r}:{z}{r}),"")')
for key, i2, pt, en in [('s1', F2(159), 'Segmento 1 — receita líquida', 'Segment 1 revenue'), ('s2', F2(160), 'Segmento 2 — receita líquida', 'Segment 2 revenue'),
                        ('s3', F2(161), 'Segmento 3 — receita líquida', 'Segment 3 revenue'), ('s4', F2(162), 'Segmento 4 — receita líquida', 'Segment 4 revenue'),
                        ('srec', F2(163), 'Ajustes de reconciliação (FTE e outros)', 'Reconciling items'),
                        ('n1', F2(164), 'Segmento 1 — lucro líquido', 'Segment 1 net income'), ('n2', F2(165), 'Segmento 2 — lucro líquido', 'Segment 2 net income'),
                        ('n3', F2(166), 'Segmento 3 — lucro líquido', 'Segment 3 net income'), ('n4', F2(167), 'Segmento 4 — lucro líquido', 'Segment 4 net income')]:
    put(MR[key], pt, en, segq(i2), segfy(MR[key]), FM(140))
put(MR['schk'], 'Check: receita líquida total – (segmentos + reconciliação)', 'Check: revenue less segments',
    lambda x, p, c: f'=IF(COUNT({x}{MR["s1"]}:{x}{MR["srec"]})=0,"",ROUND({x}{FM(140)}-SUM({x}{MR["s1"]}:{x}{MR["srec"]}),0))',
    lambda x, a, z, c: f'=IF(COUNT({x}{MR["s1"]}:{x}{MR["srec"]})=0,"",ROUND({x}{FM(140)}-SUM({x}{MR["s1"]}:{x}{MR["srec"]}),0))', FM(140))

# capital
cap_from = lambda i2: (lambda x, p, c: f"=IF({x}$14=1,IF(ISNUMBER('Inputs (2)'!{x}{i2}),'Inputs (2)'!{x}{i2},\"\"),\"\")")
put(MR['cet1'], 'Capital principal (CET1)', 'CET1 capital',
    lambda x, p, c: f"=IF(AND({x}$14=1,ISNUMBER('Inputs (2)'!{x}{F2(139)}),'Inputs (2)'!{x}{F2(139)}<>0),'Inputs (2)'!{x}{F2(139)},{x}{FM(155)}*Inputs!{x}${RD_CET1})", last(MR['cet1']), FM(157))
put(MR['cet1r'], 'Índice CET1 — Standardized', 'CET1 ratio — Standardized', lambda x, p, c: f'=IFERROR({x}{MR["cet1"]}/{x}{FM(155)},0)', lambda x, a, z, c: f'=IFERROR({x}{MR["cet1"]}/{x}{FM(155)},0)', FM(159))
put(MR['cet1r_adv'], 'Índice CET1 — Advanced (reportado)', 'CET1 ratio — Advanced (reported)', cap_from(R2['cet1r_adv']), last(MR['cet1r_adv']), FM(159))
put(MR['slr'], 'Supplementary leverage ratio — SLR (reportado)', 'SLR (reported)', cap_from(R2['slr']), last(MR['slr']), FM(159))
put(MR['t1lev'], 'Tier 1 leverage ratio (reportado)', 'Tier 1 leverage (reported)', cap_from(F2(143)), last(MR['t1lev']), FM(159))
put(MR['lcr'], 'LCR — liquidity coverage ratio (reportado)', 'LCR (reported)', cap_from(F2(144)), last(MR['lcr']), FM(159))

# rentabilidade
put(MR['roe_c'], 'ROE ordinário — lucro ao acionista comum / PL ordinário médio (% a.a.)', 'Return on common equity',
    lambda x, p, c: f'=IFERROR({x}{MR["nicom"]}*4/{x}{MR["ceq_avg"]},0)', lambda x, a, z, c: f'=IFERROR({x}{MR["nicom"]}/{x}{MR["ceq_avg"]},0)', FM(162))
put(MR['rotce'], 'ROTCE — lucro ao acionista comum / TCE médio (% a.a.)', 'Return on tangible common equity',
    lambda x, p, c: f'=IFERROR({x}{MR["nicom"]}*4/{x}{MR["tce_avg"]},0)', lambda x, a, z, c: f'=IFERROR({x}{MR["nicom"]}/{x}{MR["tce_avg"]},0)', FM(162))
put(MR['cet1r_ratio'], 'Índice CET1 — Standardized', 'CET1 ratio', lambda x, p, c: f'={x}{MR["cet1r"]}', lambda x, a, z, c: f'={x}{MR["cet1r"]}', FM(193))
put(MR['tbvps'], '="Valor patrimonial tangível por ação — TBVPS ("&Inputs!$E$200&")"', 'Tangible book value per share',
    lambda x, p, c: f'=IFERROR({x}{MR["tce"]}/{x}{FM(199)},0)', lambda x, a, z, c: f'=IFERROR({x}{MR["tce"]}/{x}{FM(199)},0)', FM(201))

# fórmulas existentes ajustadas
for c in range(5, 60):
    x = L(c); p = prevq(c); P = L(p) if p else None
    q = (c - 5) % 5 != 4
    if q:
        m.cell(FM(71), c).value = f'={x}{FM(72)}+{x}{FM(73)}+{x}{FM(74)}+{x}{FM(75)}+{x}{FM(76)}+{x}{MR["pref"]}'
        if P: m.cell(FM(76), c).value = f"=IF({x}$14=1,'Inputs (2)'!{x}{F2(64)},{P}{FM(76)}+{x}{MR['buy']})"
        m.cell(FM(84), c).value = f"=IF({x}$14=1,-N('Inputs (2)'!{x}{F2(149)})+{x}{MR['prefdiv']},-MAX(0,{x}{MR['nicom']})*Inputs!{x}${RD_PAY}+{x}{MR['prefdiv']})"
        m.cell(FM(87), c).value = f'={x}{FM(78)}-({x}{FM(82)}+{x}{FM(83)}+{x}{FM(84)}+{x}{MR["buy"]}+{x}{FM(85)}+{x}{FM(86)})'
    m.cell(FM(200), c).value = f'=IFERROR({x}{MR["nicom"]}/{x}{FM(199)},0)'
    m.cell(FM(201), c).value = f'=IFERROR({x}{MR["ceq"]}/{x}{FM(199)},0)'
    m.cell(FM(202), c).value = f'=IFERROR(-({x}{FM(84)}-{x}{MR["prefdiv"]})/{x}{FM(199)},0)'
    m.cell(FM(203), c).value = f'=IFERROR(-({x}{FM(84)}-{x}{MR["prefdiv"]}+{x}{MR["buy"]})/{x}{MR["nicom"]},0)'

# Valuation: lucro e PL ordinários, retorno total ao acionista, P/TBV e ROTCE
FYM = {2020 + i: L(9 + 5 * i) for i in range(11)}
for i, y in enumerate(range(2020, 2031)):
    vc, mc = L(5 + i), FYM[y]
    val[f'{vc}29'] = f'=Model!{mc}{MR["nicom"]}'
    val[f'{vc}30'] = f'=Model!{mc}{MR["ceq"]}'
    val[f'{vc}31'] = f'={L(4 + i)}30' if i else None
    val[f'{vc}33'] = f'=-Model!{mc}{FM(84)}+Model!{mc}{MR["prefdiv"]}-Model!{mc}{MR["buy"]}'
    val[f'{vc}143'] = f'=IFERROR($E$5/Model!{mc}{MR["tbvps"]},0)'
    val[f'{vc}144'] = f'=Model!{mc}{MR["rotce"]}'
    val[f'{vc}143'].number_format = val[f'{vc}140'].number_format; val[f'{vc}144'].number_format = '0.0%'
    for rr in (143, 144):
        if val[f'{vc}140'].has_style: val[f'{vc}{rr}']._style = copy.copy(val[f'{vc}140']._style)
val['E31'] = None
val['C143'], val['D143'] = 'P/TBV (x)', 'Price to tangible book'
val['C144'], val['D144'] = 'ROTCE do exercício', 'Return on tangible common equity'
val['E144'].number_format = '0.0%'
for c in range(5, 16): val.cell(144, c).number_format = '0.0%'
val['C97'] = '="Valor por ação via P/TBV ("&Inputs!$E$200&")"'
val['E95'] = f'=IFERROR(INDEX(Model!{FYM[2020]}{MR["tbvps"]}:{FYM[2030]}{MR["tbvps"]},0),0)'
# TBVPS do ano-alvo para o múltiplo (P/TBV é a métrica de mercado nos EUA)
val['E95'] = '=IFERROR(INDEX($E$145:$O$145,MATCH($E$25,$E$27:$O$27,0)),0)'
val['C95'] = '="TBVPS do ano-alvo ("&Inputs!$E$200&")"'
for i, y in enumerate(range(2020, 2031)):
    val[f'{L(5 + i)}145'] = f'=Model!{FYM[y]}{MR["tbvps"]}'
val['C145'], val['D145'] = '="TBVPS ("&Inputs!$E$200&")"', 'TBVPS'

# Glossário
G = [
 ('Rentabilidade', None, None, None),
 ('ROTCE — Return on Tangible Common Equity', 'Lucro ao acionista comum ÷ PL tangível ordinário médio (PL ordinário – goodwill e intangíveis ex-MSR), anualizado', f'Model, linha {MR["rotce"]}', 'Métrica-chave dos bancos dos EUA: tira o efeito de aquisições e mostra o retorno sobre o capital que de fato absorve perdas.'),
 ('ROE ordinário', 'Lucro ao acionista comum ÷ PL ordinário médio (ex-preferenciais), anualizado', f'Model, linha {MR["roe_c"]}', 'Como o ROE brasileiro, mas sem as preferenciais no denominador e com o dividendo preferencial fora do lucro.'),
 ('ROA — Return on Assets', 'Lucro líquido ÷ ativo total médio, anualizado', f'Model, linha {FM(163)}', 'Rentabilidade antes do efeito da alavancagem.'),
 ('NII e NIM', 'Receita líquida de juros ÷ ativos rentáveis médios, anualizada', f'Model, linha {FM(164)}', 'Nos EUA o resultado de trading fica fora do NII (em principal transactions). Bancos com mesa grande divulgam o NII ex-Markets, que é o número comparável.'),
 ('NII ex-Markets', 'NII total – NII da área de Markets (reportado no earnings supplement)', f'Model, linha {FM(138)}', 'Isola a margem do balanço bancário (empréstimos e depósitos) da carteira de trading, cujo NII oscila com o funding das posições.'),
 ('Eficiência', None, None, None),
 ('Overhead ratio', 'Despesas não-juros ÷ receita líquida total', f'Model, linha {FM(170)}', 'Equivalente ao índice de eficiência brasileiro, com a receita líquida (NII + tarifas + trading + outras) no denominador.'),
 ('Receita líquida total', 'NII + tarifas e comissões + principal transactions + outras receitas não-juros', f'Model, linha {FM(140)}', 'É o "total net revenue" do release.'),
 ('Qualidade de crédito (CECL)', None, None, None),
 ('CECL', 'Perda esperada para toda a vida do empréstimo, reconhecida na originação (ASC 326)', f'Model, linha {FM(145)}', 'Não há estágios como na Resolução 4.966: a provisão cobre a vida inteira desde o início.'),
 ('Nonaccrual', 'Empréstimos que deixaram de reconhecer juros ÷ empréstimos', f'Model, linha {FM(183)}', 'Equivale ao NPL. Cartão de crédito normalmente não entra em nonaccrual: é baixado com 180 dias de atraso.'),
 ('Vencidos 30-89 e 90+ ainda acumulando', 'Saldos vencidos que ainda reconhecem juros ÷ empréstimos', f'Model, linhas {FM(181)} e {FM(182)}', 'Indicadores antecedentes, usados sobretudo em cartão.'),
 ('Net charge-offs (NCO)', 'Baixas – recuperações no período ÷ empréstimos médios, anualizado', f'Model, linha {FM(190)}', 'Perda efetivamente realizada. Provisão acima do NCO = reserve build; abaixo = reserve release.'),
 ('Custo do crédito', 'Provisão para perdas de crédito ÷ empréstimos médios, anualizado', f'Model, linha {FM(189)}', 'Inclui a provisão de compromissos e de títulos quando a companhia divulga só o total.'),
 ('Cobertura', 'ACL ÷ nonaccrual; ACL ÷ empréstimos', f'Model, linhas {FM(185)} e {FM(188)}', 'Colchão de provisão frente à carteira problemática e ao total da carteira.'),
 ('Capital e liquidez (Basileia III dos EUA)', None, None, None),
 ('CET1 — Standardized e Advanced', 'Capital principal ÷ RWA pelas duas abordagens; vale a mais restritiva', f'Model, linhas {MR["cet1r"]} e {MR["cet1r_adv"]}', 'Mínimo = 4,5% + stress capital buffer (SCB) + sobretaxa G-SIB. É a restrição que define recompras.'),
 ('SLR — Supplementary leverage ratio', 'Tier 1 ÷ exposição total (inclui fora do balanço)', f'Model, linha {MR["slr"]}', 'Restrição de alavancagem dos grandes bancos (mínimo de 3% + buffer de 2% para G-SIBs).'),
 ('LCR', 'Ativos líquidos de alta qualidade ÷ saídas de caixa em 30 dias de estresse', f'Model, linha {MR["lcr"]}', 'Mínimo de 100%.'),
 ('Valuation', None, None, None),
 ('Ke — custo do capital próprio', 'Treasury 10 anos + beta × prêmio de risco de mercado dos EUA', 'Inputs, premissas de valuation', 'Em banco não se usa WACC: a dívida é matéria-prima do negócio.'),
 ('Lucro residual (Excess Return)', 'Lucro ao acionista comum – Ke × PL ordinário inicial', 'Valuation, Método 1', 'Parcela do lucro acima da remuneração exigida pelo capital ordinário.'),
 ('DDM com retorno total', 'Dividendos ordinários + recompras, descontados a Ke', 'Valuation, Método 2', 'Bancos dos EUA devolvem a maior parte do capital via recompra. Só dividendos subestimaria o valor.'),
 ('P/TBV e P/BV justificado', '(ROTCE sustentável – g) ÷ (Ke – g)', 'Valuation, Métodos 3 e 4', 'O mercado precifica bancos dos EUA por P/TBV contra ROTCE.'),
]
for r in range(6, 40):
    for c in range(3, 7): glo.cell(r, c).value = None
for i, (a, b2, c2, d2) in enumerate(G):
    r = 6 + i
    src = 6 if b2 is None else 7
    for c in range(3, 7):
        if glo.cell(src, c).has_style: glo.cell(r, c)._style = copy.copy(glo.cell(src, c)._style)
    glo.cell(r, 3).value, glo.cell(r, 4).value, glo.cell(r, 5).value, glo.cell(r, 6).value = a, b2, c2, d2
glo['C2'] = 'GLOSSÁRIO DE INDICADORES — BANCOS DOS EUA'
wb.save(OUT)
print('ok', OUT, MR, R2, {'RI_BUY': RI_BUY, 'RD_BUY': RD_BUY, 'RD_CET1': RD_CET1})
