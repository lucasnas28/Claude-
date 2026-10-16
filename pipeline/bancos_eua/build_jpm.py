"""Preenche a base de bancos dos EUA com a JPMorgan Chase (SEC XBRL, 10-Q/10-K, 1T20..2T26).
uso: build_jpm.py BASE_EUA.xlsx SAIDA.xlsx PASTA_DL   (PASTA_DL = saída do sec/dl.sh para o CIK 0000019617)
Valores gravados em US$ milhões (fator 1 na Inputs). Fluxos: trimestre isolado (3M ou YTD - YTD anterior, par consistente)."""
import sys, os, datetime as dt, collections
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'sec'))
from sx import Co, P, label, kind_of
import openpyxl
from openpyxl.utils import get_column_letter as L

BASE, OUT, DL = sys.argv[1:4]
co = Co(DL)
wb = openpyxl.load_workbook(BASE)
inp, w2, m = wb['Inputs'], wb['Inputs (2)'], wb['Model']
COL = {}
for c in range(5, 80):
    h = inp.cell(2, c).value
    if isinstance(h, str) and len(h) == 4 and h[1] == 'T': COL[f'{h[0]}Q{h[2:]}'] = c
assert '1Q20' in COL and '2Q26' in COL, COL
ROW = {inp.cell(r, 2).value: r for r in range(5, 191) if inp.cell(r, 2).value}
SRC = 'SEC XBRL (10-Q/10-K JPMorgan Chase, CIK 19617)'
M = 1e-6
written = collections.Counter()

def put(code, vals, note, scale=M, sign=1, fmt=None):
    r = ROW[code]
    inp.cell(r, 4).value = note
    for p, v in vals.items():
        if p in COL and v is not None:
            inp.cell(r, COL[p]).value = round(sign * v * scale, 6); written[code] += 1
def formula_row(code, expr, note):
    """linha derivada: expr usa {memo} -> linha de memória na mesma coluna"""
    r = ROW[code]; inp.cell(r, 4).value = note
    for p in P:
        x = L(COL[p]); inp.cell(r, COL[p]).value = '=' + expr.format(x=x)
def empty(code, why): inp.cell(ROW[code], 4).value = f'[vazio] {why}'

MEMO0 = max(r for r in range(151, 191) if inp.cell(r, 2).value) + 2
memo_r = [MEMO0]
inp.cell(MEMO0 - 1, 3).value = None
def memo(desc, vals, note, scale=M):
    r = memo_r[0]; memo_r[0] += 1
    inp.cell(r, 3).value = desc; inp.cell(r, 4).value = note
    for p, v in vals.items():
        if p in COL and v is not None: inp.cell(r, COL[p]).value = round(v * scale, 6)
    return r

# ---------------------------------------------------------------- séries
def q(*taglists, orig=False):
    """soma de várias séries trimestrais; cada item é uma lista de tags alternativas. Período só entra se todas existirem."""
    parts = [co.quarter(list(t), orig=orig) for t in taglists]
    return {p: sum(x[p] for x in parts) for p in P if all(p in x for x in parts)}
def inst(tags, members=()):
    want = tuple(sorted(mm.split(':')[-1] for mm in members)); out = {}
    for t in tags:
        got = {}
        for (n, ns, s, e, dd, v, u, a, fl, i) in co.facts:
            if n != t or s is not None: continue
            mem = tuple(sorted(mm.split(':')[-1] for ax, mm in dd))
            if mem != want: continue
            lab = label(e)
            if lab in P and (lab not in got or fl > got[lab][1]): got[lab] = (v, fl)
        for k, (v, fl) in got.items(): out.setdefault(k, v)
    return out

# ---------------------------------------------------------------- balanço (face do próprio 10-Q/10-K)
BSMAP = {
 'CashAndDueFromBanks': 'cash', 'InterestBearingDepositsInBanks': 'dep_banks', 'FederalFundsSoldAndSecuritiesPurchasedUnderAgreementsToResell': 'resale',
 'SecuritiesBorrowed': 'sec_borrowed', 'TradingAssets': 'trading',
 'AvailableForSaleSecuritiesDebtSecurities': 'afs', 'DebtSecuritiesAvailableForSaleExcludingAccruedInterest': 'afs',
 'DebtSecuritiesHeldToMaturityNetOfAllowanceForCreditLosses': 'htm', 'DebtSecuritiesHeldToMaturityExcludingAccruedInterestAfterAllowanceForCreditLoss': 'htm',
 'FinancingReceivableBeforeAllowanceForCreditLossesNetofDeferredIncome': 'loans', 'FinancingReceivableExcludingAccruedInterestBeforeAllowanceForCreditLossesNetOfDeferredIncome': 'loans',
 'FinancingReceivableAllowanceForCreditLosses': 'acl', 'FinancingReceivableAllowanceForCreditLossExcludingAccruedInterest': 'acl',
 'AccruedInterestAndAccountsReceivable': 'receivables', 'PropertyPlantAndEquipmentNet': 'premises',
 'PropertyPlantAndEquipmentAndOperatingLeaseRightOfUseAssetAfterAccumulatedDepreciationAndAmortization': 'premises',
 'GoodwillServicingAssetsatFairValueandOtherIntangibleAssets': 'gw_intang', 'GoodwillServicingAssetsAtFairValueAndOtherIntangibleAssets': 'gw_intang',
 'OtherAssets': 'other_a',
 'Deposits': '_deposits', 'FederalFundsPurchasedSecuritiesSoldUnderAgreementsToRepurchase': 'repo', 'FederalFundsPurchasedAndSecuritiesSoldUnderAgreementsToRepurchase': 'repo',
 'ShortTermBorrowings': 'st_borr', 'TradingLiabilities': 'trading_l', 'AccountsPayableAndAccruedLiabilitiesCurrentAndNoncurrent': 'other_l',
 'BeneficialInterest': 'vie', 'LongTermDebtAndCapitalLeaseObligationsIncludingCurrentMaturities': 'ltd',
 'PreferredStockIncludingAdditionalPaidInCapitalNetOfDiscount': 'pref', 'CommonStockValue': 'common_stk', 'AdditionalPaidInCapitalCommonStock': 'apic',
 'CommonStockHeldInTrust': 'apic', 'RetainedEarningsAccumulatedDeficit': 'ret_earn', 'AccumulatedOtherComprehensiveIncomeLossNetOfTax': 'aoci',
 'TreasuryStockValue': 'treasury', 'TreasuryStockCommonValue': 'treasury',
}
NEG = {'acl', 'treasury'}
NEGTAG = {'CommonStockHeldInTrust'}
SKIP = {'AvailableForSaleDebtSecuritiesAmortizedCostBasis', 'DebtSecuritiesAvailableForSaleAmortizedCostAfterAllowanceForCreditLoss',
        'DebtSecuritiesAvailableForSaleAmortizedCostExcludingAccruedInterestAfterAllowanceForCreditLoss', 'DebtSecuritiesNetCarryingAmount',
        'DebtSecuritiesHeldToMaturityAllowanceForCreditLoss', 'NotesReceivableNet', 'FinancingReceivableExcludingAccruedInterestAfterAllowanceForCreditLoss',
        'PreferredStockParOrStatedValuePerShare', 'CommonStockParOrStatedValuePerShare', 'Assets', 'Liabilities', 'StockholdersEquity', 'LiabilitiesAndStockholdersEquity'}
bs = collections.defaultdict(dict); tot = collections.defaultdict(dict); unk = []
for p in P:
    face, acc = co.face(p)
    for n, v in face:
        if n in ('Assets', 'Liabilities', 'StockholdersEquity', 'LiabilitiesAndStockholdersEquity'): tot[n][p] = v
        if n in SKIP: continue
        code = BSMAP.get(n)
        if code is None: unk.append((p, n, v)); continue
        sgn = -1 if (code in NEG or n in NEGTAG) else 1
        bs[code][p] = bs[code].get(p, 0) + sgn * abs(v) if code in NEG or n in NEGTAG else bs[code].get(p, 0) + v
assert not unk, unk[:10]
ASSETS = ['cash', 'dep_banks', 'resale', 'sec_borrowed', 'trading', 'afs', 'htm', 'loans', 'acl', 'receivables', 'premises', 'gw_intang', 'other_a']
LIABS = ['_deposits', 'repo', 'st_borr', 'trading_l', 'other_l', 'vie', 'ltd']
EQ = ['pref', 'common_stk', 'apic', 'ret_earn', 'aoci', 'treasury']
for p in P:
    a = sum(bs[k].get(p, 0) for k in ASSETS); l = sum(bs[k].get(p, 0) for k in LIABS); e = sum(bs[k].get(p, 0) for k in EQ)
    assert abs(a - tot['Assets'][p]) < 1e6 and abs(l - tot['Liabilities'][p]) < 1e6 and abs(e - tot['StockholdersEquity'][p]) < 1e6, (p, a - tot['Assets'][p], l - tot['Liabilities'][p], e - tot['StockholdersEquity'][p])
FN = f'{SRC}, balanço patrimonial consolidado (face)'
desc = {'cash': 'Cash and due from banks', 'dep_banks': 'Deposits with banks', 'resale': 'Federal funds sold and securities purchased under resale agreements',
        'sec_borrowed': 'Securities borrowed', 'trading': 'Trading assets (inclui derivativos a receber)', 'afs': 'Available-for-sale securities (valor justo)',
        'htm': 'Held-to-maturity securities (líquidos de provisão)', 'loans': 'Loans (líquidos de receitas diferidas)', 'acl': 'Allowance for loan losses (negativo)',
        'receivables': 'Accrued interest and accounts receivable', 'premises': 'Premises and equipment (inclui direito de uso)',
        'gw_intang': 'Goodwill, MSRs and other intangible assets', 'other_a': 'Other assets', 'repo': 'Federal funds purchased and securities loaned or sold under repurchase agreements',
        'st_borr': 'Short-term borrowings', 'trading_l': 'Trading liabilities (inclui derivativos a pagar)', 'other_l': 'Accounts payable and other liabilities',
        'vie': 'Beneficial interests issued by consolidated VIEs', 'ltd': 'Long-term debt', 'pref': 'Preferred stock', 'common_stk': 'Common stock',
        'apic': 'Additional paid-in capital (menos ações em trust de RSU, quando apresentadas)', 'ret_earn': 'Retained earnings', 'aoci': 'Accumulated other comprehensive income (loss)',
        'treasury': 'Treasury stock, at cost (negativo)'}
for k, d in desc.items(): put(k, bs[k], f'{FN}: {d} · fonte · moeda')
r_dep = memo('[memória] Depósitos totais (face do balanço)', bs['_deposits'], f'{FN}: Deposits · memoria · moeda')
nib = {p: v + inst(['NoninterestBearingDepositLiabilitiesForeign']).get(p, 0) for p, v in inst(['NoninterestBearingDepositLiabilitiesDomestic']).items()}
put('dep_nib', nib, f'{SRC}, nota de depósitos: noninterest-bearing (EUA + fora dos EUA) · fonte · moeda')
formula_row('dep_ib', f'{{x}}{r_dep}-N({{x}}{ROW["dep_nib"]})', f'Depósitos com juros = depósitos totais (linha {r_dep}) – depósitos sem juros · derivada · moeda')
for k, why in [('liq_acl', 'não apresentada à parte'), ('sec_acl', 'HTM já vem líquido de provisão na face; AFS a valor justo'), ('deriv_a', 'dentro de trading assets'),
               ('loans_fv', 'dentro de loans'), ('dta', 'dentro de other assets'), ('eq_inv', 'dentro de other assets'), ('inv_out', 'dentro de other assets'),
               ('dep_x1', 'não se aplica'), ('dep_time', 'dentro de depósitos com juros'), ('dep_other', 'não se aplica'), ('sec_loaned', 'dentro de repos'),
               ('ltd_x', 'não se aplica'), ('sub_debt', 'dentro de long-term debt'), ('fhlb', 'dentro de long-term debt / short-term borrowings'), ('borr_x', 'não se aplica'),
               ('brokers_l', 'dentro de accounts payable and other liabilities'), ('lrc_liab', 'dentro de accounts payable and other liabilities'), ('ins_res', 'não se aplica'),
               ('prov_other', 'dentro de accounts payable and other liabilities'), ('dtl', 'dentro de accounts payable and other liabilities'), ('nci', 'não apresentada à parte (imaterial)')]:
    empty(k, why)

# ---------------------------------------------------------------- DRE (trimestre isolado)
FI = f'{SRC}, demonstração do resultado consolidada'
r_ii = memo('[memória] Receita de juros total (total interest income)', q(['InterestIncomeOperating', 'InterestAndDividendIncomeOperating']), f'{FI} · memoria · moeda')
r_ie = memo('[memória] Despesa de juros total (total interest expense)', q(['InterestExpenseOperating', 'InterestExpense']), f'{FI} · memoria · moeda')
r_nir = memo('[memória] Receita não-juros total (total noninterest revenue)', q(['NoninterestIncome']), f'{FI} · memoria · moeda')
r_nie = memo('[memória] Despesa não-juros total (total noninterest expense)', q(['NoninterestExpense']), f'{FI} · memoria · moeda')
r_rev = memo('[memória] Receita líquida total (total net revenue)', q(['RevenuesNetOfInterestExpense', 'Revenues']), f'{FI} · memoria · moeda')
r_ni = memo('[memória] Lucro líquido (net income)', q(['NetIncomeLoss', 'ProfitLoss']), f'{FI} · memoria · moeda')
r_nic = memo('[memória] Lucro líquido ao acionista comum (net income applicable to common stockholders)', q(['NetIncomeLossAvailableToCommonStockholdersBasic']), f'{FI} · memoria · moeda')
put('ii_loans', q(['InterestAndFeeIncomeLoansAndLeases']), f'{FI}: Loans · fonte · moeda')
put('ii_sec', q(['InterestIncomeSecuritiesTaxable'], ['InterestIncomeSecuritiesTaxExempt'], ['InterestIncomeTradingAssets']), f'{FI}: Securities (taxable + tax-exempt) + trading assets · fonte · moeda')
put('ii_repo', q(['InterestIncomeFederalFundsSoldAndSecuritiesPurchasedUnderAgreementsToResell'], ['InterestIncomeSecuritiesBorrowed']), f'{FI}: Federal funds sold and resale agreements + securities borrowed · fonte · moeda')
formula_row('ii_other', f'{{x}}{r_ii}-N({{x}}{ROW["ii_loans"]})-N({{x}}{ROW["ii_sec"]})-N({{x}}{ROW["ii_repo"]})',
            f'Receita de juros total (linha {r_ii}) – empréstimos – títulos – compromissadas: depósitos em bancos e outros · derivada · moeda')
put('ie_dep', q(['InterestExpenseDeposits']), f'{FI}: Interest-bearing deposits (negativo) · fonte · moeda', sign=-1)
put('ie_ltd', q(['InterestExpenseLongTermDebt'], ['InterestExpenseBeneficialInterestsIssuedByConsolidatedVariableInterestEntities']), f'{FI}: Long-term debt + beneficial interests of VIEs (negativo) · fonte · moeda', sign=-1)
formula_row('ie_st', f'-{{x}}{r_ie}-N({{x}}{ROW["ie_dep"]})-N({{x}}{ROW["ie_ltd"]})',
            f'–Despesa de juros total (linha {r_ie}) – depósitos – dívida LP e VIEs: compromissadas, curto prazo, trading liabilities e outros · derivada · moeda')
for k in ('ie_other', 'recoveries', 'legal', 'prov_x1', 'prov_x2', 'fdic', 'eq_income', 'nonop', 'tax_def', 'nci_ni'):
    empty(k, {'ie_other': 'dentro de ie_st', 'recoveries': 'a provisão dos EUA já é líquida de recuperações (ver nco)', 'legal': 'dentro de outras despesas não-juros',
              'prov_x1': 'não se aplica', 'prov_x2': 'não se aplica', 'fdic': 'dentro de outras despesas não-juros', 'eq_income': 'dentro de outras receitas',
              'nonop': 'não se aplica', 'tax_def': 'imposto total em tax', 'nci_ni': 'imaterial; o lucro reportado já é o atribuível à companhia'}[k])
put('provision', q(['ProvisionForLoanLeaseAndOtherLosses', 'ProvisionForCreditLosses', 'ProvisionForLoanAndLeaseLosses']), f'{FI}: Provision for credit losses (negativo; inclui compromissos e títulos) · fonte · moeda', sign=-1)
FEES = [('fee_ib', [['InvestmentBankingRevenue']], 'Investment banking fees'), ('fee_lend_dep', [['LendingAndDepositRelatedFees']], 'Lending- and deposit-related fees'),
        ('fee_am', [['AssetManagementFees', 'InvestmentAdvisoryManagementAndAdministrativeFees', 'AssetManagementFees1']], 'Asset management fees'),
        ('fee_comm', [['FeesAndCommissions1', 'FeesAndCommissionsOther', 'FeesAndCommissions']], 'Commissions and other fees'),
        ('fee_card', [['FeesAndCommissionsCreditAndDebitCards1', 'FeesAndCommissionsCreditAndDebitCards']], 'Card income'),
        ('fee_mort', [['MortgageFeesAndRelatedIncome', 'FeesAndCommissionsMortgageBanking']], 'Mortgage fees and related income')]
for code, tags, d in FEES: put(code, q(*tags), f'{FI}: {d} · fonte · moeda')
formula_row('fees', '+'.join(f'N({{x}}{ROW[c]})' for c, *_ in FEES), 'Soma das tarifas abertas (fee_ib + fee_lend_dep + fee_am + fee_comm + fee_card + fee_mort) · derivada · moeda')
put('principal', q(['PrincipalTransactionsRevenue', 'TradingGainsLosses']), f'{FI}: Principal transactions (resultado de trading) · fonte · moeda')
formula_row('other_inc', f'{{x}}{r_nir}-N({{x}}{ROW["fees"]})-N({{x}}{ROW["principal"]})',
            f'Receita não-juros total (linha {r_nir}) – tarifas – principal transactions: investment securities gains (losses) + other income · derivada · moeda')
put('comp', q(['LaborAndRelatedExpense']), f'{FI}: Compensation expense (negativo) · fonte · moeda', sign=-1)
OPX = [('occupancy', [['OccupancyNet']], 'Occupancy expense'), ('tech', [['CommunicationsAndInformationTechnology', 'TechnologyCommunicationsAndEquipmentExpense', 'InformationTechnologyAndDataProcessing']], 'Technology, communications and equipment'),
       ('prof_serv', [['ProfessionalAndContractServicesExpense']], 'Professional and outside services'), ('marketing', [['MarketingAndAdvertisingExpense', 'MarketingExpense']], 'Marketing')]
for code, tags, d in OPX: put(code, q(*tags), f'{FI}: {d} (negativo) · fonte · moeda', sign=-1)
formula_row('opex', '+'.join(f'N({{x}}{ROW[c]})' for c, *_ in OPX), 'Ocupação + tecnologia + serviços profissionais + marketing · derivada · moeda')
formula_row('other_exp', f'-{{x}}{r_nie}-N({{x}}{ROW["comp"]})-N({{x}}{ROW["opex"]})',
            f'–Despesa não-juros total (linha {r_nie}) – remuneração – opex aberto: other expense (inclui jurídico e FDIC) · derivada · moeda')
put('tax', q(['IncomeTaxExpenseBenefit']), f'{FI}: Income tax expense (negativo) · fonte · moeda', sign=-1)
formula_row('div_pref', f'{{x}}{r_nic}-{{x}}{r_ni}', f'Lucro ao acionista comum (linha {r_nic}) – lucro líquido (linha {r_ni}): dividendos preferenciais e alocação a títulos participantes · derivada · moeda')
r_chk = memo_r[0]; memo_r[0] += 1
inp.cell(r_chk, 3).value = 'CHECK: soma da DRE codificada – lucro líquido reportado (deve ser zero)'
IS = ['ii_loans', 'ii_sec', 'ii_repo', 'ii_other', 'ie_dep', 'ie_ltd', 'ie_st', 'provision', 'fees', 'principal', 'other_inc', 'comp', 'opex', 'other_exp', 'tax']
for p in P:
    x = L(COL[p]); inp.cell(r_chk, COL[p]).value = f'=ROUND({"+".join(f"N({x}{ROW[c]})" for c in IS)}-{x}{r_ni},0)'
for k in ('comp_sal', 'comp_ben', 'comp_inc', 'comp_oth', 'adm_da', 'adm_oth'): empty(k, 'não aberto no XBRL')

# ---------------------------------------------------------------- crédito, capital e ações
put('loans_tot', inst(['FinancingReceivableExcludingAccruedInterestBeforeAllowanceForCreditLoss', 'FinancingReceivableBeforeAllowanceForCreditLossesNetofDeferredIncome']), f'{SRC}, nota de empréstimos: loans retained (base da ACL; exclui empréstimos mantidos para venda e a valor justo, que estão no balanço) · fonte · moeda')
for code, mem, d in [('loans_cons', 'us-gaap:ConsumerExcludingCreditCardLoanPortfolioSegmentMember', 'Consumer, excluding credit card'),
                     ('loans_card', 'CreditCardReceivablesMember', 'Credit card'), ('loans_wh', 'us-gaap:CommercialPortfolioSegmentMember', 'Wholesale')]:
    put(code, inst(['FinancingReceivableExcludingAccruedInterestBeforeAllowanceForCreditLoss'], (mem,)), f'{SRC}, nota de empréstimos: {d} (retidos) · fonte · moeda')
acl = inst(['FinancingReceivableAllowanceForCreditLossExcludingAccruedInterest', 'FinancingReceivableAllowanceForCreditLosses'])
put('acl_tot', acl, f'{SRC}: allowance for loan losses (saldo positivo) · fonte · moeda')
put('acl_cons', inst(['FinancingReceivableAllowanceForCreditLossExcludingAccruedInterest'], ('us-gaap:ConsumerExcludingCreditCardLoanPortfolioSegmentMember',)), f'{SRC}: allowance — consumer ex-card · fonte · moeda')
afs_a, htm_a = inst(['DebtSecuritiesAvailableForSaleAmortizedCostAllowanceForCreditLossExcludingAccruedInterest']), inst(['DebtSecuritiesHeldToMaturityAllowanceForCreditLossExcludingAccruedInterest'])
put('acl_secs', {p: afs_a.get(p, 0) + htm_a.get(p, 0) for p in set(afs_a) | set(htm_a)}, f'{SRC}: allowance de títulos AFS + HTM (desde 1T23) · fonte · moeda')
put('nco', q(['FinancingReceivableExcludingAccruedInterestAllowanceForCreditLossWriteoffAfterRecovery', 'FinancingReceivableAllowanceForCreditLossWriteoffAfterRecovery']),
    f'{SRC}, roll-forward da ACL: net charge-offs (baixas – recuperações) · fonte · moeda')
CARD = 'CreditCardReceivablesMember'
put('dpd90', inst(['FinancingReceivableExcludingAccruedInterestBeforeAllowanceForCreditLoss'], (CARD, 'FinancialAsset90orMoreDaysPastDueandStillAccruingMember')),
    f'{SRC}, nota de empréstimos: cartão de crédito vencido 90+ dias ainda acumulando juros (SÓ CARTÃO; consumo ex-cartão e atacado vão para nonaccrual) · fonte · moeda')
put('dpd30', inst(['FinancingReceivableExcludingAccruedInterestBeforeAllowanceForCreditLoss'], (CARD, 'FinancialAsset30To89DaysPastDueAndStillAccruingMember')),
    f'{SRC}, nota de empréstimos: cartão de crédito vencido 30-89 dias (SÓ CARTÃO) · fonte · moeda')
for k, why in [('nonaccrual', 'o XBRL só etiqueta o nonaccrual por portfólio e de forma incompleta; total no earnings supplement'), ('npa', 'earnings supplement'), ('mods', 'não incluído'), ('acl_lrc', 'earnings supplement'),
               ('nii_exmkts', 'earnings supplement (NII ex-Markets)'), ('nii_mkts', 'earnings supplement'), ('iea', 'earnings supplement (saldos médios)'), ('ibl', 'idem'),
               ('rwa_cred', 'Pillar 3'), ('rwa_mkt', 'Pillar 3'), ('rwa_op', 'Pillar 3'), ('lcr', 'earnings supplement / LCR disclosure (não etiquetado no 10-Q)'), ('nsfr', 'idem'),
               ('price', 'preencher (cotação)'), ('headcount', 'earnings supplement'), ('branches', 'earnings supplement'), ('clients', 'não divulgado de forma comparável')]:
    empty(k, why)
PC, ST, AD = 'srt:ParentCompanyMember', 'us-gaap:BaselIIIStandardizedMember', 'us-gaap:BaselIIIAdvancedMember'
CAPN = f'{SRC}, nota de capital regulatório (JPMorgan Chase & Co., consolidado)'
put('rwa', inst(['RiskWeightedAssets'], (PC, ST)), f'{CAPN}: RWA Standardized · fonte · moeda')
put('rwa_adv', inst(['RiskWeightedAssets'], (PC, AD)), f'{CAPN}: RWA Advanced · fonte · moeda')
put('total_cap', inst(['Capital'], (PC, ST)), f'{CAPN}: total capital (Standardized) · fonte · moeda')
put('tier1', inst(['TierOneRiskBasedCapital'], (PC, ST)), f'{CAPN}: Tier 1 capital · fonte · moeda')
put('cet1', inst(['CommonEquityTier1Capital'], (PC, ST)), f'{CAPN}: CET1 capital · fonte · moeda')
put('tcr_rep', inst(['CapitalToRiskWeightedAssets'], (PC, ST)), f'{CAPN}: total capital ratio Standardized · fonte · pct', scale=1)
put('t1r_rep', inst(['TierOneRiskBasedCapitalToRiskWeightedAssets'], (PC, ST)), f'{CAPN}: Tier 1 ratio Standardized · fonte · pct', scale=1)
put('cet1r_rep', inst(['CommonEquityTier1CapitaltoRiskWeightedAssets'], (PC, ST)), f'{CAPN}: CET1 ratio Standardized · fonte · pct', scale=1)
put('cet1r_adv', inst(['CommonEquityTier1CapitaltoRiskWeightedAssets'], (PC, AD)), f'{CAPN}: CET1 ratio Advanced · fonte · pct', scale=1)
t1l = inst(['TierOneLeverageCapitalToAverageAssets'], (PC,)); t1l.update({p: v for p, v in inst(['TierOneLeverageCapitalToAverageAssets'], (PC, ST)).items() if p not in t1l})
slr = inst(['SupplementaryLeverageRatio'], (PC,)); slr.update({p: v for p, v in inst(['SupplementaryLeverageRatio'], (PC, AD)).items() if p not in slr})
put('t1lev', t1l, f'{CAPN}: Tier 1 leverage ratio · fonte · pct', scale=1)
put('slr', slr, f'{CAPN}: supplementary leverage ratio · fonte · pct', scale=1)
put('msr', inst(['ServicingAssetAtFairValueAmount']), f'{SRC}, nota de MSRs: mortgage servicing rights (valor justo) · fonte · moeda')
iss, tre = inst(['CommonStockSharesIssued']), inst(['TreasuryStockCommonShares', 'TreasuryStockShares'])
dei = {}
for (n, ns, s, e, dd, v, u, a, fl, i) in co.facts:
    if n == 'EntityCommonStockSharesOutstanding' and not dd:
        rd = co.filings[a]['rd']; lab = label(rd)
        if lab in P: dei[lab] = v
shares = {p: (iss[p] - tre[p]) if (p in iss and p in tre) else dei.get(p) for p in P}
put('shares', shares, f'{SRC}: ações emitidas – tesouraria no fim do trimestre; antes de 4T21 (tesouraria não etiquetada) usa a capa do 10-Q/10-K (ações em circulação na data de arquivamento) · fonte · acoes', scale=1e-6)
put('div_common', co.dq('DividendsCommonStockCash', ['us-gaap:RetainedEarningsMember']), f'{SRC}, demonstração das mutações do PL: dividendos ordinários declarados · fonte · moeda')
put('buyback', q(['PaymentsForRepurchaseOfCommonStock'], orig=True), f'{SRC}, demonstração dos fluxos de caixa: recompra de ações ordinárias · fonte · moeda')

# ---------------------------------------------------------------- segmentos (receita líquida e lucro, base gerencial)
def seg(tag, alts): return co.dq(tag, alts)
def cib(tag):
    new = seg(tag, ['jpm:CommercialAndInvestmentBankMember'])
    old_cib = seg(tag, ['jpm:CorporateInvestmentBankMember', 'jpm:CorporateAndInvestmentBankMember'])
    old_cb = seg(tag, ['jpm:CommercialBankingMember'])
    return {p: new[p] if p in new else (old_cib[p] + old_cb[p]) for p in P if p in new or (p in old_cib and p in old_cb)}
SEGN = f'{SRC}, nota de segmentos (base gerencial, receita em base FTE)'
put('seg1_rev', seg('RevenuesNetOfInterestExpense', ['jpm:ConsumerCommunityBankingMember']), f'{SEGN}: Consumer & Community Banking · fonte · moeda')
put('seg3_rev', seg('RevenuesNetOfInterestExpense', ['jpm:AssetandWealthManagementSegmentMember']), f'{SEGN}: Asset & Wealth Management · fonte · moeda')
put('seg4_rev', seg('RevenuesNetOfInterestExpense', ['us-gaap:CorporateNonSegmentMember']), f'{SEGN}: Corporate · fonte · moeda')
put('seg_recon', seg('RevenuesNetOfInterestExpense', ['us-gaap:MaterialReconcilingItemsMember']), f'{SEGN}: reconciling items (ajuste FTE) · fonte · moeda')
put('seg1_ni', seg('NetIncomeLoss', ['jpm:ConsumerCommunityBankingMember']), f'{SEGN}: CCB net income · fonte · moeda')
r_snr = memo('[memória] Lucro líquido — itens de reconciliação dos segmentos', seg('NetIncomeLoss', ['us-gaap:MaterialReconcilingItemsMember']), f'{SEGN} · memoria · moeda')
formula_row('seg2_rev', f'{{x}}{r_rev}-N({{x}}{ROW["seg1_rev"]})-N({{x}}{ROW["seg3_rev"]})-N({{x}}{ROW["seg4_rev"]})-N({{x}}{ROW["seg_recon"]})',
            f'Commercial & Investment Bank POR DIFERENÇA: receita líquida total (linha {r_rev}) – CCB – AWM – Corporate – reconciliação. Motivo: a JPM juntou CIB e Commercial Banking em 2024 e reutilizou a etiqueta do CIB antigo para o segmento combinado; somar as etiquetas dupla-conta 2023 · derivada · moeda')
formula_row('seg2_ni', f'{{x}}{r_ni}-N({{x}}{ROW["seg1_ni"]})-N({{x}}{ROW["seg3_ni"]})-N({{x}}{ROW["seg4_ni"]})-N({{x}}{r_snr})',
            f'CIB lucro líquido POR DIFERENÇA: lucro líquido (linha {r_ni}) – CCB – AWM – Corporate – reconciliação (linha {r_snr}) · derivada · moeda')
put('seg3_ni', seg('NetIncomeLoss', ['jpm:AssetandWealthManagementSegmentMember']), f'{SEGN}: AWM net income · fonte · moeda')
put('seg4_ni', seg('NetIncomeLoss', ['us-gaap:CorporateNonSegmentMember']), f'{SEGN}: Corporate net income · fonte · moeda')
SEGNAMES = {'seg1': 'Consumer & Community Banking (CCB)', 'seg2': 'Commercial & Investment Bank (CIB)', 'seg3': 'Asset & Wealth Management (AWM)', 'seg4': 'Corporate'}
for k, nm in SEGNAMES.items():
    for suf, t in (('_rev', 'receita líquida'), ('_ni', 'lucro líquido')):
        inp.cell(ROW[k + suf], 3).value = f'{nm} — {t}'
        for r in range(5, 200):
            if w2.cell(r, 2).value == k + suf: w2.cell(r, 3).value = f'{nm} — {t}'
        for r in range(140, 175):
            if m.cell(r, 3).value == f'Segmento {k[-1]} — {t}': m.cell(r, 3).value = f'{nm} — {t}'

# ---------------------------------------------------------------- painel
if memo_r[0] <= 190:
    inp.cell(memo_r[0] + 1, 3).value = f'Espaço livre até a linha 190 (linhas {memo_r[0] + 1} a 190) para novas linhas de origem.'
inp['E197'] = 'JPMorgan Chase & Co.'; inp['E198'] = 'JPM'
inp['E199'] = 'Banco universal (G-SIB) / US GAAP (CECL). Balanço, DRE, capital e segmentos: 10-Q/10-K via SEC XBRL'
inp['E203'] = dt.datetime(2026, 6, 30); inp['E204'] = dt.datetime(2020, 3, 31)
inp['E203'].number_format = inp['E204'].number_format = 'dd/mm/yyyy'
last = '2Q26'
inp['E207'] = round(iss[last] * 1e-6, 3); inp['E209'] = -round(tre[last] * 1e-6, 3)
inp['F1'] = ('JPMorgan Chase & Co. (JPM). Fonte: instâncias XBRL dos 10-Q e 10-K arquivados na SEC (CIK 19617), 1T20 a 2T26, em US$ milhões. '
             'Trimestres de fluxo: 3M direto ou acumulado do ano menos o acumulado anterior (par consistente). Linhas "derivada" são fórmulas nesta aba. '
             'Métricas que só existem no earnings supplement (NII ex-Markets, nonaccrual total, LCR, saldos médios) ficam vazias com o motivo escrito. Modo histórico: premissas e valuation vazios.')
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('linhas gravadas', dict(written))
